"""COBRANZAS DE SEGUNDA (dueño, 02/10): las reglas, las palabras y los
archivos. Es puro: los lotes llegan armados de `lotes_de_segunda`
(app/db.py).

EL CIRCUITO, dictado por Lionel:

- Cada salida al puesto de segunda (`remitos_segunda`, destino 'puesto', no
  anulada) es un LOTE, y nace pendiente de cobro. La merma de segunda no
  entra: se tiró, no hay nada que cobrar.
- Hay UN SOLO puesto de segunda: no se elige.
- El puesto liquida LOTE POR LOTE. No hay un total de rendición que
  repartir: se carga cuánto pagó cada lote. $0 es un cobro (cobrado en
  cero), no un pendiente.
- "Quién" es el SECTOR y la hora: el sistema no tiene usuarios.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer

from core.exportar_vacios import (
    VERDE_ENCABEZADO_HEX,
    _documento_pdf,
    _estilos_pdf,
    _tabla_seccion_pdf,
)
from core.vales import TEXTO_DEL_SECTOR, hora_argentina

# Un lote con MÁS de estos días sin cobrar dispara la alerta
# `segunda_sin_cobrar`. Corridos, contra la fecha de la salida.
DIAS_SEGUNDA_SIN_COBRAR = 40

TEXTO_DEL_ESTADO = {"pendiente": "Pendientes de cobro", "cobrado": "Cobrados"}
OPCIONES_DE_ESTADO = (("todos", "Todos"), ("pendiente", "Pendientes"), ("cobrado", "Cobrados"))


def estado_del_filtro(texto: str) -> str | None:
    """'pendiente' o 'cobrado'; cualquier otra cosa es no filtrar (None)."""
    return texto if texto in TEXTO_DEL_ESTADO else None


def dias_que_lleva(lote: dict, hoy: date) -> int:
    """Días corridos desde la salida del lote."""
    return (hoy - lote["fecha"]).days


def sin_cobrar_hace_mucho(lote: dict, hoy: date) -> bool:
    """La regla de la alerta: pendiente y con MÁS de DIAS_SEGUNDA_SIN_COBRAR días."""
    return lote["importe"] is None and dias_que_lleva(lote, hoy) > DIAS_SEGUNDA_SIN_COBRAR


def importe_por_bulto(lote: dict) -> float | None:
    if lote["importe"] is None or not lote["bultos"]:
        return None
    return float(lote["importe"]) / float(lote["bultos"])


def partir_por_estado(lotes: list[dict]) -> tuple[list[dict], list[dict]]:
    """(pendientes, cobrados). Un lote cobrado en $0 es COBRADO."""
    return ([l for l in lotes if l["importe"] is None],
            [l for l in lotes if l["importe"] is not None])


def resumen(lotes: list[dict]) -> dict:
    pendientes, cobrados = partir_por_estado(lotes)
    return {
        "pendientes": len(pendientes),
        "bultos_pendientes": sum(float(l["bultos"]) for l in pendientes),
        "cobrados": len(cobrados),
        "bultos_cobrados": sum(float(l["bultos"]) for l in cobrados),
        "total_cobrado": round(sum(float(l["importe"]) for l in cobrados), 2),
    }


def texto_del_filtro(desde: date | None, hasta: date | None, articulo: str | None,
                     estado: str | None) -> str:
    """Lo que dice el encabezado del PDF y del Excel: TODOS los filtros
    aplicados, o que no hay ninguno (regla de v1063)."""
    partes = []
    if desde and hasta:
        partes.append(f"salidas del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}")
    elif desde:
        partes.append(f"salidas desde el {desde.strftime('%d/%m/%Y')}")
    elif hasta:
        partes.append(f"salidas hasta el {hasta.strftime('%d/%m/%Y')}")
    if articulo:
        partes.append(articulo)
    if estado:
        partes.append(TEXTO_DEL_ESTADO[estado].lower())
    return " · ".join(partes) or "sin filtros: todos los lotes"


def _dia(fecha) -> str:
    return fecha.strftime("%d/%m/%Y") if fecha else ""


def _moneda(valor) -> str:
    if valor is None:
        return ""
    texto = f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return "$" + (texto[:-3] if texto.endswith(",00") else texto)


def _bultos(valor) -> str:
    numero = float(valor)
    return str(int(numero)) if numero.is_integer() else f"{numero:.2f}".replace(".", ",")


def generar_excel_cobranzas(filtro: str, lotes: list[dict], hoy: date) -> bytes:
    """Una hoja por lista, con los importes como NÚMERO para poder sumarlos."""
    pendientes, cobrados = partir_por_estado(lotes)
    libro = Workbook()
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")

    def _hoja(hoja, titulo, encabezados, filas, anchos):
        hoja.cell(row=1, column=1, value=f"Cobranzas de segunda: {titulo}").font = Font(bold=True, size=14)
        hoja.cell(row=2, column=1, value=f"Filtros: {filtro}")
        for columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=4, column=columna, value=encabezado)
            celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
            celda.fill = relleno
        for numero_fila, valores in enumerate(filas, start=5):
            for columna, valor in enumerate(valores, start=1):
                hoja.cell(row=numero_fila, column=columna, value=valor)
        for indice, ancho in enumerate(anchos):
            hoja.column_dimensions[chr(ord("A") + indice)].width = ancho

    hoja = libro.active
    hoja.title = "Pendientes"
    _hoja(hoja, "pendientes de cobro", ("Lote", "Salida", "Artículo", "Bultos", "Días"),
          [(l["id"], _dia(l["fecha"]), l["articulo"], float(l["bultos"]), dias_que_lleva(l, hoy))
           for l in pendientes], (8, 12, 30, 10, 8))
    _hoja(libro.create_sheet("Cobrados"), "cobrados",
          ("Lote", "Salida", "Artículo", "Bultos", "Importe", "Por bulto", "Cobro", "Cargado desde", "Cargado el"),
          [(l["id"], _dia(l["fecha"]), l["articulo"], float(l["bultos"]), float(l["importe"]),
            round(importe_por_bulto(l), 2) if importe_por_bulto(l) is not None else None,
            _dia(l["fecha_cobro"]), TEXTO_DEL_SECTOR.get(l["sector"], "sin dato"),
            hora_argentina(l["cobro_creado_en"])) for l in cobrados],
          (8, 12, 30, 10, 14, 12, 12, 16, 18))
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_cobranzas(filtro: str, lotes: list[dict], hoy: date) -> bytes:
    """A4 vertical: los pendientes primero (es lo que hay que reclamar) y
    después los cobrados, con el filtro en el encabezado."""
    pendientes, cobrados = partir_por_estado(lotes)
    totales = resumen(lotes)
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, "Cobranzas de segunda",
                                           f"Al {hoy.strftime('%d/%m/%Y')} · Filtros: {filtro}")
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    historia = []
    if pendientes or not cobrados:
        filas = [[_p(_dia(l["fecha"])), _p(l["articulo"]), _p(_bultos(l["bultos"]), "numero"),
                  _p(f"{dias_que_lleva(l, hoy)} días")] for l in pendientes] \
            or [[_p("No hay lotes pendientes con estos filtros.", "vacio"), "", "", ""]]
        historia.append(_tabla_seccion_pdf(
            f"Pendientes de cobro: {totales['pendientes']} lotes, {_bultos(totales['bultos_pendientes'])} bultos",
            ["Salida", "Artículo", "Bultos", "Lleva"], filas, [28 * mm, 82 * mm, 30 * mm, 38 * mm], estilos))
        historia.append(Spacer(1, 6 * mm))
    if cobrados:
        filas = [[_p(_dia(l["fecha"])), _p(l["articulo"]), _p(_bultos(l["bultos"]), "numero"),
                  _p(_moneda(l["importe"]), "numero"), _p(_moneda(importe_por_bulto(l))),
                  _p(_dia(l["fecha_cobro"])), _p(TEXTO_DEL_SECTOR.get(l["sector"], "sin dato"), "dato_gris")]
                 for l in cobrados]
        historia.append(_tabla_seccion_pdf(
            f"Cobrados: {totales['cobrados']} lotes, {_moneda(totales['total_cobrado'])}",
            ["Salida", "Artículo", "Bultos", "Importe", "Por bulto", "Cobro", "Cargado desde"],
            filas, [21 * mm, 45 * mm, 16 * mm, 24 * mm, 22 * mm, 21 * mm, 29 * mm], estilos))
    documento.build(historia, onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
