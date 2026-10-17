"""CONSULTAS: el Excel de una respuesta (dueño, 17/10).

Cuando la pregunta pide "en Excel" o "pasame la lista", la IA arma la
consulta y dice, por columna, su nombre y su tipo. Acá se arma el .xlsx con
lo que trajo la base: el título arriba (y si se cortó en
FILAS_MAXIMAS_DEL_EXCEL, lo dice), los encabezados fijos, los pesos y los
kilos con su formato, y las fechas como fechas. No toca la base.
"""
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.consultas_ia import FILAS_MAXIMAS_DEL_EXCEL

# El formato de cada tipo. Excel los muestra con la coma y el punto de la
# compu que lo abre: en una argentina, $ 1.234,56.
FORMATOS = {"pesos": '"$" #,##0.00', "kilos": "#,##0.0", "fecha": "dd/mm/yyyy"}
ANCHO_MAXIMO_DE_COLUMNA = 60


def _celda(valor, tipo: str):
    """El valor como lo tiene que guardar Excel: número, fecha o texto."""
    if valor is None:
        return None
    if tipo in ("numero", "pesos", "kilos"):
        if isinstance(valor, Decimal):
            return float(valor)
        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
            return valor
        try:
            return float(str(valor).replace(",", "."))
        except ValueError:
            return str(valor)
    if tipo == "fecha":
        if isinstance(valor, datetime):
            return valor.replace(tzinfo=None)
        if isinstance(valor, date):
            return valor
        try:
            return date.fromisoformat(str(valor)[:10])
        except ValueError:
            return str(valor)
    if isinstance(valor, (dict, list)):
        return str(valor)
    return valor if isinstance(valor, (str, int, float, bool)) else str(valor)


def nombre_del_archivo(titulo: str) -> str:
    """Sin acentos ni símbolos: el nombre viaja en un encabezado y en Telegram."""
    sin_acentos = unicodedata.normalize("NFKD", titulo or "").encode("ascii", "ignore").decode()
    limpio = re.sub(r"[^A-Za-z0-9\- ]+", "", sin_acentos).strip().replace(" ", "_")
    return (limpio[:60] or "consulta") + ".xlsx"


def armar_xlsx(titulo: str, columnas: list[dict], filas: list, cortada: bool) -> bytes:
    """El .xlsx: en la fila 1 el título, en la 2 los encabezados (fijos al
    bajar) y desde la 3 los renglones. `columnas`: [{nombre, tipo}], una por
    columna de `filas`; si la IA mandó menos, las que faltan van como texto."""
    ancho = max([len(columnas)] + [len(f) for f in filas[:1]])
    columnas = list(columnas) + [{"nombre": f"Columna {n + 1}", "tipo": "texto"}
                                 for n in range(len(columnas), ancho)]
    libro = Workbook()
    hoja = libro.active
    hoja.title = (re.sub(r"[\[\]:*?/\\]", "", titulo or "") or "Consulta")[:31]
    arriba = titulo or "Consulta"
    if cortada:
        arriba += f" (CORTADA en {FILAS_MAXIMAS_DEL_EXCEL:,} renglones: hay más)".replace(",", ".")
    hoja.cell(row=1, column=1, value=arriba).font = Font(bold=True, size=13)
    for n, columna in enumerate(columnas, start=1):
        hoja.cell(row=2, column=n, value=columna["nombre"]).font = Font(bold=True)
    for r, fila in enumerate(filas, start=3):
        for n, columna in enumerate(columnas, start=1):
            valor = fila[n - 1] if n - 1 < len(fila) else None
            celda = hoja.cell(row=r, column=n, value=_celda(valor, columna["tipo"]))
            if columna["tipo"] in FORMATOS:
                celda.number_format = FORMATOS[columna["tipo"]]
    hoja.freeze_panes = "A3"
    for n, columna in enumerate(columnas, start=1):
        largos = [len(str(columna["nombre"]))] + [len(str(f[n - 1])) for f in filas[:200] if n - 1 < len(f)]
        hoja.column_dimensions[get_column_letter(n)].width = min(max(largos) + 3, ANCHO_MAXIMO_DE_COLUMNA)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
