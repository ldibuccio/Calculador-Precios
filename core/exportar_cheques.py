"""CARTERA DE CHEQUES (dueño, 13/10): el listado en PDF y en Excel, con los
mismos filtros que la pantalla y el encabezado que dice cuáles."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph

from core.cheques import NOMBRES_DE_ESTADO, total_del_listado
from core.exportar_vacios import (
    VERDE_ENCABEZADO_HEX,
    _documento_pdf,
    _estilos_pdf,
    _tabla_seccion_pdf,
)


def _dia(fecha) -> str:
    return fecha.strftime("%d/%m/%Y") if fecha else ""


def _moneda(valor) -> str:
    if valor is None:
        return ""
    texto = f"{float(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return "$" + (texto[:-3] if texto.endswith(",00") else texto)


def texto_del_filtro(partes: list[str]) -> str:
    """Las partes ya escritas por la pantalla ("pago del 01/10 al 31/10",
    "banco Galicia"...), o que no hay ninguna."""
    return " · ".join(partes) or "sin filtros: todos los cheques"


ENCABEZADOS = ("Número", "Banco", "Tipo", "De quién", "CUIT librador", "Carga", "Pago", "Importe",
               "Estado", "Concepto", "Proveedor", "Salida", "Acreditado")


def _fila(c: dict) -> tuple:
    return (c["numero"], c["banco"], "ECHEQ" if c["tipo"] == "echeq" else "Físico", c["de_quien"],
            c["cuit_librador"], _dia(c["fecha_carga"]), _dia(c["fecha_pago"]), float(c["importe"]),
            NOMBRES_DE_ESTADO.get(c["estado"], c["estado"]), c.get("concepto") or "",
            c.get("proveedor") or "", _dia(c.get("fecha_salida")), _dia(c.get("acreditado_el")))


def generar_excel_cheques(filtro: str, cheques: list[dict], hoy: date) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Cheques"
    hoja.cell(row=1, column=1, value="Cartera de cheques").font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=f"Al {hoy.strftime('%d/%m/%Y')} · Filtros: {filtro}")
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for numero_fila, cheque in enumerate(cheques, start=5):
        for columna, valor in enumerate(_fila(cheque), start=1):
            hoja.cell(row=numero_fila, column=columna, value=valor)
    fila_total = 5 + len(cheques)
    hoja.cell(row=fila_total, column=7, value="Total").font = Font(bold=True)
    hoja.cell(row=fila_total, column=8, value=total_del_listado(cheques)).font = Font(bold=True)
    for indice, ancho in enumerate((12, 18, 8, 26, 14, 11, 11, 14, 16, 18, 22, 11, 11)):
        hoja.column_dimensions[chr(ord("A") + indice)].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_cheques(filtro: str, cheques: list[dict], hoy: date) -> bytes:
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, "Cartera de cheques",
                                           f"Al {hoy.strftime('%d/%m/%Y')} · Filtros: {filtro}")
    documento.pagesize = landscape(A4)
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    filas = [[_p(c["numero"]), _p(c["banco"]), _p(c["de_quien"]), _p(_dia(c["fecha_carga"])),
              _p(_dia(c["fecha_pago"])), _p(_moneda(c["importe"]), "numero"),
              _p(NOMBRES_DE_ESTADO.get(c["estado"], c["estado"])), _p(c.get("proveedor") or c.get("concepto") or "")]
             for c in cheques] or [[_p("No hay cheques con estos filtros.", "vacio")] + [""] * 7]
    titulo = f"{len(cheques)} cheques · total {_moneda(total_del_listado(cheques))}"
    tabla = _tabla_seccion_pdf(titulo, ["Número", "Banco", "De quién", "Carga", "Pago", "Importe", "Estado",
                                        "Salió a"], filas,
                               [26 * mm, 32 * mm, 50 * mm, 22 * mm, 22 * mm, 30 * mm, 28 * mm, 50 * mm], estilos)
    documento.build([tabla], onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()


# --- CONCILIAR CON EL BANCO (dueño, 16/10) ------------------------------------

def _monto_o_nada(valor) -> str:
    return _moneda(valor) if valor is not None else "sin monto"


def _resumen_de_la_conciliacion(r: dict) -> str:
    return (f"Banco: {r['banco']['cuantos']} cheques {_moneda(r['banco']['total'])} · "
            f"Sistema: {r['sistema']['cuantos']} cheques {_moneda(r['sistema']['total'])} · "
            f"Diferencia {_moneda(r['diferencia'])}")


def _secciones_de_la_conciliacion(r: dict) -> list[tuple[str, list[str], list[tuple]]]:
    """(título con cantidad y total, encabezados, filas) de cada grupo, en el
    orden de la pantalla. Lo usan el PDF y el Excel."""
    t = r["totales"]
    secciones = [
        (f"En el sistema y NO en el banco: {len(r['solo_sistema'])} · {_moneda(t['solo_sistema'])}",
         ["Número", "Banco", "De quién", "Pago", "Importe"],
         [(c["numero"], c["banco"], c["de_quien"], _dia(c["fecha_pago"]), _moneda(c["importe"]))
          for c in r["solo_sistema"]]),
        (f"En el banco y NO en el sistema (faltan cargar): {len(r['solo_banco'])} · {_moneda(t['solo_banco'])}",
         ["Número", "Banco", "Librador", "Pago", "Importe"],
         [(l["numero"], l["banco"], l["librador"], _dia(l["fecha_pago"]), _monto_o_nada(l["importe"]))
          for l in r["solo_banco"]]),
        (f"Con diferencia: {len(r['con_diferencia'])} · banco {_moneda(t['con_diferencia_banco'])}"
         f" · sistema {_moneda(t['con_diferencia_sistema'])}",
         ["Número", "Banco", "Pago banco", "Pago sistema", "Importe banco", "Importe sistema"],
         [(p["sistema"]["numero"], p["sistema"]["banco"], _dia(p["banco"]["fecha_pago"]),
           _dia(p["sistema"]["fecha_pago"]), _monto_o_nada(p["banco"]["importe"]), _moneda(p["sistema"]["importe"]))
          for p in r["con_diferencia"]]),
        (f"Coinciden: {len(r['coinciden'])} · {_moneda(t['coinciden'])}",
         ["Número", "Banco", "De quién", "Pago", "Importe"],
         [(p["sistema"]["numero"], p["sistema"]["banco"], p["sistema"]["de_quien"], _dia(p["sistema"]["fecha_pago"]),
           _moneda(p["sistema"]["importe"])) for p in r["coinciden"]]),
    ]
    if r["solo_echeq"]:
        secciones.append((f"Físicos en cartera, aparte (no figuran en el banco): {len(r['fisicos'])} · "
                          f"{_moneda(t['fisicos'])}", ["Número", "Banco", "De quién", "Pago", "Importe"],
                          [(c["numero"], c["banco"], c["de_quien"], _dia(c["fecha_pago"]), _moneda(c["importe"]))
                           for c in r["fisicos"]]))
    return secciones


def generar_excel_conciliacion(resultado: dict, filtro: str, hoy: date) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Conciliación"
    hoja.cell(row=1, column=1, value="Cartera de cheques: conciliación con el banco").font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=f"Al {hoy.strftime('%d/%m/%Y')} · {filtro}")
    hoja.cell(row=3, column=1, value=_resumen_de_la_conciliacion(resultado)).font = Font(bold=True)
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    fila = 5
    for titulo, encabezados, filas in _secciones_de_la_conciliacion(resultado):
        hoja.cell(row=fila, column=1, value=titulo).font = Font(bold=True)
        fila += 1
        for columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=fila, column=columna, value=encabezado)
            celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
            celda.fill = relleno
        fila += 1
        for valores in filas or [("Ninguno",)]:
            for columna, valor in enumerate(valores, start=1):
                hoja.cell(row=fila, column=columna, value=valor)
            fila += 1
        fila += 1
    for indice, ancho in enumerate((14, 20, 28, 14, 16, 16)):
        hoja.column_dimensions[chr(ord("A") + indice)].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_conciliacion(resultado: dict, filtro: str, hoy: date) -> bytes:
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, "Cartera de cheques: conciliación con el banco",
                                           f"Al {hoy.strftime('%d/%m/%Y')} · {filtro}")
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    partes = [_p(_resumen_de_la_conciliacion(resultado), "titulo_tabla")]
    ancho_total = 178 * mm
    for titulo, encabezados, filas in _secciones_de_la_conciliacion(resultado):
        anchos = [ancho_total / len(encabezados)] * len(encabezados)
        tabla = [[_p(v) for v in valores] for valores in filas] or \
            [[_p("Ninguno.", "vacio")] + [""] * (len(encabezados) - 1)]
        partes.append(_tabla_seccion_pdf(titulo, encabezados, tabla, anchos, estilos))
    documento.build(partes, onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
