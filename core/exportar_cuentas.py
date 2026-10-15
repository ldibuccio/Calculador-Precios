"""CUENTAS (dueño, 15/10): el listado de una cuenta en PDF y en Excel, con los
mismos filtros que la pantalla, el encabezado que dice cuáles, el saldo de
hoy y los totales por concepto y por etiqueta de lo filtrado."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer

from core.cuentas import NOMBRE_DE_SENTIDO, NOMBRE_DE_TIPO, importe_en, saldo_en
from core.exportar_vacios import VERDE_ENCABEZADO_HEX, _documento_pdf, _estilos_pdf, _tabla_seccion_pdf

ENCABEZADOS = ("Fecha", "Ingreso o egreso", "Concepto", "Etiqueta", "Detalle", "Importe", "Saldo de la cuenta",
               "Anulado")


def _dia(fecha) -> str:
    return fecha.strftime("%d/%m/%Y") if fecha else ""


def _subtitulo(datos: dict) -> str:
    cuenta = datos["cuenta"]
    return (f"{NOMBRE_DE_TIPO.get(cuenta['tipo'], cuenta['tipo'])}{' · ' + cuenta['numero'] if cuenta['numero'] else ''}"
            f" · Al {datos['hoy'].strftime('%d/%m/%Y')} · Filtros: {datos['filtro']}")


def _texto_del_saldo(datos: dict) -> str:
    moneda = datos["cuenta"]["moneda"]
    texto = f"Saldo de hoy: {saldo_en(datos['saldo']['saldo'], moneda)}"
    if datos["saldo"]["a_futuro"]:
        texto += f" · Por acreditarse: {saldo_en(datos['saldo']['a_futuro'], moneda)}"
    return texto


def generar_excel_cuenta(datos: dict) -> bytes:
    cuenta = datos["cuenta"]
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Movimientos"
    hoja.cell(row=1, column=1, value=cuenta["nombre"]).font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=_subtitulo(datos))
    hoja.cell(row=3, column=1, value=_texto_del_saldo(datos)).font = Font(bold=True)
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=5, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for fila, m in enumerate(datos["movimientos"], start=6):
        valores = (_dia(m["fecha"]), NOMBRE_DE_SENTIDO.get(m["tipo"], m["tipo"]), m["concepto"], m["etiqueta"] or "",
                   m["detalle"] or "", float(m["importe"]) * (1 if m["tipo"] == "ingreso" else -1),
                   float(m["saldo"]), "Sí" if m["anulado"] else "")
        for columna, valor in enumerate(valores, start=1):
            hoja.cell(row=fila, column=columna, value=valor)
    for indice, ancho in enumerate((11, 15, 26, 22, 34, 16, 18, 9)):
        hoja.column_dimensions[chr(ord("A") + indice)].width = ancho
    for titulo, totales in (("Por concepto", datos["por_concepto"]), ("Por etiqueta", datos["por_etiqueta"])):
        otra = libro.create_sheet(titulo)
        otra.cell(row=1, column=1, value=f"{cuenta['nombre']} · {titulo.lower()}").font = Font(bold=True, size=13)
        otra.cell(row=2, column=1, value=_subtitulo(datos))
        for columna, encabezado in enumerate(("Nombre", "Ingresos", "Egresos", "Neto", "Movimientos"), start=1):
            celda = otra.cell(row=4, column=columna, value=encabezado)
            celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
            celda.fill = relleno
        for fila, t in enumerate(totales, start=5):
            for columna, valor in enumerate((t["nombre"], t["ingresos"], t["egresos"], t["neto"], t["cuantos"]), start=1):
                otra.cell(row=fila, column=columna, value=valor)
        otra.column_dimensions["A"].width = 30
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_cuenta(datos: dict) -> bytes:
    cuenta, moneda = datos["cuenta"], datos["cuenta"]["moneda"]
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, cuenta["nombre"], _subtitulo(datos))
    documento.pagesize = landscape(A4)
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    filas = [[_p(_dia(m["fecha"])),
              _p(NOMBRE_DE_SENTIDO.get(m["tipo"], m["tipo"]) + (" (anulado)" if m["anulado"] else "")),
              _p(m["concepto"]), _p(m["etiqueta"] or ""), _p(m["detalle"] or ""),
              _p(("−" if m["tipo"] == "egreso" else "") + importe_en(m["importe"], moneda), "numero"),
              _p(saldo_en(m["saldo"], moneda), "numero")]
             for m in datos["movimientos"]] or [[_p("No hay movimientos con estos filtros.", "vacio")] + [""] * 6]
    tabla = _tabla_seccion_pdf(f"{len(datos['movimientos'])} movimientos",
                               ["Fecha", "Ingreso o egreso", "Concepto", "Etiqueta", "Detalle", "Importe", "Saldo"],
                               filas, [22 * mm, 30 * mm, 40 * mm, 36 * mm, 64 * mm, 32 * mm, 34 * mm], estilos)

    def _totales(titulo, totales):
        cuerpo = [[_p(t["nombre"]), _p(importe_en(t["ingresos"], moneda), "numero"),
                   _p(importe_en(t["egresos"], moneda), "numero"), _p(saldo_en(t["neto"], moneda), "numero")]
                  for t in totales] or [[_p("Nada en lo filtrado.", "vacio"), "", "", ""]]
        return _tabla_seccion_pdf(titulo, ["Nombre", "Ingresos", "Egresos", "Neto"], cuerpo,
                                  [80 * mm, 40 * mm, 40 * mm, 40 * mm], estilos)

    documento.build([_p(_texto_del_saldo(datos), "numero"), Spacer(1, 4 * mm), tabla, Spacer(1, 6 * mm),
                     _totales("Por concepto", datos["por_concepto"]), Spacer(1, 6 * mm),
                     _totales("Por etiqueta", datos["por_etiqueta"])],
                    onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
