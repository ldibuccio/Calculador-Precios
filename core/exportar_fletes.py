"""La CUENTA DEL FLETERO en Excel y PDF (dueño, 05/10) — puro, sin base.

Los renglones llegan armados de `cuenta_del_fletero`, con los mismos filtros
que la pantalla, y el encabezado dice qué se filtró: el archivo y la
pantalla no pueden decir números distintos. Un renglón por viaje.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.platypus import Paragraph

from core.exportar_vacios import VERDE_ENCABEZADO_HEX, _documento_pdf, _estilos_pdf, _tabla_seccion_pdf

TITULO = "Cuenta del fletero"
ENCABEZADOS = ("Fecha", "Fletero", "Sucursal", "Camión", "Precio", "Frutamax", "Palmala", "Estado")


def _estado(fila) -> str:
    return f"Pagado el {fila['pagado_el'].strftime('%d/%m/%Y')}" if fila["pagado_el"] else "A pagar"


def _pesos(valor) -> str:
    entero = round(float(valor))
    return ("-" if entero < 0 else "") + "$" + f"{abs(entero):,}".replace(",", ".")


def generar_excel_cuenta_fletero(filtro_texto: str, filas: list[dict], totales: dict) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Cuenta del fletero"
    hoja.cell(row=1, column=1, value=TITULO).font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=filtro_texto)
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    renglon = 5
    for f in filas:
        valores = (f["fecha"], f["fletero"], f["sucursal"], f["camion"], float(f["precio"]),
                   float(f["parte_frutamax"]), float(f["parte_palmala"]), _estado(f))
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=renglon, column=columna, value=valor)
            if columna == 1:
                celda.number_format = "DD/MM/YYYY"
            elif 5 <= columna <= 7:
                celda.number_format = "#,##0.00"
        renglon += 1
    for texto, clave in (("Total", "precio"), ("A pagar", "a_pagar"), ("Pagado", "pagado")):
        hoja.cell(row=renglon, column=4, value=texto).font = Font(bold=True)
        celda = hoja.cell(row=renglon, column=5, value=float(totales[clave]))
        celda.font = Font(bold=True)
        celda.number_format = "#,##0.00"
        if clave == "precio":
            for columna, parte in ((6, "parte_frutamax"), (7, "parte_palmala")):
                celda = hoja.cell(row=renglon, column=columna, value=float(totales[parte]))
                celda.font = Font(bold=True)
                celda.number_format = "#,##0.00"
        renglon += 1
    for letra, ancho in zip("ABCDEFGH", (12, 18, 18, 14, 14, 14, 14, 20)):
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_cuenta_fletero(filtro_texto: str, filas: list[dict], totales: dict) -> bytes:
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, TITULO, filtro_texto)
    estilos = _estilos_pdf()
    ancho = documento.width
    datos = [[Paragraph(texto, estilos["dato"]) for texto in (
        f["fecha"].strftime("%d/%m/%Y"), f["fletero"], f["sucursal"], f["camion"], _pesos(f["precio"]),
        _pesos(f["parte_frutamax"]), _pesos(f["parte_palmala"]), _estado(f))] for f in filas]
    datos.append([Paragraph(t, estilos["numero"]) for t in (
        "", "", "", "Total", _pesos(totales["precio"]), _pesos(totales["parte_frutamax"]),
        _pesos(totales["parte_palmala"]), f"A pagar {_pesos(totales['a_pagar'])}")])
    anchos = [ancho * p for p in (0.135, 0.12, 0.14, 0.1, 0.115, 0.115, 0.115, 0.16)]
    elementos = [_tabla_seccion_pdf("Viajes", list(ENCABEZADOS), datos, anchos, estilos)]
    documento.build(elementos, onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
