"""PALMALA CRUCES (dueño, 14/10): el listado en PDF y en Excel, con los mismos
filtros que la pantalla, el encabezado que dice cuáles, y los tres saldos."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, Spacer

from core.exportar_vacios import VERDE_ENCABEZADO_HEX, _documento_pdf, _estilos_pdf, _tabla_seccion_pdf
from core.palmala_cruces import NOMBRE_DE_CUENTA, NOMBRE_DE_SENTIDO, importe_con_signo

ENCABEZADOS = ("Fecha", "Cuenta", "Sentido", "Concepto", "Detalle", "Importe", "Saldo de la cuenta",
               "Cheque", "Banco", "Pago del cheque", "Anulado")


def _dia(fecha) -> str:
    return fecha.strftime("%d/%m/%Y") if fecha else ""


def _fila(m: dict) -> tuple:
    return (_dia(m["fecha"]), NOMBRE_DE_CUENTA.get(m["cuenta"], m["cuenta"]),
            NOMBRE_DE_SENTIDO.get(m["sentido"], m["sentido"]), m["concepto"], m.get("detalle") or "",
            float(m["con_signo"]) if not m["anulado"] else 0.0, float(m["saldo"]), m.get("cheque_numero") or "",
            m.get("cheque_banco") or "", _dia(m.get("cheque_fecha_pago")), "Sí" if m["anulado"] else "")


def generar_excel_palmala(filtro: str, movimientos: list[dict], saldos: list[dict], hoy: date) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Palmala Cruces"
    hoja.cell(row=1, column=1, value="Palmala Cruces").font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=f"Al {hoy.strftime('%d/%m/%Y')} · Filtros: {filtro}")
    for n, s in enumerate(saldos):
        hoja.cell(row=3 + n, column=1, value=s["nombre"]).font = Font(bold=True)
        hoja.cell(row=3 + n, column=2, value=s["saldo"])
        hoja.cell(row=3 + n, column=3, value=s["texto"])
    titulos = 3 + len(saldos) + 1
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=titulos, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for numero_fila, movimiento in enumerate(movimientos, start=titulos + 1):
        for columna, valor in enumerate(_fila(movimiento), start=1):
            hoja.cell(row=numero_fila, column=columna, value=valor)
    for indice, ancho in enumerate((11, 15, 14, 20, 34, 16, 18, 14, 18, 14, 9)):
        hoja.column_dimensions[chr(ord("A") + indice)].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_palmala(filtro: str, movimientos: list[dict], saldos: list[dict], hoy: date) -> bytes:
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, "Palmala Cruces",
                                           f"Al {hoy.strftime('%d/%m/%Y')} · Filtros: {filtro}")
    documento.pagesize = landscape(A4)
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    resumen = _tabla_seccion_pdf("Saldos", ["Cuenta", "Saldo"],
                                 [[_p(s["nombre"]), _p(s["texto"])] for s in saldos], [60 * mm, 120 * mm], estilos)
    filas = [[_p(_dia(m["fecha"])), _p(NOMBRE_DE_CUENTA.get(m["cuenta"], m["cuenta"])),
              _p(NOMBRE_DE_SENTIDO.get(m["sentido"], m["sentido"]) + (" (anulado)" if m["anulado"] else "")),
              _p(m["concepto"]), _p(m.get("detalle") or ""),
              _p(importe_con_signo(m["importe"], m["cuenta"]), "numero"),
              _p(("−" if float(m["saldo"]) < 0 else "") + importe_con_signo(m["saldo"], m["cuenta"]), "numero")]
             for m in movimientos] or [[_p("No hay movimientos con estos filtros.", "vacio")] + [""] * 6]
    tabla = _tabla_seccion_pdf(f"{len(movimientos)} movimientos", ["Fecha", "Cuenta", "Sentido", "Concepto", "Detalle",
                                                                    "Importe", "Saldo"], filas,
                               [22 * mm, 30 * mm, 30 * mm, 36 * mm, 70 * mm, 32 * mm, 34 * mm], estilos)
    documento.build([resumen, Spacer(1, 6 * mm), tabla], onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
