"""PRECIOS COTIZACIONES en Excel y PDF (dueño, 07/10) — puro, sin base.

ES EL LISTADO PARA EL CLIENTE, y por eso lleva SOLO cliente, fecha, artículo,
presentación y precio. Sin costos, sin proveedores, sin compras, sin
condiciones y sin decir qué se tocó a mano: todo eso queda en la pantalla.

Las filas llegan armadas de `_filas_del_listado_de_cotizacion`: solo los
renglones con un precio que vale (ni vacío ni 0), con el precio que quedó en
el casillero de la pantalla.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.platypus import Paragraph

from core.exportar_vacios import VERDE_ENCABEZADO_HEX, _documento_pdf, _estilos_pdf, _tabla_seccion_pdf

TITULO = "Precios Cotizaciones"
ENCABEZADOS = ("Artículo", "Presentación", "Precio")
_POR_UNIDAD = {"kilo": "por kilo", "unidad": "por unidad", "cubeta": "por cubeta"}


def _pesos(valor) -> str:
    entero = round(float(valor))
    return ("-" if entero < 0 else "") + "$" + f"{abs(entero):,}".replace(",", ".")


def _precio(fila: dict) -> str:
    return f"{_pesos(fila['precio'])} {_POR_UNIDAD.get(fila['unidad_venta'], '')}".strip()


def _valores(fila: dict) -> tuple:
    return (fila["articulo"], fila["presentacion"], _precio(fila))


def _encabezado(cliente: str, fecha: str) -> str:
    return f"Cliente: {cliente} · Fecha: {fecha}"


def generar_excel_precios_cotizaciones(cliente: str, fecha: str, filas: list[dict]) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Precios Cotizaciones"
    hoja.cell(row=1, column=1, value=TITULO).font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=_encabezado(cliente, fecha))
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, titulo in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=4, column=columna, value=titulo)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for renglon, fila in enumerate(filas, start=5):
        for columna, valor in enumerate(_valores(fila), start=1):
            hoja.cell(row=renglon, column=columna, value=valor)
    for letra, ancho in zip("ABC", (32, 30, 20)):
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_precios_cotizaciones(cliente: str, fecha: str, filas: list[dict]) -> bytes:
    buffer = BytesIO()
    documento, dibujar_encabezado = _documento_pdf(buffer, TITULO, _encabezado(cliente, fecha))
    estilos = _estilos_pdf()
    ancho = documento.width
    datos = [[Paragraph(_escapar(texto), estilos["dato"]) for texto in _valores(fila)] for fila in filas]
    anchos = [ancho * p for p in (0.42, 0.33, 0.25)]
    elementos = [_tabla_seccion_pdf("Precios", list(ENCABEZADOS), datos, anchos, estilos)]
    documento.build(elementos, onFirstPage=dibujar_encabezado, onLaterPages=dibujar_encabezado)
    return buffer.getvalue()


def _escapar(texto) -> str:
    return str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
