"""RENDIMIENTOS en Excel y PDF — puro, sin tocar la base (dueño, 13/10).

Arriba de todo, el resumen en tres renglones (guías R, sin procesar y
TOTAL), igual que la pantalla. Recibe lo que ya armó
`core.rendimientos.armar_rendimientos` para la pantalla,
con los MISMOS filtros (período y artículo), y el encabezado dice qué se
filtró. Acá no se recalcula nada.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Spacer, TableStyle

from core.exportar_vacios import VERDE_ENCABEZADO_HEX, _documento_pdf, _estilos_pdf, _tabla_seccion_pdf
from core.rendimientos import texto_de_la_perdida

TITULO = "Rendimientos"
ENCABEZADOS_CON_GUIA = ("Artículo", "Kg que entraron", "% primera", "% segunda", "% merma",
                        "% falta sin registrar", "% pérdida total", "% pérdida en la ficha")
ENCABEZADOS_RESUMEN = ("", "Kg que entraron", "Aprovechado", "Segunda", "Merma", "Falta sin registrar",
                       "Pérdida total", "Pérdida al costo", "% del costo")
ENCABEZADOS_SIN_GUIA = ("Artículo", "Kg por cajón comprado", "Kg por bulto enviado", "Diferencia",
                        "Qué pasa", "% pérdida en la ficha")


def subtitulo(desde, hasta, articulo_nombre: str | None) -> str:
    """Lo que se filtró, como lo dice la pantalla."""
    texto = f"Del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}"
    return texto + (f" · {articulo_nombre}" if articulo_nombre else " · todos los artículos")


def _pct(valor) -> str:
    return "—" if valor is None else f"{valor:.1f}%".replace(".", ",")


def _kg(valor) -> str:
    return "—" if valor is None else f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _pesos(valor) -> str:
    entero = round(float(valor))
    return ("-" if entero < 0 else "") + "$" + f"{abs(entero):,}".replace(",", ".")


def _kg_entero(valor) -> str:
    return f"{round(float(valor)):,}".replace(",", ".")


def _filas_resumen(datos: dict) -> list[list[str]]:
    """Los tres renglones del resumen (dueño, 15/10), en el orden de la pantalla."""
    return [[f["nombre"], _kg_entero(f["entraron"]), _pct(f["aprovechado_pct"]), _pct(f["segunda_pct"]),
             _pct(f["merma_pct"]), _pct(f["no_explicado_pct"]), _pct(f["perdida_pct"]),
             _pesos(f["perdida_pesos"]), _pct(f["perdida_costo_pct"])]
            for f in (datos["resumen"][c] for c in ("guias_r", "sin_procesar", "total"))]


def _notas_resumen(datos: dict) -> list[str]:
    """Lo que la pantalla dice debajo de los renglones: los kilos de sin procesar
    y lo que no entra en la cuenta."""
    r = datos["resumen"]
    s = r["sin_procesar"]
    notas = [f"Sin procesar: comprado {_kg_entero(s['kg_comprados'])} kg, enviado {_kg_entero(s['kg_enviados'])} kg, "
             f"diferencia {_pct(s['diferencia_pct'])}, merma {_kg_entero(s['merma'])} kg, "
             f"segunda {_kg_entero(s['segunda'])} kg."]
    sin_cuenta = r["guias_r"]["guias_sin_cuenta"]
    if sin_cuenta:
        notas.append(f"{sin_cuenta} guía{'' if sin_cuenta == 1 else 's'} R sin pesada o con ficha por unidad "
                     f"no entra{'' if sin_cuenta == 1 else 'n'}.")
    sin_dato = s["articulos_sin_dato"]
    if sin_dato:
        notas.append(f"{sin_dato} artículo{'' if sin_dato == 1 else 's'} sin compra pesada o sin envíos "
                     f"no entra{'' if sin_dato == 1 else 'n'}.")
    for clave in ("guias_r", "sin_procesar"):
        if r[clave]["kg_sin_costo"]:
            notas.append(f"{r[clave]['nombre']}: {_kg_entero(r[clave]['kg_sin_costo'])} kg sin precio de compra "
                         "no entran en la plata.")
    return notas


def _filas_con_guia(datos: dict) -> list[list[str]]:
    return [[f["articulo"], _kg(f["entraron"]), _pct(f["primera_pct"]), _pct(f["segunda_pct"]),
             _pct(f["merma_pct"]), _pct(f["falta_pct"]), _pct(f["perdida_total_pct"]),
             texto_de_la_perdida(f["perdida_ficha"])] for f in datos["con_guia"]]


def _filas_sin_guia(datos: dict) -> list[list[str]]:
    return [[f["articulo"], _kg(f["kg_por_cajon"]), _kg(f["kg_por_bulto"]), _pct(f["diferencia_pct"]),
             f["que_pasa"] or "—", texto_de_la_perdida(f["perdida_ficha"])] for f in datos["sin_guia"]]


def generar_excel_rendimientos(datos: dict, desde, hasta, articulo_nombre: str | None) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Rendimientos"
    hoja.cell(row=1, column=1, value=TITULO).font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=subtitulo(desde, hasta, articulo_nombre))
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    fila = 4
    # PRIMERO el resumen (dueño, 15/10), destacado.
    hoja.cell(row=fila, column=1, value="Resumen del período").font = Font(bold=True, size=13)
    fila += 1
    for columna, encabezado in enumerate(ENCABEZADOS_RESUMEN, start=1):
        celda = hoja.cell(row=fila, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    fila += 1
    destacado = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    for valores in _filas_resumen(datos):
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)
            celda.fill = destacado
            celda.font = Font(bold=valores[0] == "TOTAL")
        fila += 1
    for nota in _notas_resumen(datos):
        hoja.cell(row=fila, column=1, value=nota)
        fila += 1
    fila += 1
    for titulo, encabezados, filas in (
            ("Pasan por guía R", ENCABEZADOS_CON_GUIA, _filas_con_guia(datos)),
            ("No pasan por guía R", ENCABEZADOS_SIN_GUIA, _filas_sin_guia(datos))):
        hoja.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=12)
        fila += 1
        for columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=fila, column=columna, value=encabezado)
            celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
            celda.fill = relleno
        fila += 1
        for valores in filas:
            for columna, valor in enumerate(valores, start=1):
                hoja.cell(row=fila, column=columna, value=valor)
            fila += 1
        if not filas:
            hoja.cell(row=fila, column=1, value="Nada en el período.")
            fila += 1
        fila += 1
    hoja.column_dimensions["A"].width = 26
    for letra in "BCDEFGHI":
        hoja.column_dimensions[letra].width = 16
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_rendimientos(datos: dict, desde, hasta, articulo_nombre: str | None) -> bytes:
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, TITULO, subtitulo(desde, hasta, articulo_nombre))
    estilos = _estilos_pdf()
    ancho = documento.width

    con = _filas_con_guia(datos)
    sin = _filas_sin_guia(datos)
    chico = ParagraphStyle("encabezado_chico", parent=estilos["encabezado_tabla"], fontSize=8)
    resumen = _tabla_seccion_pdf(
        "Resumen del período", list(ENCABEZADOS_RESUMEN),
        [[Paragraph(str(v), estilos["numero"]) for v in f] for f in _filas_resumen(datos)],
        [ancho * f for f in (0.153, 0.10, 0.125, 0.09, 0.09, 0.11, 0.11, 0.142, 0.08)], estilos,
        estilo_encabezado=chico)
    # Destacado (dueño, 15/10): fondo azul y el TOTAL con borde arriba.
    resumen.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                                 ("BACKGROUND", (0, 2), (-1, -1), colors.Color(0.86, 0.92, 1)),
                                 ("BOX", (0, 1), (-1, -1), 1.2, colors.Color(0.15, 0.39, 0.92)),
                                 ("LINEABOVE", (0, -1), (-1, -1), 1.2, colors.Color(0.15, 0.39, 0.92))]))
    elementos = [
        resumen,
        *[Paragraph(nota, estilos["aviso"]) for nota in _notas_resumen(datos)],
        Spacer(1, 14),
        _tabla_seccion_pdf("Pasan por guía R", list(ENCABEZADOS_CON_GUIA),
                           [[Paragraph(str(v), estilos["dato"] if i == 0 else estilos["numero"])
                             for i, v in enumerate(f)] for f in con]
                           or [[Paragraph("Nada en el período.", estilos["dato_gris"])] + [""] * 7],
                           [ancho * 0.19] + [ancho * 0.115] * 7, estilos),
        Spacer(1, 12),
        _tabla_seccion_pdf("No pasan por guía R", list(ENCABEZADOS_SIN_GUIA),
                           [[Paragraph(str(v), estilos["dato"] if i == 0 else estilos["numero"])
                             for i, v in enumerate(f)] for f in sin]
                           or [[Paragraph("Nada en el período.", estilos["dato_gris"])] + [""] * 5],
                           [ancho * 0.24, ancho * 0.15, ancho * 0.15, ancho * 0.13, ancho * 0.18, ancho * 0.15],
                           estilos),
    ]
    documento.build(elementos, onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
