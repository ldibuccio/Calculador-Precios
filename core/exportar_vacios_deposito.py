"""El stock de VACÍOS DEL DEPÓSITO en Excel y PDF — puro, sin tocar la base.

Es la lista simple que pidió el dueño el 23/09: proveedor, tipo de cajón y
cantidad — y desde el 25/09 una fila por PILA (proveedor y marca del cajón),
que es lo que se le devuelve a cada uno. Los datos llegan armados de `stock_de_vacios_deposito`, que es la
cuenta de la pantalla: acá no se recalcula nada, así el archivo y la
pantalla no pueden decir números distintos.

NO ES `core/exportar_vacios.py`, que es del circuito del PUESTO (otro
proveedor, otras tablas). De ahí se reusa SOLO la forma del PDF.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.platypus import Paragraph

from core.exportar_vacios import (
    VERDE_ENCABEZADO_HEX,
    _documento_pdf,
    _estilos_pdf,
    _tabla_seccion_pdf,
)

TITULO = "Stock de vacíos del depósito"


def filas_del_stock(proveedores: list[dict]) -> list[dict]:
    """Una fila por PILA: proveedor, marca ("sin asignar" si no tiene) y cajones.

    Ya no hay proveedores "sin conteo" que contar al pie: desde el corte del
    25/09 todos arrancan de la foto, y el que no aparece es porque está en
    cero y no se movió.
    """
    return [
        {"nombre": p["nombre"], "tipo_cajon": p.get("tipo_cajon"),
         "marca": pila["marca"] or "sin asignar", "stock": int(pila["stock"])}
        for p in proveedores
        for pila in p["pilas"]
    ]


def _texto_tipo(fila: dict) -> str:
    return fila.get("tipo_cajon") or "sin declarar"


def generar_excel_stock_vacios_deposito(fecha: date, proveedores: list[dict]) -> bytes:
    filas = filas_del_stock(proveedores)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Stock de vacíos"
    hoja.cell(row=1, column=1, value=TITULO).font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=f"Al {fecha.strftime('%d/%m/%Y')}")
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(("Proveedor", "Marca", "Tipo de cajón", "Cajones"), start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    fila_actual = 5
    for p in filas:
        hoja.cell(row=fila_actual, column=1, value=p["nombre"])
        hoja.cell(row=fila_actual, column=2, value=p["marca"])
        hoja.cell(row=fila_actual, column=3, value=_texto_tipo(p))
        hoja.cell(row=fila_actual, column=4, value=p["stock"])
        fila_actual += 1
    hoja.cell(row=fila_actual, column=3, value="Total").font = Font(bold=True)
    hoja.cell(row=fila_actual, column=4, value=sum(p["stock"] for p in filas)).font = Font(bold=True)
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 20
    hoja.column_dimensions["C"].width = 22
    hoja.column_dimensions["D"].width = 10
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_stock_vacios_deposito(fecha: date, proveedores: list[dict]) -> bytes:
    filas = filas_del_stock(proveedores)
    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, TITULO, f"Al {fecha.strftime('%d/%m/%Y')}")
    estilos = _estilos_pdf()
    ancho = documento.width
    datos = [
        [Paragraph(p["nombre"], estilos["dato"]), Paragraph(p["marca"], estilos["dato"]),
         Paragraph(_texto_tipo(p), estilos["dato_gris"]),
         Paragraph(str(p["stock"]), estilos["numero"])]
        for p in filas
    ]
    datos.append([Paragraph("", estilos["dato"]), Paragraph("", estilos["dato"]),
                  Paragraph("Total", estilos["numero"]),
                  Paragraph(str(sum(p["stock"] for p in filas)), estilos["numero"])])
    elementos = [_tabla_seccion_pdf("Cajones en el galpón", ["Proveedor", "Marca", "Tipo de cajón", "Cajones"],
                                    datos, [ancho * 0.36, ancho * 0.24, ancho * 0.26, ancho * 0.14], estilos)]
    documento.build(elementos, onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
