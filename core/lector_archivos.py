"""Conversión de archivos (PDF, Excel) a algo que la IA de core/lector_comandas.py pueda leer.

Este módulo NO habla con la API de Claude ni interpreta nada: solo
convierte el archivo tal como llegó a imágenes (PDF, una por página) o a
texto plano (Excel, celda por celda). La lectura con IA — que es la que
decide qué es un artículo y qué un precio — vive en core/lector_comandas.py.
"""

import io

import openpyxl
import pypdfium2 as pdfium
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas as canvas_pdf

ESCALA_RENDER_PDF = 2.0  # 72 * 2 ≈ 144 DPI: legible sin generar imágenes gigantes.
CALIDAD_JPEG_PDF = 85

# Para el PDF que se GUARDA en Storage: mismo criterio que las fotos de
# comandas (máx 1000px de lado, JPEG calidad 60) — alcanza para releerlo
# a ojo, y baja un escaneo de varios MB a ~100-150 KB por página.
LADO_MAXIMO_PDF_GUARDADO = 1000
CALIDAD_JPEG_PDF_GUARDADO = 60


def comprimir_pdf(bytes_pdf: bytes) -> bytes:
    """Rearma el PDF con cada página como imagen comprimida, para guardarlo chico en Storage.

    Un PDF escaneado trae cada página como foto en alta resolución: acá se
    renderiza cada página, se achica y comprime como las fotos de comandas,
    y se arma un PDF nuevo con esas imágenes. Se pierde el texto
    seleccionable (si lo había) — no importa: este archivo se guarda solo
    como respaldo visual de dónde salió cada precio, la lectura con IA ya
    se hizo con el original.

    Si algo falla, o si el resultado no achica nada, devuelve el original:
    comprimir es una mejora, nunca un motivo para perder el archivo.
    """
    try:
        documento = pdfium.PdfDocument(bytes_pdf)
        try:
            buffer_pdf = io.BytesIO()
            lienzo = None
            for pagina in documento:
                bitmap = pagina.render(scale=ESCALA_RENDER_PDF)
                try:
                    imagen = bitmap.to_pil().convert("RGB")
                finally:
                    bitmap.close()
                pagina.close()
                imagen.thumbnail((LADO_MAXIMO_PDF_GUARDADO, LADO_MAXIMO_PDF_GUARDADO))
                buffer_jpeg = io.BytesIO()
                imagen.save(buffer_jpeg, format="JPEG", quality=CALIDAD_JPEG_PDF_GUARDADO)
                buffer_jpeg.seek(0)

                # Página del PDF nuevo del tamaño exacto de la imagen (en
                # puntos, a ~144 DPI: la mitad de los píxeles).
                ancho_pt, alto_pt = imagen.width / 2, imagen.height / 2
                if lienzo is None:
                    lienzo = canvas_pdf.Canvas(buffer_pdf, pagesize=(ancho_pt, alto_pt))
                else:
                    lienzo.setPageSize((ancho_pt, alto_pt))
                lienzo.drawImage(ImageReader(buffer_jpeg), 0, 0, width=ancho_pt, height=alto_pt)
                lienzo.showPage()
            if lienzo is None:
                return bytes_pdf
            lienzo.save()
        finally:
            documento.close()
        bytes_comprimidos = buffer_pdf.getvalue()
        return bytes_comprimidos if len(bytes_comprimidos) < len(bytes_pdf) else bytes_pdf
    except Exception:
        return bytes_pdf


def imagenes_desde_pdf(bytes_pdf: bytes) -> list[bytes]:
    """Convierte cada página de un PDF en una imagen JPEG (bytes), en el mismo orden.

    Sirve tanto para un PDF con texto real como para uno escaneado (fotos
    metidas en un PDF): en los dos casos la página se renderiza tal cual se
    ve, y es la IA (en modo imagen) la que la lee — no hace falta distinguir
    de antemano cuál es cuál.

    Lanza ValueError con un mensaje claro si el archivo no es un PDF válido
    o no se pudo abrir, para que quien llame pueda mostrar un error
    entendible en vez de un stack trace de la librería.
    """
    try:
        documento = pdfium.PdfDocument(bytes_pdf)
    except Exception as error:
        raise ValueError(f"No se pudo abrir el PDF: {error}") from error

    try:
        imagenes = []
        for pagina in documento:
            bitmap = pagina.render(scale=ESCALA_RENDER_PDF)
            try:
                imagen_pil = bitmap.to_pil().convert("RGB")
            finally:
                bitmap.close()
            buffer = io.BytesIO()
            imagen_pil.save(buffer, format="JPEG", quality=CALIDAD_JPEG_PDF)
            imagenes.append(buffer.getvalue())
            pagina.close()
        return imagenes
    finally:
        documento.close()


def texto_desde_excel(bytes_excel: bytes) -> str:
    """Vuelca el contenido de todas las hojas de un Excel a texto plano, fila por fila.

    Cada fila se representa como sus valores separados por " | ", una fila
    por línea — no hace falta más estructura que esa: la IA es la que
    interpreta cuál columna es el artículo y cuál el precio (ver
    PROMPT_LISTADO_PRECIOS en core/lector_comandas.py). Filas y celdas
    vacías no aportan nada, se saltean. Si el archivo tiene varias hojas,
    se vuelcan todas juntas, una atrás de la otra.

    Lanza ValueError con un mensaje claro si el archivo no es un Excel
    válido (.xlsx) o no se pudo abrir.
    """
    try:
        libro = openpyxl.load_workbook(io.BytesIO(bytes_excel), data_only=True, read_only=True)
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Excel: {error}") from error

    try:
        lineas = []
        for hoja in libro.worksheets:
            for fila in hoja.iter_rows(values_only=True):
                valores = [str(valor).strip() for valor in fila if valor is not None and str(valor).strip()]
                if valores:
                    lineas.append(" | ".join(valores))
        return "\n".join(lineas)
    finally:
        libro.close()
