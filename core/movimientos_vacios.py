"""Los MOVIMIENTOS de vacíos del depósito: textos, ventana y Excel — puro.

Las filas llegan armadas de `movimientos_de_vacios` (app/db.py), que es la
MISMA consulta del historial del detalle de un proveedor. Acá no se recalcula
nada: se decide cómo se DICE cada fila y qué ventana se pide.

"POR DÓNDE ENTRÓ" Y NO "QUIÉN" (dueño, 29/09): el sistema no tiene usuarios,
solo sectores. Lo dice el SECTOR que cargó la fila.
"""

from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.exportar_vacios import VERDE_ENCABEZADO_HEX

TEXTO_DEL_TIPO = {
    "arranque": "Conteo físico",
    "entrada": "Compra con seña",
    "devolucion": "Devolución",
    "ajuste": "Ajuste",
    "asignacion": "Pase",
}

# NULL es una devolución cargada antes del 29/09: no hay de dónde saberlo, y
# se dice así en vez de inventarle una puerta.
TEXTO_DE_LA_PUERTA = {
    "conteo": "Conteo físico",
    "recepcion": "Recepción",
    "deposito": "Depósito",
    "administracion": "Administración",
    "compras": "Compras",
    None: "sin dato",
}

DIAS_POR_DEFECTO = 30
DIAS_COMO_MAXIMO = 90


def ventana(desde: str, hasta: str, hoy: date) -> tuple[date, date, str | None]:
    """La ventana pedida, con los dos días incluidos. Por defecto, los últimos
    30 días hasta hoy; como máximo 90 (dueño, 29/09).

    Devuelve `(desde, hasta, error)`. Con error, las fechas son las que se
    muestran en el formulario y NO se consulta nada.
    """
    try:
        fin = date.fromisoformat(hasta) if hasta.strip() else hoy
        inicio = (date.fromisoformat(desde) if desde.strip()
                  else fin - timedelta(days=DIAS_POR_DEFECTO - 1))
    except ValueError:
        fin = hoy
        return fin - timedelta(days=DIAS_POR_DEFECTO - 1), fin, "Una de las fechas no se entiende."
    if inicio > fin:
        return inicio, fin, "La fecha de inicio es posterior a la de fin."
    if (fin - inicio).days + 1 > DIAS_COMO_MAXIMO:
        return inicio, fin, f"Hasta {DIAS_COMO_MAXIMO} días por búsqueda: achicá las fechas."
    return inicio, fin, None


def texto_de_cantidad(movimiento: dict) -> str:
    """El número con su signo: suma o resta de la pila. El PASE no lleva signo
    —resta de una marca y suma a otra— y el CONTEO es lo contado."""
    cantidad = int(movimiento["cantidad"])
    if movimiento["tipo"] in ("arranque", "asignacion"):
        return str(cantidad)
    return f"+{cantidad}" if cantidad > 0 else str(cantidad)


def texto_de_la_marca(movimiento: dict) -> str:
    desde = movimiento["marca"] or "Sin marca"
    if movimiento["tipo"] == "asignacion":
        return f"{desde} → {movimiento['marca_hasta'] or 'Sin marca'}"
    return desde


def estado(movimiento: dict) -> str:
    if movimiento["anulada"]:
        return "anulado"
    if movimiento["antes_del_arranque"]:
        return "antes del conteo: no cuenta"
    return ""


def generar_excel_movimientos_vacios_deposito(desde: date, hasta: date, filtro: str,
                                     movimientos: list[dict]) -> bytes:
    """Lo filtrado, fila por fila, con las mismas palabras que la pantalla."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Movimientos de vacíos"
    hoja.cell(row=1, column=1, value="Movimientos de vacíos del depósito").font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1,
              value=f"Del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')} · {filtro}")
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    encabezados = ("Fecha", "Tipo", "Proveedor", "Marca", "Cajones", "Por dónde entró",
                   "Compra", "Motivo", "Importe", "Estado")
    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for fila, m in enumerate(movimientos, start=5):
        valores = (
            m["fecha"].strftime("%d/%m/%Y") if m["fecha"] else "",
            TEXTO_DEL_TIPO[m["tipo"]],
            m["proveedor"],
            texto_de_la_marca(m),
            # El número con signo como NÚMERO, para que se pueda sumar; el pase
            # va sin signo, igual que en la pantalla.
            int(m["cantidad"]),
            TEXTO_DE_LA_PUERTA.get(m["cargada_desde"], m["cargada_desde"]),
            m["compra_id"] if m["tipo"] == "entrada" else None,
            m["motivo"] if m["tipo"] == "ajuste" else None,
            float(m["importe"]) if m["tipo"] == "devolucion" and m["importe"] is not None else None,
            estado(m),
        )
        for columna, valor in enumerate(valores, start=1):
            hoja.cell(row=fila, column=columna, value=valor)
    for letra, ancho in zip("ABCDEFGHIJ", (12, 16, 30, 28, 10, 16, 10, 30, 12, 26)):
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
