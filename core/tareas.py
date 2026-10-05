"""TAREAS (dueño, 02/10): las reglas, las palabras y los archivos. Es puro.

Gerencia le carga a un sector (Compras, Administración o Gerencia) tareas
puntuales: de una sola vez, con vencimiento, o repetitivas. El sector las
marca como hechas y queda cuándo y quién (el sector: el sistema no tiene
usuarios). No tiene nada que ver con las alertas, salvo una cosa: una tarea
vencida y no hecha sale en las alertas de Gerencia.

CADA SECTOR CARGA LAS SUYAS (dueño, 02/10). Compras, Administración y
Gerencia crean tareas para su propio sector; Gerencia, para cualquiera. Cada
tarea dice quién la CREÓ (`creada_por`), y un sector edita, pausa o da de
baja SOLO las que creó él. La mensual puede salir en VARIOS días del mes.

LAS OCURRENCIAS. Lo que el sector ve y marca es una OCURRENCIA: cada vez que
la tarea sale. La de una sola vez sale al crearla. La repetitiva sale el día
que le toca (y vence ese día), y NO SE ACUMULA: si llega la fecha de la
siguiente y la anterior sigue pendiente, la anterior queda "no hecha" en el
registro y la nueva sale marcada como atrasada. Una sola en el recuadro.
"""

import calendar
from datetime import date, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate

from core.exportar_vacios import VERDE_ENCABEZADO_HEX, _estilos_pdf, _tabla_seccion_pdf
from core.vales import hora_argentina

SECTORES = {"compras": "Compras", "administracion": "Administración", "gerencia": "Gerencia"}
# DOS FORMAS DE REPETIR (dueño, 05/10): a fecha FIJA (cada X días, semanal,
# mensual en varios días, anual) o RELATIVA: vuelve a salir X días después de
# marcada hecha (`despues_de_hecha`, con los días en `cada_dias`).
TIPOS = {"una_vez": "Una sola vez", "cada_dias": "Cada X días", "semanal": "Semanal", "mensual": "Mensual",
         "anual": "Una vez por año", "despues_de_hecha": "X días después de hecha"}
MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre")
DIAS_DE_LA_SEMANA = ("lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo")
# "Programada" (dueño, 05/10): pendiente cuyo día todavía no llegó. No sale
# en el recuadro del hub hasta ese día; sí en la pantalla de Tareas.
ESTADOS_VISIBLES = {"programada": "Programada", "pendiente": "Pendiente", "vencida": "Vencida", "hecha": "Hecha",
                    "no_hecha": "No hecha", "eliminada": "Eliminada"}
_ZONA = ZoneInfo("America/Argentina/Buenos_Aires")
ESTADOS_DE_LA_TAREA = {"activa": "Activa", "pausada": "Pausada", "baja": "Eliminada"}


def fechas_que_tocan(tarea: dict, despues_de: date, hasta: date) -> list[date]:
    """Las fechas de una repetitiva en (despues_de, hasta], en orden. Nunca
    antes de `desde`. Mensual: cada día de `dias_mes`, y uno que el mes no
    tiene (31 en septiembre) cae el último día del mes. Si dos caen el mismo
    día (30 y 31 en febrero) sale una sola vez."""
    inicio = max(despues_de + timedelta(days=1), tarea["desde"])
    if inicio > hasta:
        return []
    fechas = []
    if tarea["tipo"] == "cada_dias":
        paso = int(tarea["cada_dias"])
        saltos = max(0, -(-(inicio - tarea["desde"]).days // paso))
        dia = tarea["desde"] + timedelta(days=saltos * paso)
        while dia <= hasta:
            fechas.append(dia)
            dia += timedelta(days=paso)
    elif tarea["tipo"] == "semanal":
        dia = inicio + timedelta(days=(int(tarea["dia_semana"]) - inicio.weekday()) % 7)
        while dia <= hasta:
            fechas.append(dia)
            dia += timedelta(days=7)
    elif tarea["tipo"] == "mensual":
        anio, mes = inicio.year, inicio.month
        while date(anio, mes, 1) <= hasta:
            ultimo = calendar.monthrange(anio, mes)[1]
            for dia in sorted({date(anio, mes, min(int(d), ultimo)) for d in tarea["dias_mes"] or []}):
                if inicio <= dia <= hasta:
                    fechas.append(dia)
            anio, mes = (anio + 1, 1) if mes == 12 else (anio, mes + 1)
    elif tarea["tipo"] == "anual":
        mes = int(tarea["anual_mes"])
        for anio in range(inicio.year, hasta.year + 1):
            dia = date(anio, mes, min(int(tarea["anual_dia"]), calendar.monthrange(anio, mes)[1]))
            if inicio <= dia <= hasta:
                fechas.append(dia)
    # `despues_de_hecha` no tiene calendario: la siguiente la saca el "Hecha"
    # (`siguiente_despues_de_hecha`), no el paso de los días.
    return fechas


def siguiente_despues_de_hecha(tarea: dict, hecha_el: date) -> date | None:
    """La RELATIVA (dueño, 05/10): vuelve a salir X días después del día en
    que se marcó hecha. Las demás no tienen siguiente por acá."""
    if tarea["tipo"] != "despues_de_hecha":
        return None
    return hecha_el + timedelta(days=int(tarea["cada_dias"]))


def estado_visible(ocurrencia: dict, hoy: date) -> str:
    """'programada' (pendiente y su día no llegó), 'pendiente' (es hoy),
    'vencida' (pasó su día), 'hecha', 'no_hecha' o 'eliminada'."""
    if ocurrencia["estado"] == "pendiente" and ocurrencia["vence_el"] < hoy:
        return "vencida"
    if ocurrencia["estado"] == "pendiente" and ocurrencia["vence_el"] > hoy:
        return "programada"
    return ocurrencia["estado"]


def dias_de_atraso(ocurrencia: dict, hoy: date) -> int | None:
    """Días entre el vencimiento y cuando se hizo (o hoy si sigue pendiente,
    o el día en que quedó no hecha). Cero o negativo es None: no hubo atraso."""
    if ocurrencia["estado"] == "hecha":
        hasta = ocurrencia["hecha_el"].astimezone(_ZONA).date()
    elif ocurrencia["estado"] == "no_hecha":
        hasta = ocurrencia["no_hecha_el"]
    elif ocurrencia["estado"] == "eliminada":
        return None
    else:
        hasta = hoy
    dias = (hasta - ocurrencia["vence_el"]).days
    return dias if dias > 0 else None


def texto_de_la_regla(tarea: dict) -> str:
    if tarea["tipo"] == "una_vez":
        return f"una sola vez, vence el {tarea['vence_el'].strftime('%d/%m/%Y')}"
    if tarea["tipo"] == "cada_dias":
        n = int(tarea["cada_dias"])
        return "todos los días" if n == 1 else f"cada {n} días"
    if tarea["tipo"] == "semanal":
        return f"todos los {DIAS_DE_LA_SEMANA[int(tarea['dia_semana'])]}"
    if tarea["tipo"] == "anual":
        return f"todos los años, el {int(tarea['anual_dia'])} de {MESES[int(tarea['anual_mes']) - 1]}"
    if tarea["tipo"] == "despues_de_hecha":
        n = int(tarea["cada_dias"])
        return f"{n} día{'s' if n != 1 else ''} después de hecha"
    dias = [str(int(d)) for d in sorted(tarea["dias_mes"] or [])]
    if len(dias) == 1:
        return f"el día {dias[0]} de cada mes"
    return f"los días {', '.join(dias[:-1])} y {dias[-1]} de cada mes"


def dias_del_mes(texto) -> list[int] | None:
    """Los días del mes elegidos (las casillas del 1 al 31, o escritos "1, 15"),
    ordenados y sin repetir. None si no hay ninguno o alguno no es un día del
    1 al 31. Casillas y no texto (05/10): el teclado numérico del iPhone no
    tiene coma, y "1, 15" no se podía escribir."""
    if isinstance(texto, (list, tuple)):
        texto = " ".join(str(t) for t in texto)
    partes = [p for p in (texto or "").replace(",", " ").replace(";", " ").split() if p]
    if not partes or not all(p.isdigit() and 1 <= int(p) <= 31 for p in partes):
        return None
    return sorted({int(p) for p in partes})


def puede_manejar(tarea: dict, sector: str) -> bool:
    """Editar, pausar o eliminar (dueño, 02/10): Gerencia todas; un sector
    SOLO las que creó él. La escritura pregunta lo mismo en su WHERE."""
    return sector == "gerencia" or tarea["creada_por"] == sector


def texto_del_filtro(sector: str | None, desde: date | None, hasta: date | None, estado: str | None,
                     creada_por: str | None = None) -> str:
    """Lo que dice el encabezado del PDF y del Excel: TODOS los filtros
    aplicados, o que no hay ninguno (regla de v1063)."""
    partes = []
    if sector:
        partes.append(f"sector {SECTORES[sector]}")
    if desde and hasta:
        partes.append(f"vencen del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}")
    elif desde:
        partes.append(f"vencen desde el {desde.strftime('%d/%m/%Y')}")
    elif hasta:
        partes.append(f"vencen hasta el {hasta.strftime('%d/%m/%Y')}")
    if estado:
        partes.append(f"estado {ESTADOS_VISIBLES[estado].lower()}")
    if creada_por:
        partes.append(f"creadas por {SECTORES[creada_por]}")
    return " · ".join(partes) or "sin filtros: todas las tareas"


def _fila(o: dict, hoy: date) -> tuple:
    return (
        o["vence_el"].strftime("%d/%m/%Y"), SECTORES[o["sector"]], SECTORES[o["creada_por"]], o["titulo"],
        o["detalle"] or "",
        ESTADOS_VISIBLES[estado_visible(o, hoy)] + (" (atrasada)" if o["atrasada"] else ""),
        hora_argentina(o["hecha_el"]), SECTORES.get(o["hecha_por"], ""), o["nota"] or "",
        dias_de_atraso(o, hoy),
    )


ENCABEZADOS = ("Vence", "Sector", "Creada por", "Tarea", "Detalle", "Estado", "Marcada el", "Quién", "Nota", "Días de atraso")


def generar_excel_tareas(filtro: str, ocurrencias: list[dict], hoy: date) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tareas"
    hoja.cell(row=1, column=1, value="Tareas").font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=f"Al {hoy.strftime('%d/%m/%Y')} · Filtros: {filtro}")
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(ENCABEZADOS, start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for numero_fila, o in enumerate(ocurrencias, start=5):
        for columna, valor in enumerate(_fila(o, hoy), start=1):
            hoja.cell(row=numero_fila, column=columna, value=valor)
    for letra, ancho in zip("ABCDEFGHIJ", (12, 15, 15, 30, 36, 18, 18, 15, 30, 10)):
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_tareas(filtro: str, ocurrencias: list[dict], hoy: date) -> bytes:
    """A4 apaisado: son diez columnas."""
    buffer = BytesIO()
    documento = SimpleDocTemplate(buffer, pagesize=landscape(A4), topMargin=14 * mm,
                                  leftMargin=12 * mm, rightMargin=12 * mm, bottomMargin=12 * mm)
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str("" if texto is None else texto)
        texto = texto.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    filas = [[_p(v) for v in _fila(o, hoy)] for o in ocurrencias] \
        or [[_p("No hay tareas con estos filtros.", "vacio")] + [""] * (len(ENCABEZADOS) - 1)]
    titulo = f"Tareas al {hoy.strftime('%d/%m/%Y')} — Filtros: {filtro}"
    anchos = [23 * mm, 22 * mm, 22 * mm, 38 * mm, 43 * mm, 24 * mm, 26 * mm, 20 * mm, 34 * mm, 18 * mm]
    documento.build([_tabla_seccion_pdf(titulo, list(ENCABEZADOS), filas, anchos, estilos)])
    return buffer.getvalue()
