"""VALES A COBRAR: cómo se dice cada cosa, los filtros y el Excel — puro.

Los vales llegan armados de `listar_vales` y `movimientos_de_vales`
(app/db.py). Acá no se recalcula nada: se decide cómo se DICE cada estado y
cada movimiento, y la pantalla y el Excel usan las mismas palabras.

"QUIÉN" ES EL SECTOR Y LA HORA (dueño, 30/09): el sistema no tiene usuarios.
"""

from datetime import date
from zoneinfo import ZoneInfo
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.exportar_vacios import VERDE_ENCABEZADO_HEX

TEXTO_DEL_ESTADO = {
    "en_cartera": "En cartera",
    "cobrado": "Cobrado",
    "cruzado": "Cruzado con el proveedor",
    "anulado": "Anulado",
    "devolucion_anulada": "Se anuló la devolución",
}

# El filtro de estado de la pantalla: "en_cartera" por defecto, "todos" es
# no filtrar.
OPCIONES_DE_ESTADO = (("en_cartera", "En cartera"), ("todos", "Todos")) + tuple(
    (k, v) for k, v in TEXTO_DEL_ESTADO.items() if k != "en_cartera")

TEXTO_DEL_ORIGEN = {
    "devolucion": "Devolución de vacíos",
    "anterior_al_sistema": "Anterior al sistema",
    "carga_manual": "Carga manual",
}

# El filtro de origen de la pantalla y de Movimientos: vacío es todos.
OPCIONES_DE_ORIGEN = (("", "Todos"),) + tuple(TEXTO_DEL_ORIGEN.items())

TEXTO_DEL_CAMPO = {"importe": "Importe", "numero": "Número", "fecha": "Fecha", "proveedor": "Proveedor"}


def origen_del_filtro(texto: str) -> str | None:
    """El origen pedido en la URL, o None (todos)."""
    return texto if texto in TEXTO_DEL_ORIGEN else None


def texto_del_valor_corregido(campo: str, valor: str | None) -> str:
    """Un valor del historial de correcciones como se lee en la pantalla."""
    if valor is None:
        return "sin número" if campo == "numero" else "—"
    if campo == "fecha":
        return date.fromisoformat(valor).strftime("%d/%m/%Y")
    if campo == "importe":
        return "$" + _miles(float(valor))
    return valor

TEXTO_DEL_SECTOR = {
    "administracion": "Administración",
    "gerencia": "Gerencia",
    "deposito": "Depósito",
    "compras": "Compras",
    None: "sin dato",
}

TEXTO_DEL_MOVIMIENTO = {"entrada": "Entró a cartera", **{
    k: v for k, v in TEXTO_DEL_ESTADO.items() if k != "en_cartera"}}


def estado_del_filtro(texto: str) -> str | None:
    """El estado pedido en la URL. Vacío es "en cartera"; "todos" es None."""
    if texto == "todos":
        return None
    return texto if texto in TEXTO_DEL_ESTADO else "en_cartera"


def dias_del_filtro(texto: str) -> int | None:
    """"Más de X días": un entero mayor o igual a cero, o nada."""
    texto = (texto or "").strip()
    return int(texto) if texto.isdigit() else None


def fecha_del_filtro(texto: str) -> date | None:
    try:
        return date.fromisoformat(texto) if (texto or "").strip() else None
    except ValueError:
        return None


def total_de(vales: list[dict]) -> float:
    """La suma de los importes de una lista, la misma cuenta para el total de
    arriba y para el de un filtro."""
    return round(sum(v["importe"] or 0 for v in vales), 2)


def texto_de_la_foto(foto: dict) -> str:
    """De dónde viene una foto del vale (dueño, 30/09). La original va con su
    fecha; la anexada con día, hora y el sector que la subió."""
    if foto["que"] == "devolucion":
        return f"Foto de la devolución ({_dia_argentino(foto['creado_en'])})"
    if foto["que"] == "papel":
        return f"Foto del vale en papel (cargado el {_dia_argentino(foto['creado_en'])})"
    cuando = foto["creado_en"].astimezone(ZoneInfo("America/Argentina/Buenos_Aires")).strftime("%d/%m %H:%M")
    return f"Anexada el {cuando} desde {TEXTO_DEL_SECTOR.get(foto['sector'], 'sin dato')}"


def _dia_argentino(instante) -> str:
    return instante.astimezone(ZoneInfo("America/Argentina/Buenos_Aires")).strftime("%d/%m/%Y")


def texto_de_la_salida(vale: dict) -> str:
    """Una línea con lo que pasó en la salida, o vacía si sigue en cartera."""
    estado = vale["estado"]
    if estado == "cobrado":
        texto = f"Cobrado ${_miles(vale['importe_cobrado'])}"
        if vale.get("ingreso_a_caja"):
            texto += f" · ingresó: {vale['ingreso_a_caja']}"
        return texto
    if estado == "cruzado":
        return f"Cruzado · {vale['referencia']}"
    if estado == "anulado":
        return f"Anulado · {vale['motivo']}"
    if estado == "devolucion_anulada":
        return "Se anuló la devolución de vacíos"
    return ""


def hora_argentina(instante) -> str:
    """Un instante de la base, en hora argentina ("30/09/2026 14:35"). Lo hace
    la que muestra, aunque llegue convertido (corolario 21)."""
    if instante is None:
        return ""
    return instante.astimezone(ZoneInfo("America/Argentina/Buenos_Aires")).strftime("%d/%m/%Y %H:%M")


def _miles(valor) -> str:
    entero = int(round(float(valor or 0)))
    return ("-" if entero < 0 else "") + f"{abs(entero):,}".replace(",", ".")


def generar_excel_movimientos_vales(desde: date, hasta: date, filtro: str,
                                    movimientos: list[dict]) -> bytes:
    """Lo filtrado, fila por fila, con las mismas palabras que la pantalla.
    Los importes van como NÚMERO, para poder sumarlos."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Vales a cobrar"
    hoja.cell(row=1, column=1, value="Vales a cobrar: movimientos").font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1,
              value=f"Del {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')} · {filtro}")
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    encabezados = ("Fecha", "Movimiento", "Vale", "Número", "Proveedor", "Origen", "Importe",
                   "Detalle", "Sector", "Cargado el")
    for columna, encabezado in enumerate(encabezados, start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    for fila, m in enumerate(movimientos, start=5):
        v = m["vale"]
        if m["que"] == "entrada":
            detalle = (f"{v['cajones']} cajones" if v.get("cajones") else "")
            sector = "carga por SQL" if v["origen"] == "anterior_al_sistema" \
                else TEXTO_DEL_SECTOR.get(v.get("cargada_desde"), "sin dato")
            cargado = v["creado_en"]
        else:
            detalle = texto_de_la_salida(v)
            sector = TEXTO_DEL_SECTOR.get(v.get("salida_sector"), "") if v.get("salida_sector") else ""
            cargado = v.get("salida_creado_en") or v.get("devolucion_anulada_el")
        valores = (
            m["fecha"].strftime("%d/%m/%Y"),
            TEXTO_DEL_MOVIMIENTO[m["que"]],
            v["id"],
            v["numero"] or "",
            v["proveedor"],
            TEXTO_DEL_ORIGEN[v["origen"]],
            m["importe"],
            detalle,
            sector,
            hora_argentina(cargado),
        )
        for columna, valor in enumerate(valores, start=1):
            hoja.cell(row=fila, column=columna, value=valor)
    for letra, ancho in zip("ABCDEFGHIJ", (12, 26, 8, 12, 30, 22, 14, 36, 16, 18)):
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# VALES POR PROVEEDOR (dueño, 09/10): los vales EN CARTERA agrupados, un
# renglón por proveedor con cuántos y cuánto, de mayor a menor importe, y el
# total general abajo. Recibe los vales de `resumen_de_la_cartera`, la MISMA
# lista que suma el cuadro del Panel de control: no hay otra consulta, así
# que el total general no puede dar distinto del panel.

def vales_por_proveedor(vales: list[dict]) -> dict:
    grupos: dict = {}
    for vale in vales:
        grupo = grupos.setdefault(vale["proveedor_id"], {
            "proveedor_id": vale["proveedor_id"], "proveedor": vale["proveedor"], "cantidad": 0, "total": 0.0})
        grupo["cantidad"] += 1
        grupo["total"] += float(vale["importe"] or 0)
    filas = sorted(({**g, "total": round(g["total"], 2)} for g in grupos.values()),
                   key=lambda g: (-g["total"], g["proveedor"].lower()))
    return {"filas": filas, "cantidad": len(vales), "total": total_de(vales)}


TITULO_POR_PROVEEDOR = "Vales por proveedor"
QUE_SE_MUESTRA_POR_PROVEEDOR = "En cartera: ni cobrados ni aplicados a una liquidación · todos los proveedores"


def generar_excel_vales_por_proveedor(hoy: date, agrupados: dict) -> bytes:
    """Lo mismo que la pantalla: un renglón por proveedor y el total abajo.
    Los importes van como NÚMERO, para poder sumarlos."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Vales por proveedor"
    hoja.cell(row=1, column=1, value=TITULO_POR_PROVEEDOR).font = Font(bold=True, size=14)
    hoja.cell(row=2, column=1, value=f"Al {hoy.strftime('%d/%m/%Y')} · {QUE_SE_MUESTRA_POR_PROVEEDOR}")
    relleno = PatternFill(start_color="DEEFE3", end_color="DEEFE3", fill_type="solid")
    for columna, encabezado in enumerate(("Proveedor", "Vales", "Importe"), start=1):
        celda = hoja.cell(row=4, column=columna, value=encabezado)
        celda.font = Font(bold=True, color=VERDE_ENCABEZADO_HEX)
        celda.fill = relleno
    fila = 5
    for grupo in agrupados["filas"]:
        for columna, valor in enumerate((grupo["proveedor"], grupo["cantidad"], grupo["total"]), start=1):
            hoja.cell(row=fila, column=columna, value=valor)
        fila += 1
    for columna, valor in enumerate(("Total", agrupados["cantidad"], agrupados["total"]), start=1):
        hoja.cell(row=fila, column=columna, value=valor).font = Font(bold=True)
    for letra, ancho in zip("ABC", (34, 10, 16)):
        hoja.column_dimensions[letra].width = ancho
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def generar_pdf_vales_por_proveedor(hoy: date, agrupados: dict) -> bytes:
    """A4 vertical, con lo que se muestra en el encabezado (la pantalla no
    tiene filtros: son todos los vales en cartera)."""
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph

    from core.exportar_vacios import _documento_pdf, _estilos_pdf, _tabla_seccion_pdf

    buffer = BytesIO()
    documento, encabezado = _documento_pdf(buffer, TITULO_POR_PROVEEDOR,
                                           f"Al {hoy.strftime('%d/%m/%Y')} · {QUE_SE_MUESTRA_POR_PROVEEDOR}")
    estilos = _estilos_pdf()

    def _p(texto, estilo="dato"):
        texto = str(texto).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(texto, estilos[estilo])

    filas = [[_p(g["proveedor"]), _p(g["cantidad"], "numero"), _p(f"${_miles(g['total'])}", "numero")]
             for g in agrupados["filas"]] or [[_p("No hay vales en cartera.", "vacio"), "", ""]]
    filas.append([_p("Total", "numero"), _p(agrupados["cantidad"], "numero"),
                  _p(f"${_miles(agrupados['total'])}", "numero")])
    tabla = _tabla_seccion_pdf(f"{len(agrupados['filas'])} proveedores con vales en cartera",
                               ["Proveedor", "Vales", "Importe"], filas, [104 * mm, 30 * mm, 44 * mm], estilos)
    documento.build([tabla], onFirstPage=encabezado, onLaterPages=encabezado)
    return buffer.getvalue()
