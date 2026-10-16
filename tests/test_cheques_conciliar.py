"""CARTERA DE CHEQUES: conciliar con el banco (dueño, 16/10).

El listado del banco (ECHEQ en cartera o en custodia) contra los cheques EN
CARTERA del sistema, en cuatro grupos. La IA lee cada foto (acá se reemplaza
por lo que habría devuelto: lo que se prueba es el cruce).

EL GALPÓN (nombres inventados; "P" es la fecha de pago de todos):
  ECHEQ  100 EJEMPLO Banco $10.000   → el banco lo lee igual (y en dos archivos)  → coincide
  ECHEQ  200 EJEMPLO Banco $20.000   → el banco dice $20.500                     → con diferencia (monto)
  ECHEQ  300 EJEMPLO Banco $30.000   → el banco no lo tiene                       → en el sistema y no en el banco
  ECHEQ  400 Galicia       $40.000   → el banco dice P+1                          → con diferencia (fecha)
  ECHEQ  500 Galicia       $50.000   → el banco lee un 500 de OTRO banco (el RIVAL): no es el mismo
  ECHEQ  900 EJEMPLO Banco $9.000    → el banco lo lee SIN banco: por número y monto → coincide
  FÍSICO 600 EJEMPLO Banco $60.000   → no está en el banco: aparte, no falta
  el banco además lee el 800 ($80.000), que no está en el sistema
"""
import io
import os
import re
import sys
from datetime import timedelta
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.conciliar_cartera import cruzar_con_la_cartera, juntar_lo_leido  # noqa: E402
from tests.test_administracion_reordenada import _cliente  # noqa: E402
from tests.test_cheques import CUIT_A, _hoy, d  # noqa: E402,F401
from tests.test_cheques_pantallas import _foto, _sin_estilo, m  # noqa: E402,F401


def _cheque(d, numero, banco, importe, pago, tipo="echeq"):
    return d.cargar_cheque(numero=numero, banco_id=banco, tipo=tipo, fecha_carga=_hoy(d) - timedelta(days=40),
                           fecha_pago=pago, importe=importe, origen="cliente", cliente_id=1, nombre_a_mano=None,
                           cuit_librador=CUIT_A, fotos=[])


def _galpon(d):
    pago = _hoy(d) + timedelta(days=10)
    galicia = d.agregar_banco("Galicia")
    ids = {n: _cheque(d, n, b, i, pago) for n, b, i in (
        ("100", 1, 10000.0), ("200", 1, 20000.0), ("300", 1, 30000.0), ("400", galicia, 40000.0),
        ("500", galicia, 50000.0), ("900", 1, 9000.0))}
    ids["600"] = _cheque(d, "600", 1, 60000.0, pago, tipo="fisico")
    return ids, pago


def _leido(numero, banco, importe, pago):
    return {"numero": numero, "banco": banco, "importe": importe, "fecha_pago": pago.isoformat(),
            "cuit_librador": "30712345671", "librador": "EJEMPLO Librador SA", "tipo": "echeq"}


def _archivos(pago):
    uno = [_leido("0100", "Banco EJEMPLO S.A.", 10000, pago), _leido("200", "Banco EJEMPLO S.A.", 20500, pago),
           _leido("400", "BANCO DE GALICIA Y BUENOS AIRES", 40000, pago + timedelta(days=1)),
           _leido("500", "Banco EJEMPLO S.A.", 50000, pago), _leido("800", "Banco EJEMPLO S.A.", 80000, pago)]
    # El segundo trae otra vez el 100, igual: se toma una sola vez.
    dos = [_leido("0100", "Banco EJEMPLO S.A.", 10000, pago), _leido("900", "", 9000, pago)]
    return [{"cheques": uno}, {"cheques": dos}]


def _subir(m, lecturas, solo_echeq=True):
    archivos = [("archivos", (f"banco{i + 1}.jpg", _foto(), "image/jpeg")) for i in range(len(lecturas))]
    with patch("app.main.extraer_cheques_de_imagenes", side_effect=lecturas) as lector:
        pagina = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar", files=archivos,
                                              data={"solo_echeq": "1"} if solo_echeq else {})
    assert lector.call_count == len(lecturas)
    return pagina


def _grupo(html, nombre):
    marcado = _sin_estilo(html)
    bloque = marcado.split(f'data-grupo="{nombre}"')[1].split('data-grupo="')[0]
    return {"cuantos": int(re.search(r"data-cuantos>(\d+)<", bloque).group(1)),
            "total": re.search(r"data-total>([^<]+)<", bloque).group(1),
            "numeros": re.findall(r'class="fila" data-numero="([^"]+)"', bloque), "html": bloque}


def _campos(html, accion):
    formulario = _sin_estilo(html).split(f'action="{accion}"')[1].split("</form>")[0]
    datos = {n: v for n, v in re.findall(r'<input type="hidden" name="([a-z_]+)" value="([^"]*)"', formulario)}
    import html as h
    return {k: h.unescape(v) for k, v in datos.items()}


def test_los_CUATRO_GRUPOS_con_un_archivo_REPETIDO(m, d):
    ids, pago = _galpon(d)
    pagina = _subir(m, _archivos(pago))
    assert pagina.status_code == 200, pagina.text[:300]
    marcado = _sin_estilo(pagina.text)
    assert re.findall(r'<div class="aviso" data-aviso>([^<]*)</div>', marcado) == [
        "1 cheque repetido entre archivos: se tomó una sola vez."]
    # Arriba el resumen; el formulario para subir otro listado, al final.
    assert marcado.index("data-resumen>") < marcado.index('data-grupo="coinciden"') < marcado.index("data-subir")
    assert re.sub(r"<[^>]+>", "", re.search(r"data-resumen>(.*?)</p>", marcado).group(1)) == (
        "Banco: 6 cheques $209.500 · Sistema: 6 cheques $159.000 · Diferencia $50.500")
    grupos = {g: _grupo(pagina.text, g) for g in ("coinciden", "solo_banco", "solo_sistema", "con_diferencia",
                                                   "fisicos")}
    assert {g: (v["cuantos"], sorted(v["numeros"]), v["total"]) for g, v in grupos.items()} == {
        "coinciden": (2, ["100", "900"], "$19.000"),
        "solo_banco": (2, ["500", "800"], "$130.000"),
        "solo_sistema": (2, ["300", "500"], "$80.000"),
        "con_diferencia": (2, ["200", "400"], "banco $60.500 · sistema $60.000"),
        "fisicos": (1, ["600"], "$60.000")}
    # Coinciden: verde y plegado. Con diferencia: las dos versiones, lado a lado.
    assert '<details class="tarjeta grupo verde" data-grupo="coinciden">' in marcado
    dif = grupos["con_diferencia"]["html"]
    assert re.findall(r"difiere ([a-z ]+)\n", dif) == ["monto", "fecha de pago"]
    assert len(re.findall(r'data-version="banco"', dif)) == len(re.findall(r'data-version="sistema"', dif)) == 2
    # En el sistema y no en el banco: Aplicar y Rechazado, cada uno con SU cheque.
    falta = grupos["solo_sistema"]["html"]
    assert re.findall(r'href="/gerencia/cheques/aplicar\?cheque_id=(\d+)" data-aplicar', falta) == [
        str(ids["300"]), str(ids["500"])]
    assert re.findall(r'href="/gerencia/cheques/(\d+)#rechazar" data-rechazado', falta) == [
        str(ids["300"]), str(ids["500"])]


def test_sin_SOLO_ECHEQ_el_fisico_cuenta_como_que_no_esta_en_el_banco(m, d):
    ids, pago = _galpon(d)
    pagina = _subir(m, _archivos(pago), solo_echeq=False)
    assert sorted(_grupo(pagina.text, "solo_sistema")["numeros"]) == ["300", "500", "600"]
    assert 'data-grupo="fisicos"' not in _sin_estilo(pagina.text)
    # Y volver a cruzar con "Solo ECHEQ" lo saca, sin volver a subir nada.
    campos = _campos(pagina.text, "/gerencia/cheques/conciliar/cruzar")
    otra = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar/cruzar",
                                        data={"leidos": campos["leidos"], "solo_echeq": "1"})
    assert sorted(_grupo(otra.text, "solo_sistema")["numeros"]) == ["300", "500"]
    assert _grupo(otra.text, "fisicos")["numeros"] == ["600"]


def test_CARGARLOS_abre_Revisar_cheques_con_los_que_faltan(m, d):
    _galpon(d)
    pagina = _subir(m, _archivos(_hoy(d) + timedelta(days=10)))
    campos = _campos(pagina.text, "/gerencia/cheques/conciliar/cargar")
    assert campos["solo_echeq"] == "1"
    revisar = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar/cargar", data=campos)
    assert revisar.status_code == 200
    marcado = _sin_estilo(revisar.text)
    numeros = re.findall(r'name="numero_(\d+)"[^>]*value="([^"]*)"', marcado)
    assert sorted(v for _, v in numeros) == ["500", "800"]
    assert "2 cheques del banco que no estaban en el sistema" in revisar.text


def test_PDF_y_EXCEL_salen_con_el_MISMO_cruce_y_lo_dicen(m, d):
    from openpyxl import load_workbook
    _galpon(d)
    pagina = _subir(m, _archivos(_hoy(d) + timedelta(days=10)))
    campos = _campos(pagina.text, "/gerencia/cheques/conciliar/pdf")
    pdf = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar/pdf", data=campos)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    excel = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar/excel", data=campos)
    hoja = load_workbook(io.BytesIO(excel.content)).active
    textos = [str(c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None]
    assert textos[1].endswith("6 cheques leídos del banco contra los ECHEQ en cartera (los físicos, aparte)")
    assert textos[2] == "Banco: 6 cheques $209.500 · Sistema: 6 cheques $159.000 · Diferencia $50.500"
    titulos = [t for t in textos if re.match(r"^(En el|Con diferencia|Coinciden|Físicos)", t)]
    assert titulos == ["En el sistema y NO en el banco: 2 · $80.000",
                       "En el banco y NO en el sistema (faltan cargar): 2 · $130.000",
                       "Con diferencia: 2 · banco $60.500 · sistema $60.000", "Coinciden: 2 · $19.000",
                       "Físicos en cartera, aparte (no figuran en el banco): 1 · $60.000"]
    # Sin "Solo ECHEQ", el encabezado lo dice y los físicos no van aparte.
    sin = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar/excel", data={"leidos": campos["leidos"]})
    textos = [str(c.value) for fila in load_workbook(io.BytesIO(sin.content)).active.iter_rows()
              for c in fila if c.value is not None]
    assert textos[1].endswith("contra los ECHEQ y físicos en cartera")
    assert not any(t.startswith("Físicos en cartera") for t in textos)


def test_un_archivo_que_NO_SE_LEE_no_frena_los_demas_y_otro_error_no_se_esconde(m, d):
    _galpon(d)
    pago = _hoy(d) + timedelta(days=10)
    pagina = _subir(m, [RuntimeError("la API no contestó"), _archivos(pago)[0]])
    assert re.findall(r'<div class="aviso" data-aviso>([^<]*)</div>', _sin_estilo(pagina.text)) == [
        "El archivo 1 («banco1.jpg») no se pudo leer."]
    assert _grupo(pagina.text, "coinciden")["numeros"] == ["100"]
    with pytest.raises(KeyError):
        _subir(m, [KeyError("otra cosa")])
    sin_nada = _cliente(m, "gerencia").post("/gerencia/cheques/conciliar", data={"solo_echeq": "1"})
    assert sin_nada.status_code == 400 and "Subí el listado del banco" in sin_nada.text


def test_la_PANTALLA_de_inicio_y_el_boton_en_la_Cartera_solo_Gerencia(m, d):
    cliente = _cliente(m, "gerencia")
    inicio = _sin_estilo(cliente.get("/gerencia/cheques/conciliar").text)
    assert '<input type="file" name="archivos" accept="image/*,.pdf,.xlsx,.xls,.csv" multiple>' in inicio
    assert '<input type="checkbox" name="solo_echeq" value="1" checked>' in inicio     # por defecto
    assert 'href="/gerencia/cheques/conciliar">Conciliar con el banco<' in _sin_estilo(cliente.get("/gerencia/cheques").text)
    sin_clave = _cliente(m).get("/gerencia/cheques/conciliar")
    assert 'name="clave"' in sin_clave.text and 'name="archivos"' not in sin_clave.text


# --- SIN BASE ----------------------------------------------------------------------

def test_REPETIDOS_solo_entre_archivos_distintos():
    fila = {"numero": "1", "banco": "X", "importe": 5, "fecha_pago": "2026-11-01"}
    juntos, repetidos = juntar_lo_leido([[fila, fila], [dict(fila, numero="0001")], [dict(fila, importe=6)]])
    # Dentro de un archivo, cada renglón es un cheque; el del segundo es el mismo; el tercero no.
    assert (len(juntos), repetidos) == (3, 1)


def test_el_cruce_sin_banco_es_por_MONTO_y_con_banco_no_mira_el_monto():
    from datetime import date
    cartera = [{"id": 1, "numero": "7", "banco": "EJEMPLO Banco", "importe": 100, "fecha_pago": date(2026, 11, 1),
                "estado": "en_cartera", "tipo": "echeq"}]
    leidos, _ = juntar_lo_leido([[{"numero": "7", "banco": "", "importe": 101, "fecha_pago": "2026-11-01"}]])
    r = cruzar_con_la_cartera(leidos, cartera)
    assert (len(r["solo_banco"]), len(r["solo_sistema"])) == (1, 1)            # sin banco y otro monto: no es
    leidos, _ = juntar_lo_leido([[{"numero": "7", "banco": "EJEMPLO", "importe": 101, "fecha_pago": "2026-11-01"}]])
    r = cruzar_con_la_cartera(leidos, cartera)
    assert [p["diferencias"] for p in r["con_diferencia"]] == [["monto"]]   # con banco: es, con diferencia
