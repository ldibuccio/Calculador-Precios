"""CARTERA DE CHEQUES (dueño, 13/10): las pantallas, contra Postgres.

Solo Gerencia, con su clave. La lectura por foto se reemplaza por lo que la IA
habría devuelto (el lector no se prueba acá: se prueba qué hace la pantalla
con lo leído). Storage tampoco: la subida devuelve una ruta inventada.

EL GALPÓN, con nombres inventados: cliente 1 EJEMPLO Verdulería, proveedor 1
EJEMPLO Puesto, banco 1 EJEMPLO Banco. Hoy, para las pantallas, es el de la
base (la vista lee la fecha de Argentina).
"""
import io
import os
import re
import sys
from datetime import timedelta
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from tests.test_administracion_reordenada import CLAVES, _cliente  # noqa: E402
from tests.test_cheques import CUIT_A, CUIT_B, _cargar, _hoy, d  # noqa: E402,F401


@pytest.fixture
def m(d, monkeypatch):
    for clave, valor in CLAVES.items():
        monkeypatch.setenv(clave, valor)
    import app.main as modulo
    return modulo


def _foto() -> bytes:
    from PIL import Image
    salida = io.BytesIO()
    Image.new("RGB", (20, 10), "white").save(salida, format="JPEG")
    return salida.getvalue()


def _error(respuesta) -> str:
    """El cartel de error de la pantalla, para que un assert diga el motivo."""
    texto = respuesta.text
    return texto.split("data-error>")[1].split("<")[0] if "data-error>" in texto else texto[:200]


def _sin_estilo(html: str) -> str:
    return html.split("</style>")[-1]


def test_SOLO_GERENCIA_sin_la_clave_pide_la_clave(m):
    pagina = _cliente(m).get("/gerencia/cheques")
    assert 'name="clave"' in pagina.text and "Cartera de cheques" not in pagina.text.split("<title>")[1][:40]
    assert _cliente(m).post("/gerencia/cheques/cargar", data={}, follow_redirects=False).status_code in (303, 401, 403)


def test_la_LISTA_dice_lo_de_arriba_y_el_listado_filtrado(m, d):
    hoy = _hoy(d)
    _cargar(d, "1", importe=1000.0, pago=hoy + timedelta(days=2))
    _cargar(d, "2", importe=2000.0, pago=hoy - timedelta(days=24))        # por vencer
    _cargar(d, "3", banco=2, importe=4000.0, pago=hoy + timedelta(days=20))
    pagina = _cliente(m, "gerencia").get("/gerencia/cheques")
    assert pagina.status_code == 200, pagina.text[:300]
    marcado = _sin_estilo(pagina.text)
    assert re.search(r'<p class="total" data-en-cartera>\$7\.000</p>', marcado)
    assert "Entra en 7 días: <strong>$1.000</strong> (1)" in marcado
    assert re.search(r'<p class="total-dato rojo" data-en-rojo>\s*1 por vencer</p>', marcado)
    assert re.findall(r'data-cheque="(\d+)"', marcado) == ["2", "1", "3"]          # por fecha de pago
    filtrada = _sin_estilo(_cliente(m, "gerencia").get("/gerencia/cheques", params={"banco_id": 2}).text)
    assert re.findall(r'data-cheque="(\d+)"', filtrada) == ["3"]
    assert re.search(r"<h2>1 cheque · <span class=\"importe\">\$4\.000</span></h2>", filtrada)
    # Arriba sigue diciendo TODA la cartera, no lo filtrado.
    assert re.search(r'data-en-cartera>\$7\.000</p>', filtrada)


def test_las_EXPORTACIONES_salen_con_los_filtros_de_la_pantalla(m, d):
    hoy = _hoy(d)
    _cargar(d, "1", importe=1000.0)
    _cargar(d, "2", banco=2, importe=4000.0)
    cliente = _cliente(m, "gerencia")
    pagina = cliente.get("/gerencia/cheques", params={"banco_id": 2, "pago_desde": hoy.isoformat()}).text
    links = dict(re.findall(r'<a href="(/gerencia/cheques/exportar-(pdf|excel)\?[^"]*)">', pagina))
    links = {tipo: url.replace("&amp;", "&") for url, tipo in links.items()}
    assert set(links) == {"pdf", "excel"}
    for url in links.values():
        assert "banco_id=2" in url and f"pago_desde={hoy.isoformat()}" in url
    excel = cliente.get(links["excel"])
    assert excel.headers["content-type"].startswith("application/vnd.openxmlformats")
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(excel.content)).active
    assert "banco EJEMPLO Otro Banco" in hoja.cell(row=2, column=1).value
    assert f"pago desde el {hoy.strftime('%d/%m/%Y')}" in hoja.cell(row=2, column=1).value
    numeros = [hoja.cell(row=f, column=1).value for f in range(5, hoja.max_row + 1) if hoja.cell(row=f, column=1).value]
    assert numeros == ["2"]
    pdf = cliente.get(links["pdf"])
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")


def _cargar_por_pantalla(m, filas, **comunes):
    datos = {"cantidad": str(len(filas)), "fecha_carga": comunes.get("fecha_carga", "2026-10-01"),
             "origen": comunes.get("origen", "cliente"), "cliente_id": comunes.get("cliente_id", "1"),
             "nombre_a_mano": comunes.get("nombre_a_mano", "")}
    for i, fila in enumerate(filas):
        for campo, valor in fila.items():
            datos[f"{campo}_{i}"] = valor
    return _cliente(m, "gerencia").post("/gerencia/cheques/cargar", data=datos, follow_redirects=False)


def _fila(numero, cuit, banco="EJEMPLO Banco", importe="1500.50"):
    return {"numero": numero, "banco": banco, "tipo": "fisico", "fecha_pago": "2026-10-20",
            "importe": importe, "cuit_librador": cuit}


def test_CARGAR_por_pantalla_los_REPETIDOS_son_banco_numero_y_CUIT(m, d):
    """Dos libradores con el mismo banco y número entran juntos; el mismo CUIT
    otra vez rebota con el motivo, y no queda nada a medias."""
    respuesta = _cargar_por_pantalla(m, [_fila("555", CUIT_A), _fila("555", CUIT_B)])
    assert respuesta.status_code == 303, _error(respuesta)
    assert {c["cuit_librador"] for c in d.buscar_cheques({"numero": "555"})} == {CUIT_A, CUIT_B}
    otra = _cargar_por_pantalla(m, [_fila("777", CUIT_A), _fila("0555", CUIT_A)])
    assert otra.status_code == 400
    assert "Cheque 0555: Ese cheque ya está cargado: mismo banco, número y CUIT del librador." in otra.text
    assert d.buscar_cheques({"numero": "777"}) == []                               # todo o nada


def test_CARGAR_un_banco_nuevo_se_agrega_y_el_importe_tipeado_con_coma(m, d):
    respuesta = _cargar_por_pantalla(m, [_fila("9", CUIT_A, banco="  EJEMPLO   Nuevo ", importe="1.234,50")],
                                     origen="a_mano", nombre_a_mano="EJEMPLO Persona", cliente_id="")
    assert respuesta.status_code == 303, _error(respuesta)
    cheque, = d.buscar_cheques({"numero": "9"})
    assert (cheque["banco"], float(cheque["importe"]), cheque["de_quien"]) == ("EJEMPLO Nuevo", 1234.5, "EJEMPLO Persona")


def test_CARGAR_valida_antes_de_tocar_la_base(m, d):
    respuesta = _cargar_por_pantalla(m, [_fila("1", "20-1")], origen="a_mano", nombre_a_mano="")
    assert respuesta.status_code == 400
    assert "Escribí de quién es el cheque." in respuesta.text
    assert "el CUIT del librador son 11 números." in respuesta.text
    assert d.buscar_cheques({}) == []
    # 11 números con el dígito verificador mal (13/10): también rebota.
    mal = _cargar_por_pantalla(m, [_fila("1", "20111111111")], origen="canje")
    assert mal.status_code == 400 and "el CUIT del librador no es válido" in mal.text


def test_CARGAR_por_FOTO_muestra_lo_leido_para_corregir_y_no_guarda(m, d):
    leido = {"fecha": None, "cheques": [{"numero": "00042", "banco": "Banco EJEMPLO Otro Banco",
                                         "fecha_pago": "2026-11-02", "importe": 85000.5,
                                         "cuit_librador": "20-11111111-1", "tipo": "fisico", "librador": "EJEMPLO SA"}]}
    with patch("app.main.extraer_cheques_de_imagenes", return_value=leido) as lector:
        pagina = _cliente(m, "gerencia").post("/gerencia/cheques/cargar/leer",
                                              files={"archivo": ("cheque.jpg", _foto(), "image/jpeg")})
    assert lector.call_count == 1 and pagina.status_code == 200, pagina.text[:300]
    marcado = _sin_estilo(pagina.text)
    assert 'name="numero_0" inputmode="numeric" value="00042"' in marcado
    assert 'name="banco_0" list="bancos" placeholder="Ej. Galicia" value="EJEMPLO Otro Banco"' in marcado
    assert 'name="cuit_librador_0" inputmode="numeric" placeholder="Ej. 20123456789" value="20111111111"' in marcado
    assert 'name="importe_0" step="0.01" min="0.01" inputmode="decimal" value="85000.5"' in marcado
    assert re.search(r'name="foto_leida" value="data:image/jpeg;base64,', marcado)
    assert d.buscar_cheques({}) == []


def test_CARGAR_la_foto_leida_queda_como_FRENTE(m, d):
    from tests.test_cheques_pantallas import _foto as foto
    import base64
    data_uri = "data:image/jpeg;base64," + base64.standard_b64encode(foto()).decode()
    datos = {"cantidad": "1", "fecha_carga": "2026-10-01", "origen": "canje", "foto_leida": data_uri,
             **{f"{k}_0": v for k, v in _fila("31", CUIT_A).items()}}
    with patch("app.main.subir_foto_comanda", return_value="cheques/2026-10-13/frente.jpg") as subir:
        respuesta = _cliente(m, "gerencia").post("/gerencia/cheques/cargar", data=datos, follow_redirects=False)
    assert respuesta.status_code == 303 and subir.call_count == 1
    cheque_id = int(respuesta.headers["location"].split("/gerencia/cheques/")[1].split("?")[0])
    assert [(f["lado"], f["foto_ruta"]) for f in d.obtener_cheque(cheque_id)["fotos"]] == \
        [("frente", "cheques/2026-10-13/frente.jpg")]


def test_APLICAR_varios_juntos_y_por_FOTO_tilda_los_que_encuentra(m, d):
    uno, dos, tres = _cargar(d, "11"), _cargar(d, "12", cuit=CUIT_B), _cargar(d, "13")
    leido = {"fecha": "2026-10-12", "cheques": [{"numero": "0011", "cuit_librador": None},
                                                 {"numero": "12"}, {"numero": "99"}]}
    cliente = _cliente(m, "gerencia")
    with patch("app.main.extraer_cheques_de_imagenes", return_value=leido):
        pagina = cliente.post("/gerencia/cheques/aplicar/leer",
                              files={"archivo": ("boleta.jpg", _foto(), "image/jpeg")}).text
    tildados = re.findall(r'name="cheque_id" value="(\d+)" checked', pagina)
    assert sorted(map(int, tildados)) == [uno, dos]
    assert "No están en cartera: 99." in pagina
    assert 'name="fecha" value="2026-10-12"' in pagina
    respuesta = cliente.post("/gerencia/cheques/aplicar", data={"cheque_id": [str(uno), str(dos)],
                             "fecha": _hoy(d).isoformat(), "concepto_id": "3"}, follow_redirects=False)
    assert respuesta.status_code == 303, _error(respuesta)
    estados = {c["id"]: c["estado"] for c in d.buscar_cheques({})}
    assert estados == {uno: "depositado", dos: "depositado", tres: "en_cartera"}


def test_APLICAR_un_pago_sin_proveedor_rebota_y_un_concepto_nuevo_se_crea(m, d):
    cheque = _cargar(d, "21")
    cliente = _cliente(m, "gerencia")
    sin = cliente.post("/gerencia/cheques/aplicar", data={"cheque_id": str(cheque), "fecha": _hoy(d).isoformat(),
                                                          "concepto_id": "1"}, follow_redirects=False)
    assert sin.status_code == 400 and "Un pago a proveedor tiene que decir a qué proveedor." in sin.text
    nuevo = cliente.post("/gerencia/cheques/aplicar", data={
        "cheque_id": str(cheque), "fecha": _hoy(d).isoformat(), "concepto_nuevo": "EJEMPLO Pago de flete",
        "destino_nuevo": "entregado", "proveedor_id": "1"}, follow_redirects=False)
    assert nuevo.status_code == 303, nuevo.text[:300]
    ficha = d.obtener_cheque(cheque)
    assert (ficha["estado"], ficha["concepto"], ficha["proveedor"]) == ("entregado", "EJEMPLO Pago de flete",
                                                                         "EJEMPLO Puesto")


def test_la_FICHA_rechazar_resolver_y_el_historial(m, d):
    cheque = _cargar(d, "41")
    cliente = _cliente(m, "gerencia")
    hoy = _hoy(d).isoformat()
    ficha = _sin_estilo(cliente.get(f"/gerencia/cheques/{cheque}").text)
    assert f'href="/gerencia/cheques/aplicar?cheque_id={cheque}' in ficha
    assert cliente.post(f"/gerencia/cheques/{cheque}/rechazar", data={"fecha": hoy, "motivo": "Sin fondos"},
                        follow_redirects=False).status_code == 303
    ficha = _sin_estilo(cliente.get(f"/gerencia/cheques/{cheque}").text)
    assert '<span class="rojo" data-estado="rechazado">Rechazado</span>' in ficha
    assert "aplicar?cheque_id=" not in ficha                                   # un rechazado no se aplica
    assert f'action="/gerencia/cheques/{cheque}/resolver' in ficha
    assert cliente.post(f"/gerencia/cheques/{cheque}/resolver", data={"fecha": hoy, "resolucion": "Pagó"},
                        follow_redirects=False).status_code == 303
    ficha = _sin_estilo(cliente.get(f"/gerencia/cheques/{cheque}").text)
    assert 'data-estado="reclamo_resuelto">Reclamo resuelto</span>' in ficha
    historial = ficha.split("data-historial>")[1].split("</ul>")[0]
    assert historial.count("<li>") == 3 and "lo marcó rechazado: Sin fondos" in historial


def test_la_FICHA_corregir_solo_en_cartera(m, d):
    cheque = _cargar(d, "51", importe=1000.0)
    cliente = _cliente(m, "gerencia")
    ficha = d.obtener_cheque(cheque)
    datos = {"numero": "51", "banco_id": "1", "tipo": "fisico", "fecha_carga": ficha["fecha_carga"].isoformat(),
             "fecha_pago": ficha["fecha_pago"].isoformat(), "importe": "1100", "origen": "cliente",
             "cliente_id": "1", "nombre_a_mano": "", "cuit_librador": CUIT_A}
    assert cliente.post(f"/gerencia/cheques/{cheque}/corregir", data=datos,
                        follow_redirects=False).status_code == 303
    assert float(d.obtener_cheque(cheque)["importe"]) == 1100.0
    d.aplicar_cheque(cheque, fecha=_hoy(d), concepto_id=3, proveedor_id=None, foto_ruta=None)
    rebote = cliente.post(f"/gerencia/cheques/{cheque}/corregir", data=dict(datos, importe="1200"),
                          follow_redirects=False)
    assert rebote.status_code == 400 and "ya no está en cartera" in rebote.text


def test_el_PANEL_cuadro_Cheques_rojo_si_hay_por_vencer(m, d):
    hoy = _hoy(d)
    _cargar(d, "61", importe=500.0, pago=hoy + timedelta(days=3))
    datos = m._panel_cheques(hoy)
    assert (datos["en_cartera"], datos["entran_7_dias"], datos["en_rojo"]) == (500.0, 500.0, 0)
    _cargar(d, "62", importe=700.0, pago=hoy - timedelta(days=23))
    assert m._panel_cheques(hoy)["en_rojo"] == 1
    with patch.dict(os.environ, CLAVES), patch.object(m, "_hoy_argentina", return_value=hoy):
        detalle = _sin_estilo(_cliente(m, "gerencia").get("/gerencia/panel/cheques").text)
    assert '<p data-ajustar class="grande rojo">$1.200</p>' in detalle
    assert 'class="dato rojo" data-por-vencer>Por vencer: 1</p>' in detalle
