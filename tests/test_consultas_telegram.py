"""CONSULTAS POR TELEGRAM, EXCEL Y FOTOS (dueño, 17/10).

El mismo motor que Gerencia → Consultas: la IA es de mentira (contesta lo que
contestaría la API), Telegram también (se mira lo que la app le pide a su
API) y el SQL corre de verdad contra Postgres. "Pesada" y "balanza" son la
misma foto: la de la balanza al recibir; la otra es la comanda de la guía.
"""
import ast
import io
import os
import sys
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.consultas_ia import FILAS_MAXIMAS_DEL_EXCEL, instrucciones, preguntar  # noqa: E402
from tests.test_consultas_ia import ClienteDeMentira, _pide, _texto  # noqa: E402
from tests.test_panel_control_contra_la_base import _base_con, _cliente_gerencia  # noqa: E402

SECRETO = "secreto-de-prueba-123"
LIONEL, EXTRANO, OTRO = 111, 222, 333


# --- LO PURO ------------------------------------------------------------------------

def test_el_EXCEL_tiene_titulo_encabezados_fijos_y_cada_tipo_con_su_formato():
    from openpyxl import load_workbook

    from core.consultas_excel import armar_xlsx, nombre_del_archivo
    columnas = [{"nombre": "Proveedor", "tipo": "texto"}, {"nombre": "Kilos", "tipo": "kilos"},
                {"nombre": "Importe", "tipo": "pesos"}, {"nombre": "Fecha", "tipo": "fecha"}]
    filas = [["EJEMPLO Puesto", 160.5, 100000, "2026-09-10"], ["EJEMPLO Otro", None, "12,5", date(2026, 9, 11)]]
    hoja = load_workbook(io.BytesIO(armar_xlsx("Compras de Tomate", columnas, filas, False))).active
    assert hoja["A1"].value == "Compras de Tomate" and hoja.freeze_panes == "A3"
    assert [c.value for c in hoja[2]] == ["Proveedor", "Kilos", "Importe", "Fecha"]
    assert (hoja["B3"].value, hoja["B3"].number_format) == (160.5, "#,##0.0")
    assert (hoja["C3"].value, hoja["C3"].number_format) == (100000, '"$" #,##0.00')
    assert hoja["C4"].value == 12.5                                     # un número con coma, como número
    assert hoja["D3"].value.date() == date(2026, 9, 10) and hoja["D3"].number_format == "dd/mm/yyyy"
    cortada = load_workbook(io.BytesIO(armar_xlsx("Lista", columnas, filas, True))).active
    assert "CORTADA en 5.000 renglones" in cortada["A1"].value and FILAS_MAXIMAS_DEL_EXCEL == 5000
    assert nombre_del_archivo("Compras de Tomate: sept.") == "Compras_de_Tomate_sept.xlsx"


def test_la_IA_arma_un_EXCEL_y_pide_FOTOS_con_las_herramientas_y_queda_todo_anotado():
    cliente = ClienteDeMentira(
        _pide(("armar_excel", {"sql": "select x", "titulo": "Lista", "columnas": [{"nombre": "X", "tipo": "numero"}]}),
              ("mostrar_fotos", {"compra_id": 7, "que": "pesaje"}),
              ("mostrar_fotos", {"compra_id": 8, "que": "comanda"})),
        _texto("Te mandé la lista y la pesada."))
    renglones = [[n] for n in range(12)]

    def fotos(compra_id, que):
        if compra_id == 8:
            raise ValueError("no existe la compra 8")
        return {"cantidad": 2, "descripcion": "10/09/2026 · EJEMPLO"}

    resultado = preguntar(cliente, "?", [], sistema="S", correr_sql=None, leer_doc=None,
                          nombre_del_modelo="claude-sonnet-5-5",
                          armar_excel=lambda sql: {"columnas": ["x"], "filas": renglones, "cortada": False, "ms": 1},
                          mostrar_fotos=fotos)
    excel, pesaje, mala = resultado.consultas
    assert (excel["tipo"], excel["titulo"], len(excel["filas"])) == ("excel", "Lista", 12)   # todos, para el archivo
    assert (pesaje["tipo"], pesaje["compra_id"], pesaje["que"], pesaje["cantidad"]) == ("fotos", 7, "pesaje", 2)
    assert mala["error"] == "no existe la compra 8"
    vuelta = cliente.pedidos[1]["messages"][-1]["content"]
    assert '"renglones": 12' in vuelta[0]["content"] and '"primeros": [[0], [1], [2], [3], [4]]' in vuelta[0]["content"]
    assert [r["is_error"] for r in vuelta] == [False, False, True]
    # Sin la función, la herramienta no existe (la IA lo lee como error).
    sin = preguntar(ClienteDeMentira(_pide(("mostrar_fotos", {"compra_id": 1, "que": "comanda"})), _texto("ok")),
                    "?", [], sistema="S", correr_sql=None, leer_doc=None, nombre_del_modelo="claude-sonnet-5-5")
    assert sin.consultas == [] and sin.texto == "ok"


def test_por_TELEGRAM_contesta_corto_y_las_dos_saben_de_EXCEL_y_FOTOS():
    web = instrucciones("EJEMPLO", "t: a int", "idx", date(2026, 10, 17), "web")
    tele = instrucciones("EJEMPLO", "t: a int", "idx", date(2026, 10, 17), "telegram")
    assert "Telegram" in tele and "MUY corta" in tele and "Telegram" not in web
    for texto in (web, tele):
        assert "armar_excel" in texto and "mostrar_fotos" in texto and "guias_compra.id = 953" in texto
        assert "MÁS DE UNA compra posible" in texto and "la misma foto" in texto


def test_el_SECRETO_del_webhook():
    from core import telegram
    with patch.dict(os.environ, {"TELEGRAM_WEBHOOK_SECRET": SECRETO}):
        assert telegram.secreto_valido(SECRETO) and not telegram.secreto_valido("otro")
        assert not telegram.secreto_valido(None)
    with patch.dict(os.environ, {"TELEGRAM_WEBHOOK_SECRET": ""}):
        assert not telegram.secreto_valido("")                       # sin secreto configurado no pasa nada


def test_COMO_LO_CALCULE_por_telegram_escapa_el_html_y_si_se_corta_cierra_las_etiquetas():
    import app.main as m
    corto = m._como_lo_calcule_por_telegram({"consultas": [
        {"para_que": "contar <todo>", "sql": "select a < b", "columnas": ["n"], "filas": [[3]], "cantidad_filas": 1},
        {"tipo": "fotos", "para_que": "Fotos: pesaje de la compra 1", "descripcion": "EJ", "cantidad": 2}]})
    assert "<b>contar &lt;todo&gt;</b>" in corto and "<pre>select a &lt; b</pre>" in corto and "n: 3" in corto
    assert "EJ (2 fotos)" in corto
    largo = m._como_lo_calcule_por_telegram({"consultas": [
        {"para_que": "x", "sql": "select " + "a, " * 3000, "columnas": ["n"], "filas": [], "cantidad_filas": 0}]})
    assert len(largo) <= 4096 and largo.count("<pre>") == largo.count("</pre>") and "registro" in largo


def test_a_nadie_se_lo_AUTORIZA_desde_Telegram():
    """Lo único que escribe telegram_autorizados es el formulario de Gerencia."""
    arbol = ast.parse(open(os.path.join(RAIZ, "app/main.py"), encoding="utf-8").read())
    quien = {f.name for f in ast.walk(arbol) if isinstance(f, (ast.FunctionDef, ast.AsyncFunctionDef))
             for n in ast.walk(f) if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)
             and n.func.id == "guardar_autorizado_de_telegram"}
    assert quien == {"autorizar_en_telegram"}


# --- CONTRA LA BASE --------------------------------------------------------------------

SIEMBRA = """
insert into articulos (id, nombre, grupo) overriding system value values (1, 'EJEMPLO Tomate', 'hortaliza');
insert into proveedores (id, nombre, codigo_puesto) overriding system value values (1, 'EJEMPLO Puesto', 'N01P01');
insert into guias_compra (id, fecha_operacion, proveedor_id) overriding system value values (953, '2026-09-10', 1);
insert into compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones, contenido_por_cajon,
                     cantidad_kilos, importe, estado, guia_id, guia_punto) overriding system value
  values (1, 1, 1, '2026-09-10', 10, 16, 160, 100000, 'pendiente', 953, 1),
         (2, 1, 1, '2026-09-10', 5, 16, 80, 50000, 'pendiente', 953, 2);
insert into fotos_guia (guia_id, foto_ruta) values (953, 'comandas/EJ-guia-953.jpg');
insert into fotos_recepcion (compra_id, foto_ruta) values (1, 'balanza/EJ-1-a.jpg'), (1, 'balanza/EJ-1-b.jpg'),
                                                          (2, 'balanza/EJ-2-RIVAL.jpg');
"""


@pytest.fixture(scope="module")
def m():
    url = _base_con(SIEMBRA)
    with patch.dict(os.environ, {"DATABASE_URL": url, "CONSULTAS_IA_DATABASE_URL": url,
                                 "ANTHROPIC_API_KEY": "sk-de-mentira", "TELEGRAM_BOT_TOKEN": "1:EJ",
                                 "TELEGRAM_WEBHOOK_SECRET": SECRETO}):
        import app.main as modulo
        import app.db as d
        d.guardar_autorizado_de_telegram(LIONEL, "Lionel")
        yield modulo


class TelegramDeMentira:
    """Anota lo que la app le pide a la API de Telegram y contesta ok."""

    def __init__(self):
        self.pedidos = []

    def __call__(self, metodo, *, datos=None, archivos=None):
        self.pedidos.append((metodo, datos or {}, archivos or {}))
        if metodo == "getWebhookInfo":
            return {"url": "https://testserver/telegram/webhook"}
        if metodo == "getMe":
            return {"username": "ejemplo_consultas_bot"}
        return {}

    def de(self, metodo):
        return [d | {"_archivos": a} for met, d, a in self.pedidos if met == metodo]


def _mensaje(update_id, de, texto=None, chat="private", **mas):
    mensaje = {"message_id": update_id, "chat": {"id": de, "type": chat},
               "from": {"id": de, "first_name": "EJEMPLO", "last_name": "Persona", "username": "ejemplo"}, **mas}
    if texto is not None:
        mensaje["text"] = texto
    return {"update_id": update_id, "message": mensaje}


def _al_webhook(m, update, ia=None, secreto=SECRETO):
    from fastapi.testclient import TestClient
    tele = TelegramDeMentira()
    with patch("core.telegram._llamar", tele), patch("anthropic.Anthropic", return_value=ia or ClienteDeMentira()), \
            patch("app.main.obtener_url_foto", side_effect=lambda ruta: f"https://firmada/{ruta}"):
        respuesta = TestClient(m.app).post("/telegram/webhook", json=update,
                                           headers={"X-Telegram-Bot-Api-Secret-Token": secreto})
    return respuesta, tele


def _preguntas(pregunta):
    import app.db as d
    return [p for p in d.ultimas_preguntas_de_consultas() if p["pregunta"] == pregunta]


def test_el_WEBHOOK_sin_el_secreto_no_se_lee_y_el_que_no_esta_autorizado_queda_anotado(m):
    import app.db as d
    falso, tele = _al_webhook(m, _mensaje(1, EXTRANO, "¿cuánto vendimos?"), secreto="otro")
    assert falso.status_code == 403 and tele.pedidos == [] and d.intentos_de_telegram() == []
    intruso, tele = _al_webhook(m, _mensaje(2, EXTRANO, "¿cuánto vendimos?"))
    assert intruso.status_code == 200
    assert [x["text"] for x in tele.de("sendMessage")] == ["No autorizado."]
    (intento,) = d.intentos_de_telegram()
    assert (intento["telegram_id"], intento["nombre"], intento["usuario"], intento["texto"]) == (
        EXTRANO, "EJEMPLO Persona", "ejemplo", "¿cuánto vendimos?")
    assert _preguntas("¿cuánto vendimos?") == []                         # no se le preguntó a la IA
    # En un grupo, ni siquiera eso.
    _r, tele = _al_webhook(m, _mensaje(3, LIONEL, "hola", chat="group"))
    assert tele.pedidos == []


def test_un_AUTORIZADO_pregunta_y_recibe_la_respuesta_el_EXCEL_y_las_FOTOS_de_ESA_compra(m):
    ia = ClienteDeMentira(
        _pide(("correr_sql", {"sql": "select id from compras where guia_id = 953 and guia_punto = 1",
                              "para_que": "la compra de la guía 953.1"})),
        _pide(("mostrar_fotos", {"compra_id": 1, "que": "pesaje"}),
              ("armar_excel", {"sql": "select id, cantidad_kilos, importe from compras order by id",
                               "titulo": "Compras de la guía 953",
                               "columnas": [{"nombre": "Compra", "tipo": "numero"}, {"nombre": "Kilos", "tipo": "kilos"},
                                            {"nombre": "Importe", "tipo": "pesos"}]})),
        _texto("Te mando la pesada de la 953.1 y la lista."))
    respuesta, tele = _al_webhook(m, _mensaje(10, LIONEL, "mostrame la pesada de la guía 953.1"), ia)
    assert respuesta.status_code == 200
    (anotada,) = _preguntas("mostrame la pesada de la guía 953.1")
    assert (anotada["quien"], anotada["canal"], anotada["respuesta"]) == (
        "Lionel", "telegram", "Te mando la pesada de la 953.1 y la lista.")
    assert "Telegram" in ia.pedidos[0]["system"][0]["text"]
    (texto,) = tele.de("sendMessage")
    assert texto["text"] == "Te mando la pesada de la 953.1 y la lista." and texto["chat_id"] == LIONEL
    assert texto["reply_markup"]["inline_keyboard"][0][0] == {"text": "Cómo lo calculé",
                                                              "callback_data": f"calc:{anotada['id']}"}
    # Las DOS fotos de balanza de la compra 1 (la de la compra 2 es el rival y no va).
    (fotos,) = tele.de("sendMediaGroup")
    assert [f["media"] for f in fotos["media"]] == ["https://firmada/balanza/EJ-1-a.jpg",
                                                     "https://firmada/balanza/EJ-1-b.jpg"]
    assert "guía 953.1" in fotos["media"][0]["caption"]
    (documento,) = tele.de("sendDocument")
    nombre, contenido = documento["_archivos"]["document"]
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(contenido)).active
    assert nombre == "Compras_de_la_guia_953.xlsx" and [c.value for c in hoja[2]] == ["Compra", "Kilos", "Importe"]
    assert [[c.value for c in f] for f in hoja.iter_rows(min_row=3)] == [[1, 160, 100000], [2, 80, 50000]]


def test_el_CONTEXTO_dura_hasta_NUEVA_y_COMO_LO_CALCULE_es_solo_de_quien_pregunto(m):
    import app.db as d
    ia = ClienteDeMentira(_texto("Son 2 compras."))
    _al_webhook(m, _mensaje(20, LIONEL, "¿cuántas compras hay?"), ia)
    segunda = ClienteDeMentira(_texto("De Tomate."))
    _al_webhook(m, _mensaje(21, LIONEL, "¿de qué?"), segunda)
    mensajes = segunda.pedidos[0]["messages"]
    assert [x["content"] for x in mensajes[-3:]] == ["¿cuántas compras hay?", "Son 2 compras.", "¿de qué?"]
    _r, tele = _al_webhook(m, _mensaje(22, LIONEL, "/nueva"))
    assert [x["text"] for x in tele.de("sendMessage")] == ["Listo, empezamos de cero."]
    tercera = ClienteDeMentira(_texto("Hola."))
    _al_webhook(m, _mensaje(23, LIONEL, "¿y ahora?"), tercera)
    assert [x["content"] for x in tercera.pedidos[0]["messages"]] == ["¿y ahora?"]
    # Voz o foto: solo texto, y no se anota.
    _r, tele = _al_webhook(m, _mensaje(24, LIONEL, voice={"file_id": "x"}))
    assert [x["text"] for x in tele.de("sendMessage")] == ["Por ahora leo solo texto."]
    # El mismo update dos veces (Telegram reintenta): se contesta una.
    _r, tele = _al_webhook(m, _mensaje(24, LIONEL, voice={"file_id": "x"}))
    assert tele.pedidos == []
    # El botón: el que preguntó ve el SQL; otro (aunque esté autorizado) no.
    (anotada,) = _preguntas("mostrame la pesada de la guía 953.1")
    d.guardar_autorizado_de_telegram(OTRO, "EJEMPLO Otro")
    boton = {"update_id": 30, "callback_query": {"id": "b1", "from": {"id": OTRO}, "data": f"calc:{anotada['id']}",
                                                 "message": {"chat": {"id": OTRO}}}}
    _r, tele = _al_webhook(m, boton)
    assert tele.de("sendMessage") == [] and tele.de("answerCallbackQuery")[0]["text"] == "No autorizado."
    boton["update_id"], boton["callback_query"]["from"]["id"] = 31, LIONEL
    boton["callback_query"]["message"]["chat"]["id"] = LIONEL
    _r, tele = _al_webhook(m, boton)
    (detalle,) = tele.de("sendMessage")
    assert detalle["parse_mode"] == "HTML" and "guia_punto = 1</pre>" in detalle["text"]
    assert "Fotos: pesaje de la compra 1" in detalle["text"]


def test_en_la_WEB_el_EXCEL_se_baja_y_las_FOTOS_se_ven_solo_si_estan_en_la_respuesta(m):
    cliente_http, claves = _cliente_gerencia(m)
    ia = ClienteDeMentira(
        _pide(("armar_excel", {"sql": "select id, importe from compras order by id", "titulo": "Compras",
                               "columnas": [{"nombre": "Compra", "tipo": "numero"},
                                            {"nombre": "Importe", "tipo": "pesos"}]}),
              ("mostrar_fotos", {"compra_id": 1, "que": "comanda"})),
        _texto("Ahí va."))
    with patch.dict(os.environ, claves), patch("anthropic.Anthropic", return_value=ia), \
            patch("app.main.obtener_url_foto", side_effect=lambda ruta: f"https://firmada/{ruta}"):
        llegada = cliente_http.post("/gerencia/consultas/preguntar", json={"pregunta": "pasame las compras en Excel"})
        datos = cliente_http.get(f"/gerencia/consultas/pregunta/{llegada.json()['id']}").json()
        pid = datos["id"]
        excel = cliente_http.get(f"/gerencia/consultas/pregunta/{pid}/excel/0")
        visor = cliente_http.get(f"/gerencia/consultas/pregunta/{pid}/fotos/1")
        foto = cliente_http.get(f"/gerencia/consultas/pregunta/{pid}/fotos/1/0", follow_redirects=False)
        fuera = [cliente_http.get(u, follow_redirects=False).status_code for u in (
            f"/gerencia/consultas/pregunta/{pid}/fotos/1/1",       # la comanda tiene UNA foto
            f"/gerencia/consultas/pregunta/{pid}/fotos/0",         # el paso 0 es un Excel
            f"/gerencia/consultas/pregunta/{pid}/excel/1",         # y el 1, fotos
            f"/gerencia/consultas/pregunta/{pid}/excel/9")]
    assert [c["tipo"] for c in datos["consultas"]] == ["excel", "fotos"]
    assert excel.status_code == 200 and excel.headers["content-disposition"] == 'attachment; filename="Compras.xlsx"'
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(excel.content)).active
    assert [[c.value for c in f] for f in hoja.iter_rows(min_row=3)] == [[1, 100000], [2, 50000]]
    assert visor.status_code == 200 and visor.text.count(f'src="/gerencia/consultas/pregunta/{pid}/fotos/1/0"') == 1
    assert "Comanda · 10/09/2026 · EJEMPLO Puesto · EJEMPLO Tomate · guía 953.1" in visor.text
    assert (foto.status_code, foto.headers["location"]) == (307, "https://firmada/comandas/EJ-guia-953.jpg")
    assert fuera == [404, 404, 404, 404]


def test_GERENCIA_autoriza_desactiva_y_conecta_el_bot(m):
    import app.db as d
    cliente_http, claves = _cliente_gerencia(m)
    tele = TelegramDeMentira()
    with patch.dict(os.environ, claves), patch("core.telegram._llamar", tele):
        pagina = cliente_http.get("/gerencia/consultas/telegram").text
        mal = cliente_http.post("/gerencia/consultas/telegram/autorizar", data={"telegram_id": "abc", "nombre": "X"},
                                follow_redirects=False)
        cliente_http.post("/gerencia/consultas/telegram/autorizar", data={"telegram_id": str(EXTRANO),
                                                                          "nombre": "  EJEMPLO   Nuevo "})
        assert d.autorizado_de_telegram(EXTRANO) == {"telegram_id": EXTRANO, "nombre": "EJEMPLO Nuevo"}
        assert d.intentos_de_telegram() == []                           # autorizado: sale de Intentos
        cliente_http.post(f"/gerencia/consultas/telegram/{EXTRANO}/activo", data={"activo": "no"})
        assert d.autorizado_de_telegram(EXTRANO) is None
        cliente_http.post("/gerencia/consultas/telegram/conectar")
        sin_clave = __import__("fastapi.testclient", fromlist=["TestClient"]).TestClient(m.app).post(
            "/gerencia/consultas/telegram/autorizar", data={"telegram_id": "999", "nombre": "Nadie"},
            follow_redirects=False)
    assert "@ejemplo_consultas_bot" in pagina and "data-conectado" in pagina
    assert "error=" in mal.headers["location"]
    (conectar,) = tele.de("setWebhook")
    assert (conectar["url"], conectar["secret_token"], conectar["allowed_updates"]) == (
        "https://testserver/telegram/webhook", SECRETO, ["message", "callback_query"])
    assert sin_clave.status_code in (303, 401, 403) and d.autorizado_de_telegram(999) is None


def test_si_un_ADJUNTO_no_sale_los_otros_si_y_se_avisa(m):
    """Telegram rechaza las fotos (por ejemplo, una muy pesada): el Excel
    igual llega y la persona se entera de que faltan las fotos."""
    ia = ClienteDeMentira(
        _pide(("mostrar_fotos", {"compra_id": 1, "que": "pesaje"}),
              ("armar_excel", {"sql": "select id from compras order by id", "titulo": "Compras",
                               "columnas": [{"nombre": "Compra", "tipo": "numero"}]})),
        _texto("Ahí van."))
    tele = TelegramDeMentira()

    def telegram_que_rechaza_fotos(metodo, *, datos=None, archivos=None):
        if metodo == "sendMediaGroup":
            from core.telegram import TelegramNoAnda
            raise TelegramNoAnda("Telegram contestó 400: wrong file")
        return tele(metodo, datos=datos, archivos=archivos)

    from fastapi.testclient import TestClient
    with patch("core.telegram._llamar", telegram_que_rechaza_fotos), patch("anthropic.Anthropic", return_value=ia), \
            patch("app.main.obtener_url_foto", side_effect=lambda ruta: f"https://firmada/{ruta}"):
        TestClient(m.app).post("/telegram/webhook", json=_mensaje(40, LIONEL, "la pesada y la lista"),
                               headers={"X-Telegram-Bot-Api-Secret-Token": SECRETO})
    assert [x["text"] for x in tele.de("sendMessage")] == [
        "Ahí van.", "No pude mandar las fotos (Telegram contestó 400: wrong file)."]
    assert [x["_archivos"]["document"][0] for x in tele.de("sendDocument")] == ["Compras.xlsx"]

