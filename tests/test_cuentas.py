"""CUENTAS (dueño, 15/10): Caja, Banco e Inversiones con UN SOLO MOTOR, en las
dos apps, contra Postgres.

EL GALPÓN de Frutamax se arma corriendo los bloques 1 a 3 de la migración
(db/cuentas_1..3), como en producción: Caja Lionel, Banco Credicoop (CC$
191-220-000247/7) e Inversiones Credicoop, con sus 31 conceptos. El de Palmala
es la base sin cuentas, y la cuenta se crea por la pantalla. Nombres de
personas y empresas inventados; fechas fijas de 2025, que no pueden ser hoy
(las de los cheques, relativas a hoy de la base: la vista las lee así).
"""
import io
import os
import re
import sys
from datetime import date, timedelta
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.cuentas import (  # noqa: E402
    CONCEPTOS_POR_TIPO,
    cuenta_por_numero,
    problemas_del_movimiento,
    saldo_de_la_cuenta,
    totales_por,
)
from tests.test_administracion_reordenada import _cliente  # noqa: E402
from tests.test_cheques import _cargar, _hoy, d  # noqa: E402,F401
from tests.test_cheques_pantallas import _foto, _sin_estilo, m  # noqa: E402,F401

NUMERO = "CC$ 191-220-000247/7"


def _sembrar_frutamax(d):
    """Los bloques de la migración que siembran (1 a 3), como los corrió Lionel."""
    conexion = d.obtener_conexion()
    with conexion.cursor() as cursor:
        for archivo in ("cuentas_1_cuentas.sql", "cuentas_2_conceptos_caja.sql",
                        "cuentas_3_conceptos_banco_e_inversiones.sql"):
            cursor.execute(io.open(os.path.join(RAIZ, "db", archivo), encoding="utf-8").read())
    conexion.commit()
    conexion.close()
    return {c["nombre"]: c["id"] for c in d.listar_cuentas()}


def _concepto(d, cuenta_id, nombre) -> int:
    return next(c["id"] for c in d.conceptos_de_la_cuenta(cuenta_id) if c["nombre"] == nombre)


def _mov(d, cuenta_id, fecha, tipo, importe, concepto, etiqueta=None, detalle=None):
    return d.guardar_movimientos_de_cuenta(cuenta_id, [{
        "fecha": date.fromisoformat(fecha), "tipo": tipo, "importe": importe,
        "concepto_id": _concepto(d, cuenta_id, concepto), "etiqueta": etiqueta, "detalle": detalle}],
        como_se_cargo="a_mano")[0]


def _gerencia(m):
    return _cliente(m, "gerencia")


def _filas(html) -> list[tuple]:
    """(origen-id, fecha, importe con signo, concepto, saldo) de cada renglón del listado."""
    marcado = _sin_estilo(html).split("data-listado>")[1]
    return [(o, f, i, c.strip(), s) for o, f, i, c, s in re.findall(
        r'data-movimiento="([a-z]+-\d+)">.*?<span class="fila-que dia">([^<]*)</span>\s*'
        r'<span class="importe[^"]*">([^<]*)</span>.*?data-concepto>([^<]*)</div>.*?data-saldo-al>saldo ([^<]*)</span>',
        marcado, re.S)]


def _saldo(html) -> str:
    return re.search(r"data-saldo-hoy>([^<]*)<", _sin_estilo(html)).group(1)


# --- LO PURO ------------------------------------------------------------------

def test_el_NUMERO_se_reconoce_por_los_digitos_y_solo_en_un_banco_activo():
    cuentas = [{"id": 1, "tipo": "caja", "numero": None, "activa": True},
               {"id": 2, "tipo": "banco", "numero": NUMERO, "activa": True},
               {"id": 3, "tipo": "banco", "numero": "999 888 777", "activa": False},
               # El RIVAL: una inversión con el mismo número no recibe depósitos.
               {"id": 4, "tipo": "inversion", "numero": "1912200002477", "activa": True}]
    assert cuenta_por_numero("191 220 000247 7", cuentas)["id"] == 2
    assert cuenta_por_numero(" CC$ 191-220-000247/7 ", cuentas)["id"] == 2
    assert cuenta_por_numero("999888777", cuentas) is None          # dada de baja
    assert cuenta_por_numero("191-220", cuentas) is None             # un pedazo no es el número
    assert cuenta_por_numero("", cuentas) is None


def test_el_SALDO_de_hoy_deja_aparte_lo_que_se_acredita_despues_y_no_suma_anulados():
    hoy = date(2025, 5, 10)
    movimientos = [{"fecha": date(2025, 5, 1), "tipo": "ingreso", "importe": 1000, "anulado": False},
                   {"fecha": date(2025, 5, 10), "tipo": "egreso", "importe": 300, "anulado": False},
                   {"fecha": date(2025, 5, 11), "tipo": "ingreso", "importe": 500, "anulado": False},
                   {"fecha": date(2025, 5, 2), "tipo": "ingreso", "importe": 9999, "anulado": True}]
    assert saldo_de_la_cuenta(movimientos, hoy) == {"saldo": 700.0, "a_futuro": 500.0}


def test_los_TOTALES_por_concepto_y_por_etiqueta_sin_anulados():
    movimientos = [{"tipo": "ingreso", "importe": 100, "concepto": "Cobro", "etiqueta": "EJ Uno", "anulado": False},
                   {"tipo": "egreso", "importe": 30, "concepto": "Cobro", "etiqueta": None, "anulado": False},
                   {"tipo": "egreso", "importe": 500, "concepto": "Retiro", "etiqueta": "EJ Uno", "anulado": False},
                   {"tipo": "ingreso", "importe": 7, "concepto": "Cobro", "etiqueta": "EJ Uno", "anulado": True}]
    assert [(t["nombre"], t["ingresos"], t["egresos"], t["neto"], t["cuantos"])
            for t in totales_por(movimientos, "concepto")] == [("Retiro", 0, 500, -500, 1), ("Cobro", 100, 30, 70, 2)]
    assert [(t["nombre"], t["neto"]) for t in totales_por(movimientos, "etiqueta")] == [
        ("EJ Uno", -400), ("Sin etiqueta", -30)]


def test_los_PROBLEMAS_de_un_movimiento_antes_de_guardar():
    conceptos = [{"id": 1, "nombre": "Cobro Mecafront", "tipo": "ingreso", "activo": True},
                 {"id": 2, "nombre": "Cruce banco", "tipo": "egreso", "activo": False}]
    bien = {"fecha": "2025-05-01", "tipo": "ingreso", "importe": "100", "concepto_id": 1}
    assert problemas_del_movimiento(bien, conceptos) == []
    assert problemas_del_movimiento(dict(bien, tipo="egreso"), conceptos) == ["Cobro Mecafront es un ingreso"]
    assert problemas_del_movimiento(dict(bien, concepto_id=2, tipo="egreso"), conceptos) == [
        "el concepto Cruce banco ya no se usa"]
    assert problemas_del_movimiento({"fecha": "", "tipo": "", "importe": "0", "concepto_id": None}, conceptos) == [
        "falta la fecha", "decí si es ingreso o egreso", "falta el importe", "elegí el concepto"]


# --- FRUTAMAX: las tres cuentas sembradas, y los botones -----------------------

def test_los_BOTONES_de_Gerencia_salen_de_las_cuentas_ACTIVAS(m, d):
    cuentas = _sembrar_frutamax(d)
    hub = _sin_estilo(_gerencia(m).get("/gerencia").text)
    assert re.findall(r'href="/gerencia/cuentas/(\d+)"><svg[^>]*>.*?</svg><span>([^<]*)</span>', hub, re.S) == [
        (str(cuentas[n]), n) for n in ("Banco Credicoop", "Inversiones Credicoop", "Caja Lionel")]
    assert hub.count('href="/gerencia/cuentas"') == 1
    baja = _gerencia(m).post(f"/gerencia/cuentas/{cuentas['Caja Lionel']}/activa", data={"activa": ""},
                             follow_redirects=False)
    assert baja.status_code == 303
    hub = _sin_estilo(_gerencia(m).get("/gerencia").text)
    assert f'href="/gerencia/cuentas/{cuentas["Caja Lionel"]}"' not in hub
    assert f'href="/gerencia/cuentas/{cuentas["Banco Credicoop"]}"' in hub


def test_sin_la_CLAVE_de_Gerencia_no_se_ve_nada(m, d):
    cuentas = _sembrar_frutamax(d)
    pagina = _cliente(m).get(f"/gerencia/cuentas/{cuentas['Caja Lionel']}")
    assert 'name="clave"' in pagina.text and "Caja Lionel" not in _sin_estilo(pagina.text)


def test_CARGAR_a_mano_el_saldo_acumulado_y_la_ETIQUETA_que_se_arma_sola(m, d):
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    cliente = _gerencia(m)
    concepto = _concepto(d, caja, "Cobro Mecafront")
    uno = cliente.post(f"/gerencia/cuentas/{caja}/cargar", data={
        "fecha": "2025-05-01", "tipo": "ingreso", "importe": "1000,50", "concepto_id": str(concepto),
        "etiqueta": "  EJEMPLO   Mecafront ", "detalle": "factura 1"}, follow_redirects=False)
    assert uno.status_code == 303
    # La misma etiqueta escrita distinto es la MISMA: no se duplica.
    cliente.post(f"/gerencia/cuentas/{caja}/cargar", data={
        "fecha": "2025-05-03", "tipo": "egreso", "importe": "200", "concepto_id": str(_concepto(d, caja, "BP")),
        "etiqueta": "ejemplo mecafront"})
    assert [e["nombre"] for e in d.etiquetas_de_cuentas()] == ["EJEMPLO Mecafront"]
    # Un cobro cargado como egreso rebota antes de tocar la base.
    mal = cliente.post(f"/gerencia/cuentas/{caja}/cargar", data={
        "fecha": "2025-05-04", "tipo": "egreso", "importe": "5", "concepto_id": str(concepto)})
    assert mal.status_code == 400 and "Cobro Mecafront es un ingreso" in mal.text
    pagina = cliente.get(f"/gerencia/cuentas/{caja}")
    assert _saldo(pagina.text) == "$800,50"
    assert [(f, i, c, s) for _, f, i, c, s in _filas(pagina.text)] == [
        ("03/05/2025", "−$200", "BP · EJEMPLO Mecafront", "$800,50"),
        ("01/05/2025", "+$1.000,50", "Cobro Mecafront · EJEMPLO Mecafront", "$1.000,50")]


def test_el_concepto_de_OTRA_cuenta_y_el_que_es_solo_del_HISTORICO_los_frena_la_base(d):
    cuentas = _sembrar_frutamax(d)
    banco = cuentas["Banco Credicoop"]
    with pytest.raises(d.CuentaNoSePuede, match="es de otra cuenta"):
        d.guardar_movimientos_de_cuenta(banco, [{"fecha": date(2025, 5, 1), "tipo": "ingreso", "importe": 1,
                                                "concepto_id": _concepto(d, cuentas["Caja Lionel"], "Cobro Mecafront")}],
                                        como_se_cargo="a_mano")
    cruce = _concepto(d, cuentas["Caja Lionel"], "Cruce banco")
    with pytest.raises(d.CuentaNoSePuede, match="ya no se usa"):
        d.guardar_movimientos_de_cuenta(cuentas["Caja Lionel"], [
            {"fecha": date(2025, 5, 1), "tipo": "egreso", "importe": 1, "concepto_id": cruce}], como_se_cargo="a_mano")
    # El histórico sí lo usa (los 668 de la Caja traen "Cruce banco").
    conexion = d.obtener_conexion()
    with conexion.cursor() as cursor:
        cursor.execute("INSERT INTO cuentas_movimientos (cuenta_id, fecha, tipo, importe, concepto_id, como_se_cargo, "
                       "fila_historico, sector) VALUES (%s, '2021-03-01', 'egreso', 1, %s, 'historico', 7, 'gerencia')",
                       (cuentas["Caja Lionel"], cruce))
    conexion.commit()
    conexion.close()


# --- TRANSFERENCIAS: una fila, dos patas -----------------------------------------

def test_una_SUSCRIPCION_se_carga_una_vez_y_se_ve_en_las_dos_cuentas(m, d):
    cuentas = _sembrar_frutamax(d)
    banco, inversion = cuentas["Banco Credicoop"], cuentas["Inversiones Credicoop"]
    _mov(d, banco, "2025-05-01", "ingreso", 10000, "Saldo inicial")
    cliente = _gerencia(m)
    nueva = cliente.get(f"/gerencia/cuentas/{banco}/transferir").text
    assert "Inversiones Credicoop (si sale: Suscripción; si entra: Rescate)" in nueva
    assert "Banco Credicoop (si sale" not in nueva                     # no transfiere consigo misma
    respuesta = cliente.post(f"/gerencia/cuentas/{banco}/transferir", data={
        "sentido": "sale", "otra_cuenta_id": str(inversion), "fecha": "2025-05-02", "importe": "3000"},
        follow_redirects=False)
    assert respuesta.status_code == 303
    cliente.post(f"/gerencia/cuentas/{inversion}/transferir", data={
        "sentido": "sale", "otra_cuenta_id": str(banco), "fecha": "2025-05-05", "importe": "500"})
    _mov(d, inversion, "2025-05-06", "ingreso", 45, "Rendimiento")
    filas_banco = [(f, i, c, s) for _, f, i, c, s in _filas(cliente.get(f"/gerencia/cuentas/{banco}").text)]
    assert filas_banco == [("05/05/2025", "+$500", "Rescate · Inversiones Credicoop", "$7.500"),
                           ("02/05/2025", "−$3.000", "Suscripción · Inversiones Credicoop", "$7.000"),
                           ("01/05/2025", "+$10.000", "Saldo inicial", "$10.000")]
    filas_inversion = [(f, i, c, s) for _, f, i, c, s in _filas(cliente.get(f"/gerencia/cuentas/{inversion}").text)]
    assert filas_inversion == [("06/05/2025", "+$45", "Rendimiento", "$2.545"),
                               ("05/05/2025", "−$500", "Rescate · Banco Credicoop", "$2.500"),
                               ("02/05/2025", "+$3.000", "Suscripción · Banco Credicoop", "$3.000")]


def test_CORREGIR_y_ANULAR_una_transferencia_cambia_las_DOS_patas_con_historial(m, d):
    cuentas = _sembrar_frutamax(d)
    caja, banco = cuentas["Caja Lionel"], cuentas["Banco Credicoop"]
    tid = d.guardar_transferencia(fecha=date(2025, 5, 2), origen_id=caja, destino_id=banco, importe=1000,
                                  detalle=None, foto_ruta=None)
    cliente = _gerencia(m)
    sin_motivo = cliente.post(f"/gerencia/cuentas/transferencias/{tid}/corregir", data={
        "fecha": "2025-05-02", "importe": "1200", "cuenta_origen_id": str(caja), "cuenta_destino_id": str(banco)},
        follow_redirects=False)
    assert "error=" in sin_motivo.headers["location"]
    cliente.post(f"/gerencia/cuentas/transferencias/{tid}/corregir", data={
        "fecha": "2025-05-02", "importe": "1200", "cuenta_origen_id": str(caja), "cuenta_destino_id": str(banco),
        "motivo": "eran 1200"})
    assert _saldo(cliente.get(f"/gerencia/cuentas/{caja}").text) == "−$1.200"
    assert _saldo(cliente.get(f"/gerencia/cuentas/{banco}").text) == "$1.200"
    cliente.post(f"/gerencia/cuentas/transferencias/{tid}/anular", data={"anulado": "1", "motivo": "no fue"})
    assert _saldo(cliente.get(f"/gerencia/cuentas/{caja}").text) == "$0"
    assert _saldo(cliente.get(f"/gerencia/cuentas/{banco}").text) == "$0"
    ficha = _sin_estilo(cliente.get(f"/gerencia/cuentas/transferencias/{tid}").text)
    assert "<title>Depósito de efectivo</title>" in cliente.get(f"/gerencia/cuentas/transferencias/{tid}").text
    assert re.findall(r"data-correccion>[^·]*· ([^<]*)</li>", ficha) == [
        "Importe: 1000.00 → 1200.00 · eran 1200", "Anulado: no → sí · no fue"]


# --- LA CUENTA DE PALMALA: se crea por la pantalla ---------------------------------

def test_PALMALA_crea_su_cuenta_con_los_CONCEPTOS_de_su_tipo(m, d):
    assert d.listar_cuentas() == []                                   # la base de Palmala, sin cuentas
    assert not re.search(r'href="/gerencia/cuentas/\d+"', _gerencia(m).get("/gerencia").text)
    cliente = _gerencia(m)
    sin_tipo = cliente.post("/gerencia/cuentas", data={"nombre": "EJEMPLO Banco Palmala", "moneda": "pesos"})
    assert sin_tipo.status_code == 400 and "si es caja, banco o inversión" in sin_tipo.text
    hecha = cliente.post("/gerencia/cuentas", data={"nombre": " EJEMPLO  Banco Palmala ", "tipo": "banco",
                                                    "moneda": "pesos", "numero": "CA$ 111-222-333/4"},
                         follow_redirects=False)
    assert hecha.status_code == 303
    banco = d.listar_cuentas()[0]
    assert (banco["nombre"], banco["tipo"], banco["numero"]) == ("EJEMPLO Banco Palmala", "banco", "CA$ 111-222-333/4")
    assert [(c["nombre"], c["tipo"]) for c in d.conceptos_de_la_cuenta(banco["id"])] == list(CONCEPTOS_POR_TIPO["banco"])
    cliente.post("/gerencia/cuentas", data={"nombre": "EJEMPLO Caja Palmala", "tipo": "caja", "moneda": "pesos"})
    caja = d.listar_cuentas()[1]
    assert [c["nombre"] for c in d.conceptos_de_la_cuenta(caja["id"])] == [n for n, _ in CONCEPTOS_POR_TIPO["caja"]]
    # Los de solo histórico de la Caja de Frutamax NO vienen.
    assert "Cruce banco" not in [c["nombre"] for c in d.conceptos_de_la_cuenta(caja["id"])]
    repetida = cliente.post("/gerencia/cuentas", data={"nombre": "ejemplo banco palmala", "tipo": "banco",
                                                       "moneda": "pesos"})
    assert repetida.status_code == 400 and "Ya hay una cuenta con ese nombre." in repetida.text
    mismo_numero = cliente.post("/gerencia/cuentas", data={"nombre": "EJEMPLO Otro", "tipo": "banco",
                                                           "moneda": "pesos", "numero": "111 222 333 4"})
    assert mismo_numero.status_code == 400 and "Ya hay una cuenta con ese número." in mismo_numero.text
    hub = _sin_estilo(cliente.get("/gerencia").text)
    assert re.findall(r'href="/gerencia/cuentas/\d+"><svg[^>]*>.*?</svg><span>([^<]*)</span>', hub, re.S) == [
        "EJEMPLO Banco Palmala", "EJEMPLO Caja Palmala"]


# --- CHEQUES DEPOSITADOS ------------------------------------------------------------

def test_un_DEPOSITO_desde_la_Cartera_va_al_banco_con_la_fecha_de_ACREDITACION_y_su_RECHAZO(m, d):
    cuentas = _sembrar_frutamax(d)
    banco = cuentas["Banco Credicoop"]
    hoy = _hoy(d)
    cheque = _cargar(d, "777", importe=5000.0, pago=hoy - timedelta(days=5))
    rival = _cargar(d, "778", importe=9000.0, pago=hoy - timedelta(days=5))      # sale como pago: no va al banco
    from tests.test_cheques_aplicar import _formulario, _revisar, _seguir
    cliente = _gerencia(m)
    salio = (hoy - timedelta(days=3)).isoformat()
    # Tildado a mano y pasado a depósito: con un solo banco, ya viene elegido.
    datos = _formulario(_seguir(m, elegidos=[cheque]).text)
    revisada = _revisar(m, dict(datos, concepto_id_0="3", fecha_0=salio), "revisar")
    assert re.search(rf'<select name="cuenta_id_0"><option value="">Elegí</option>\s*<option value="{banco}" selected>',
                     _sin_estilo(revisada.text))
    hecho = _revisar(m, _formulario(revisada.text), "confirmar")
    assert hecho.status_code == 303
    # El RIVAL sale como pago: aunque traiga la cuenta, no va al banco.
    datos = _formulario(_seguir(m, elegidos=[rival]).text)
    assert _revisar(m, dict(datos, concepto_id_0="1", proveedor_id_0="1", cuenta_id_0=str(banco), fecha_0=salio),
                    "confirmar").status_code == 303
    acreditado = hoy - timedelta(days=2)
    filas = _filas(cliente.get(f"/gerencia/cuentas/{banco}").text)
    assert filas == [(f"deposito-{cheque}", acreditado.strftime("%d/%m/%Y"), "+$5.000",
                      "Depósito de cheque · EJEMPLO Verduleria", "$5.000")]
    assert d.listar_cuentas() and not _filas(cliente.get(f"/gerencia/cuentas/{cuentas['Caja Lionel']}").text)
    ficha = _sin_estilo(cliente.get(f"/gerencia/cheques/{cheque}").text)
    assert re.search(rf'data-cuenta><a href="/gerencia/cuentas/{banco}">Banco Credicoop</a>', ficha)
    listado = _sin_estilo(cliente.get("/gerencia/cheques").text)
    assert listado.count("data-cuenta>Banco Credicoop<") == 1
    d.rechazar_cheque(cheque, fecha=hoy - timedelta(days=1), motivo="Sin fondos")
    filas = _filas(cliente.get(f"/gerencia/cuentas/{banco}").text)
    assert [(o, i, c, s) for o, _, i, c, s in filas] == [
        (f"rechazo-{cheque}", "−$5.000", "Cheque rechazado · EJEMPLO Verduleria", "$0"),
        (f"deposito-{cheque}", "+$5.000", "Depósito de cheque · EJEMPLO Verduleria", "$5.000")]


def test_un_cheque_que_se_ACREDITA_MANANA_va_aparte_del_saldo_de_hoy(m, d):
    banco = _sembrar_frutamax(d)["Banco Credicoop"]
    hoy = _hoy(d)
    cheque = _cargar(d, "900", importe=700.0, pago=hoy + timedelta(days=3))
    d.aplicar_cheques([cheque], fecha=hoy, concepto_id=3, concepto_nuevo=None, destino_nuevo=None,
                      proveedor_id=None, foto_ruta=None, cuenta_id=banco)
    pagina = _gerencia(m).get(f"/gerencia/cuentas/{banco}").text
    assert _saldo(pagina) == "$0"
    assert re.search(r"data-a-futuro>Por acreditarse: <strong>\$700</strong>", _sin_estilo(pagina))


def test_un_cheque_no_se_deposita_en_una_CAJA(d):
    cuentas = _sembrar_frutamax(d)
    hoy = _hoy(d)
    cheque = _cargar(d, "901", importe=1.0, pago=hoy)
    with pytest.raises(d.ChequeNoSePuede, match="cuenta de banco"):
        d.aplicar_cheques([cheque], fecha=hoy, concepto_id=3, concepto_nuevo=None, destino_nuevo=None,
                          proveedor_id=None, foto_ruta=None, cuenta_id=cuentas["Caja Lionel"])
    assert d.buscar_cheques({"numero": "901"})[0]["estado"] == "en_cartera"       # no salió: todo o nada


def test_la_CAPTURA_de_un_deposito_reconoce_la_cuenta_por_el_NUMERO(m, d):
    from tests.test_cheques_aplicar import DEPOSITO_REAL, _formulario, _galpon, _los_reales, _revisar, _seguir
    cuentas = _sembrar_frutamax(d)
    # El RIVAL: otro banco activo, para que no gane por ser el único.
    d.crear_cuenta(nombre="EJEMPLO Otro Banco", tipo="banco", moneda="pesos", numero="CC$ 999-000-111/2",
                   madre_id=None, conceptos=[])
    _galpon(d)
    reales = _los_reales(d)
    pagina = _seguir(m, [DEPOSITO_REAL])
    datos = _formulario(pagina.text)
    assert [datos[f"cuenta_id_{i}"] for i in range(2)] == [str(cuentas["Banco Credicoop"])] * 2
    with patch("app.main.subir_foto_comanda", return_value="cheques/deposito.jpg"):
        final = _revisar(m, datos, "confirmar")
    assert final.status_code == 303
    assert {c["id"]: c["cuenta"] for c in d.buscar_cheques({}) if c["id"] in reales.values()} == {
        reales["00932970"]: "Banco Credicoop", reales["00020758"]: "Banco Credicoop"}


def test_un_numero_que_NO_es_de_ninguna_cuenta_pide_elegirla(m, d):
    from tests.test_cheques_aplicar import DEPOSITO_REAL, _galpon, _los_reales, _seguir
    _sembrar_frutamax(d)
    d.crear_cuenta(nombre="EJEMPLO Otro Banco", tipo="banco", moneda="pesos", numero=None, madre_id=None,
                   conceptos=[])
    _galpon(d)
    _los_reales(d)
    pagina = _seguir(m, [dict(DEPOSITO_REAL, cuenta_deposito="CC$ 555-555-555/5")])
    assert re.findall(r"data-problemas>([^<]*)</p>", _sin_estilo(pagina.text)) == ["elegí la cuenta"] * 2


# --- FILTROS, TOTALES Y EXPORTACIONES -----------------------------------------------

def test_los_FILTROS_los_TOTALES_y_el_PDF_y_EXCEL_con_los_mismos_filtros(m, d):
    from openpyxl import load_workbook
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    _mov(d, caja, "2025-04-01", "ingreso", 5000, "Aporte Lionel", "EJEMPLO Lionel")
    _mov(d, caja, "2025-05-01", "ingreso", 1000, "Cobro Mecafront", "EJEMPLO Mecafront")
    _mov(d, caja, "2025-05-02", "egreso", 300, "Retiro Lionel", "EJEMPLO Lionel")
    anulado = _mov(d, caja, "2025-05-03", "egreso", 77, "BP", "EJEMPLO Lionel")
    d.corregir_movimiento_de_cuenta(anulado, {"anulado": True}, "repetido")
    cliente = _gerencia(m)
    filtro = {"desde": "2025-05-01", "etiqueta": "EJEMPLO Lionel"}
    pagina = cliente.get(f"/gerencia/cuentas/{caja}", params=filtro).text
    assert _saldo(pagina) == "$5.700"                                   # el saldo es de TODO
    assert [(f, i) for _, f, i, _, _ in _filas(pagina)] == [("03/05/2025", "−$77"), ("02/05/2025", "−$300")]
    marcado = _sin_estilo(pagina)
    conceptos = marcado.split('data-totales="concepto"')[1].split('data-totales="etiqueta"')[0]
    assert re.findall(r'<span class="nombre">([^<]*)</span>', conceptos) == ["Retiro Lionel"]   # el anulado no suma
    excel = cliente.get(f"/gerencia/cuentas/{caja}/exportar-excel", params=filtro)
    hoja = load_workbook(io.BytesIO(excel.content))
    movimientos = hoja["Movimientos"]
    assert "Filtros: desde el 01/05/2025 · etiqueta EJEMPLO Lionel" in movimientos["A2"].value
    assert movimientos["A3"].value == "Saldo de hoy: $5.700"
    assert [[c.value for c in fila][:3] for fila in movimientos.iter_rows(min_row=6)] == [
        ["03/05/2025", "Egreso", "BP"], ["02/05/2025", "Egreso", "Retiro Lionel"]]
    assert [[c.value for c in fila][:4] for fila in hoja["Por concepto"].iter_rows(min_row=5)] == [
        ["Retiro Lionel", 0, 300, -300]]
    pdf = cliente.get(f"/gerencia/cuentas/{caja}/exportar-pdf", params=filtro)
    import pypdfium2 as pdfium
    texto = " ".join(" ".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(pdf.content)).split())
    assert "Filtros: desde el 01/05/2025 · etiqueta EJEMPLO Lionel" in texto and "Retiro Lionel" in texto
    assert "Cobro Mecafront" not in texto


def test_CORREGIR_un_movimiento_deja_el_historial_con_NOMBRES(m, d):
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    mid = _mov(d, caja, "2025-05-01", "ingreso", 1000, "Cobro Mecafront", "EJEMPLO Mecafront")
    cliente = _gerencia(m)
    cliente.post(f"/gerencia/cuentas/movimientos/{mid}/corregir", data={
        "fecha": "2025-05-01", "tipo": "ingreso", "importe": "1000", "concepto_id": str(_concepto(d, caja, "Cobro de terceros")),
        "etiqueta": "EJEMPLO Otro", "detalle": "", "motivo": "era de otro"})
    ficha = _sin_estilo(cliente.get(f"/gerencia/cuentas/movimientos/{mid}").text)
    assert re.findall(r"data-correccion>[^·]*· ([^<]*)</li>", ficha) == [
        "Concepto: Cobro Mecafront → Cobro de terceros · era de otro",
        "De quién: EJEMPLO Mecafront → EJEMPLO Otro · era de otro"]


def test_CAMBIAR_la_cuenta_de_un_deposito_lo_pasa_de_un_banco_al_otro(m, d):
    cuentas = _sembrar_frutamax(d)
    otro = d.crear_cuenta(nombre="EJEMPLO Otro Banco", tipo="banco", moneda="pesos", numero=None, madre_id=None,
                          conceptos=[])
    hoy = _hoy(d)
    cheque = _cargar(d, "950", importe=300.0, pago=hoy - timedelta(days=4))
    d.aplicar_cheques([cheque], fecha=hoy - timedelta(days=3), concepto_id=3, concepto_nuevo=None, destino_nuevo=None,
                      proveedor_id=None, foto_ruta=None, cuenta_id=cuentas["Banco Credicoop"])
    cliente = _gerencia(m)
    mal = cliente.post(f"/gerencia/cheques/{cheque}/cuenta", data={"cuenta_id": str(cuentas["Caja Lionel"])})
    assert mal.status_code == 400 and "Elegí la cuenta." in mal.text
    cliente.post(f"/gerencia/cheques/{cheque}/cuenta", data={"cuenta_id": str(otro)})
    assert _filas(cliente.get(f"/gerencia/cuentas/{cuentas['Banco Credicoop']}").text) == []
    assert [o for o, *_ in _filas(cliente.get(f"/gerencia/cuentas/{otro}").text)] == [f"deposito-{cheque}"]


def test_si_la_base_NO_CONTESTA_el_hub_abre_y_dice_que_faltan_los_botones(m, d):
    import psycopg2
    with patch("app.main.listar_cuentas", side_effect=psycopg2.OperationalError("caída")):
        hub = _gerencia(m).get("/gerencia")
    assert hub.status_code == 200
    assert "No se pudieron leer las cuentas: sus botones no están." in hub.text
    sana = _gerencia(m).get("/gerencia").text
    assert "data-cuentas-sin-leer" not in sana


# --- CARGA POR EXCEL Y POR FOTOS, con la revisión (dueño, 15/10) -----------------

def _xlsx(filas) -> bytes:
    from openpyxl import Workbook
    libro = Workbook()
    for fila in filas:
        libro.active.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _formulario_de_revision(html: str, cuenta_id: int) -> dict:
    from tests.test_palmala_cruces import _formulario
    return _formulario(html, f"/gerencia/cuentas/{cuenta_id}/revisar")


def _problemas(html) -> list[str]:
    return re.findall(r"data-problemas>([^<]*)</span>", _sin_estilo(html))


def test_la_PLANILLA_reconoce_ingreso_y_egreso_y_saltea_los_TOTALES():
    from core.cuentas import crudos_de_la_planilla, renglones_para_revisar
    filas = [["Movimientos de la caja", None, None, None],
             ["Fecha", "Ingreso", "Egreso", "Detalle"],
             ["01/05/2025", "1.000,50", None, "uno"],
             ["02/05/2025", None, 300, "dos"],
             [None, None, None, None],
             ["Total", 1000.5, 300, None]]                     # el RIVAL: no es un movimiento
    crudos = crudos_de_la_planilla(filas)
    renglones = renglones_para_revisar(crudos, [], hoy=date(2025, 6, 1))
    assert [(r["fecha"], r["tipo"], r["importe"], r["detalle"]) for r in renglones] == [
        ("2025-05-01", "ingreso", "1000.5", "uno"), ("2025-05-02", "egreso", "300", "dos")]
    negativo = renglones_para_revisar(crudos_de_la_planilla([["Fecha", "Importe"], ["03/05/2025", "-45"]]), [],
                                      hoy=date(2025, 6, 1))
    assert [(r["tipo"], r["importe"]) for r in negativo] == [("egreso", "45")]
    assert crudos_de_la_planilla([["Día", "Concepto"], ["01/05/2025", "x"]]) is None


def test_un_EXCEL_se_revisa_PARA_TODOS_y_se_guardan_los_que_estan_BIEN(m, d):
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    cliente = _gerencia(m)
    planilla = _xlsx([["Fecha", "Importe", "Tipo", "Concepto", "Etiqueta", "Detalle"],
                      ["01/05/2025", 1000, "Ingreso", "Cobro Mecafront", "EJEMPLO Mecafront", "fact 1"],
                      ["02/05/2025", 300, "Egreso", "BP", "", ""],
                      ["03/05/2025", 50, "Egreso", "Algo que no existe", "", ""],
                      ["Total", 1350, None, None, None, None]])
    pagina = cliente.post(f"/gerencia/cuentas/{caja}/cargar/excel",
                          files={"archivo": ("caja.xlsx", planilla, "application/octet-stream")})
    assert pagina.status_code == 200
    assert re.search(r"data-conteo>3 movimientos · <span data-bien>2 bien</span> · <span class=\"rojo\" "
                     r"data-con-problema>1 con problema", _sin_estilo(pagina.text))
    assert _problemas(pagina.text) == ["elegí el concepto («Algo que no existe» no está en la lista)"]
    datos = _formulario_de_revision(pagina.text, caja)
    # "Para todos" con el renglón con problema TILDADO: solo ése cambia.
    malo = next(k.split("_")[-1] for k, v in datos.items() if k.startswith("concepto_leido_") and v == "Algo que no existe")
    otro = _concepto(d, caja, "Retiro Lionel")
    pagina = cliente.post(f"/gerencia/cuentas/{caja}/revisar", data=dict(
        datos, accion="a_todos", concepto_todos=str(otro), etiqueta_todos="EJEMPLO Lionel", **{f"tildado_{malo}": "1"}))
    assert "Se cambió 1 movimiento." in pagina.text and _problemas(pagina.text) == []
    final = cliente.post(f"/gerencia/cuentas/{caja}/revisar", data=dict(_formulario_de_revision(pagina.text, caja),
                                                                          accion="guardar"), follow_redirects=False)
    assert final.status_code == 303 and "3+movimientos+cargados" in final.headers["location"]
    filas = [(f, i, c) for _, f, i, c, _ in _filas(cliente.get(f"/gerencia/cuentas/{caja}").text)]
    assert filas == [("03/05/2025", "−$50", "Retiro Lionel · EJEMPLO Lionel"), ("02/05/2025", "−$300", "BP"),
                     ("01/05/2025", "+$1.000", "Cobro Mecafront · EJEMPLO Mecafront")]
    assert {mov["como_se_cargo"] for mov in (d.obtener_movimiento_de_cuenta(x["origen_id"])
                                             for x in d.movimientos_de_la_cuenta(caja))} == {"excel"}


def test_GUARDAR_deja_en_la_pagina_los_que_tienen_PROBLEMA(m, d):
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    cliente = _gerencia(m)
    planilla = _xlsx([["Fecha", "Importe", "Tipo", "Concepto"],
                      ["01/05/2025", 1000, "Ingreso", "Cobro Mecafront"],
                      ["02/05/2025", 300, "Ingreso", "BP"]])          # BP es un egreso: no entra
    pagina = cliente.post(f"/gerencia/cuentas/{caja}/cargar/excel",
                          files={"archivo": ("caja.xlsx", planilla, "application/octet-stream")})
    final = cliente.post(f"/gerencia/cuentas/{caja}/revisar",
                         data=dict(_formulario_de_revision(pagina.text, caja), accion="guardar"))
    assert "Se guardaron 1. Quedan 1 con problemas" in final.text and _problemas(final.text) == ["BP es un egreso"]
    assert len(d.movimientos_de_la_cuenta(caja)) == 1


def test_las_FOTOS_un_renglon_por_movimiento_con_SU_foto_y_la_que_no_se_lee_avisa(m, d):
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    leidas = [{"movimientos": [{"fecha": "2025-05-01", "tipo": "ingreso", "importe": 2500, "quien": "EJEMPLO Mecafront",
                                "concepto": "Cobro Mecafront", "detalle": "recibo 7"},
                               {"fecha": "2025-05-01", "tipo": "egreso", "importe": 100, "quien": None,
                                "concepto": "BP", "detalle": None}]},
              RuntimeError("no se entiende"),
              {"movimientos": [{"fecha": "2025-05-02", "tipo": "egreso", "importe": 70, "quien": "EJEMPLO Lionel",
                                "concepto": "Retiro Lionel", "detalle": None}]}]
    archivos = [("fotos", (f"f{i}.jpg", _foto(), "image/jpeg")) for i in range(3)]
    with patch("core.lector_cuentas.extraer_movimientos_de_imagenes", side_effect=leidas) as lector:
        pagina = _gerencia(m).post(f"/gerencia/cuentas/{caja}/cargar/fotos", files=archivos)
    assert lector.call_count == 3
    # A la IA se le pasan los conceptos ACTIVOS de esta cuenta, no los de solo histórico.
    assert "Cruce banco" not in lector.call_args.kwargs["conceptos"]
    assert "Cobro Mecafront" in lector.call_args.kwargs["conceptos"]
    assert "La foto 2 no se pudo leer: «f1.jpg» no se entiende como un comprobante" in pagina.text
    datos = _formulario_de_revision(pagina.text, caja)
    assert sorted((datos[f"foto_{i}"], datos[f"importe_{i}"]) for i in range(3)) == [
        ("0", "100"), ("0", "2500"), ("2", "70")]
    with patch("app.main.subir_foto_comanda", side_effect=["cuentas/f0.jpg", "cuentas/f2.jpg"]) as subida:
        final = _gerencia(m).post(f"/gerencia/cuentas/{caja}/revisar", data=dict(datos, accion="guardar"),
                                  follow_redirects=False)
    assert final.status_code == 303 and subida.call_count == 2           # una subida por FOTO, no por renglón
    fotos = {float(x["importe"]): x["foto_ruta"] for x in d.movimientos_de_la_cuenta(caja)}
    assert fotos == {2500.0: "cuentas/f0.jpg", 100.0: "cuentas/f0.jpg", 70.0: "cuentas/f2.jpg"}


# --- CONCILIACIÓN (dueño, 15/10) -------------------------------------------------

def test_los_DIAS_HABILES_y_el_CRUCE_uno_a_uno_por_importe_con_signo():
    from core.conciliacion import cruzar, dias_habiles_entre
    assert dias_habiles_entre(date(2025, 5, 9), date(2025, 5, 12)) == 1      # viernes → lunes
    assert dias_habiles_entre(date(2025, 5, 12), date(2025, 5, 19)) == 5
    assert dias_habiles_entre(date(2025, 5, 19), date(2025, 5, 12)) == 5
    renglones = [{"id": 1, "fecha": date(2025, 5, 12), "importe": 500.0},
                 {"id": 2, "fecha": date(2025, 5, 14), "importe": 500.0},
                 {"id": 3, "fecha": date(2025, 5, 12), "importe": -100.0},
                 {"id": 4, "fecha": date(2025, 5, 12), "importe": 70.0}]
    sistema = [{"origen": "movimiento", "origen_id": 10, "fecha": date(2025, 5, 13), "tipo": "ingreso", "importe": 500},
               {"origen": "deposito", "origen_id": 7, "fecha": date(2025, 5, 14), "tipo": "ingreso", "importe": 500},
               # El RIVAL del -100: un INGRESO de 100 no es un débito de 100.
               {"origen": "movimiento", "origen_id": 11, "fecha": date(2025, 5, 12), "tipo": "ingreso", "importe": 100},
               # A 6 días hábiles: fuera del margen.
               {"origen": "movimiento", "origen_id": 12, "fecha": date(2025, 5, 20), "tipo": "ingreso", "importe": 70}]
    assert cruzar(renglones, sistema) == [
        {"renglon_id": 1, "origen": "movimiento", "origen_id": 10, "a_revisar": True},
        {"renglon_id": 2, "origen": "deposito", "origen_id": 7, "a_revisar": True}]
    # UNO A UNO: dos renglones de 500 y un solo movimiento de 500 hacen un par, no dos.
    assert [p["renglon_id"] for p in cruzar(renglones[:2], sistema[:1])] == [1]
    # Solo, sin otro candidato, no queda para revisar.
    assert cruzar(renglones[:1], sistema[:1]) == [{"renglon_id": 1, "origen": "movimiento", "origen_id": 10,
                                                   "a_revisar": False}]


def test_el_CONCEPTO_SUGERIDO_sale_de_la_descripcion_del_banco():
    from core.conciliacion import concepto_sugerido
    conceptos = [{"id": 1, "nombre": "Impuestos y comisiones", "activo": True},
                 {"id": 2, "nombre": "Transferencia recibida", "activo": True},
                 {"id": 3, "nombre": "Transferencia enviada", "activo": True},
                 {"id": 4, "nombre": "Otros ingresos", "activo": True}, {"id": 5, "nombre": "Otros egresos", "activo": True}]
    assert concepto_sugerido("COMISION MANT. CUENTA", -1500, conceptos) == 1
    assert concepto_sugerido("IMP.LEY 25413 DEB", -30, conceptos) == 1
    assert concepto_sugerido("TRANSF. RECIBIDA CBU 0191", 900, conceptos) == 2
    assert concepto_sugerido("TRANSFERENCIA A TERCEROS", -900, conceptos) == 3
    assert concepto_sugerido("DEP CHEQ 30-7XXXX", 5000, conceptos) == 4
    assert concepto_sugerido("ALGO RARO", -5, conceptos) == 5


EXTRACTO_MAYO = {"desde": "2025-05-01", "hasta": "2025-05-31", "saldo_final": 9198.25, "renglones": [
    {"fecha": "2025-05-02", "descripcion": "TRANSF RECIBIDA", "importe": 500, "saldo": None},       # antes del corte
    {"fecha": "2025-05-09", "descripcion": "PAGO PROV 0001", "importe": -300, "saldo": None},
    {"fecha": "2025-05-12", "descripcion": "DEP EFECTIVO", "importe": 1000, "saldo": None},
    {"fecha": "2025-05-12", "descripcion": "COMISION MANT. CUENTA", "importe": -1500.75, "saldo": None},
    {"fecha": None, "descripcion": "ilegible", "importe": 12, "saldo": None}]}
# El segundo extracto se PISA con el primero en el pago de 300: no entra dos veces.
EXTRACTO_PISADO = {"desde": "2025-05-09", "hasta": "2025-05-09", "saldo_final": None, "renglones": [
    {"fecha": "2025-05-09", "descripcion": "PAGO PROV 0001", "importe": -300, "saldo": None}]}


def _banco_para_conciliar(d):
    cuentas = _sembrar_frutamax(d)
    banco, caja = cuentas["Banco Credicoop"], cuentas["Caja Lionel"]
    _mov(d, banco, "2025-05-05", "ingreso", 10000, "Saldo inicial")
    pago = _mov(d, banco, "2025-05-07", "egreso", 300, "Pago a proveedor", "EJEMPLO Puesto")
    d.guardar_transferencia(fecha=date(2025, 5, 8), origen_id=caja, destino_id=banco, importe=1000, detalle=None,
                            foto_ruta=None)
    suelto = _mov(d, banco, "2025-05-12", "ingreso", 77, "Otros ingresos")
    return banco, pago, suelto


def _subir(m, banco, *leidos):
    archivos = [("extractos", (f"mayo{i}.pdf", b"%PDF-1.4", "application/pdf")) for i in range(len(leidos))]
    with patch("core.lector_cuentas.extraer_extracto_de_imagenes", side_effect=list(leidos)), \
            patch("app.main.imagenes_desde_pdf", return_value=[b"x"]):
        return _gerencia(m).post(f"/gerencia/cuentas/{banco}/conciliar/extractos", files=archivos)


def _secciones(html) -> dict:
    marcado = _sin_estilo(html)
    return {"solo_extracto": re.findall(r'data-renglon="\d+">\s*<div class="renglon-cabeza"><span class="dia">([^<]*)</span>'
                                        r'\s*<span class="importe[^"]*">([^<]*)</span></div>\s*<div class="descripcion">([^<]*)<',
                                        marcado),
            "solo_sistema": re.findall(r'data-sistema="([a-z]+-\d+)"', marcado),
            "conciliados": re.findall(r'data-conciliado="\d+".*?<div class="descripcion">([^<]*)<', marcado, re.S),
            "a_revisar": marcado.count("data-a-revisar>")}


def test_CONCILIAR_sube_dos_extractos_que_se_pisan_cruza_y_deja_lo_suelto_a_la_vista(m, d):
    banco, _, suelto = _banco_para_conciliar(d)
    pagina = _subir(m, banco, EXTRACTO_MAYO, EXTRACTO_PISADO)
    assert pagina.status_code == 200
    assert "«mayo0.pdf»: 4 renglones, 1 sin fecha o sin importe · «mayo1.pdf»: 0 renglones (1 ya estaban). " \
           "Se conciliaron 2." in pagina.text
    s = _secciones(pagina.text)
    assert s["conciliados"] == ["PAGO PROV 0001 ↔ Pago a proveedor · EJEMPLO Puesto (07/05)",
                                "DEP EFECTIVO ↔ Depósito de efectivo · Caja Lionel (08/05)"]
    assert s["solo_extracto"] == [("12/05/2025", "−$1.500,75", "COMISION MANT. CUENTA")]
    assert s["solo_sistema"] == [f"movimiento-{suelto}"]           # ni el saldo inicial ni lo conciliado
    assert "1 renglón del extracto anterior al saldo inicial: no se cargan." in _sin_estilo(pagina.text)
    # La comisión trae sugerido su concepto.
    comision = _sin_estilo(pagina.text).split("COMISION MANT. CUENTA")[1].split("</form>")[0]
    assert re.search(r'<option value="(\d+)" selected>Impuestos y comisiones</option>', comision)
    # Diferencia al cierre: banco 9.198,25 contra el sistema 10.000 − 300 + 1.000 + 77 = 10.777.
    assert "Banco $9.198,25 · sistema $10.777 · diferencia −$1.578,75" in _sin_estilo(pagina.text)


def test_CARGAR_de_un_toque_lo_del_extracto_y_SEPARAR_un_par(m, d):
    banco, pago, _ = _banco_para_conciliar(d)
    _subir(m, banco, EXTRACTO_MAYO)
    cliente = _gerencia(m)
    renglones = {r["descripcion"]: r["id"] for r in d.renglones_de_extractos(banco)}
    cliente.post(f"/gerencia/cuentas/{banco}/conciliar/{renglones['COMISION MANT. CUENTA']}/cargar",
                 data={"concepto_id": str(_concepto(d, banco, "Impuestos y comisiones")), "etiqueta": "EJEMPLO Credicoop"})
    s = _secciones(cliente.get(f"/gerencia/cuentas/{banco}/conciliar").text)
    assert s["solo_extracto"] == [] and len(s["conciliados"]) == 3
    cargado = next(x for x in d.movimientos_de_la_cuenta(banco) if x["concepto"] == "Impuestos y comisiones")
    assert (cargado["fecha"], cargado["tipo"], float(cargado["importe"]), cargado["etiqueta"], cargado["conciliado"]) == (
        date(2025, 5, 12), "egreso", 1500.75, "EJEMPLO Credicoop", True)
    cliente.post(f"/gerencia/cuentas/{banco}/conciliar/{renglones['PAGO PROV 0001']}/desconciliar")
    s = _secciones(cliente.get(f"/gerencia/cuentas/{banco}/conciliar").text)
    assert ("09/05/2025", "−$300", "PAGO PROV 0001") in s["solo_extracto"] and f"movimiento-{pago}" in s["solo_sistema"]
    cliente.post(f"/gerencia/cuentas/{banco}/conciliar/cruzar")
    assert len(_secciones(cliente.get(f"/gerencia/cuentas/{banco}/conciliar").text)["conciliados"]) == 3


def test_ANULAR_un_movimiento_conciliado_suelta_su_renglon(m, d):
    banco, pago, _ = _banco_para_conciliar(d)
    _subir(m, banco, EXTRACTO_MAYO)
    d.corregir_movimiento_de_cuenta(pago, {"anulado": True}, "no era")
    s = _secciones(_gerencia(m).get(f"/gerencia/cuentas/{banco}/conciliar").text)
    assert ("09/05/2025", "−$300", "PAGO PROV 0001") in s["solo_extracto"]
    assert f"movimiento-{pago}" not in s["solo_sistema"]                    # anulado: no se revisa


def test_dos_del_MISMO_importe_quedan_PARA_REVISAR_y_se_confirman(m, d):
    banco, _, _ = _banco_para_conciliar(d)
    _mov(d, banco, "2025-05-13", "egreso", 300, "Pago a proveedor", "EJEMPLO Otro Puesto")
    pagina = _subir(m, banco, EXTRACTO_MAYO)
    assert _secciones(pagina.text)["a_revisar"] == 1
    renglon = next(r["id"] for r in d.renglones_de_extractos(banco) if r["a_revisar"])
    _gerencia(m).post(f"/gerencia/cuentas/{banco}/conciliar/{renglon}/confirmar")
    assert _secciones(_gerencia(m).get(f"/gerencia/cuentas/{banco}/conciliar").text)["a_revisar"] == 0


def test_solo_un_BANCO_se_concilia(m, d):
    caja = _sembrar_frutamax(d)["Caja Lionel"]
    assert _gerencia(m).get(f"/gerencia/cuentas/{caja}/conciliar").status_code == 404
    assert "/conciliar" not in _gerencia(m).get(f"/gerencia/cuentas/{caja}").text


def test_lo_ANTERIOR_al_saldo_inicial_no_se_cruza_ni_se_carga(m, d):
    banco, _, _ = _banco_para_conciliar(d)
    # El RIVAL: un ingreso de 500 a 2 días hábiles del renglón del 02/05, que es anterior al saldo inicial.
    rival = _mov(d, banco, "2025-05-06", "ingreso", 500, "Otros ingresos")
    s = _secciones(_subir(m, banco, EXTRACTO_MAYO).text)
    assert f"movimiento-{rival}" in s["solo_sistema"]
    assert ("02/05/2025", "+$500", "TRANSF RECIBIDA") not in s["solo_extracto"]


def test_lo_YA_CONCILIADO_no_se_vuelve_a_cruzar_con_un_extracto_nuevo(m, d):
    banco, _, _ = _banco_para_conciliar(d)
    _subir(m, banco, EXTRACTO_MAYO)
    otro = {"desde": "2025-05-13", "hasta": "2025-05-13", "saldo_final": None, "renglones": [
        {"fecha": "2025-05-13", "descripcion": "DEP EFECTIVO 2", "importe": 1000, "saldo": None}]}
    pagina = _subir(m, banco, otro)
    assert pagina.status_code == 200 and "Se conciliaron 0." in pagina.text
    assert ("13/05/2025", "+$1.000", "DEP EFECTIVO 2") in _secciones(pagina.text)["solo_extracto"]


# --- EL HUB DE GERENCIA: "Bancos y cajas" (dueño, 16/10) ---------------------------

def _bancos_y_cajas(html) -> list[tuple]:
    from tests.test_administracion_reordenada import _grupos_con_color
    return [(c, t) for c, _h, t in _grupos_con_color(_sin_estilo(html))["Bancos y cajas"]]


def test_BANCOS_Y_CAJAS_en_FRUTAMAX_va_en_el_orden_del_dueno(m, d):
    _sembrar_frutamax(d)
    with patch.dict(m.templates.env.globals, {"ES_FRUTAMAX": True}):
        hub = _gerencia(m).get("/gerencia").text
    assert _bancos_y_cajas(hub) == [("5", "Banco Credicoop"), ("5", "Inversiones Credicoop"),
                                    ("5", "Cartera de cheques"), ("5", "Caja Lionel"),
                                    ("5", "Palmala Cruces"), ("5", "Cuentas")]
    # Ninguno queda repetido en otro recuadro.
    marcado = _sin_estilo(hub)
    for href in ("/gerencia/cheques", "/gerencia/palmala-cruces", "/gerencia/cuentas", "/gerencia/panel"):
        assert marcado.count(f'href="{href}"') == 1, href


def test_BANCOS_Y_CAJAS_en_PALMALA_con_sus_cuentas_y_sin_Palmala_Cruces(m, d):
    cliente = _gerencia(m)
    for nombre, tipo in (("EJEMPLO Caja Palmala", "caja"), ("EJEMPLO Banco Palmala", "banco")):
        cliente.post("/gerencia/cuentas", data={"nombre": nombre, "tipo": tipo, "moneda": "pesos"})
    banco = next(c["id"] for c in d.listar_cuentas() if c["tipo"] == "banco")
    cliente.post("/gerencia/cuentas", data={"nombre": "EJEMPLO Plazo fijo", "tipo": "inversion", "moneda": "pesos",
                                            "madre_id": str(banco)})
    with patch.dict(m.templates.env.globals, {"ES_FRUTAMAX": False}):
        hub = cliente.get("/gerencia").text
    # El banco con su inversión adentro, aunque la caja se haya creado antes.
    assert _bancos_y_cajas(hub) == [("5", "EJEMPLO Banco Palmala"), ("5", "EJEMPLO Plazo fijo"),
                                    ("5", "Cartera de cheques"), ("5", "EJEMPLO Caja Palmala"), ("5", "Cuentas")]
