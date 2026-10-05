"""Devoluciones de mercadería al proveedor (dueño, 30/09), contra Postgres.

Dos caminos, cada uno registrado por separado:

  · POR RECHAZO: lo que Día rechazó y se le devuelve al proveedor. Pide la
    compra, y cancela el costo al costo por bulto de ESA compra cuando lo que
    volvió iba en su cajón.
  · DESDE DEPÓSITO: lo que queda en el piso. Sale del lote de SU compra y de
    ningún otro, y nunca más de lo que queda.

En los dos, si la compra dejó seña, sus cajones vuelven llenos y salen de
Vacíos; y los dos aparecen en Movimientos del depósito.

Con la base mockeada el reparto lo decide el fixture, así que esto corre
contra el esquema real (corolario 89). Los nombres son de EJEMPLO.

EL GALPÓN QUE SE PLANTA, con el RIVAL puesto (corolario 11: el lote que no
tiene que ganar):

  compra 11  EJ Uno  05/09  10 cajones a $100, CON seña     <- la más vieja
  compra 12  EJ Uno  06/09   8 cajones a $120, sin seña
  compra 13  EJ Dos  07/09   5 cajones a $90

Una devolución de la 12 tiene que salir de la 12 aunque la 11 sea más vieja:
el FIFO puro la mandaría a la 11, que es exactamente lo que no pasó.
"""
import os
import re
import urllib.parse
import sys
from datetime import date

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from scripts.humo import hay_postgres, preparar_base  # noqa: E402

OBLIGATORIO = os.environ.get("HUMO_OBLIGATORIO") == "1"

SIEMBRA = """
insert into clientes (id, nombre) overriding system value values (1, 'EJEMPLO Dia');
insert into articulos (id, nombre) overriding system value values (1, 'EJEMPLO Fruta');
insert into proveedores (id, nombre, codigo_puesto) overriding system value
  values (1, 'EJ Uno', 'N91P01'), (2, 'EJ Dos', 'N91P02');
insert into envases (id, nombre) overriding system value values (1, 'EJEMPLO Caja Dia');
insert into fichas_logistica (id, cliente_id, articulo_id, contenido_caja, unidad_venta, envase_id)
  overriding system value values (1, 1, 1, 16, 'kilo', null), (2, 1, 1, 6, 'kilo', 1);
insert into compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
                     contenido_por_cajon, cantidad_kilos, importe, sena, estado, procesada_el,
                     cantidad_cajones_real, contenido_por_cajon_real)
  overriding system value values
  (11, 1, 1, '2026-09-05', 10, 16, 160, 100, 500, 'recepcionado', '2026-09-05 18:00-03', 10, 16),
  (12, 1, 1, '2026-09-06',  8, 16, 128, 120, null, 'recepcionado', '2026-09-06 18:00-03',  8, 16),
  (13, 2, 1, '2026-09-07',  5, 16,  80,  90, null, 'recepcionado', '2026-09-07 18:00-03',  5, 16);
insert into pedidos (id, cliente_id, fecha_operacion, origen) overriding system value
  values (1, 1, '2026-09-08', 'texto');
-- Un renglón SIN caja nuestra (ficha 1) y otro EN CAJA DE DÍA (ficha 2).
insert into pedidos_renglones (id, pedido_id, articulo_id, ficha_id, sucursal, cantidad,
                               cantidad_armada, armado_el, kilos_enviados)
  overriding system value values
  (1, 1, 1, 1, 'EJ Suc', 4, 4, '2026-09-08 10:00-03', 64),
  (2, 1, 1, 2, 'EJ Suc', 4, 4, '2026-09-08 10:00-03', 24);
"""


@pytest.fixture
def base(monkeypatch):
    if not hay_postgres():
        if OBLIGATORIO:
            pytest.fail("HUMO_OBLIGATORIO=1 y no hay Postgres: las devoluciones no se verificaron")
        pytest.skip("sin Postgres local")
    url = preparar_base()
    import psycopg2
    conexion = psycopg2.connect(url)
    with conexion.cursor() as cursor:
        cursor.execute(SIEMBRA)
    conexion.commit()
    conexion.close()
    monkeypatch.setenv("DATABASE_URL", url)
    import app.db as d

    def sql(consulta, parametros=None):
        conexion = d.obtener_conexion()
        try:
            with conexion.cursor() as cursor:
                cursor.execute(consulta, parametros)
                filas = cursor.fetchall() if cursor.description else None
            conexion.commit()
            return filas
        finally:
            conexion.close()

    def rechazo(renglon_id, cantidad, compra_id=None, proveedor_id=None, costo=77):
        sql("""INSERT INTO movimientos_stock (articulo_id, tipo, cantidad, motivo, cliente_id,
                   fecha_operacion, stock_sistema, pedido_renglon_id, costo_por_bulto,
                   destino_rechazo, compra_devolucion_id, proveedor_devolucion_id)
               VALUES (1, 'reingreso_rechazo', %s, 'EJ rechazo', 1, '2026-09-09', 0, %s, %s,
                       'devolucion_proveedor', %s, %s)""",
            (cantidad, renglon_id, costo, compra_id, proveedor_id))
    return d, sql, rechazo


def _lotes(d):
    """{compra_id: restante} del reparto de ahora, y lo que quedó sin lote."""
    from core.stock import repartir_fifo, salidas_para_reparto
    entradas, salidas = d.entradas_y_salidas_stock_articulo(1)
    reparto = repartir_fifo(entradas, salidas_para_reparto(salidas))
    restos = {l["origen_id"]: l["restante"] for l in reparto["lotes"] if l["tipo_lote"] == "guia"}
    return restos, reparto["sin_lote"]


def _pila(d, proveedor_id):
    return next((p["stock"] for p in d.stock_de_vacios_deposito() if p["id"] == proveedor_id), 0)


# --- desde depósito: sale de SU compra --------------------------------------

def test_la_devolucion_desde_deposito_sale_de_SU_compra_y_no_de_la_mas_vieja(base):
    d, sql, _ = base
    antes, sin_lote_antes = _lotes(d)
    stock_antes = d.stock_deposito_de_articulo(1)
    d.crear_devolucion_deposito(12, 3, "EJ se puso fea", date(2026, 9, 10), cargada_desde="deposito")
    despues, sin_lote = _lotes(d)
    # El armado en cajón consumió la 11 (FIFO); la devolución sale de la 12.
    # El de caja de Día ya estaba sin lote (espera su guía R): no se mueve.
    assert despues[12] == antes[12] - 3
    assert despues[11] == antes[11], "la compra más vieja no se toca"
    assert sin_lote == sin_lote_antes
    assert d.stock_deposito_de_articulo(1) == stock_antes - 3


def test_lo_que_su_compra_no_cubre_queda_SIN_LOTE_y_no_sale_de_otra(base):
    """La guarda de "no más de lo que queda" está en la ruta; si igual llega
    de más (un día cargado tarde), el sobrante queda a la vista sin lote en vez
    de llevarse mercadería de otra compra."""
    d, sql, _ = base
    antes, sin_lote_antes = _lotes(d)
    d.crear_devolucion_deposito(13, 7, "EJ de más", date(2026, 9, 10), cargada_desde="deposito")
    despues, sin_lote = _lotes(d)
    assert despues.get(13, 0) == 0
    assert despues[11] == antes[11] and despues[12] == antes[12]
    assert sin_lote == sin_lote_antes + 2


def test_la_devolucion_lleva_sus_FOTOS_a_las_de_la_compra(base):
    d, sql, _ = base
    d.crear_devolucion_deposito(12, 1, "EJ", date(2026, 9, 10), fotos_pesada=["pesaje/x/a.jpg"], cargada_desde="deposito")
    assert [f["foto_ruta"] for f in d.listar_fotos_de_recepcion(12)] == ["pesaje/x/a.jpg"]


def test_una_compra_que_NO_esta_recibida_no_se_puede_devolver(base):
    d, sql, _ = base
    sql("UPDATE compras SET estado = 'pendiente' WHERE id = 13")
    with pytest.raises(ValueError, match="no está recibida"):
        d.crear_devolucion_deposito(13, 1, "EJ", date(2026, 9, 10), cargada_desde="deposito")
    assert sql("SELECT count(*) FROM movimientos_stock")[0][0] == 0


# --- vacíos: los cajones que vuelven llenos -----------------------------------

def test_con_SENA_los_cajones_que_vuelven_llenos_salen_de_Vacios(base):
    d, sql, rechazo = base
    assert _pila(d, 1) == 10, "la compra 11 dejó seña por 10 cajones"
    d.crear_devolucion_deposito(11, 3, "EJ", date(2026, 9, 10), cargada_desde="deposito")
    assert _pila(d, 1) == 7
    # Sin seña (la 12) no hay cajón que devolver.
    d.crear_devolucion_deposito(12, 2, "EJ", date(2026, 9, 10), cargada_desde="deposito")
    assert _pila(d, 1) == 7


def test_el_RECHAZO_en_cajon_resta_y_el_que_iba_en_CAJA_DE_DIA_no(base):
    d, sql, rechazo = base
    rechazo(1, 2, compra_id=11)        # ficha sin caja nuestra: vuelve el cajón
    assert _pila(d, 1) == 8
    rechazo(2, 2, compra_id=11)        # iba en caja de Día: su cajón quedó en el galpón
    assert _pila(d, 1) == 8
    rechazo(1, 1, proveedor_id=1)      # los viejos, sin compra: no se sabe la marca
    assert _pila(d, 1) == 8


def test_la_LISTA_de_vacios_y_la_PILA_dicen_lo_mismo_con_las_devoluciones(base):
    d, sql, rechazo = base
    d.crear_devolucion_deposito(11, 3, "EJ", date(2026, 9, 10), cargada_desde="deposito")
    rechazo(1, 2, compra_id=11)
    movimientos = [m for m in d.movimientos_de_vacios(1, limite=1000) if not m["anulada"]]
    llenas = [m["cantidad"] for m in movimientos if m["tipo"] == "devolucion_llena"]
    assert sorted(llenas) == [-3, -2]
    assert sum(int(m["cantidad"]) for m in movimientos) == _pila(d, 1)


# --- el costo en la Rentabilidad Real ------------------------------------------

def test_el_VALOR_de_la_devolucion_es_el_PRECIO_DE_SU_COMPRA_en_cajon_y_en_caja_de_Dia(base):
    """Regla de Lionel (01/10): vale EXACTAMENTE el precio por cajón de la
    compra (`compras.importe`), y no el costo del armado (77, el congelado).
    En caja de Día con compra atada, también (opción B, 02/10). Sin compra
    queda el congelado. Una compra SIN PRECIO da None, no el congelado: caer
    ahí sería la regla vieja."""
    d, sql, rechazo = base
    sql("""INSERT INTO compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
               contenido_por_cajon, cantidad_kilos, importe, estado, procesada_el,
               cantidad_cajones_real, contenido_por_cajon_real) OVERRIDING SYSTEM VALUE
           VALUES (14, 1, 1, '2026-09-07', 3, 16, 48, NULL, 'recepcionado', '2026-09-07 18:00-03', 3, 16)""")
    rechazo(1, 2, compra_id=12)        # cajón: $120 de la compra 12
    rechazo(2, 2, compra_id=12)        # caja de Día: también $120 (opción B)
    rechazo(1, 1, proveedor_id=1)      # viejo, sin compra: el congelado
    rechazo(1, 1, compra_id=14)        # cajón, compra sin precio: None
    filas = d.devoluciones_vinculadas_por_rango(1, date(2026, 9, 1), date(2026, 9, 30))
    def clave(t):
        return tuple(-1 if v is None else float(v) for v in t)
    obtenido = sorted(((f["ficha_id"], f["valor_por_bulto"], f["vale_la_compra"]) for f in filas), key=clave)
    assert obtenido == sorted([(1, 120, True), (2, 120, True), (1, 77, False), (1, None, True)], key=clave)


def test_MOVIMIENTOS_y_la_PLANILLA_valen_el_precio_de_la_compra_y_SIN_PRECIO_no_cae_al_armado(base):
    d, sql, rechazo = base
    sql("""INSERT INTO compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
               contenido_por_cajon, cantidad_kilos, importe, estado, procesada_el,
               cantidad_cajones_real, contenido_por_cajon_real) OVERRIDING SYSTEM VALUE
           VALUES (14, 1, 1, '2026-09-07', 3, 16, 48, NULL, 'recepcionado', '2026-09-07 18:00-03', 3, 16)""")
    rechazo(1, 2, compra_id=12)
    rechazo(2, 1, compra_id=12)
    rechazo(1, 1, compra_id=14)
    filas = d.movimientos_del_deposito(date(2026, 9, 1), date(2026, 9, 30), tipo="rechazo")
    assert sorted(((m["compra_id"], m["bultos"], m["valor"]) for m in filas),
                  key=lambda t: (t[0], t[1])) == [
        (12, -2.0, -240.0),      # cajón: 2 × $120, no 2 × $77
        (12, -1.0, -120.0),      # caja de Día: también el precio (opción B)
        (14, -1.0, None),        # sin precio: "sin precio", no $77
    ]


def test_CHERRY_en_caja_de_Dia_vale_2_cajones_de_su_compra_y_no_toca_VACIOS(base):
    """Caso real de Frutamax (movimiento 199): 2 bultos de Tomate Cherry en
    caja de Día, compra 826 a $30.000, costo del armado $30.523,26. Con la
    opción B (Lionel, 02/10) vale 2 × $30.000 = $60.000 en Movimientos y en
    la planilla, y la Rentabilidad nombra la diferencia contra el armado:
    2 × (30.000 − 30.523,26) = −1.046,52.

    Y la caja de Día NO es un cajón del proveedor: aunque la compra dejó seña,
    ese rechazo no sale de Vacíos. El valor cambió; los cajones, no."""
    from core.costo_real import calcular_rentabilidad_real
    from tests.test_costo_real import MARGEN, _armado, _datos, _devolucion
    d, sql, rechazo = base
    sql("""INSERT INTO compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
               contenido_por_cajon, cantidad_kilos, importe, sena, estado, procesada_el,
               cantidad_cajones_real, contenido_por_cajon_real) OVERRIDING SYSTEM VALUE
           VALUES (826, 1, 1, '2026-09-07', 4, 16, 64, 30000, 500, 'recepcionado',
                   '2026-09-07 18:00-03', 4, 16)""")
    pila_antes = _pila(d, 1)
    rechazo(2, 2, compra_id=826, costo=30523.26)
    assert _pila(d, 1) == pila_antes, "la caja de Día no es un cajón del proveedor"
    assert not [m for m in d.movimientos_de_vacios(1, limite=1000)
                if m["tipo"] == "devolucion_llena" and not m["anulada"]]
    movs = [m for m in d.movimientos_del_deposito(date(2026, 9, 1), date(2026, 9, 30), tipo="rechazo")
            if m["compra_id"] == 826]
    assert [(m["bultos"], m["valor"], m["sena"]) for m in movs] == [(-2.0, -60000.0, None)]
    from app.main import _devoluciones_para_pagar, _fila_de_devolucion
    para_pagar = [_fila_de_devolucion(m) for m in _devoluciones_para_pagar(
        date(2026, 9, 1), date(2026, 9, 30), None, None) if m["compra_id"] == 826]
    assert [(r["total"], r["importe"], r["total_sena"], r["total_a_depositar"]) for r in para_pagar] == [
        (-60000.0, 30000.0, None, -60000.0)]
    fila, = [f for f in d.devoluciones_vinculadas_por_rango(1, date(2026, 9, 1), date(2026, 9, 30))
             if float(f["valor_por_bulto"] or 0) == 30000]
    assert fila["vale_la_compra"] is True and float(fila["costo_por_bulto"]) == 30523.26
    fecha = date(2026, 8, 25)
    devolucion = dict(_devolucion(float(fila["bultos"]), fecha, costo_por_bulto=float(fila["costo_por_bulto"]),
                                  destino=fila["destino_rechazo"]),
                      valor_por_bulto=float(fila["valor_por_bulto"]), vale_la_compra=fila["vale_la_compra"])
    renta = calcular_rentabilidad_real(
        _datos([_armado(fecha, 25, 500.0)]), {fecha: {901: dict(MARGEN)}}, 1, fecha, fecha,
        devoluciones=[devolucion],
    )["grupos"][0]["filas"][0]
    assert round(renta["diferencia_devolucion_proveedor"], 2) == -1046.52


def test_la_Rentabilidad_lee_el_valor_de_la_MISMA_regla_que_Movimientos():
    """Escrita una vez: la consulta de la Rentabilidad y la de Movimientos
    nombran el MISMO fragmento, y la vieja (que caía al congelado con un
    COALESCE) no está en ningún lado."""
    import inspect
    import app.db as d
    assert d._SQL_VALOR_POR_BULTO_DE_LA_DEVOLUCION in d._SQL_MOVIMIENTOS_DEL_DEPOSITO
    fuente = inspect.getsource(d.devoluciones_vinculadas_por_rango)
    assert "_SQL_VALOR_POR_BULTO_DE_LA_DEVOLUCION + \"\"\" AS valor_por_bulto" in fuente
    assert "_SQL_DEVOLUCION_VALE_LA_COMPRA + \"\"\" AS vale_la_compra" in fuente
    assert not hasattr(d, "_SQL_COSTO_DE_LA_COMPRA_DEVUELTA")
    # Opción B (02/10): la regla del VALOR no pregunta por la caja de Día.
    assert "envase_id" not in d._SQL_DEVOLUCION_VALE_LA_COMPRA


# --- Movimientos del depósito --------------------------------------------------

def test_MOVIMIENTOS_del_deposito_trae_entradas_devoluciones_y_segunda(base):
    d, sql, rechazo = base
    d.crear_devolucion_deposito(12, 3, "EJ se puso fea", date(2026, 9, 10), cargada_desde="deposito")
    rechazo(1, 2, compra_id=11)
    sql("INSERT INTO remitos_segunda (articulo_id, bultos, fecha_operacion) VALUES (1, 4, '2026-09-11')")
    sql("INSERT INTO remitos_segunda (articulo_id, bultos, fecha_operacion, anulado_el) "
        "VALUES (1, 9, '2026-09-11', now())")
    filas = d.movimientos_del_deposito(date(2026, 9, 1), date(2026, 9, 30))
    resumen = [(m["tipo"], m["proveedor"], m["compra_id"], m["bultos"], m["valor"]) for m in filas]
    assert resumen == [
        ("entrada", "EJ Dos", 13, 5.0, 450.0),
        ("entrada", "EJ Uno", 11, 10.0, 1000.0),
        ("entrada", "EJ Uno", 12, 8.0, 960.0),
        ("rechazo", "EJ Uno", 11, -2.0, -200.0),
        ("deposito", "EJ Uno", 12, -3.0, -360.0),
        ("segunda", None, None, -4.0, None),
    ]
    # Los filtros.
    assert {m["tipo"] for m in d.movimientos_del_deposito(
        date(2026, 9, 1), date(2026, 9, 30), tipo="deposito")} == {"deposito"}
    por_proveedor = d.movimientos_del_deposito(date(2026, 9, 1), date(2026, 9, 30), proveedor_id=2)
    assert [m["compra_id"] for m in por_proveedor] == [13]
    assert d.movimientos_del_deposito(date(2026, 9, 8), date(2026, 9, 9), tipo="entrada") == []


def test_la_CUENTA_de_un_proveedor_cierra_como_la_dice_el_duenio(base):
    """"El 01 entraron 50; el 02 salieron 10 por rechazo y 40 desde depósito."""
    from core.movimientos_deposito import agrupar
    d, sql, rechazo = base
    rechazo(1, 2, compra_id=11)
    d.crear_devolucion_deposito(12, 3, "EJ", date(2026, 9, 10), cargada_desde="deposito")
    grupos = agrupar(d.movimientos_del_deposito(date(2026, 9, 1), date(2026, 9, 30), proveedor_id=1))
    uno, = grupos
    assert (uno["entraron"], uno["salieron"], uno["neto"]) == (18, 5, 13)
    assert (uno["valor_entrado"], uno["valor_salido"]) == (1960, 560)


# --- las RUTAS, contra la misma base -------------------------------------------
# El tope de "lo que queda" se calcula en la ruta con el reparto de ahora. Con
# la base mockeada ese número lo pondría el fixture, así que va contra Postgres.

HOY = date(2026, 9, 10)


def _cliente():
    from fastapi.testclient import TestClient
    from app.main import app
    return TestClient(app, base_url="https://testserver")


def test_la_pantalla_de_DEVOLVER_ofrece_solo_las_compras_con_lo_que_QUEDA(base):
    from unittest.mock import patch
    d, sql, _ = base
    with patch("app.main._hoy_argentina", return_value=HOY):
        respuesta = _cliente().get("/deposito/devolver?proveedor_id=1")
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    # La 11 la consumió el armado en cajón (4 de 10: quedan 6); la 12 entera.
    assert re.search(r'value="11"[^>]*>.*?quedan <span class="resto">6</span> de 10', marcado, re.S)
    assert re.search(r'value="12"[^>]*>.*?quedan <span class="resto">8</span> de 8', marcado, re.S)
    # La 13 es de OTRO proveedor.
    assert 'value="13"' not in marcado


def test_DEVOLVER_mas_de_lo_que_queda_rebota_y_NO_escribe(base):
    from unittest.mock import patch
    d, sql, _ = base
    with patch("app.main._hoy_argentina", return_value=HOY):
        respuesta = _cliente().post("/deposito/devolver", data={
            "proveedor_id": "1", "compra_id": "11", "cantidad": "7", "motivo": "EJ fea"},
            follow_redirects=False)
    assert respuesta.status_code == 400
    assert "quedan 6 bultos" in respuesta.text
    assert sql("SELECT count(*) FROM movimientos_stock WHERE tipo = 'devolucion_deposito'") == [(0,)]


@pytest.mark.parametrize("datos, error", [
    ({"proveedor_id": "1", "compra_id": "13", "cantidad": "1", "motivo": "EJ"}, "Elegí de qué compra"),
    ({"proveedor_id": "1", "compra_id": "12", "cantidad": "1", "motivo": "  "}, "El motivo es obligatorio"),
    ({"proveedor_id": "1", "compra_id": "12", "cantidad": "0", "motivo": "EJ"}, ""),
])
def test_DEVOLVER_una_compra_AJENA_sin_motivo_o_en_cero_rebota(base, datos, error):
    from unittest.mock import patch
    d, sql, _ = base
    with patch("app.main._hoy_argentina", return_value=HOY):
        respuesta = _cliente().post("/deposito/devolver", data=datos, follow_redirects=False)
    assert respuesta.status_code == 400
    assert error in respuesta.text
    assert sql("SELECT count(*) FROM movimientos_stock WHERE tipo = 'devolucion_deposito'") == [(0,)]


def test_DEVOLVER_lo_que_queda_guarda_con_su_FOTO_y_avisa_la_sena(base):
    import io
    from unittest.mock import patch
    from PIL import Image
    d, sql, _ = base
    jpeg = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 100, 50)).save(jpeg, format="JPEG")
    with patch("app.main._hoy_argentina", return_value=HOY), \
         patch("app.main.subir_foto_comanda", return_value="pesaje/2026-09-10/ej.jpg") as subir:
        respuesta = _cliente().post(
            "/deposito/devolver",
            data={"proveedor_id": "1", "compra_id": "11", "cantidad": "6", "motivo": "EJ fea"},
            files=[("fotos", ("a.jpg", jpeg.getvalue(), "image/jpeg"))],
            follow_redirects=False)
    assert respuesta.status_code == 303, respuesta.text[-400:]
    assert subir.call_count == 1
    assert "esos cajones salen de Vacíos" in urllib.parse.unquote_plus(respuesta.headers["location"])
    assert sql("""SELECT compra_devolucion_id, cantidad, motivo, fecha_operacion FROM movimientos_stock
                  WHERE tipo = 'devolucion_deposito'""") == [(11, -6, "EJ fea", HOY)]
    assert sql("SELECT foto_ruta FROM fotos_recepcion WHERE compra_id = 11") == [("pesaje/2026-09-10/ej.jpg",)]


def test_MOVIMIENTOS_en_la_pantalla_y_el_EXCEL_de_Administracion(base, monkeypatch):
    import io
    import openpyxl
    from unittest.mock import patch
    from app.main import PUERTA_ADMINISTRACION
    d, sql, rechazo = base
    d.crear_devolucion_deposito(12, 3, "EJ se puso fea", date(2026, 9, 10), cargada_desde="deposito")
    monkeypatch.setenv("CLAVE_ADMINISTRACION", "admin-secreta")
    cliente = _cliente()
    cliente.cookies.set(PUERTA_ADMINISTRACION.cookie, PUERTA_ADMINISTRACION.firma("admin-secreta"))
    with patch("app.main._hoy_argentina", return_value=date(2026, 9, 12)):
        pantalla = cliente.get("/administracion/ingresos?proveedor_id=1")
        excel = cliente.get("/administracion/ingresos/movimientos-excel?proveedor_id=1&tipo=deposito")
    assert pantalla.status_code == 200
    marcado = pantalla.text.split("</style>")[-1]
    assert "Entraron 18 · se devolvieron 3 · neto 15 bultos" in marcado
    assert marcado.count('<div class="fila">') == 3
    assert "Devolución desde depósito · EJEMPLO Fruta" in marcado
    assert 'href="/administracion/ingresos/pagar?' in marcado
    assert excel.status_code == 200
    hoja = openpyxl.load_workbook(io.BytesIO(excel.content)).active
    filas = [tuple(c.value for c in fila) for fila in hoja.iter_rows(min_row=5)]
    assert filas == [("10/09/2026", "Devolución desde depósito", "EJ Uno", 12, "EJEMPLO Fruta",
                      -3, -360, "EJ se puso fea", None, None, "Depósito")]


def test_MOVIMIENTOS_dicen_que_ENTRO_con_sena_y_VOLVIO_con_sena(base, monkeypatch):
    """Dueño, 30/09: la devolución de mercadería con seña no genera un vale,
    pero Movimientos tiene que mostrar los cajones y la seña en la entrada y
    en la devolución. En un rechazo, solo si iba en el cajón."""
    import io
    import openpyxl
    from unittest.mock import patch
    from app.main import PUERTA_ADMINISTRACION
    d, sql, rechazo = base
    d.crear_devolucion_deposito(11, 2, "EJ se puso fea", date(2026, 9, 10), cargada_desde="deposito")
    rechazo(1, 1, compra_id=11)          # iba en el cajón: vuelve con seña
    rechazo(2, 1, compra_id=11)          # iba en caja de Día: no
    movs = {(m["tipo"], m["bultos"], m["sena"]) for m in d.movimientos_del_deposito(
        date(2026, 9, 1), date(2026, 9, 30), proveedor_id=1)
        if m["compra_id"] == 11}
    assert movs == {("entrada", 10.0, 500.0), ("deposito", -2.0, 500.0),
                    ("rechazo", -1.0, 500.0), ("rechazo", -1.0, None)}
    # Nada de esto crea un vale.
    assert sql("SELECT count(*) FROM vales_a_cobrar") == [(0,)]
    monkeypatch.setenv("CLAVE_ADMINISTRACION", "admin-secreta")
    cliente = _cliente()
    cliente.cookies.set(PUERTA_ADMINISTRACION.cookie, PUERTA_ADMINISTRACION.firma("admin-secreta"))
    with patch("app.main._hoy_argentina", return_value=date(2026, 9, 12)):
        pantalla = cliente.get("/administracion/ingresos?proveedor_id=1")
        excel = cliente.get("/administracion/ingresos/movimientos-excel?proveedor_id=1&tipo=deposito")
    marcado = pantalla.text.split("</style>")[-1]
    assert "entró con seña: 10 cajones × $500" in marcado
    assert "volvió con seña: 2 cajones × $500" in marcado
    assert marcado.count("con seña:") == 3
    hoja = openpyxl.load_workbook(io.BytesIO(excel.content)).active
    filas = [tuple(c.value for c in fila) for fila in hoja.iter_rows(min_row=5)]
    assert [f[8:] for f in filas] == [(500, "volvió con seña: 2 cajones × $500", "Depósito")]


# --- DEVOLVER también desde ADMINISTRACIÓN (dueño, 01/10) --------------------

def _cliente_administracion(monkeypatch):
    from app.main import PUERTA_ADMINISTRACION
    monkeypatch.setenv("CLAVE_ADMINISTRACION", "admin-secreta")
    cliente = _cliente()
    cliente.cookies.set(PUERTA_ADMINISTRACION.cookie, PUERTA_ADMINISTRACION.firma("admin-secreta"))
    return cliente


def test_la_MISMA_pantalla_en_ADMINISTRACION_se_queda_en_Administracion(base, monkeypatch):
    from unittest.mock import patch
    d, sql, _ = base
    with patch("app.main._hoy_argentina", return_value=HOY):
        respuesta = _cliente_administracion(monkeypatch).get("/administracion/devolver?proveedor_id=1")
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    # Las mismas compras con lo que queda, por la misma consulta.
    assert re.search(r'value="11"[^>]*>.*?quedan <span class="resto">6</span> de 10', marcado, re.S)
    # La barra, los dos formularios y el volver son de Administración: ninguna
    # mitad del camino manda a Depósito (corolario 63).
    assert marcado.count('action="/administracion/devolver"') == 2
    assert "/deposito" not in marcado
    # El atrás de la barra vive ANTES del último </style> (corolario 50): se
    # mira sobre la página entera, anclado en el elemento (corolario 57).
    assert 'href="/administracion" aria-label="Volver atrás"' in respuesta.text
    assert 'href="/deposito"' not in respuesta.text
    assert 'href="/administracion">Volver a Administración' in marcado


def test_DEVOLVER_desde_ADMINISTRACION_guarda_su_SECTOR_y_vuelve_a_Administracion(base, monkeypatch):
    from unittest.mock import patch
    d, sql, _ = base
    with patch("app.main._hoy_argentina", return_value=HOY):
        respuesta = _cliente_administracion(monkeypatch).post(
            "/administracion/devolver",
            data={"proveedor_id": "1", "compra_id": "11", "cantidad": "2", "motivo": "EJ fea"},
            follow_redirects=False)
        deposito = _cliente().post(
            "/deposito/devolver",
            data={"proveedor_id": "1", "compra_id": "12", "cantidad": "1", "motivo": "EJ fea"},
            follow_redirects=False)
    assert respuesta.status_code == 303 and respuesta.headers["location"].startswith("/administracion/devolver?")
    assert deposito.status_code == 303 and deposito.headers["location"].startswith("/deposito/devolver?")
    assert sql("""SELECT compra_devolucion_id, cantidad, cargada_desde FROM movimientos_stock
                  WHERE tipo = 'devolucion_deposito' ORDER BY id""") == [
        (11, -2, "administracion"), (12, -1, "deposito")]


def test_desde_ADMINISTRACION_tampoco_se_devuelve_MAS_de_lo_que_queda(base, monkeypatch):
    from unittest.mock import patch
    d, sql, _ = base
    with patch("app.main._hoy_argentina", return_value=HOY):
        respuesta = _cliente_administracion(monkeypatch).post("/administracion/devolver", data={
            "proveedor_id": "1", "compra_id": "11", "cantidad": "7", "motivo": "EJ fea"})
    assert respuesta.status_code == 400 and "quedan 6 bultos" in respuesta.text
    assert sql("SELECT count(*) FROM movimientos_stock WHERE tipo = 'devolucion_deposito'") == [(0,)]


def test_ADMINISTRACION_sin_su_clave_no_devuelve(base, monkeypatch):
    d, sql, _ = base
    monkeypatch.setenv("CLAVE_ADMINISTRACION", "admin-secreta")
    _cliente().post("/administracion/devolver", data={
        "proveedor_id": "1", "compra_id": "11", "cantidad": "1", "motivo": "EJ"})
    assert sql("SELECT count(*) FROM movimientos_stock WHERE tipo = 'devolucion_deposito'") == [(0,)]


def test_sin_SECTOR_no_se_escribe_y_la_base_tampoco_lo_deja(base):
    d, sql, _ = base
    import psycopg2
    with pytest.raises(TypeError):
        d.crear_devolucion_deposito(12, 1, "EJ", HOY)
    with pytest.raises(ValueError):
        d.crear_devolucion_deposito(12, 1, "EJ", HOY, cargada_desde="compras")
    with pytest.raises(psycopg2.errors.CheckViolation):
        sql("""INSERT INTO movimientos_stock (articulo_id, tipo, cantidad, motivo, fecha_operacion,
                                              stock_sistema, compra_devolucion_id)
               VALUES (1, 'devolucion_deposito', -1, 'EJ', '2026-09-10', 0, 12)""")


def test_MOVIMIENTOS_dicen_el_SECTOR_y_se_FILTRAN_por_el(base, monkeypatch):
    import io
    import openpyxl
    from unittest.mock import patch
    d, sql, _ = base
    d.crear_devolucion_deposito(11, 2, "EJ uno", date(2026, 9, 10), cargada_desde="administracion")
    d.crear_devolucion_deposito(12, 1, "EJ dos", date(2026, 9, 10), cargada_desde="deposito")
    assert [(m["motivo"], m["sector"]) for m in d.movimientos_del_deposito(
        date(2026, 9, 1), date(2026, 9, 30), sector="administracion")] == [("EJ uno", "administracion")]
    cliente = _cliente_administracion(monkeypatch)
    with patch("app.main._hoy_argentina", return_value=date(2026, 9, 12)):
        todo = cliente.get("/administracion/ingresos?proveedor_id=1").text.split("</style>")[-1]
        filtrado = cliente.get("/administracion/ingresos?sector=deposito").text.split("</style>")[-1]
        excel = cliente.get("/administracion/ingresos/movimientos-excel?sector=administracion")
    assert '<span class="sector">cargada desde Administración</span>' in todo
    assert '<span class="sector">cargada desde Depósito</span>' in todo
    assert "EJ dos" in filtrado and "EJ uno" not in filtrado
    assert '<option value="deposito" selected>Depósito</option>' in filtrado
    hoja = openpyxl.load_workbook(io.BytesIO(excel.content)).active
    filas = [tuple(c.value for c in fila) for fila in hoja.iter_rows(min_row=5)]
    assert [(f[7], f[10]) for f in filas] == [("EJ uno", "Administración")]
    assert "cargadas desde Administración" in hoja.cell(row=2, column=1).value


def test_el_HUB_de_Administracion_tiene_Devolver_mercaderia_y_Deposito_sigue_igual():
    import io as _io
    admin = _io.open("templates/administracion.html", encoding="utf-8").read()
    deposito = _io.open("templates/deposito.html", encoding="utf-8").read()
    assert re.search(r'href="/administracion/devolver">[^<]*<span>Devolver mercadería</span></a>', admin)
    # En la tarjeta de stock (dueño, 01/10; desde el 04/10 se llama "Stock"),
    # no en Facturación.
    tarjetas = admin.split('<div class="tarjeta">')
    control = next(c for c in tarjetas if "<h2>Stock</h2>" in c)
    facturacion = next(c for c in tarjetas if "<h2>Facturación y cobranzas</h2>" in c)
    assert 'href="/administracion/devolver"' in control
    assert 'href="/administracion/devolver"' not in facturacion
    assert 'href="/deposito/devolver">' in deposito


def test_el_RESUMEN_PROVEEDORES_resta_las_devoluciones_del_dia_en_que_se_devolvieron(base, monkeypatch):
    """Dueño, 01/10: al pagarle al proveedor se ve lo que entró menos lo que se
    le devolvió, cada devolución como renglón NEGATIVO del día en que se devolvió."""
    from fastapi.testclient import TestClient
    from app.main import PUERTA_ADMINISTRACION, app
    d, sql, rechazo = base
    rechazo(1, 2, compra_id=11)                     # 09/09, en cajón, compra con seña $500
    d.crear_devolucion_deposito(12, 3, "EJ fea", date(2026, 9, 10), cargada_desde="administracion")
    monkeypatch.setenv(PUERTA_ADMINISTRACION.env_var, "clave-adm")
    cliente = TestClient(app, base_url="https://testserver")
    cliente.cookies.set(PUERTA_ADMINISTRACION.cookie, PUERTA_ADMINISTRACION.firma("clave-adm"))
    respuesta = cliente.get("/administracion/ingresos/pagar",
                            params={"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-30"})
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    tarjeta_uno = marcado.split('<p class="proveedor-encabezado">EJ Uno')[1].split('<div class="tarjeta">')[0]
    filas = re.findall(r"<tr[^>]*>(.*?)</tr>", tarjeta_uno, re.S)[1:]         # sin el encabezado
    celdas = [[re.sub(r"<[^>]+>", "", c).strip() for c in re.findall(r"<td[^>]*>(.*?)</td>", f, re.S)]
              for f in filas]
    # En orden de DÍA: 05/09 y 06/09 entran; 09/09 y 10/09 se devuelven.
    assert [c[0][:5] for c in celdas] == ["05/09", "06/09", "09/09", "10/09"]
    # El artículo lleva abajo la marca del cajón (dueño, 05/10); la 11 no tiene.
    assert celdas[2][1:4] == ["Compra 11", "EJEMPLO FrutaSin marca", "-2 bultos"]
    assert celdas[2][7] == "Devolución por rechazo"
    assert celdas[3][1] == "Compra 12" and celdas[3][7] == "Devolución desde depósito"
    # Entró 10×($100+$500) + 8×$120 = $6.960; se devolvieron 2×($100+$500) + 3×$120 = $1.560.
    assert re.search(r"Subtotal EJ Uno \(entró menos lo devuelto\)</span>\s*<span>\$5\.400</span>",
                     tarjeta_uno)
    # El estado "No ingresó" es para controlar y no lleva devoluciones.
    solo_rechazadas = cliente.get("/administracion/ingresos/pagar", params={
        "fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-30", "estado": "rechazado"})
    assert "Devolución" not in solo_rechazadas.text.split("</style>")[-1]
