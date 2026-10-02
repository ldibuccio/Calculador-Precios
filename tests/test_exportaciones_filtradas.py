"""Toda exportación sale con los MISMOS filtros que la pantalla (01/10, dueño).

Lionel filtró la planilla para pagar por FRUTAMAX S.R.L. y el PDF bajó todos
los proveedores. La regla: el PDF y el Excel salen con exactamente lo que la
pantalla muestra en ese momento, y su encabezado dice qué se filtró.

Dos mitades:

  · ESTRUCTURA (sin base): toda ruta que devuelve un PDF o un Excel está en
    `EXPORTACIONES` —el conjunto ENCONTRADO contra el DECIDIDO (corolario
    60)—, y cada link de exportar lleva TODOS los campos del formulario de su
    pantalla.
  · CONTENIDO (contra Postgres): se abre la pantalla filtrada, se sigue el
    link de exportar QUE LA PANTALLA DIBUJA (no una URL armada en el test: el
    error de Consultar precios estaba en el link), y se lee el archivo: el
    otro proveedor, cliente o artículo no puede aparecer, y el filtro sí.

Los nombres son de EJEMPLO.
"""
import io
import os
import re
import sys

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from scripts.humo import hay_postgres, preparar_base  # noqa: E402

OBLIGATORIO = os.environ.get("HUMO_OBLIGATORIO") == "1"

# --- ESTRUCTURA -----------------------------------------------------------------

# Cada ruta que devuelve un PDF o un Excel, con cómo se cuida que salga con los
# filtros de su pantalla. "contenido" = la cubre un test de abajo contra la base.
EXPORTACIONES = {
    "/compras/buscar/exportar-pdf": "contenido",
    "/compras/buscar/exportar-excel": "contenido",
    "/logistica/consultar/exportar-pdf": "contenido",
    "/logistica/consultar/exportar-excel": "contenido",
    "/administracion/ingresos/exportar-pdf": "contenido",
    "/administracion/ingresos/exportar-excel": "contenido",
    "/administracion/ingresos/movimientos-excel": "contenido",
    "/administracion/vacios/movimientos/excel": "contenido",
    "/gerencia/vales/movimientos-excel": "contenido",
    "/precios/consultar/exportar-pdf": "contenido",
    "/precios/consultar/exportar-excel": "contenido",
    "/precios/vigencias/exportar-excel": "contenido",
    "/administracion/precios-por-periodo/exportar-excel": "contenido",
    "/gerencia/rentabilidad/exportar-pdf": "contenido",
    "/gerencia/rentabilidad/exportar-excel": "contenido",
    "/gerencia/rentabilidad-real/exportar-pdf": "contenido",
    "/gerencia/rentabilidad-real/exportar-excel": "contenido",
    "/administracion/pedidos/buscar/exportar-pdf": "contenido",
    "/administracion/pedidos/buscar/exportar-excel": "contenido",
    # Cobranzas de segunda (02/10): lo cuida tests/test_cobranzas_segunda.py,
    # siguiendo el link que dibuja la pantalla. La de Administración es la
    # misma función con otra puerta.
    "/gerencia/cobranzas-segunda/exportar-pdf": "tests/test_cobranzas_segunda.py",
    "/gerencia/cobranzas-segunda/exportar-excel": "tests/test_cobranzas_segunda.py",
    "/administracion/cobranzas-segunda/exportar-pdf": "misma función que la de Gerencia",
    "/administracion/cobranzas-segunda/exportar-excel": "misma función que la de Gerencia",
    # El mismo Excel que la de Gerencia, con la misma función y otra puerta.
    "/administracion/vales/movimientos-excel": "misma función que /gerencia/vales/movimientos-excel",
    # El filtro por tipo lo cuida tests/test_remanente_por_tipo.py.
    "/administracion/stock/remanente/exportar-excel": "tests/test_remanente_por_tipo.py",
    "/administracion/stock/remanente/exportar-pdf": "tests/test_remanente_por_tipo.py",
    # La pantalla no tiene filtros: el stock de todas las pilas.
    "/administracion/vacios/stock/excel": "sin filtros en la pantalla",
    "/administracion/vacios/stock/pdf": "sin filtros en la pantalla",
    # La pantalla filtra solo por fecha, y el link la lleva.
    "/puesto/envases/stock/exportar-pdf": "solo fecha",
    "/puesto/envases/stock/exportar-excel": "solo fecha",
    "/puesto/envases/movimientos/exportar-pdf": "solo fechas",
    "/puesto/envases/movimientos/exportar-excel": "solo fechas",
    # Exportan lo GUARDADO, que es lo que la pantalla muestra: guardan primero.
    "/compras/que-comprar/pdf": "lo guardado (tildes)",
    "/compras/disponibles/guardar-y-exportar-excel": "lo guardado",
    "/precios/resultado-negociacion": "una negociación puntual, sin filtros",
    # "Guardar y generar listado": la lista entera del cliente recién guardada.
    # La pantalla de carga no filtra artículos.
    "/precios/cargar/guardar-y-exportar-pdf": "lista entera, la pantalla no filtra",
    "/precios/cargar/guardar-y-exportar-excel": "lista entera, la pantalla no filtra",
    "/precios/cargar-foto/guardar-y-exportar-pdf": "lista entera, la pantalla no filtra",
    "/precios/cargar-foto/guardar-y-exportar-excel": "lista entera, la pantalla no filtra",
}

# Rutas que devuelven una imagen o un archivo subido, no una exportación.
NO_SON_EXPORTACIONES = {"/compras/{compra_id}/fotos", "/deposito/pedido/{pedido_id}/fotos"}


def _rutas_que_exportan():
    """Las rutas cuyo cuerpo devuelve un PDF o un Excel, más las que lo hacen
    llamando a un ayudante (su URL lo dice: exportar, excel, pdf). Un
    decorador apilado cuenta las dos URLs; un `def` sin decorador corta."""
    texto = io.open(os.path.join(RAIZ, "app", "main.py"), encoding="utf-8").read().split("\n")
    rutas, pendientes, actuales = set(), [], []
    for linea in texto:
        m = re.match(r'@app\.(?:get|post)\("([^"]+)"', linea)
        if m:
            pendientes.append(m.group(1))
            continue
        if re.match(r"(async )?def ", linea):
            actuales, pendientes = pendientes, []
            rutas |= {r for r in actuales if re.search(r"exportar|excel|/pdf$", r)}
        if ("application/pdf" in linea or "spreadsheetml" in linea) and actuales:
            rutas |= set(actuales)
    return rutas - NO_SON_EXPORTACIONES


def test_TODA_exportacion_esta_DECIDIDA():
    encontradas = _rutas_que_exportan()
    assert encontradas - set(EXPORTACIONES) == set(), "exportación nueva sin decidir cómo sigue los filtros"
    assert set(EXPORTACIONES) - encontradas == set(), "sacala de EXPORTACIONES: ya no exporta"


def test_cada_link_de_exportar_lleva_TODOS_los_campos_del_formulario_de_su_pantalla():
    """Las pantallas cuyo link se escribe a mano en la plantilla: cada campo
    del formulario GET tiene que viajar en el link. Las que arman el link en
    el server (`consulta`, `query_exportar`) las mira la mitad de contenido."""
    revisadas = 0
    for nombre in sorted(os.listdir(os.path.join(RAIZ, "templates"))):
        texto = io.open(os.path.join(RAIZ, "templates", nombre), encoding="utf-8").read()
        marcado = texto.split("</style>")[-1]
        links = re.findall(r'href="([^"]*(?:exportar-pdf|exportar-excel|movimientos-excel)\?[^"]*)"', marcado)
        links = [l for l in links if "{{ consulta }}" not in l and "{{ query_exportar }}" not in l]
        if not links:
            continue
        formularios = re.findall(r'<form[^>]*method="get"[^>]*>(.*?)</form>', marcado, re.S)
        campos = {c for f in formularios for c in re.findall(r'name="([^"]+)"', f)}
        for link in links:
            parametros = set(re.findall(r'[?&](?:amp;)?([a-z_]+)=', link))
            assert campos <= parametros, (nombre, campos - parametros)
            revisadas += 1
    assert revisadas >= 14, revisadas


# --- CONTENIDO, contra Postgres ---------------------------------------------------

SIEMBRA = """
insert into clientes (id, nombre) overriding system value values (1, 'EJEMPLO Dia'), (2, 'EJEMPLO Coto');
insert into articulos (id, nombre, grupo) overriding system value
  values (1, 'EJEMPLO Fruta', 'fruta'), (2, 'EJEMPLO Verdura', 'hortaliza');
insert into proveedores (id, nombre, codigo_puesto) overriding system value
  values (1, 'EJ Uno', 'N91P01'), (2, 'EJ Dos', 'N91P02');
insert into fichas_logistica (id, cliente_id, articulo_id, contenido_caja, unidad_venta, nombre_cliente)
  overriding system value values
  (1, 1, 1, 16, 'kilo', 'EJ Ficha Fruta Dia'),
  (2, 1, 2, 16, 'kilo', 'EJ Ficha Verdura Dia'),
  (3, 2, 1, 16, 'kilo', 'EJ Ficha Fruta Coto');
insert into precios_venta_historial (articulo_id, cliente_id, ficha_id, precio, vigente_desde)
  values (1, 1, 1, 100, '2026-09-01'), (2, 1, 2, 200, '2026-09-01'), (1, 2, 3, 300, '2026-09-01');
insert into compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
                     contenido_por_cajon, cantidad_kilos, importe, sena, estado, procesada_el,
                     cantidad_cajones_real, contenido_por_cajon_real)
  overriding system value values
  (11, 1, 1, '2026-09-05', 10, 16, 160, 50, 500, 'recepcionado', '2026-09-05 18:00-03', 10, 16),
  (12, 2, 2, '2026-09-06',  8, 16, 128, 60, 500, 'recepcionado', '2026-09-06 18:00-03',  8, 16);
insert into pedidos (id, cliente_id, fecha_operacion, origen) overriding system value
  values (1, 1, '2026-09-08', 'texto'), (2, 2, '2026-09-08', 'texto');
insert into pedidos_sucursales (pedido_id, sucursal) values (1, 'EJ DiaSuc'), (2, 'EJ CotoSuc');
insert into pedidos_renglones (id, pedido_id, articulo_id, ficha_id, sucursal, cantidad,
                               cantidad_armada, armado_el, kilos_enviados)
  overriding system value values
  (1, 1, 1, 1, 'EJ DiaSuc', 4, 4, '2026-09-08 10:00-03', 64),
  (2, 1, 2, 2, 'EJ DiaSuc', 3, 3, '2026-09-08 10:00-03', 48),
  (3, 2, 1, 3, 'EJ CotoSuc', 2, 2, '2026-09-08 10:00-03', 32);
insert into vales_a_cobrar (origen, proveedor_id, fecha, importe, numero)
  values ('anterior_al_sistema', 1, '2026-09-10', 1000, 'V-UNO'),
         ('anterior_al_sistema', 2, '2026-09-10', 2000, 'V-DOS');
"""


@pytest.fixture
def cliente(monkeypatch):
    if not hay_postgres():
        if OBLIGATORIO:
            pytest.fail("HUMO_OBLIGATORIO=1 y no hay Postgres: las exportaciones no se verificaron")
        pytest.skip("sin Postgres local")
    url = preparar_base()
    import psycopg2
    conexion = psycopg2.connect(url)
    with conexion.cursor() as cursor:
        cursor.execute(SIEMBRA)
    conexion.commit()
    conexion.close()
    monkeypatch.setenv("DATABASE_URL", url)
    from fastapi.testclient import TestClient
    from app.main import PUERTA_ADMINISTRACION, PUERTA_COMPRAS, PUERTA_GERENCIA, app
    navegador = TestClient(app, base_url="https://testserver")
    for puerta in (PUERTA_ADMINISTRACION, PUERTA_COMPRAS, PUERTA_GERENCIA):
        monkeypatch.setenv(puerta.env_var, "clave-ej")
        navegador.cookies.set(puerta.cookie, puerta.firma("clave-ej"))
    return navegador


def _texto_del_archivo(respuesta) -> str:
    assert respuesta.status_code == 200, respuesta.text[-300:]
    if respuesta.content.startswith(b"%PDF"):
        import pypdfium2 as pdfium
        documento = pdfium.PdfDocument(respuesta.content)
        return " ".join(documento[i].get_textpage().get_text_range() for i in range(len(documento)))
    import openpyxl
    libro = openpyxl.load_workbook(io.BytesIO(respuesta.content))
    return " ".join(str(c.value) for hoja in libro.worksheets
                    for fila in hoja.iter_rows() for c in fila if c.value is not None)


DESDE, HASTA = "2026-09-01", "2026-09-30"

# (pantalla filtrada, lo que tiene que estar, lo que NO puede estar, lo que
# dice el encabezado del filtro)
CASOS = [
    pytest.param(f"/compras/buscar?fecha_desde={DESDE}&fecha_hasta={HASTA}&proveedor_id=1",
                 "EJ Uno", "EJ Dos", "proveedor EJ Uno", id="buscar_compras"),
    pytest.param(f"/logistica/consultar?fecha_desde={DESDE}&fecha_hasta={HASTA}&proveedor_id=1&estado=todos",
                 "EJ Uno", "EJ Dos", "EJ Uno", id="consultar_retiros"),
    pytest.param(f"/administracion/ingresos/pagar?fecha_desde={DESDE}&fecha_hasta={HASTA}&proveedor_id=1",
                 "EJ Uno", "EJ Dos", "proveedor EJ Uno", id="planilla_para_pagar"),
    pytest.param(f"/administracion/ingresos?desde={DESDE}&hasta={HASTA}&proveedor_id=1",
                 "EJ Uno", "EJ Dos", "proveedor EJ Uno", id="movimientos_del_deposito"),
    pytest.param(f"/administracion/vacios/movimientos?desde={DESDE}&hasta={HASTA}&proveedor_id=1",
                 "EJ Uno", "EJ Dos", "EJ Uno", id="movimientos_de_vacios"),
    pytest.param(f"/gerencia/vales/movimientos?desde={DESDE}&hasta={HASTA}&proveedor_id=1",
                 "V-UNO", "V-DOS", "proveedor EJ Uno", id="movimientos_de_vales"),
    pytest.param("/precios/consultar?cliente_id=1&fecha=2026-09-20&ficha_id=1",
                 "EJ Ficha Fruta Dia", "EJ Ficha Verdura Dia", "EJ Ficha Fruta Dia", id="consultar_precios"),
    pytest.param(f"/precios/vigencias?cliente_id=1&desde={DESDE}&hasta={HASTA}",
                 "EJ Ficha Fruta Dia", "EJ Ficha Fruta Coto", "EJEMPLO Dia", id="precios_por_periodo"),
    pytest.param(f"/administracion/precios-por-periodo?cliente_id=1&desde={DESDE}&hasta={HASTA}",
                 "EJ Ficha Fruta Dia", "EJ Ficha Fruta Coto", "EJEMPLO Dia", id="precios_por_periodo_adm"),
    pytest.param(f"/gerencia/rentabilidad?cliente_id=1&fecha_desde={DESDE}&fecha_hasta={HASTA}&articulo_id=1",
                 "EJEMPLO Fruta", "EJEMPLO Verdura", "EJEMPLO Fruta", id="rentabilidad"),
    pytest.param(f"/gerencia/rentabilidad-real?cliente_id=1&fecha_desde={DESDE}&fecha_hasta={HASTA}&articulo_id=1",
                 "EJEMPLO Fruta", "EJEMPLO Verdura", "EJEMPLO Fruta", id="rentabilidad_real"),
    pytest.param(f"/administracion/pedidos/buscar?cliente_id=1&fecha_desde={DESDE}&fecha_hasta={HASTA}",
                 "EJ DiaSuc", "EJ CotoSuc", "EJEMPLO Dia", id="armar_remito"),
]


@pytest.mark.parametrize("pantalla, esta, no_esta, encabezado", CASOS)
def test_la_EXPORTACION_sale_con_los_filtros_de_la_PANTALLA(cliente, pantalla, esta, no_esta, encabezado):
    from unittest.mock import patch
    from datetime import date
    with patch("app.main._hoy_argentina", return_value=date(2026, 9, 20)):
        respuesta = cliente.get(pantalla)
        assert respuesta.status_code == 200, respuesta.text[-300:]
        marcado = respuesta.text.split("</style>")[-1]
        links = [l.replace("&amp;", "&") for l in re.findall(
            r'href="([^"]*(?:exportar-pdf|exportar-excel|movimientos-excel|movimientos/excel)\?[^"]*)"', marcado)]
        assert links, "la pantalla no dibujó ningún link de exportar"
        for link in links:
            texto = _texto_del_archivo(cliente.get(link))
            assert esta in texto, (link, esta)
            assert no_esta not in texto, (link, no_esta)
            assert encabezado in texto, (link, encabezado)


def test_la_PLANILLA_PARA_PAGAR_abierta_desde_MOVIMIENTOS_conserva_el_proveedor(cliente):
    """El camino que perdía el filtro: Movimientos del depósito filtrado por
    proveedor → "Planilla para pagar" llegaba con las fechas y sin el
    proveedor, y su PDF bajaba todos."""
    respuesta = cliente.get(f"/administracion/ingresos?desde={DESDE}&hasta={HASTA}&proveedor_id=1&articulo_id=1")
    link = re.search(r'href="(/administracion/ingresos/pagar\?[^"]*)"', respuesta.text).group(1)
    link = link.replace("&amp;", "&")
    assert "proveedor_id=1" in link and "articulo_id=1" in link
    planilla = cliente.get(link).text.split("</style>")[-1]
    assert "EJ Dos" not in planilla.split("</form>")[-1]
