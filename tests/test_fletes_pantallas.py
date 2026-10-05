"""Las pantallas de Fletes (dueño, 05/10), contra Postgres.

Administración → Pedidos → Fletes, tres pestañas: Fleteros y camiones ·
Flete del día · Cuenta del fletero. El flete se corrige el mismo día desde
Administración; después, solo Gerencia, con historial. La cuenta exporta
PDF y Excel con los filtros de la pantalla. Y Rentabilidad Real resta la
parte de la empresa en una línea "Flete" aparte. Nombres de EJEMPLO.
"""
import io
import os
import re
import sys
from datetime import date

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from scripts.humo import hay_postgres, preparar_base  # noqa: E402
from tests.test_administracion_reordenada import CLAVES, _cliente  # noqa: E402
from tests.test_fletes_contra_la_base import SIEMBRA  # noqa: E402

OBLIGATORIO = os.environ.get("HUMO_OBLIGATORIO") == "1"
PALLETS = {"fm_VL": "10", "pm_VL": "8", "fm_BZ": "8", "pm_BZ": "0", "fm_GR": "0", "pm_GR": ""}
MENSAJE = ("Hola EJ Juan, para el lunes 09/03 necesito: Vicente López: 1 Grande + 1 Mediano (18 pallets). "
           "Burzaco: 1 Mediano (8 pallets). Gracias.")


@pytest.fixture
def base(monkeypatch):
    if not hay_postgres():
        if OBLIGATORIO:
            pytest.fail("HUMO_OBLIGATORIO=1 y no hay Postgres: las pantallas de fletes no se verificaron")
        pytest.skip("sin Postgres local")
    url = preparar_base()
    import psycopg2
    conexion = psycopg2.connect(url)
    with conexion.cursor() as cursor:
        cursor.execute(SIEMBRA)
    conexion.commit()
    conexion.close()
    monkeypatch.setenv("DATABASE_URL", url)
    for variable, valor in CLAVES.items():
        monkeypatch.setenv(variable, valor)
    import app.main as m

    def sql(consulta, parametros=None):
        conexion = psycopg2.connect(url)
        try:
            with conexion.cursor() as cursor:
                cursor.execute(consulta, parametros)
                filas = cursor.fetchall() if cursor.description else None
            conexion.commit()
            return filas
        finally:
            conexion.close()
    return m, sql


def _confirmar(cliente, fecha="2026-03-09", pallets=None, camino="/administracion/fletes/dia", extra=None):
    """Arma y confirma como la persona: lo que la pantalla de armar trae en
    los campos es lo que se confirma (con `extra` pisando alguno)."""
    datos = {"fecha": fecha, "cliente_id": "1", "fletero_id": "1", **(pallets or PALLETS)}
    armado = cliente.post(f"{camino}/armar", data=datos)
    assert armado.status_code == 200, armado.text[:300]
    for nombre, valor in re.findall(r'<input type="number" name="(camion_[A-Z]+_\d+)"[^>]*value="(\d+)"', armado.text):
        datos[nombre] = valor
    datos.update(extra or {})
    return armado, cliente.post(f"{camino}/confirmar", data=datos)


def test_el_boton_FLETES_esta_en_PEDIDOS_y_las_TRES_pestanias_abren(base):
    m, _ = base
    cliente = _cliente(m, "administracion")
    hub = cliente.get("/administracion").text
    pedidos = hub.split("<h2>Pedidos</h2>")[1].split("</div>")[0]
    assert re.search(r'<a class="boton color-2" href="/administracion/fletes/dia"><svg[^>]*>.*?</svg><span>Fletes</span></a>',
                     pedidos, re.S)
    for ruta, activa in (("/administracion/fletes/fleteros", "Fleteros y camiones"),
                         ("/administracion/fletes/dia", "Flete del día"),
                         ("/administracion/fletes/cuenta", "Cuenta del fletero")):
        respuesta = cliente.get(ruta)
        assert respuesta.status_code == 200, ruta
        assert re.search(r'<a href="' + re.escape(ruta) + r'" aria-current="page">' + activa + "</a>", respuesta.text)


def test_ARMAR_propone_lo_mas_barato_y_CONFIRMAR_guarda_y_da_el_MENSAJE(base):
    m, sql = base
    cliente = _cliente(m, "administracion")
    vacio = cliente.get("/administracion/fletes/dia?fecha=2026-03-09&cliente_id=1&fletero_id=1").text
    assert '<input type="number" name="fm_VL"' in vacio and "Garín" in vacio
    armado, confirmado = _confirmar(cliente)
    assert "lo más barato: $240.000" in armado.text
    assert confirmado.status_code == 200 and "Flete confirmado" in confirmado.text
    assert f'<p class="mensaje" id="mensaje-whatsapp">{MENSAJE}</p>' in confirmado.text
    assert '<button type="button" id="copiar-whatsapp">Copiar para WhatsApp</button>' in confirmado.text
    assert sql("select sucursal, pallets_frutamax, pallets_palmala from fletes_sucursales order by id") == [
        ("VL", 10, 8), ("BZ", 8, 0)]


def test_lo_EDITADO_antes_de_confirmar_es_lo_que_se_guarda(base):
    m, sql = base
    cliente = _cliente(m, "administracion")
    # a Burzaco, en vez del Mediano propuesto, dos Chicos
    _, confirmado = _confirmar(cliente, extra={"camion_BZ_12": "0", "camion_BZ_13": "2"})
    assert "Burzaco: 2 Chico (8 pallets)" in confirmado.text
    assert sql("select count(*) from fletes_viajes v join fleteros_camiones c on c.id = v.camion_id "
               "where c.nombre = 'Chico'") == [(2,)]


def test_FLOTA_INSUFICIENTE_lo_dice_y_dice_DONDE_se_pasa(base):
    m, _ = base
    cliente = _cliente(m, "administracion")
    armado = cliente.post("/administracion/fletes/dia/armar", data={
        "fecha": "2026-03-09", "cliente_id": "1", "fletero_id": "1",
        "fm_VL": "30", "pm_VL": "0", "fm_BZ": "20", "pm_BZ": "0", "fm_GR": "0", "pm_GR": "0"}).text
    assert "La flota de EJ Juan no alcanza" in armado
    assert re.search(r"se pasa en: Grande \(usa \d+, queda 1\)", armado)


def test_NO_ENTRAN_no_confirma_y_vuelve_con_lo_cargado(base):
    m, sql = base
    cliente = _cliente(m, "administracion")
    _, confirmado = _confirmar(cliente, extra={"camion_VL_11": "0", "camion_VL_12": "1"})
    assert confirmado.status_code == 400
    assert "A Vicente López no le entran: son 18 pallets y los camiones elegidos llevan 8." in confirmado.text
    assert sql("select count(*) from fletes") == [(0,)]


def test_CORREGIR_el_mismo_dia_desde_ADMINISTRACION_y_despues_solo_GERENCIA_con_historial(base):
    m, sql = base
    cliente = _cliente(m, "administracion", "gerencia")
    _confirmar(cliente)
    (fid,), = sql("select id from fletes")
    hoy = cliente.get(f"/administracion/fletes/{fid}/corregir").text
    assert "Volver a armar" in hoy
    _, corregido = _confirmar(cliente, camino=f"/administracion/fletes/{fid}/corregir",
                              pallets={"fm_VL": "12", "pm_VL": "0", "fm_BZ": "0", "pm_BZ": "0",
                                       "fm_GR": "4", "pm_GR": "0"})
    assert "Vicente López: 1 Grande (12 pallets). Garín: 1 Chico (4 pallets)." in corregido.text
    assert sql("select sector from fletes_correcciones") == [("administracion",)]

    # PASÓ EL DÍA: Administración ya no corrige, ni por la pantalla ni por el POST
    sql("update fletes set confirmado_el = now() - interval '2 days'")
    ayer = cliente.get(f"/administracion/fletes/{fid}/corregir").text
    assert "Pasó el día en que se confirmó: lo corrige Gerencia." in ayer and "Volver a armar" not in ayer
    assert f'href="/gerencia/fletes/{fid}/corregir"' in ayer
    rebote = cliente.post(f"/administracion/fletes/{fid}/corregir/confirmar",
                          data={"fm_VL": "1", "pm_VL": "0", "camion_VL_13": "1"}, follow_redirects=False)
    assert rebote.status_code == 303 and "lo+corrige+Gerencia" in rebote.headers["location"]
    assert sql("select count(*) from fletes_correcciones") == [(1,)]

    gerencia = cliente.get(f"/gerencia/fletes/{fid}/corregir").text
    assert "Volver a armar" in gerencia and 'action="/gerencia/fletes/' in gerencia
    _, corregido = _confirmar(cliente, camino=f"/gerencia/fletes/{fid}/corregir",
                              pallets={"fm_VL": "0", "pm_VL": "0", "fm_BZ": "4", "pm_BZ": "0",
                                       "fm_GR": "0", "pm_GR": "0"})
    assert "Flete corregido" in corregido.text
    assert sql("select sector from fletes_correcciones order by id") == [("administracion",), ("gerencia",)]


def test_GERENCIA_sin_su_clave_no_corrige(base):
    m, sql = base
    _confirmar(_cliente(m, "administracion"))
    (fid,), = sql("select id from fletes")
    sin_clave = _cliente(m, "administracion")
    assert 'type="password"' in sin_clave.get(f"/gerencia/fletes/{fid}/corregir").text
    sin_clave.post(f"/gerencia/fletes/{fid}/corregir/confirmar", data={"fm_BZ": "4", "camion_BZ_13": "1"})
    assert sql("select count(*) from fletes_correcciones") == [(0,)]


def test_la_CUENTA_marca_pagados_y_un_flete_pagado_ya_no_se_corrige(base):
    m, sql = base
    cliente = _cliente(m, "administracion", "gerencia")
    _confirmar(cliente)
    cuenta = cliente.get("/administracion/fletes/cuenta?fletero_id=1&desde=2026-03-01&hasta=2026-03-31").text
    ids = re.findall(r'<input type="checkbox" name="viaje_id" value="(\d+)"', cuenta)
    assert len(ids) == 3
    pagado = cliente.post("/administracion/fletes/cuenta/pagar", data={
        "fletero_id": "1", "desde": "2026-03-01", "hasta": "2026-03-31", "pagado_el": "2026-03-20",
        "viaje_id": ids[:2]}).text
    assert "2 viajes pagados el 20/03/2026." in pagado
    assert pagado.count("Pagado el 20/03/2026") == 2
    assert len(re.findall(r'name="viaje_id"', pagado)) == 1
    (fid,), = sql("select id from fletes")
    for camino in ("/administracion", "/gerencia"):
        assert "Tiene viajes pagados: ya no se corrige." in cliente.get(f"{camino}/fletes/{fid}/corregir").text


def _texto_del_pdf(contenido: bytes) -> str:
    import pypdfium2 as pdfium
    return "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(contenido))


def test_PDF_y_EXCEL_de_la_cuenta_salen_con_los_FILTROS_de_la_pantalla(base):
    """Se sigue el link QUE DIBUJA LA PANTALLA. El RIVAL (otro fletero, otra
    fecha) no puede aparecer, y el encabezado dice qué se filtró."""
    m, sql = base
    cliente = _cliente(m, "administracion")
    _confirmar(cliente)
    _confirmar(cliente, fecha="2026-04-02")
    sql("insert into fletes (id, fecha, cliente_id, fletero_id) overriding system value values (90, '2026-03-09', 1, 2)")
    sql("insert into fletes_sucursales (id, flete_id, sucursal, pallets_frutamax, pallets_palmala) "
        "overriding system value values (90, 90, 'GR', 1, 0)")
    sql("insert into fletes_viajes (flete_sucursal_id, camion_id, precio, parte_frutamax, parte_palmala) "
        "values (90, 21, 777, 777, 0)")
    pantalla = cliente.get("/administracion/fletes/cuenta?fletero_id=1&desde=2026-03-01&hasta=2026-03-31").text
    assert "Fletero: EJ Juan · del 01/03/2026 al 31/03/2026" in pantalla
    links = dict(re.findall(r'<a href="(/administracion/fletes/cuenta/(pdf|excel)\?[^"]+)">', pantalla))
    links = {formato: url.replace("&amp;", "&") for url, formato in links.items()}
    assert set(links) == {"pdf", "excel"}

    texto = " ".join(_texto_del_pdf(cliente.get(links["pdf"]).content).split())
    assert "Fletero: EJ Juan · del 01/03/2026 al 31/03/2026" in texto and "09/03/2026" in texto
    assert "Vicente López" in texto and "$240.000" in texto
    assert "EJ Rival" not in texto and "02/04/2026" not in texto

    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(cliente.get(links["excel"]).content)).active
    celdas = [c for fila in hoja.iter_rows(values_only=True) for c in fila if c is not None]
    assert hoja["A2"].value == "Fletero: EJ Juan · del 01/03/2026 al 31/03/2026"
    assert "EJ Rival" not in celdas and 777 not in celdas
    assert sum(1 for c in celdas if c in ("Vicente López", "Burzaco")) == 3


def test_a_313px_el_MENSAJE_se_ve_ENTERO_y_nada_se_sale(base):
    pytest.importorskip("playwright", reason="lo que se ve lo decide el navegador")
    from playwright.sync_api import sync_playwright
    from scripts.medir_layout import CHROMIUM
    m, _ = base
    cliente = _cliente(m, "administracion")
    armado, confirmado = _confirmar(cliente)
    paginas = {"armar": armado.text, "confirmado": confirmado.text,
               "cuenta": cliente.get("/administracion/fletes/cuenta?desde=2026-03-01&hasta=2026-03-31").text,
               "fleteros": cliente.get("/administracion/fletes/fleteros").text}
    medidas = {}
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(executable_path=CHROMIUM)
        try:
            for nombre, html in paginas.items():
                pagina = navegador.new_page(viewport={"width": 313, "height": 700})
                pagina.set_content(html)
                medidas[nombre] = pagina.evaluate("""() => {
                  const m = document.getElementById('mensaje-whatsapp');
                  const r = m && m.getBoundingClientRect();
                  return {
                    desborde: document.documentElement.scrollWidth - document.documentElement.clientWidth,
                    mensaje: m && {cortado: m.scrollHeight - m.clientHeight, derecha: r.right, texto: m.textContent},
                    chicos: [...document.querySelectorAll('main button, form button, .acciones a, summary, .pestanas a, input, select')]
                      .filter(b => b.offsetParent && !b.closest('.barra-navegacion')
                              && b.type !== 'checkbox' && b.type !== 'hidden'
                              && b.getBoundingClientRect().height < 44).map(b => b.name || b.textContent.trim()),
                  }; }""")
                pagina.close()
        finally:
            navegador.close()
    for nombre, medida in medidas.items():
        assert medida["desborde"] == 0 and medida["chicos"] == [], (nombre, medida)
    mensaje = medidas["confirmado"]["mensaje"]
    assert mensaje["texto"] == MENSAJE and mensaje["cortado"] == 0 and mensaje["derecha"] <= 313, mensaje


# --- Rentabilidad Real: la línea "Flete" ---------------------------------------

def test_REAL_el_flete_resta_de_la_renta_TOTAL_y_no_de_los_grupos():
    from core.costo_real import calcular_rentabilidad_real
    sin = calcular_rentabilidad_real([], {}, 1, date(2026, 3, 1), date(2026, 3, 31))
    con = calcular_rentabilidad_real([], {}, 1, date(2026, 3, 1), date(2026, 3, 31), fletes=[
        {"fecha": date(2026, 3, 9), "sucursal": "Vicente López", "pesos": 94444.45},
        {"fecha": date(2026, 3, 9), "sucursal": "Burzaco", "pesos": 70000}])
    assert con["totales"]["flete"] == pytest.approx(164444.45)
    assert con["totales"]["renta_pesos"] == pytest.approx(sin["totales"]["renta_pesos"] - 164444.45)
    assert sin["totales"]["flete"] == 0 and sin["fletes"] == []
    assert [f["sucursal"] for f in con["fletes"]] == ["Vicente López", "Burzaco"]


def test_REAL_en_pantalla_PDF_y_EXCEL_con_la_parte_de_ESTA_empresa_y_sin_ella_con_filtro(base):
    m, sql = base
    cliente = _cliente(m, "administracion", "gerencia")
    _confirmar(cliente)
    consulta = "cliente_id=1&fecha_desde=2026-03-01&fecha_hasta=2026-03-31"
    pantalla = cliente.get(f"/gerencia/rentabilidad-real?{consulta}").text
    assert "<h2>Flete · $164.444</h2>" in pantalla
    assert re.search(r'<div class="renglon-flete"><span>09/03 · Vicente López</span><span>\$94.444</span></div>', pantalla)
    assert re.search(r'<div class="renglon-flete"><span>09/03 · Burzaco</span><span>\$70.000</span></div>', pantalla)
    # la parte de Palmala (75.556) no es de esta base
    assert "$75.556" not in pantalla.split("</style>")[-1]

    texto = _texto_del_pdf(cliente.get(f"/gerencia/rentabilidad-real/exportar-pdf?{consulta}").content)
    assert "Flete $164.444" in " ".join(texto.split()) and "Burzaco $70.000" in " ".join(texto.split())
    from openpyxl import load_workbook
    hoja = load_workbook(io.BytesIO(cliente.get(f"/gerencia/rentabilidad-real/exportar-excel?{consulta}").content)).active
    celdas = [c for fila in hoja.iter_rows(values_only=True) for c in fila if c is not None]
    assert "Total flete" in celdas and 164444.45 in celdas

    # con un grupo elegido, el flete no se resta y se dice
    filtrada = cliente.get(f"/gerencia/rentabilidad-real?{consulta}&grupo=fruta").text
    assert "Flete ·" not in filtrada and "el flete no se resta: no es de ningún artículo" in filtrada
