import io

import openpyxl
import pypdfium2 as pdfium
import pytest

from core.lector_archivos import comprimir_pdf, imagenes_desde_pdf, texto_desde_excel


def _pdf_de_prueba(cantidad_paginas: int = 1) -> bytes:
    documento = pdfium.PdfDocument.new()
    for _ in range(cantidad_paginas):
        documento.new_page(200, 200)
    buffer = io.BytesIO()
    documento.save(buffer)
    documento.close()
    return buffer.getvalue()


def _excel_de_prueba(hojas: dict[str, list[list]]) -> bytes:
    libro = openpyxl.Workbook()
    primera = True
    for nombre_hoja, filas in hojas.items():
        hoja = libro.active if primera else libro.create_sheet(nombre_hoja)
        if primera:
            hoja.title = nombre_hoja
            primera = False
        for fila in filas:
            hoja.append(fila)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


# --- comprimir_pdf ---


def test_comprimir_pdf_achica_un_escaneo_y_conserva_las_paginas():
    # Un "escaneo": PDF de una página con una foto grande adentro. El
    # comprimido tiene que pesar bastante menos, seguir siendo un PDF
    # válido y conservar la cantidad de páginas.
    from PIL import Image

    imagen = Image.new("RGB", (2400, 3200))
    for x in range(0, 2400, 7):
        for y in range(0, 3200, 11):
            imagen.putpixel((x, y), ((x * 7) % 256, (y * 5) % 256, (x + y) % 256))
    buffer_jpeg = io.BytesIO()
    imagen.save(buffer_jpeg, format="JPEG", quality=95)

    documento = pdfium.PdfDocument.new()
    objeto_imagen = pdfium.PdfImage.new(documento)
    objeto_imagen.load_jpeg(io.BytesIO(buffer_jpeg.getvalue()))
    objeto_imagen.set_matrix(pdfium.PdfMatrix().scale(600, 800))
    pagina = documento.new_page(600, 800)
    pagina.insert_obj(objeto_imagen)
    pagina.gen_content()
    buffer_pdf = io.BytesIO()
    documento.save(buffer_pdf)
    documento.close()
    original = buffer_pdf.getvalue()

    comprimido = comprimir_pdf(original)

    assert comprimido.startswith(b"%PDF")
    assert len(comprimido) < len(original) / 2
    relectura = pdfium.PdfDocument(comprimido)
    assert len(relectura) == 1
    relectura.close()


def test_comprimir_pdf_invalido_devuelve_el_original_sin_romper():
    # Comprimir es una mejora, nunca un motivo para perder el archivo.
    assert comprimir_pdf(b"esto no es un PDF") == b"esto no es un PDF"


def test_comprimir_pdf_que_no_achica_devuelve_el_original():
    # Un PDF ya chiquito (páginas vacías) rearmado como imágenes pesa MÁS:
    # en ese caso se guarda el original.
    original = _pdf_de_prueba(cantidad_paginas=2)
    assert comprimir_pdf(original) == original


# --- imagenes_desde_pdf ---


def test_imagenes_desde_pdf_devuelve_una_imagen_por_pagina():
    resultado = imagenes_desde_pdf(_pdf_de_prueba(cantidad_paginas=3))

    assert len(resultado) == 3
    for imagen in resultado:
        assert imagen.startswith(b"\xff\xd8\xff")  # firma JPEG


def test_imagenes_desde_pdf_una_sola_pagina():
    resultado = imagenes_desde_pdf(_pdf_de_prueba(cantidad_paginas=1))

    assert len(resultado) == 1


def test_imagenes_desde_pdf_archivo_invalido_lanza_value_error():
    with pytest.raises(ValueError, match="No se pudo abrir el PDF"):
        imagenes_desde_pdf(b"esto no es un PDF")


# --- texto_desde_excel ---


def test_texto_desde_excel_junta_las_filas_con_valores():
    datos = _excel_de_prueba({"Precios": [["Articulo", "Precio"], ["Tomate Cherry", 500], ["Mango", 350.5]]})

    resultado = texto_desde_excel(datos)

    assert "Articulo | Precio" in resultado
    assert "Tomate Cherry | 500" in resultado
    assert "Mango | 350.5" in resultado


def test_texto_desde_excel_ignora_filas_y_celdas_vacias():
    datos = _excel_de_prueba({"Precios": [["Tomate Cherry", 500], [None, None], ["Mango", None]]})

    resultado = texto_desde_excel(datos)
    lineas = resultado.splitlines()

    assert lineas == ["Tomate Cherry | 500", "Mango"]


def test_texto_desde_excel_junta_varias_hojas():
    datos = _excel_de_prueba({"Precios": [["Tomate Cherry", 500]], "Precios2": [["Palta", 800]]})

    resultado = texto_desde_excel(datos)

    assert "Tomate Cherry | 500" in resultado
    assert "Palta | 800" in resultado


def test_texto_desde_excel_archivo_invalido_lanza_value_error():
    with pytest.raises(ValueError, match="No se pudo abrir el Excel"):
        texto_desde_excel(b"esto no es un Excel")
