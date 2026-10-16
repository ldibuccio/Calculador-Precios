"""PALMALA CRUCES (dueño, 14/10): la cuenta corriente interna con Palmala,
contra Postgres. Solo Gerencia, y SOLO en Frutamax: en la base de Palmala las
tablas no existen y la app de Palmala no las nombra.

La lectura de las fotos se reemplaza por lo que la IA habría devuelto, y la
subida a Storage por una ruta inventada: lo que se prueba es qué hace el
sistema con lo leído.

EL GALPÓN, con nombres inventados: banco 1 EJEMPLO Banco; los 13 conceptos del
esquema. Fechas fijas de 2025, que no pueden ser hoy.
"""
import io
import os
import re
import sys
from datetime import date
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.palmala_cruces import (  # noqa: E402
    crudos_de_la_planilla,
    renglones_para_revisar,
    texto_del_saldo,
)
from tests.test_administracion_reordenada import _cliente  # noqa: E402
from tests.test_cheques import d  # noqa: E402,F401
from tests.test_cheques_pantallas import _foto, _sin_estilo, m  # noqa: E402,F401


def _concepto(d, nombre) -> int:
    return next(c["id"] for c in d.conceptos_de_palmala() if c["nombre"] == nombre)


def _mov(d, fecha, cuenta, sentido, importe, concepto="Comisiones Día", **otros):
    return d.guardar_movimientos_de_palmala([{
        "fecha": date.fromisoformat(fecha), "cuenta": cuenta, "sentido": sentido, "importe": importe,
        "concepto_id": _concepto(d, concepto), **otros}], como_se_cargo="a_mano")[0]


def _gerencia(m):
    return _cliente(m, "gerencia")


def _xlsx(filas) -> bytes:
    from openpyxl import Workbook
    libro = Workbook()
    for fila in filas:
        libro.active.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _formulario(html: str, action: str) -> dict:
    """Lo que el navegador mandaría: los campos con su valor, los tildes
    marcados y la opción elegida de cada lista (los repetidos, en lista)."""
    marcado = _sin_estilo(html).split(f'action="{action}"')[1].split("</form>")[0]
    datos: dict = {}
    for etiqueta in re.findall(r"<input[^>]*>", marcado):
        nombre = re.search(r'name="([^"]*)"', etiqueta)
        tipo = (re.search(r'type="([^"]*)"', etiqueta) or [None, "text"])[1]
        if not nombre or tipo == "file" or (tipo == "checkbox" and " checked" not in etiqueta):
            continue
        valor = re.search(r'value="([^"]*)"', etiqueta)
        datos.setdefault(nombre[1], []).append(valor[1] if valor else "")
    for nombre, opciones in re.findall(r'<select name="([^"]*)"[^>]*>(.*?)</select>', marcado, re.S):
        elegida = re.search(r'<option value="([^"]*)"[^>]*? selected', opciones)
        datos[nombre] = [elegida[1] if elegida else ""]
    return {k: (v if len(v) > 1 else v[0]) for k, v in datos.items()}


def _saldos(html) -> dict:
    return dict(re.findall(r'data-saldo="([a-z_]+)">\s*<p class="cuenta">[^<]*</p>\s*<p class="texto">([^<]*)</p>',
                           _sin_estilo(html)))


# --- LOS SALDOS -------------------------------------------------------------------

def test_los_TRES_SALDOS_dicen_quien_le_debe_a_quien_y_el_ANULADO_no_suma(m, d):
    _mov(d, "2025-03-01", "pesos", "debe", 1000)
    _mov(d, "2025-03-02", "pesos", "pago", 1580.29, "Pago en efectivo")
    anulado = _mov(d, "2025-03-03", "pesos", "debe", 999999)                  # el RIVAL: no cuenta
    d.corregir_movimiento_de_palmala(anulado, {"anulado": True}, "cargado dos veces")
    _mov(d, "2025-03-01", "dolares_socios", "debe", 100)
    _mov(d, "2025-03-05", "dolares_socios", "pago", 100, "Pago en efectivo")
    _mov(d, "2025-03-01", "dolares_ndt", "debe", 3000)
    pagina = _gerencia(m).get("/gerencia/palmala-cruces")
    assert pagina.status_code == 200, pagina.text[:300]
    assert _saldos(pagina.text) == {"pesos": "Le debés a Palmala $580,29", "dolares_socios": "Están a mano",
                                    "dolares_ndt": "Palmala te debe US$3.000"}
    # El saldo acumulado de cada renglón es el de su cuenta, en rojo si es negativo.
    marcado = _sin_estilo(pagina.text)
    saldos = re.findall(r'<span class="saldo-al ([a-z]*)" data-saldo-al>saldo ([^<]*)</span>', marcado)
    assert ("rojo", "−$580,29") in saldos and ("", "$1.000") in saldos
    assert texto_del_saldo("pesos", 0) == "Están a mano"


def test_SOLO_GERENCIA_y_SOLO_EN_FRUTAMAX(m, d):
    assert 'name="clave"' in _cliente(m).get("/gerencia/palmala-cruces").text
    hub = _gerencia(m).get("/gerencia").text
    assert hub.count('href="/gerencia/palmala-cruces"') == 1
    # En la app de Palmala: el botón no está y la ruta contesta 404 SIN tocar la base.
    with patch.object(m, "ES_FRUTAMAX", False), \
            patch.dict(m.templates.env.globals, {"ES_FRUTAMAX": False}), \
            patch("app.main.buscar_movimientos_de_palmala", side_effect=AssertionError("consultó la base")), \
            patch("app.main.saldos_de_palmala", side_effect=AssertionError("consultó la base")):
        assert _gerencia(m).get("/gerencia/palmala-cruces").status_code == 404
        assert _gerencia(m).post("/gerencia/palmala-cruces/cargar/seguir", data={}).status_code == 404
        assert 'href="/gerencia/palmala-cruces"' not in _gerencia(m).get("/gerencia").text


def test_las_FOTOS_de_respaldo_nombran_la_tabla_solo_si_existe(d):
    _mov(d, "2025-03-01", "pesos", "pago", 10, "Pago con cheque", cheque_numero="1", cheque_banco="EJEMPLO Banco",
         cheque_fecha_pago=date(2025, 4, 1), foto_ruta="palmala/EJ.jpg")
    assert [(f["tipo"], f["ruta"]) for f in d.fotos_de_respaldo()] == [("palmala", "palmala/EJ.jpg")]
    # La base de Palmala, sin las tablas: la lista sale igual, sin ellas.
    conexion = d.obtener_conexion()
    with conexion.cursor() as cursor:
        cursor.execute("DROP VIEW palmala_movimientos_con_saldo; DROP TABLE palmala_correcciones; "
                       "DROP TABLE palmala_movimientos")
    conexion.commit()
    conexion.close()
    assert d.fotos_de_respaldo() == []


# --- EL LISTADO, LOS FILTROS Y LAS EXPORTACIONES -----------------------------------

def test_los_FILTROS_y_el_PDF_y_el_Excel_con_los_mismos(m, d):
    _mov(d, "2025-03-01", "pesos", "debe", 1000, detalle="EJEMPLO marzo")
    _mov(d, "2025-04-01", "pesos", "pago", 300, "Pago en efectivo")
    _mov(d, "2025-04-02", "dolares_socios", "debe", 50)
    _mov(d, "2025-05-01", "pesos", "debe", 70, "Intereses")
    consulta = "?cuenta=pesos&desde=2025-03-15&sentido=debe"
    pagina = _gerencia(m).get("/gerencia/palmala-cruces" + consulta)
    marcado = _sin_estilo(pagina.text)
    assert len(re.findall(r'data-movimiento="\d+"', marcado)) == 1                 # el de Intereses
    assert re.search(r'data-filtrado>Lo filtrado: Pesos: Palmala debe \$70</p>', marcado)
    # Los tres saldos de arriba son de TODO, no del filtro.
    assert _saldos(pagina.text)["pesos"] == "Palmala te debe $770"
    # El PDF y el Excel: los links que dibuja la pantalla, con sus filtros.
    link_excel = re.search(r'href="(/gerencia/palmala-cruces/exportar-excel\?[^"]*)"', marcado).group(1)
    link_pdf = re.search(r'href="(/gerencia/palmala-cruces/exportar-pdf\?[^"]*)"', marcado).group(1)
    excel = _gerencia(m).get(link_excel.replace("&amp;", "&"))
    import openpyxl
    hoja = openpyxl.load_workbook(io.BytesIO(excel.content)).active
    assert "Filtros: del 15/03/2025 al hoy · cuenta Pesos · Palmala debe" in hoja.cell(row=2, column=1).value
    filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=8) if fila[0].value]
    assert [(f[0], f[3], f[5]) for f in filas] == [("01/05/2025", "Intereses", 70.0)]
    pdf = _gerencia(m).get(link_pdf.replace("&amp;", "&"))
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")


# --- CARGAR (dueño, 16/10): el flujo de dos pasos está en test_carga_en_dos_pasos ---

def test_a_MANO_un_efectivo_no_lleva_cheque(m, d):
    pagina = _gerencia(m).post("/gerencia/palmala-cruces/cargar/seguir", data={"a_mano": "1"})
    datos = dict(_formulario(pagina.text, "/gerencia/palmala-cruces/revisar"), fecha_0="2025-06-01",
                 sentido_0="pago", importe_0="1500,50", concepto_id_0=str(_concepto(d, "Pago en efectivo")),
                 cheque_numero_0="9", accion="confirmar")
    efectivo = _gerencia(m).post("/gerencia/palmala-cruces/revisar", data=datos)
    assert "el concepto Pago en efectivo no lleva cheque" in efectivo.text
    assert d.buscar_movimientos_de_palmala({}) == []


def test_los_CONCEPTOS_se_agregan_y_se_dan_de_baja(m, d):
    _gerencia(m).post("/gerencia/palmala-cruces/conceptos", data={"nombre": "EJEMPLO Flete", "pide_cheque": ""})
    nuevo = _concepto(d, "EJEMPLO Flete")
    _gerencia(m).post(f"/gerencia/palmala-cruces/conceptos/{nuevo}/activo", data={"activo": ""})
    pagina = _gerencia(m).post("/gerencia/palmala-cruces/cargar/seguir", data={"a_mano": "1"})
    elegir = re.search(r'<select name="concepto_id_0"[^>]*>(.*?)</select>', _sin_estilo(pagina.text), re.S)[1]
    assert "EJEMPLO Flete" not in elegir and "Comisiones Día" in elegir
    with pytest.raises(d.PalmalaNoSePuede, match="ya no se usa"):
        _mov(d, "2025-03-01", "pesos", "debe", 1, "EJEMPLO Flete")


# --- VARIAS FOTOS DE CHEQUES --------------------------------------------------------

def test_varias_FOTOS_de_cheques_un_renglon_por_cheque_y_el_banco_nuevo_una_vez(m, d):
    cheque = lambda numero, banco, importe, tipo: {  # noqa: E731
        "fecha": None, "sentido": None, "importe": importe, "concepto": None, "detalle": None,
        "cheque": {"numero": numero, "banco": banco, "fecha_pago": "2025-08-01", "tipo": tipo}}
    leidas = [{"movimientos": [cheque("100", "BANCO DE GALICIA", 1000, "fisico"), cheque("101", "Galicia", 2000, "fisico")]},
              {"movimientos": [cheque("E-5", "EJEMPLO Banco", 500, "echeq")]}]
    archivos = [("archivos", (f"c{i}.jpg", _foto(), "image/jpeg")) for i in range(2)]
    with patch("app.main.extraer_movimientos_de_palmala", side_effect=leidas):
        pagina = _gerencia(m).post("/gerencia/palmala-cruces/cargar/seguir", data={"a_mano": "0"}, files=archivos)
    assert pagina.status_code == 200, pagina.text[:300]
    marcado = _sin_estilo(pagina.text)
    assert len(re.findall(r'data-renglon="\d+"', marcado)) == 3
    assert re.findall(r'data-falta-banco><span><strong>([^<]*)</strong> no está en la lista \((\d+) cheques?\)',
                      marcado) == [("BANCO DE GALICIA", "2")]
    datos = _formulario(pagina.text, "/gerencia/palmala-cruces/revisar")
    con, echeq = str(_concepto(d, "Pago con cheque")), str(_concepto(d, "Pago con echeq"))
    assert {datos[f"cheque_numero_{i}"]: (datos[f"concepto_id_{i}"], datos[f"sentido_{i}"], datos[f"cuenta_{i}"],
                                           datos[f"archivo_{i}"]) for i in range(3)} == {
        "100": (con, "pago", "pesos", "0"), "101": (con, "pago", "pesos", "0"), "E-5": (echeq, "pago", "pesos", "1")}
    agregado = _gerencia(m).post("/gerencia/palmala-cruces/revisar", data=dict(datos, agregar_banco="BANCO DE GALICIA"))
    assert "data-falta-banco" not in agregado.text
    rutas = iter(f"palmala/foto-{n}.jpg" for n in range(5))
    with patch("app.main.subir_foto_comanda", side_effect=lambda *a, **k: next(rutas)) as subir:
        final = _gerencia(m).post("/gerencia/palmala-cruces/revisar", follow_redirects=False,
                                  data=dict(_formulario(agregado.text, "/gerencia/palmala-cruces/revisar"),
                                            accion="confirmar"))
    assert final.status_code == 303 and "3+movimientos+cargados" in final.headers["location"]
    assert subir.call_count == 2                                                  # una por foto
    assert sorted((x["cheque_numero"], x["cheque_banco"], x["foto_ruta"], x["como_se_cargo"])
                  for x in d.buscar_movimientos_de_palmala({})) == [
        ("100", "BANCO DE GALICIA", "palmala/foto-0.jpg", "foto"), ("101", "BANCO DE GALICIA", "palmala/foto-0.jpg", "foto"),
        ("E-5", "EJEMPLO Banco", "palmala/foto-1.jpg", "foto")]


# --- UN EXCEL GRANDE ---------------------------------------------------------------

def test_un_EXCEL_de_300_movimientos_con_Debe_y_Pago_y_PARA_TODOS(m, d):
    filas = [["Fecha", "Concepto", "Detalle", "Debe", "Pagó"]] + [
        [f"{1 + i % 28:02d}/0{1 + i % 9}/2024", "Comisiones Día" if i % 2 else "pago en EFECTIVO",
         f"EJEMPLO {i}", 100 + i if i % 2 else None, None if i % 2 else 50 + i] for i in range(300)]
    filas.append(["Total", None, None, 1, 1])                                      # el pie no entra
    pagina = _gerencia(m).post("/gerencia/palmala-cruces/cargar/seguir", data={"a_mano": "0"},
                               files={"archivos": ("historico.xlsx", _xlsx(filas))})
    assert pagina.status_code == 200, pagina.text[:300]
    assert re.search(r"300 movimientos · \$[\d.]+ · <span class=\"verde\">300 bien</span>", _sin_estilo(pagina.text))
    datos = _formulario(pagina.text, "/gerencia/palmala-cruces/revisar")
    assert len(datos) > 2000                                       # más del tope de 1000 campos de Starlette
    assert (datos["sentido_0"], datos["importe_0"], datos["sentido_1"], datos["importe_1"]) == (
        "pago", "50", "debe", "101")
    # Para todos: los dos tildados pasan a Dólares NDT; el resto queda en Pesos.
    cambiado = _gerencia(m).post("/gerencia/palmala-cruces/revisar",
                                 data=dict(datos, tildado_0="1", tildado_1="1", accion="a_todos",
                                           todos_cuenta="dolares_ndt"))
    assert "Se cambió en 2 movimientos." in cambiado.text
    final = _gerencia(m).post("/gerencia/palmala-cruces/revisar", follow_redirects=False,
                              data=dict(_formulario(cambiado.text, "/gerencia/palmala-cruces/revisar"),
                                        accion="confirmar"))
    assert final.status_code == 303 and "300+movimientos+cargados" in final.headers["location"]
    guardados = d.buscar_movimientos_de_palmala({})
    assert len(guardados) == 300 and sum(1 for g in guardados if g["cuenta"] == "dolares_ndt") == 2
    assert {g["como_se_cargo"] for g in guardados} == {"excel"}


def test_la_PLANILLA_sin_entender_no_inventa_y_el_concepto_desconocido_se_marca():
    assert crudos_de_la_planilla([["a", "b"], [1, 2]]) is None
    renglones = renglones_para_revisar(
        crudos_de_la_planilla([["Fecha", "Importe", "Concepto", "Cuenta"], ["01/02/2024", "-1.500,50", "Raro", "euros"]]),
        [{"id": 1, "nombre": "Ajuste", "activo": True}], hoy=date(2025, 1, 1))
    assert [(r["importe"], r["sentido"], r["concepto_id"], r["cuenta"]) for r in renglones] == [
        ("1500.5", "pago", None, "")]


# --- CORREGIR Y ANULAR, CON HISTORIAL ------------------------------------------------

def test_CORREGIR_y_ANULAR_dejan_historial_y_cambian_el_saldo(m, d):
    movimiento = _mov(d, "2025-03-01", "pesos", "debe", 1000)
    ficha = _gerencia(m).get(f"/gerencia/palmala-cruces/{movimiento}")
    datos = _formulario(ficha.text, f"/gerencia/palmala-cruces/{movimiento}/corregir")
    sin_motivo = _gerencia(m).post(f"/gerencia/palmala-cruces/{movimiento}/corregir", data=dict(datos, importe="1200"))
    assert "Escribí por qué se corrige." in sin_motivo.text
    igual = _gerencia(m).post(f"/gerencia/palmala-cruces/{movimiento}/corregir", data=dict(datos, motivo="nada"))
    assert "No cambió nada." in igual.text                                       # 1000 es 1000.00
    _gerencia(m).post(f"/gerencia/palmala-cruces/{movimiento}/corregir",
                      data=dict(datos, importe="1200", motivo="EJEMPLO mal tipeado"))
    _gerencia(m).post(f"/gerencia/palmala-cruces/{movimiento}/anular", data={"anulado": "1", "motivo": "EJEMPLO doble"})
    ficha = _sin_estilo(_gerencia(m).get(f"/gerencia/palmala-cruces/{movimiento}").text)
    assert [h.split(" · ", 1)[1] for h in re.findall(r'<li data-correccion>([^<]*)</li>', ficha)] == [
        "importe: 1000.00 → 1200.00 · EJEMPLO mal tipeado", "anulado: False → True · EJEMPLO doble"]
    assert "data-anulado" in ficha
    assert _saldos(_gerencia(m).get("/gerencia/palmala-cruces").text)["pesos"] == "Están a mano"
    _gerencia(m).post(f"/gerencia/palmala-cruces/{movimiento}/anular", data={"anulado": "", "motivo": "EJEMPLO no era"})
    assert _saldos(_gerencia(m).get("/gerencia/palmala-cruces").text)["pesos"] == "Palmala te debe $1.200"


def test_la_BASE_frena_el_cheque_incompleto_salvo_en_el_historico(d):
    with pytest.raises(d.PalmalaNoSePuede, match="lleva número, banco y fecha de pago"):
        _mov(d, "2025-03-01", "pesos", "pago", 10, "Pago con cheque", cheque_numero="1")
    with pytest.raises(d.PalmalaNoSePuede, match="no lleva cheque"):
        _mov(d, "2025-03-01", "pesos", "pago", 10, "Pago en efectivo", cheque_numero="1")
    with pytest.raises(d.PalmalaNoSePuede, match="Decí si Palmala debe o pagó"):
        _mov(d, "2025-03-01", "pesos", "a_veces", 10)
    conexion = d.obtener_conexion()
    with conexion.cursor() as cursor:
        cursor.execute("INSERT INTO palmala_movimientos (fecha, cuenta, sentido, importe, concepto_id, cheque_numero, "
                       "como_se_cargo, fila_historico, sector) VALUES ('2021-01-01', 'pesos', 'pago', 5, %s, '9', "
                       "'historico', 9, 'gerencia')", (_concepto(d, "Pago con cheque"),))
    conexion.commit()
    conexion.close()
    assert [x["cheque_numero"] for x in d.buscar_movimientos_de_palmala({})] == ["9"]
