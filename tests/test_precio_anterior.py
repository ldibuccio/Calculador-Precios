"""El "Precio anterior" del listado de precios (dueño, 05/10), contra Postgres.

Es el precio VIGENTE AL CIERRE DEL DÍA ANTERIOR a la fecha del listado. Si
ayer valía lo mismo, las dos columnas dicen lo mismo; si ayer no tenía
precio, la celda queda vacía. La pantalla, el PDF y el Excel dicen lo mismo:
el "Nuevo precio" es lo que difiere de ese anterior.

Listado del 10/09. Los tres casos pedidos y dos rivales:

  A  cambió hoy          100 (01/09) -> 120 (10/09)   anterior 100, nuevo
  B  no cambió           200 (01/09)                  anterior 200
  C  no existía ayer     300 (10/09)                  anterior vacío, nuevo
  D  cambió hace días     50 (01/09) ->  60 (05/09)   anterior 60 (la regla
                                                      vieja decía 50)
  E  recargado igual     400 (01/09) -> 400 (10/09)   anterior 400, NO nuevo
"""
import io
import os
import re
import sys
from datetime import date
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from scripts.humo import hay_postgres, preparar_base  # noqa: E402

OBLIGATORIO = os.environ.get("HUMO_OBLIGATORIO") == "1"
FECHA = date(2026, 9, 10)

SIEMBRA = """
insert into clientes (id, nombre) overriding system value values (1, 'EJEMPLO Super');
insert into articulos (id, nombre, grupo) overriding system value
  values (1, 'EJ Anana', 'Fruta'), (2, 'EJ Banana', 'Fruta'), (3, 'EJ Cereza', 'Fruta'),
         (4, 'EJ Durazno', 'Fruta'), (5, 'EJ Enebro', 'Fruta');
insert into fichas_logistica (id, cliente_id, articulo_id, contenido_caja, unidad_venta)
  overriding system value
  values (1, 1, 1, 10, 'kilo'), (2, 1, 2, 10, 'kilo'), (3, 1, 3, 10, 'kilo'),
         (4, 1, 4, 10, 'kilo'), (5, 1, 5, 10, 'kilo');
insert into precios_venta_historial (articulo_id, cliente_id, ficha_id, precio, vigente_desde) values
  (1, 1, 1, 100, '2026-09-01'), (1, 1, 1, 120, '2026-09-10'),
  (2, 1, 2, 200, '2026-09-01'),
  (3, 1, 3, 300, '2026-09-10'),
  (4, 1, 4, 50, '2026-09-01'), (4, 1, 4, 60, '2026-09-05'),
  (5, 1, 5, 400, '2026-09-01'), (5, 1, 5, 400, '2026-09-10');
"""

ESPERADO = {  # nombre: (anterior, vigente, nuevo)
    "EJ Anana": (100.0, 120.0, True),
    "EJ Banana": (200.0, 200.0, False),
    "EJ Cereza": (None, 300.0, True),
    "EJ Durazno": (60.0, 60.0, False),
    "EJ Enebro": (400.0, 400.0, False),
}


@pytest.fixture
def cliente(monkeypatch):
    if not hay_postgres():
        if OBLIGATORIO:
            pytest.fail("HUMO_OBLIGATORIO=1 y no hay Postgres")
        pytest.skip("sin Postgres local")
    url = preparar_base()
    import psycopg2
    conexion = psycopg2.connect(url)
    with conexion.cursor() as cursor:
        cursor.execute(SIEMBRA)
    conexion.commit()
    conexion.close()
    monkeypatch.setenv("DATABASE_URL", url)
    from fastapi.testclient import TestClient

    from app.main import app
    with patch("app.main._hoy_argentina", return_value=date(2026, 9, 12)):
        yield TestClient(app, base_url="https://testserver")


def test_el_EXCEL_trae_el_precio_vigente_al_cierre_del_dia_anterior(cliente):
    import openpyxl
    respuesta = cliente.get(f"/precios/consultar/exportar-excel?cliente_id=1&fecha={FECHA.isoformat()}")
    assert respuesta.status_code == 200
    hoja = openpyxl.load_workbook(io.BytesIO(respuesta.content)).active
    filas = {f[0].value: f for f in hoja.iter_rows(min_row=3) if f[0].value in ESPERADO}

    assert set(filas) == set(ESPERADO)
    naranja = filas["EJ Anana"][2].fill.start_color.rgb
    for nombre, (anterior, vigente, nuevo) in ESPERADO.items():
        assert (filas[nombre][1].value, filas[nombre][2].value) == (anterior, vigente), nombre
        assert (filas[nombre][2].fill.start_color.rgb == naranja) is nuevo, nombre


def test_la_PANTALLA_marca_lo_mismo_y_dice_el_anterior(cliente):
    respuesta = cliente.get(f"/precios/consultar?cliente_id=1&fecha={FECHA.isoformat()}")
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    filas = re.findall(r'<div class="fila-precio ?(nueva)?">\s*<span>([^<]+)</span>(.*?)</div>', marcado, re.S)
    nuevos = {nombre: bool(clase) for clase, nombre, _ in filas}

    assert nuevos == {nombre: nuevo for nombre, (_, _, nuevo) in ESPERADO.items()}
    resto = {nombre: " ".join(re.sub(r"<[^>]+>", " ", r).split()) for _, nombre, r in filas}
    assert "Nuevo precio (antes $100)" in resto["EJ Anana"]
    assert "Nuevo precio (ayer no tenía)" in resto["EJ Cereza"]


def test_el_PDF_marca_los_mismos(cliente):
    pytest.importorskip("pypdfium2")
    import pypdfium2 as pdfium
    respuesta = cliente.get(f"/precios/consultar/exportar-pdf?cliente_id=1&fecha={FECHA.isoformat()}")
    assert respuesta.status_code == 200
    texto = "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(respuesta.content))
    renglones = [r for r in texto.splitlines() if any(n in r for n in ESPERADO)]

    marcados = {n for n in ESPERADO for r in renglones if n in r and "Nuevo precio" in r}
    assert len(renglones) == len(ESPERADO)
    assert marcados == {n for n, (_, _, nuevo) in ESPERADO.items() if nuevo}
