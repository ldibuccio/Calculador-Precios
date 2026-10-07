# -*- coding: utf-8 -*-
"""PRECIOS COTIZACIONES (dueño, 07/10): el precio de cada ficha de un cliente
NUEVO, sin ninguna operación, con el costo de lo que ya se compró.

LA CUENTA NO ES NUEVA. Es `calcular_listado_para_negociar_precios`, la de
Márgenes y Cargar precios manuales, que arranca de las fichas y de las
compras de todos (nunca de las ventas). Lo que agrega esta pantalla es la
ficha que esa cuenta deja afuera —sin compras en los últimos 15 días—, que
acá sale "sin costo" en vez de desaparecer.

Contra Postgres, como el humo. El GALPÓN:

  EJEMPLO Nuevo (cliente 1): utilidad 20%, IVA +10%, flete −10%, sin una venta.
    ficha 1 Tomate: kilo, caja de 10 kg ($120).
      compras del día -3: 10 cajones a $1600 (160 kg) y una SIN PRECIO;
      del día -4: 10 cajones a $2080 (160 kg) → entra (la ventana es F1 y el
      día anterior); del día -6: $9999 → RIVAL, fuera de la ventana.
      Costo = (16000 + 20800) / 320 = $115 por kilo; envase 120/10 = $12;
      sugerido = (115 × 1,20 + 12) / (1 + 0,10 − 0,10) = $150.
    ficha 2 Zapallo: su única compra es de hace 30 días → "sin costo".
  RIVAL cliente 2: utilidad 90% y una ficha de Tomate. No puede mezclarse.
  RIVAL de vigencia: una utilidad del 50% para el cliente 1 que rige en 5 días.

EL INTERRUPTOR ES EL DEL HUMO: sin Postgres se saltea; con
HUMO_OBLIGATORIO=1 dejar de poder correrlo FALLA.
"""
import io
import os
import re
import sys
from unittest.mock import patch

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from scripts.humo import hay_postgres, preparar_base  # noqa: E402

OBLIGATORIO = os.environ.get("HUMO_OBLIGATORIO") == "1"

SIEMBRA = """
insert into clientes (id, nombre) overriding system value values (1, 'EJEMPLO Nuevo'), (2, 'EJEMPLO Rival');
insert into articulos (id, nombre, grupo) overriding system value
  values (1, 'EJEMPLO Tomate', 'hortaliza'), (2, 'EJEMPLO Zapallo', 'hortaliza');
insert into proveedores (id, nombre, codigo_puesto) overriding system value
  values (1, 'EJEMPLO Puesto Uno', 'N01P01'), (2, 'EJEMPLO Puesto Dos', 'N01P02');
insert into envases (id, nombre) overriding system value values (1, 'EJEMPLO Caja');
insert into envases_costo_historial (envase_id, costo, vigente_desde) values (1, 120, current_date - 60);
insert into fichas_logistica (id, cliente_id, articulo_id, contenido_caja, unidad_venta, envase_id)
  overriding system value values (1, 1, 1, 10, 'kilo', 1), (2, 1, 2, 10, 'kilo', 1), (3, 2, 1, 10, 'kilo', 1);
insert into clientes_parametros_historial (cliente_id, nombre_parametro, valor, vigente_desde, tipo) values
  (1, 'utilidad_objetivo', 0.20, current_date - 30, 'utilidad'),
  (1, 'IVA', 0.10, current_date - 30, 'suma'),
  (1, 'flete', 0.10, current_date - 30, 'resta'),
  (1, 'utilidad_objetivo', 0.50, current_date + 5, 'utilidad'),
  (2, 'utilidad_objetivo', 0.90, current_date - 30, 'utilidad');
insert into compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
                     contenido_por_cajon, cantidad_kilos, importe, estado)
  overriding system value values
  (1, 1, 1, current_date - 3, 10, 16, 160, 1600, 'recepcionado'),
  (2, 2, 1, current_date - 4, 10, 16, 160, 2080, 'recepcionado'),
  (3, 1, 1, current_date - 6, 10, 16, 160, 9999, 'recepcionado'),
  (4, 2, 1, current_date - 3,  5, 16,  80, null, 'recepcionado'),
  (5, 1, 2, current_date - 30, 10, 20, 200, 3000, 'recepcionado');
"""


@pytest.fixture(scope="module")
def base():
    if not hay_postgres():
        if OBLIGATORIO:
            pytest.fail("HUMO_OBLIGATORIO=1 y no hay un Postgres usable: Precios Cotizaciones no se verificó.")
        pytest.skip("sin Postgres local: se corre con el humo antes de desplegar")
    url = preparar_base()
    if url is None:
        pytest.fail("no se pudo preparar la base con db/esquema_completo.sql")
    import psycopg2
    conexion = psycopg2.connect(url)
    with conexion.cursor() as cursor:
        cursor.execute(SIEMBRA)
    conexion.commit()
    conexion.close()

    def sql(consulta, parametros=None):
        con = psycopg2.connect(url)
        try:
            with con.cursor() as cursor:
                cursor.execute(consulta, parametros)
                filas = cursor.fetchall() if cursor.description else None
            con.commit()
            return filas
        finally:
            con.close()

    with patch.dict(os.environ, {"DATABASE_URL": url}):
        import app.main as m
        yield m, sql


def _filas(m):
    from app.costeo import calcular_precios_sugeridos
    return {f["ficha_id"]: f for f in calcular_precios_sugeridos(1)}


def test_la_CUENTA_es_la_de_MARGENES_y_sale_de_las_compras_de_la_ventana(base):
    m, _ = base
    from app.costeo import calcular_listado_para_negociar_precios
    tomate = _filas(m)[1]
    assert tomate["costo_actual"] == pytest.approx(115)
    assert tomate["costo_envase_unidad_venta"] == pytest.approx(12)
    assert tomate["precio_sugerido"] == pytest.approx(150)
    # La MISMA cuenta que Márgenes, no otra que da lo mismo hoy.
    de_margenes = {f["ficha_id"]: f for f in calcular_listado_para_negociar_precios(1)}
    assert tomate["precio_sugerido"] == de_margenes[1]["precio_sugerido"]
    # De qué compras sale: las de la ventana, con la que no entra marcada; el
    # RIVAL del día -6 no aparece.
    compras = [(c["compra_id"], c["entra"], c["motivo"], c["proveedor_nombre"]) for c in tomate["compras_del_costo"]]
    assert compras == [(2, True, None, "EJEMPLO Puesto Dos"), (1, True, None, "EJEMPLO Puesto Uno"),
                       (4, False, "sin precio de compra", "EJEMPLO Puesto Dos")]


def test_una_ficha_SIN_COMPRAS_RECIENTES_sale_SIN_COSTO_en_vez_de_desaparecer(base):
    m, _ = base
    from app.costeo import calcular_listado_para_negociar_precios
    filas = _filas(m)
    assert set(filas) == {1, 2}                                   # el RIVAL cliente 2 no entra
    zapallo = filas[2]
    assert (zapallo["costo_actual"], zapallo["precio_sugerido"]) == (None, None)
    assert zapallo["sin_costo"] == "sin compras en los últimos 15 días"
    # Es justo lo que Márgenes deja afuera: por eso esta pantalla lo agrega.
    assert 2 not in {f["ficha_id"] for f in calcular_listado_para_negociar_precios(1)}


def test_la_PANTALLA_muestra_compras_condiciones_y_los_TRES_CASILLEROS(base):
    from fastapi.testclient import TestClient
    m, _ = base
    pagina = TestClient(m.app).get("/precios/cotizaciones", params={"cliente_id": 1})
    assert pagina.status_code == 200
    # El botón, la pestaña y la barra se llaman igual (dueño, 07/10).
    assert pagina.text.count('<div class="barra-titulo">Precios Cotizaciones</div>') == 1
    assert pagina.text.count('<title>Precios Cotizaciones</title>') == 1
    marcado = pagina.text.split("</style>")[-1]
    assert "Utilidad 20% sobre la mercadería · IVA +10% · flete −10%" in marcado
    tomate = marcado.split('data-ficha="1"')[1].split('data-ficha="2"')[0]
    # El cajón de las compras: (1600 + 2080) / 2 = $1840, 16 kilos -> $115 el kilo.
    assert re.search(r'name="costo_1"[^>]*value="1840"', tomate)
    assert re.search(r'name="cantidad_1"[^>]*value="16"', tomate)
    assert "¿Cuántos kilos trae el cajón?" in tomate
    assert "Costo: <strong data-costo-unidad>$115</strong> por kilo" in tomate
    assert "Envase: <span data-envase-unidad>$12</span> por kilo" in tomate
    assert "Precio sugerido: <span data-sugerido>$150</span> por kilo" in tomate
    assert re.search(r'name="pendiente_precio_1"[^>]*value="150"', tomate)
    assert "compra #2 del " in tomate and "compra #1 del " in tomate
    assert "EJEMPLO Puesto Dos (no entra: sin precio de compra)" in tomate
    assert "$9.999" not in marcado and "compra #3 " not in marcado     # el RIVAL fuera de la ventana
    # Sin compra: los tres casilleros arrancan VACÍOS y se pueden cargar.
    zapallo = marcado.split('data-ficha="2"')[1]
    assert "Sin costo: sin compras en los últimos 15 días" in zapallo
    for campo in ("costo_2", "cantidad_2", "pendiente_precio_2"):
        assert re.search(r'name="%s"[^>]*value=""' % campo, zapallo), campo
    for jerga in ("ficha_id", "costo_actual", "utilidad_objetivo", "None"):
        visible = re.sub(r"<[^>]+>", " ", re.sub(r"<script.*?</script>", " ", marcado, flags=re.S))
        assert jerga not in visible, jerga


def _calcular(m, **datos):
    from fastapi.testclient import TestClient
    respuesta = TestClient(m.app).post("/precios/cotizaciones/calcular", data={"cliente_id": "1", **datos})
    assert respuesta.status_code == 200, respuesta.text
    return respuesta.json()


def test_COSTO_A_MANO_en_un_articulo_SIN_COMPRA_da_su_sugerido(base):
    """Zapallo no tiene compra: $3000 el cajón de 20 kilos -> $150 el kilo,
    envase 120/10 = $12, sugerido = (150 × 1,20 + 12) / 1 = $192."""
    m, sql = base
    antes = sql("select count(*), sum(importe) from compras")
    respuesta = _calcular(m, ficha_id="2", costo="3000", cantidad="20")
    assert respuesta == {"precio_sugerido": 192, "texto": "$192", "costo_unidad": "$150", "envase_unidad": "$12"}
    # Vive solo en la cotización: no toca compras ni precios.
    assert sql("select count(*), sum(importe) from compras") == antes
    assert sql("select count(*) from precios_venta_historial") == [(0,)]


def test_KILOS_EDITADOS_recalculan_con_la_MISMA_cuenta(base):
    m, _ = base
    from app.costeo import calcular_listado_para_negociar_precios
    # Con el cajón de la compra, da EXACTO lo de Márgenes: es la misma cuenta.
    de_margenes = {f["ficha_id"]: f for f in calcular_listado_para_negociar_precios(1)}[1]["precio_sugerido"]
    assert _calcular(m, ficha_id="1", costo="1840", cantidad="16")["precio_sugerido"] == round(de_margenes) == 150
    # 20 kilos en vez de 16: $92 el kilo -> (92 × 1,20 + 12) / 1 = $122,40.
    assert _calcular(m, ficha_id="1", costo="1840", cantidad="20") == {
        "precio_sugerido": 122, "texto": "$122", "costo_unidad": "$92", "envase_unidad": "$12"}
    # Sin uno de los dos números no hay cuenta, y no inventa.
    assert _calcular(m, ficha_id="1", costo="1840", cantidad="")["texto"] == ""
    # Una ficha de OTRO cliente no se calcula con estas condiciones.
    from fastapi.testclient import TestClient
    otra = TestClient(m.app).post("/precios/cotizaciones/calcular",
                                  data={"cliente_id": "1", "ficha_id": "3", "costo": "1840", "cantidad": "16"})
    assert otra.status_code == 404


def test_el_PRECIO_A_MANO_no_se_pisa_y_el_COSTO_A_MANO_se_marca(base):
    """En un navegador de verdad: el recálculo es de la pantalla."""
    pytest.importorskip("playwright", reason="lo que hace la pantalla lo decide el navegador")
    from fastapi.testclient import TestClient
    from playwright.sync_api import sync_playwright
    from scripts.medir_layout import CHROMIUM
    m, _ = base
    cliente = TestClient(m.app)
    origen = "http://cotizacion.test"

    def atender(ruta, pedido):
        camino = pedido.url[len(origen):]
        if pedido.method == "POST":
            r = cliente.post(camino, content=pedido.post_data_buffer,
                             headers={"content-type": pedido.headers.get("content-type", "")})
        else:
            r = cliente.get(camino)
        ruta.fulfill(status=r.status_code, body=r.content,
                     headers={"content-type": r.headers.get("content-type", "text/html")})

    with sync_playwright() as pw:
        navegador = pw.chromium.launch(executable_path=CHROMIUM)
        try:
            pagina = navegador.new_page(viewport={"width": 360, "height": 800})
            pagina.route(origen + "/**", atender)
            pagina.goto(origen + "/precios/cotizaciones?cliente_id=1")

            def fila(n):
                return pagina.locator('.ficha[data-ficha="%d"]' % n)

            def esperar_recalculo(n, vez):
                pagina.wait_for_function(
                    "([n, v]) => document.querySelector(`.ficha[data-ficha='${n}']`).dataset.recalculado === v",
                    arg=[n, str(vez)])

            def marca_visible(n):
                return fila(n).locator("[data-costo-a-mano]").evaluate("e => getComputedStyle(e).display !== 'none'")

            tomate = fila(1)
            assert marca_visible(1) is False
            # El precio se toca a mano y DESPUÉS los kilos: el precio no se pisa.
            tomate.locator('[data-campo="precio"]').fill("170")
            tomate.locator('[data-campo="cantidad"]').fill("20")
            esperar_recalculo(1, 1)
            assert tomate.locator("[data-sugerido]").inner_text() == "$122"
            assert tomate.locator('[data-campo="precio"]').input_value() == "170"
            assert marca_visible(1) is True
            # Vuelve al cajón de la compra: la marca se va.
            tomate.locator('[data-campo="cantidad"]').fill("16")
            esperar_recalculo(1, 2)
            assert marca_visible(1) is False
            assert tomate.locator("[data-sugerido]").inner_text() == "$150"
            assert tomate.locator('[data-campo="precio"]').input_value() == "170"

            # Sin compra y SIN tocar el precio: el precio sigue al sugerido.
            zapallo = fila(2)
            zapallo.locator('[data-campo="costo"]').fill("3000")
            zapallo.locator('[data-campo="cantidad"]').fill("20")
            esperar_recalculo(2, 1)
            assert zapallo.locator("[data-sugerido]").inner_text() == "$192"
            assert zapallo.locator('[data-campo="precio"]').input_value() == "192"
            assert marca_visible(2) is True
        finally:
            navegador.close()


def _texto_del_pdf(contenido: bytes) -> str:
    import pypdfium2 as pdfium
    return "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(contenido))


# Lo que es de adentro y no puede llegarle al cliente.
DATOS_INTERNOS = ("$115", "$1.840", "$12", "Puesto", "compra", "Costo", "costo", "IVA", "flete", "Utilidad",
                  "a mano", "Envase:", "sugerido")


def test_PDF_y_EXCEL_son_el_LISTADO_PARA_EL_CLIENTE_sin_datos_internos_ni_renglones_en_0(base):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    m, _ = base
    cliente = TestClient(m.app)
    formulario = {"cliente_id": "1", "costo_1": "1840", "cantidad_1": "20", "pendiente_precio_1": "165",
                  "costo_2": "3000", "cantidad_2": "20", "pendiente_precio_2": "0"}
    excel = cliente.post("/precios/cotizaciones/exportar-excel", data=formulario)
    assert excel.status_code == 200
    hoja = load_workbook(io.BytesIO(excel.content)).active
    celdas = [tuple(c for c in fila if c is not None) for fila in hoja.iter_rows(values_only=True)]
    celdas = [fila for fila in celdas if fila]
    assert hoja.title == "Precios Cotizaciones"
    assert celdas[0] == ("Precios Cotizaciones",)
    assert re.fullmatch(r"Cliente: EJEMPLO Nuevo · Fecha: \d\d/\d\d/\d{4}", celdas[1][0])
    assert celdas[2:] == [("Artículo", "Presentación", "Precio"),
                          ("EJEMPLO Tomate", "EJEMPLO Caja de 10 kilos", "$165 por kilo")]   # Zapallo en 0: afuera
    todo = " ".join(str(c) for fila in celdas for c in fila)
    for interno in DATOS_INTERNOS:
        assert interno not in todo, interno

    formulario["pendiente_precio_2"] = ""                                   # vacío: también afuera
    pdf = cliente.post("/precios/cotizaciones/exportar-pdf", data=formulario)
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    texto = " ".join(_texto_del_pdf(pdf.content).split())
    assert texto.startswith("Precios Cotizaciones"), texto[:60]
    for esperado in ("Cliente: EJEMPLO Nuevo", "EJEMPLO Tomate", "EJEMPLO Caja de 10 kilos", "$165 por kilo"):
        assert esperado in texto, esperado
    assert "Zapallo" not in texto
    for interno in DATOS_INTERNOS:
        assert interno not in texto, interno


def test_GUARDAR_los_deja_como_precios_vigentes_DESDE_HOY_y_solo_los_que_valen(base):
    from fastapi.testclient import TestClient
    m, sql = base
    compras_antes = sql("select count(*), sum(importe) from compras")
    respuesta = TestClient(m.app).post(
        "/precios/cotizaciones/guardar",
        data={"cliente_id": "1", "costo_1": "1840", "cantidad_1": "20", "pendiente_precio_1": "150",
              "costo_2": "3000", "cantidad_2": "20", "pendiente_precio_2": "0"},
        follow_redirects=False)
    try:
        assert respuesta.status_code == 303
        assert respuesta.headers["location"].startswith("/precios/cotizaciones?cliente_id=1&aviso=Se+guardaron+1+precio+de")
        hoy = m._hoy_argentina()
        assert sql("select cliente_id, ficha_id, articulo_id, precio::float, vigente_desde "
                   "from precios_venta_historial") == [(1, 1, 1, 150.0, hoy)]
        from app.db import listar_precios_vigentes_por_cliente
        assert [(p["ficha_id"], float(p["precio"])) for p in listar_precios_vigentes_por_cliente(1, hoy)] == [(1, 150.0)]
        # El costo y los kilos de la cotización no tocaron las compras.
        assert sql("select count(*), sum(importe) from compras") == compras_antes
    finally:
        sql("delete from precios_venta_historial")
