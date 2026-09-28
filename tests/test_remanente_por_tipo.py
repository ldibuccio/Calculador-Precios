"""El Stock del Depósito filtrado por tipo de mercadería (dueño, 28/09).

Lo que cuidan, en orden de cuánto cuesta que se rompa:

1. LOS TRES TIPOS SALEN DE UN SOLO LUGAR (core/remanente_por_tipo.py): la
   pantalla, el Excel y el PDF filtran con la misma función. Si cada uno
   decidiera qué es "segunda", el papel que se baja al depósito diría otra
   cosa que la pantalla.
2. LO QUE SE EXPORTA ES LO FILTRADO, con el filtro en el título.
3. EL TOTAL DEL TIPO VA ARRIBA y con signo; los negativos en rojo.
4. EL PDF TRAE LA COLUMNA "Contado" VACÍA para anotar a mano.

El fixture es el de test_app (REMANENTE_FILAS): Mandarina tiene sueltos,
caja y segunda, así que los tres tipos salen del MISMO artículo — un filtro
que filtrara por artículo en vez de por porción no pasa.
"""
import re
from io import BytesIO
from unittest.mock import patch

import pypdfium2 as pdfium
from openpyxl import load_workbook

import tests.test_app as ta
from core import exportar_remanente_pdf
from core import remanente_por_tipo as rt

URL = "/administracion/stock/remanente"

# Sueltos: Mandarina 20 (35 − 15 en caja), Pomelo 16, Mzn Gob 20 → 56.
# Segunda: Mandarina 3. Procesada: las tres cajas, 15 + 15 + 20 → 50.
TOTAL_SUELTA, TOTAL_SEGUNDA, TOTAL_PROCESADA = 56, 3, 50


def _porcion(**cambios):
    base = {"nombre": "EJEMPLO", "bultos": 1.0, "ficha_id": None, "es_segunda": False}
    base.update(cambios)
    return base


# --- 1. el filtro puro ------------------------------------------------------

def test_el_tipo_sale_de_la_CLAVE_de_la_porcion():
    assert rt.tipo_de_porcion(_porcion()) == "suelta"
    assert rt.tipo_de_porcion(_porcion(ficha_id=7)) == "procesada"
    assert rt.tipo_de_porcion(_porcion(es_segunda=True)) == "segunda"


def test_Todo_no_es_un_tipo_y_los_TRES_juntos_son_Todo():
    assert rt.tipos_elegidos([]) == ()
    assert rt.tipos_elegidos(["todo"]) == ()
    assert rt.tipos_elegidos(["suelta", "segunda", "procesada"]) == ()
    # Tildó Segunda sin destildar Todo: gana lo que marcó a propósito.
    assert rt.tipos_elegidos(["todo", "segunda"]) == ("segunda",)
    # El orden es el fijo, no el de la URL.
    assert rt.tipos_elegidos(["segunda", "suelta"]) == ("suelta", "segunda")
    # Un valor que no es un tipo no deja la pantalla vacía.
    assert rt.tipos_elegidos(["cualquiera"]) == ()


def test_las_secciones_llevan_su_total_CON_SIGNO_y_la_vacia_sale_igual():
    porciones = [_porcion(bultos=5.0, es_segunda=True), _porcion(bultos=-7.0, es_segunda=True),
                 _porcion(bultos=4.0)]
    secciones = rt.secciones(porciones, ("segunda", "procesada"))
    assert [(s["clave"], s["total"], len(s["porciones"])) for s in secciones] == [
        ("segunda", -2.0, 2), ("procesada", 0, 0)]


def test_el_nombre_del_filtro():
    assert rt.nombre_del_filtro(()) == "Todo"
    assert rt.nombre_del_filtro(("suelta", "segunda")) == "Suelta + Segunda"
    assert rt.para_el_archivo(("suelta", "segunda")) == "Suelta_Segunda"


# --- 2. la pantalla ---------------------------------------------------------

def _marcado(url):
    return ta._remanente(url=url).text.split("</style>")[-1]


def _secciones_en_pantalla(marcado):
    """[(tipo, total, [nombres])] de lo que dibujó la pantalla."""
    resultado = []
    for tipo, cuerpo in re.findall(r'<section class="seccion-tipo" data-tipo="(\w+)">(.*?)</section>',
                                   marcado, re.S):
        total = re.search(r'<p class="total-tipo">(.*?)</p>', cuerpo, re.S).group(1)
        total = re.sub(r"<[^>]+>", "", total).strip()
        nombres = [n for n, _ in ta._porciones_en_pantalla(cuerpo)]
        resultado.append((tipo, total, nombres))
    return resultado


def test_Segunda_muestra_SOLO_la_segunda_con_su_total_arriba():
    secciones = _secciones_en_pantalla(_marcado(f"{URL}?tipo=segunda"))
    assert secciones == [("segunda", f"Segunda: {TOTAL_SEGUNDA} bultos en total",
                          ["Mandarina Segunda"])]


def test_se_marcan_VARIOS_a_la_vez_y_salen_en_el_orden_fijo():
    marcado = _marcado(f"{URL}?tipo=segunda&tipo=suelta")
    secciones = _secciones_en_pantalla(marcado)
    assert [(t, total) for t, total, _ in secciones] == [
        ("suelta", f"Suelta: {TOTAL_SUELTA} bultos en total"),
        ("segunda", f"Segunda: {TOTAL_SEGUNDA} bultos en total")]
    # Denominador: las sueltas son las tres, ninguna caja se coló.
    assert secciones[0][2] == ["Mandarina", "Mzn Gob", "Pomelo"]
    assert "Caja" not in " ".join(secciones[0][2] + secciones[1][2])
    # Los dos tildes vuelven marcados y Todo no.
    assert 'value="suelta" checked' in marcado and 'value="segunda" checked' in marcado
    assert 'value="todo" checked' not in marcado
    assert 'value="procesada" checked' not in marcado


def test_Procesada_son_las_cajas_armadas():
    secciones = _secciones_en_pantalla(_marcado(f"{URL}?tipo=procesada"))
    assert len(secciones) == 1 and secciones[0][1] == f"Procesada: {TOTAL_PROCESADA} bultos en total"
    assert len(secciones[0][2]) == 3


def test_Todo_es_la_lista_de_siempre_sin_secciones():
    marcado = _marcado(URL)
    assert 'class="seccion-tipo"' not in marcado
    assert len(ta._porciones_en_pantalla(marcado)) == 7
    assert 'value="todo" checked' in marcado


def test_los_exportar_llevan_la_fecha_Y_el_filtro():
    marcado = _marcado(f"{URL}?tipo=segunda&tipo=suelta&fecha=2026-09-06")
    for ruta in ("exportar-excel", "exportar-pdf"):
        assert (f'href="/administracion/stock/remanente/{ruta}?fecha=2026-09-06'
                f'&amp;tipo=suelta&amp;tipo=segunda"') in marcado


def test_el_negativo_del_tipo_va_en_ROJO_en_el_renglon_Y_en_el_total():
    filas = [{"articulo_id": 1, "nombre": "EJEMPLO Palta", "stock": 5.0, "segunda": -2.0,
              "grupo": "fruta"}]
    marcado = ta._remanente(filas=filas, cajas={}, url=f"{URL}?tipo=segunda").text.split("</style>")[-1]
    seccion = re.search(r'<section class="seccion-tipo" data-tipo="segunda">(.*?)</section>',
                        marcado, re.S).group(1)
    assert '<p class="total-tipo">Segunda: <span class="numero-negativo">-2</span>' in seccion
    assert ta._porciones_en_pantalla(seccion) == [("EJEMPLO Palta Segunda", "-2")]


# --- 3. el Excel ------------------------------------------------------------

def test_el_Excel_exporta_LO_FILTRADO_con_el_filtro_en_el_titulo():
    respuesta = ta._remanente(url=f"{URL}/exportar-excel?tipo=segunda")
    hoja = load_workbook(BytesIO(respuesta.content)).active
    assert hoja["A1"].value == "Stock del Depósito — Segunda"
    leido = ta._leer_excel_remanente(url=f"{URL}/exportar-excel?tipo=segunda")
    assert [n for n, _b, _c in leido["porciones"]] == ["Mandarina Segunda"]
    assert leido["total"][1] == TOTAL_SEGUNDA
    assert 'filename="Stock_del_Deposito_Segunda_' in respuesta.headers["content-disposition"]


def test_el_Excel_con_Todo_trae_todo_y_el_nombre_de_siempre():
    respuesta = ta._remanente(url=f"{URL}/exportar-excel")
    assert load_workbook(BytesIO(respuesta.content)).active["A1"].value == "Stock del Depósito — Todo"
    assert 'filename="Stock_del_Deposito_06_09_2026.xlsx"' in respuesta.headers["content-disposition"]
    assert len(ta._leer_excel_remanente()["porciones"]) == 7


# --- 4. el PDF --------------------------------------------------------------

def _texto_pdf(contenido):
    documento = pdfium.PdfDocument(contenido)
    return "\n".join(pagina.get_textpage().get_text_range() for pagina in documento)


def test_el_PDF_exporta_lo_filtrado_con_el_filtro_en_el_titulo_y_la_columna_Contado():
    respuesta = ta._remanente(url=f"{URL}/exportar-pdf?tipo=segunda&tipo=suelta")
    assert respuesta.status_code == 200
    assert respuesta.headers["content-type"] == "application/pdf"
    assert 'inline; filename="Stock_del_Deposito_Suelta_Segunda_06_09_2026.pdf"' in \
        respuesta.headers["content-disposition"]
    texto = _texto_pdf(respuesta.content)
    assert "Stock del Depósito — Suelta + Segunda" in texto
    assert f"Suelta: {TOTAL_SUELTA} bultos en total" in texto
    assert f"Segunda: {TOTAL_SEGUNDA} bultos en total" in texto
    assert "Contado" in texto
    assert "Mandarina Segunda" in texto and "Pomelo" in texto
    # Lo que no se pidió no está: ni la sección ni las cajas.
    assert "Procesada" not in texto
    assert "Caja" not in texto


def test_el_PDF_con_Todo_trae_los_TRES_tipos_con_su_total():
    texto = _texto_pdf(ta._remanente(url=f"{URL}/exportar-pdf").content)
    assert "Stock del Depósito — Todo" in texto
    for rotulo, total in (("Suelta", TOTAL_SUELTA), ("Segunda", TOTAL_SEGUNDA),
                          ("Procesada", TOTAL_PROCESADA)):
        assert f"{rotulo}: {total} bultos en total" in texto


def test_en_el_PDF_la_columna_Contado_va_VACIA_y_el_negativo_en_ROJO():
    """Se leen los Paragraph que arma el generador: el color no sale del texto del PDF."""
    escritos = []
    original = exportar_remanente_pdf.Paragraph

    def registrar(texto, estilo):
        escritos.append(texto)
        return original(texto, estilo)

    secciones = rt.secciones([_porcion(nombre="EJEMPLO Palta de segunda", bultos=-2.0,
                                       es_segunda=True),
                              _porcion(nombre="EJEMPLO Lima de segunda", bultos=4.0,
                                       es_segunda=True)], ("segunda",))
    import core.exportar_vacios as vacios
    # El título de la tabla lo arma _tabla_seccion_pdf, en otro módulo.
    with patch.object(exportar_remanente_pdf, "Paragraph", side_effect=registrar), \
            patch.object(vacios, "Paragraph", side_effect=registrar):
        exportar_remanente_pdf.generar_pdf_remanente(
            ta.date(2026, 9, 28), secciones, "Segunda", lambda n: f"{float(n):g}")
    rojo = exportar_remanente_pdf.ROJO_NEGATIVO_HEX
    assert f"<font color='{rojo}'><b>-2</b></font>" in escritos  # el renglón
    assert "4" in escritos and f"<font color='{rojo}'><b>4</b></font>" not in escritos
    # El total es 2, positivo: no va en rojo.
    assert any(t.startswith("Segunda: 2 bultos en total") for t in escritos)


# --- 5. en el navegador: el efecto, no el atributo (corolario 32) ------------

def test_en_el_NAVEGADOR_los_tildes_miden_44_no_desbordan_y_Todo_excluye_a_los_tipos():
    import pytest
    pytest.importorskip("playwright", reason="lo que se ve lo decide el navegador")
    from playwright.sync_api import sync_playwright
    from scripts.medir_layout import CHROMIUM

    html = ta._remanente(url=f"{URL}?tipo=segunda").text
    with sync_playwright() as p:
        navegador = p.chromium.launch(executable_path=CHROMIUM)
        pagina = navegador.new_page(viewport={"width": 390, "height": 844})
        pagina.set_content(html)
        medido = pagina.evaluate("""() => ({
            opciones: [...document.querySelectorAll('.opcion-tipo')].map(o => o.getBoundingClientRect().height),
            desborde: document.documentElement.scrollWidth - document.documentElement.clientWidth,
        })""")

        def tildados():
            return pagina.evaluate("""() => [...document.querySelectorAll('.opciones-tipo input')]
                .filter(i => i.checked).map(i => i.value)""")

        assert tildados() == ["segunda"]
        pagina.click("text=Todo")
        despues_de_todo = tildados()
        pagina.click("text=Suelta")
        despues_de_suelta = tildados()
        navegador.close()
    assert len(medido["opciones"]) == 4 and min(medido["opciones"]) >= 44
    assert medido["desborde"] == 0
    assert despues_de_todo == ["todo"]
    assert despues_de_suelta == ["suelta"]
