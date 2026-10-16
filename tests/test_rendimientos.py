# -*- coding: utf-8 -*-
"""RENDIMIENTOS (dueño, 13/10): Comercial → Rendimientos, contra Postgres.

Por artículo: los que pasan por guía R con la cuenta del cuadro del panel
(por guía), y los que no, con kg por cajón comprado contra kg por bulto
enviado. Al lado, el % de pérdida cargado en las fichas. PDF y Excel con los
mismos filtros.

EL GALPÓN, con nombres inventados (período del 06/10 al 12/10):

  EJEMPLO Morrón: guía R del 08/10, 10 cajones de 10 kg → 8 cajas de 10 kg y
    1 de segunda: primera 80%, segunda 10%, falta 10%, pérdida total 20%.
    Su ficha tiene 18% de pérdida cargada.
  EJEMPLO Uva: compra del 07/10, 10 cajones pesados de 7,6 kg; enviados 10
    bultos con 85 kg → 8,5 por bulto: +11,8%, se factura de más.
  EJEMPLO Pera: 10 cajones de 18,6 kg; enviados 10 bultos con 180 kg → −3,2%,
    se regala.
  PRECIOS por cajón: Morrón $30.000 (10 kg), Uva $20.000 (7,6 kg), Pera
    $18.600 (18,6 kg: $1.000 el kilo).
  MERMAS de Pera en el período: 1 bulto tirado y 2 pasados a segunda.

  RESUMEN (dueño, 15/10), ponderado por kilos:
    Guías R: entraron 100 kg, 80 primera, 10 segunda → 20% de pérdida;
      $300.000 de costo, $60.000 perdidos (20%).
    Sin procesar: Uva 10 × 7,6 = 76 kg contra 85 enviados; Pera 10 × 18,6 =
      186 contra 180, más 18,6 kg de merma y 37,2 de segunda. Entraron 317,8,
      salieron 265 → 52,8 kg, 16,6%. Al costo: Uva sobra 9 kg ($−23.684),
      Pera pierde 61,8 kg ($61.800): $38.116 de $441.800, 8,6%.
    TOTAL: 417,8 kg, 72,8 perdidos → 17,4%. El promedio SIMPLE de los dos
      renglones daría 18,3%: el test lo distingue. $98.116 de $741.800, 13,2%.

  RIVALES: una compra de Uva del 01/10 (fuera del período) de 20 kg el cajón,
    un pedido anulado y un renglón dado de baja de Pera con kilos locos, y un
    envío de Morrón (que pasa por guía R: no va en la mitad de abajo). De las
  mermas: una anulada y una fuera del período de Pera, y una del Morrón.
"""
import io
import os
import re
import sys
from datetime import date

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from core.rendimientos import armar_rendimientos, diferencia_por_bulto, texto_de_la_perdida  # noqa: E402
from scripts.humo import hay_postgres, preparar_base  # noqa: E402

OBLIGATORIO = os.environ.get("HUMO_OBLIGATORIO") == "1"

SIEMBRA = """
insert into clientes (id, nombre) overriding system value values (1, 'Día EJEMPLO');
insert into articulos (id, nombre, unidad_compra) overriding system value
  values (1, 'EJEMPLO Morron', 'kilo'), (2, 'EJEMPLO Uva', 'kilo'), (3, 'EJEMPLO Pera', 'kilo');
insert into proveedores (id, nombre, codigo_puesto) overriding system value
  values (1, 'EJEMPLO Puesto', 'N01P01');
insert into envases (id, nombre) overriding system value values (1, 'EJEMPLO Caja');
insert into fichas_logistica (id, cliente_id, articulo_id, contenido_caja, unidad_venta, envase_id)
  overriding system value values (1, 1, 1, 10, 'kilo', 1), (2, 1, 2, 8, 'kilo', null), (3, 1, 3, 18, 'kilo', null);
insert into fichas_perdida_historial (ficha_id, perdida, vigente_desde) values (1, 0.18, '2026-10-01');
insert into compras (id, proveedor_id, articulo_id, fecha_operacion, cantidad_cajones,
                     contenido_por_cajon, cantidad_kilos, importe, estado, procesada_el,
                     cantidad_cajones_real, contenido_por_cajon_real)
  overriding system value
  values (21, 1, 1, '2026-10-07', 10, 10, 100, 30000, 'recepcionado', '2026-10-07 18:00-03', 10, 10),
         (22, 1, 2, '2026-10-07', 10, 8, 80, 20000, 'recepcionado', '2026-10-07 18:00-03', 10, 7.6),
         (23, 1, 3, '2026-10-07', 10, 18, 180, 18600, 'recepcionado', '2026-10-07 18:00-03', 10, 18.6),
         (24, 1, 2, '2026-10-01', 10, 20, 200, 99999, 'recepcionado', '2026-10-01 18:00-03', 10, 20);
insert into reprocesos (id, articulo_id, fecha_operacion, bultos_tomados, bultos_primera,
                        bultos_segunda, bultos_merma, cliente_id, ficha_id, envase_id, lleva_caja_nuestra)
  overriding system value values (1, 1, '2026-10-08', 10, 8, 1, 0, 1, 1, 1, true);
insert into reprocesos_consumos (reproceso_id, origen, compra_id, origen_id, bultos)
  values (1, 'compra', 21, 21, 10);
insert into pedidos (id, cliente_id, fecha_operacion, origen, anulado_el) overriding system value
  values (1, 1, '2026-10-09', 'texto', null), (2, 1, '2026-10-10', 'texto', '2026-10-10 12:00-03');
insert into pedidos_sucursales (pedido_id, sucursal) values (1, 'VL'), (2, 'VL');
insert into pedidos_renglones (id, pedido_id, sucursal, articulo_id, ficha_id, cantidad,
                               cantidad_armada, armado_el, kilos_enviados, anulado_el)
  overriding system value
  values (1, 1, 'VL', 2, 2, 10, 10, '2026-10-09 10:00-03', 85, null),
         (2, 1, 'VL', 3, 3, 10, 10, '2026-10-09 10:00-03', 180, null),
         (3, 1, 'VL', 3, 3, 10, 10, '2026-10-09 10:00-03', 900, '2026-10-09 11:00-03'),
         (4, 2, 'VL', 3, 3, 10, 10, '2026-10-10 10:00-03', 900, null),
         (5, 1, 'VL', 1, 1, 8, 8, '2026-10-09 10:00-03', 80, null);
insert into movimientos_stock (articulo_id, fecha_operacion, tipo, cantidad, motivo, stock_sistema,
                               anulado_el, bultos_segunda)
  values (3, '2026-10-10', 'merma', -1, 'podrida', 0, null, null),
         (3, '2026-10-11', 'pase_a_segunda', -2, 'golpeada', 0, null, 2),
         (3, '2026-10-10', 'merma', -5, 'anulada', 0, '2026-10-10 12:00-03', null),
         (3, '2026-10-01', 'merma', -7, 'fuera del período', 0, null, null),
         (1, '2026-10-11', 'merma', -1, 'del Morrón, que pasa por guía R', 0, null, null);
"""
PERIODO = {"fecha_desde": "2026-10-06", "fecha_hasta": "2026-10-12"}


def test_el_ejemplo_de_UVA_y_PERA():
    assert diferencia_por_bulto(7.6, 8.5) == 11.8
    assert diferencia_por_bulto(18.6, 18.0) == -3.2
    datos = armar_rendimientos([], [
        {"articulo_id": 2, "articulo": "Uva", "kg_por_cajon": 7.6, "kg_por_bulto": 8.5},
        {"articulo_id": 3, "articulo": "Pera", "kg_por_cajon": 18.6, "kg_por_bulto": 18.0},
        {"articulo_id": 4, "articulo": "Kiwi", "kg_por_cajon": None, "kg_por_bulto": 9.0}], {3: [0.05]})
    assert [(f["articulo"], f["diferencia_pct"], f["que_pasa"]) for f in datos["sin_guia"]] == [
        ("Uva", 11.8, "se factura de más"), ("Pera", -3.2, "se regala"), ("Kiwi", None, None)]
    assert [texto_de_la_perdida(v) for v in ([], [0.18], [0.0, 0.125])] == ["—", "18%", "0% a 12,5%"]


@pytest.fixture
def base(monkeypatch):
    if not hay_postgres():
        if OBLIGATORIO:
            pytest.fail("HUMO_OBLIGATORIO=1 y no hay Postgres")
        pytest.skip("sin Postgres local")
    url = preparar_base()
    monkeypatch.setenv("DATABASE_URL", url)
    from tests.test_administracion_reordenada import CLAVES
    for clave, valor in CLAVES.items():
        monkeypatch.setenv(clave, valor)
    import app.db as d
    import app.main as m
    conexion = d.obtener_conexion()
    with conexion.cursor() as cursor:
        cursor.execute(SIEMBRA)
    conexion.commit()
    conexion.close()
    from tests.test_administracion_reordenada import _cliente
    return _cliente(m)


def _articulos(texto, seccion):
    parte = texto.split(f'data-seccion="{seccion}"')[1].split('data-seccion=')[0]
    return re.findall(r'<div class="articulo" data-articulo="(\d+)">\s*<h3>([^<]*)', parte), parte


def test_la_PANTALLA_con_las_dos_mitades_y_la_perdida_de_la_ficha(base):
    pagina = base.get("/comercial/rendimientos", params=PERIODO)
    assert pagina.status_code == 200
    con, parte_con = _articulos(pagina.text, "con-guia")
    assert [a for a, _ in con] == ["1"]
    assert "20,0% de pérdida total" in parte_con
    for dato in ("Entraron <strong>100 kg</strong>", "Primera <strong>80,0%</strong>",
                 "Segunda <strong>10,0%</strong>", "Merma <strong>0,0%</strong>"):
        assert dato in parte_con, dato
    assert re.search(r'Falta sin registrar <strong class="rojo">10,0%</strong>', parte_con)
    assert "En la ficha: 18%" in parte_con
    sin, parte_sin = _articulos(pagina.text, "sin-guia")
    assert [a for a, _ in sin] == ["2", "3"]                 # el Morrón no, aunque se envió
    assert "11,8% se factura de más" in parte_sin and "-3,2% se regala" in parte_sin
    assert "Cajón comprado <strong>7.6 kg</strong>" in parte_sin             # no el rival de 20 kg
    assert "Bulto enviado <strong>18 kg</strong>" in parte_sin                # sin el anulado ni la baja
    assert parte_sin.count("En la ficha: 0%") == 2


def test_el_FILTRO_por_articulo_y_las_fechas_malas(base):
    pagina = base.get("/comercial/rendimientos", params={**PERIODO, "articulo_id": "3"})
    assert [a for a, _ in _articulos(pagina.text, "sin-guia")[0]] == ["3"]
    assert _articulos(pagina.text, "con-guia")[0] == []
    assert "Del 06/10/2026 al 12/10/2026 · EJEMPLO Pera" in pagina.text
    mala = base.get("/comercial/rendimientos", params={"fecha_desde": "2026-10-12", "fecha_hasta": "2026-10-06"})
    assert mala.status_code == 400 and "La fecha desde no puede ser posterior" in mala.text


def test_PDF_y_EXCEL_salen_con_los_mismos_filtros(base):
    from openpyxl import load_workbook
    excel = base.get("/comercial/rendimientos/excel", params={**PERIODO, "articulo_id": "2"})
    assert excel.status_code == 200
    hoja = load_workbook(io.BytesIO(excel.content)).active
    valores = [[c.value for c in fila] for fila in hoja.iter_rows()]
    assert valores[1][0] == "Del 06/10/2026 al 12/10/2026 · EJEMPLO Uva"
    uva = next(f for f in valores if f[0] == "EJEMPLO Uva")
    assert uva[:5] == ["EJEMPLO Uva", "7,60", "8,50", "11,8%", "se factura de más"]
    assert not any(f[0] == "EJEMPLO Pera" for f in valores)
    pdf = base.get("/comercial/rendimientos/pdf", params=PERIODO)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    pantalla = base.get("/comercial/rendimientos", params={**PERIODO, "articulo_id": "2"}).text
    assert 'href="/comercial/rendimientos/pdf?fecha_desde=2026-10-06&amp;fecha_hasta=2026-10-12&amp;articulo_id=2"' in pantalla


def test_COMERCIAL_tiene_el_boton(base):
    hub = base.get("/comercial").text
    assert hub.count('href="/comercial/rendimientos"') == 1


@pytest.mark.parametrize("sector", ["comercial", "administracion", "compras"])
def test_ADMINISTRACION_y_COMPRAS_tienen_el_boton_y_la_pantalla_se_QUEDA_en_su_sector(base, sector):
    """Dueño, 16/10. Las cuatro mitades del camino (corolario 63): la barra,
    "Atrás", el formulario y las descargas, con el sector del prefijo."""
    import app.main as m
    from tests.test_administracion_reordenada import _cliente
    cliente = _cliente(m, *([] if sector == "comercial" else [sector]))     # Comercial no tiene clave
    hub = cliente.get(f"/{sector}").text
    assert hub.count(f'href="/{sector}/rendimientos"') == 1
    pagina = cliente.get(f"/{sector}/rendimientos", params={**PERIODO, "articulo_id": "2"})
    assert pagina.status_code == 200
    marcado = pagina.text.split("</style>")[-1]
    destinos = re.findall(r'(?:href|action)="(/[a-z]+)/rendimientos', marcado)
    assert destinos == [f"/{sector}"] * 3                       # el formulario, el PDF y el Excel
    # La barra (el sector y "Atrás") no lleva a ningún otro sector.
    # (la barra trae su propio <style>: se mira la página entera)
    assert set(re.findall(r'<a [^>]*href="/(comercial|administracion|compras)"', pagina.text)) == {sector}
    assert f'<a class="barra-boton" href="/{sector}" aria-label="Volver atrás">' in pagina.text
    consulta = "fecha_desde=2026-10-06&amp;fecha_hasta=2026-10-12&amp;articulo_id=2"
    assert f'href="/{sector}/rendimientos/pdf?{consulta}"' in marcado
    for formato, empieza in (("pdf", b"%PDF"), ("excel", b"PK")):
        archivo = cliente.get(f"/{sector}/rendimientos/{formato}", params=PERIODO)
        assert archivo.status_code == 200 and archivo.content.startswith(empieza), (sector, formato)


@pytest.mark.parametrize("sector", ["administracion", "compras"])
def test_desde_ADMINISTRACION_y_COMPRAS_pide_la_CLAVE_de_su_sector(base, sector):
    for ruta in (f"/{sector}/rendimientos", f"/{sector}/rendimientos/pdf", f"/{sector}/rendimientos/excel"):
        respuesta = base.get(ruta, params=PERIODO, follow_redirects=False)
        assert respuesta.status_code != 200 or 'name="clave"' in respuesta.text, ruta
    # Y en Comercial sigue abierta, sin clave.
    assert base.get("/comercial/rendimientos", params=PERIODO).status_code == 200


# ---------------------------------------------------------------- resumen

from core.rendimientos import resumen_de_rendimientos  # noqa: E402


def test_el_RESUMEN_puro_deja_afuera_y_cuenta_lo_que_no_tiene_dato():
    guia = {"consumos": [(10.0, 10.0)], "importes": [30000.0], "primera": 8.0, "segunda": 1.0,
            "merma": 0.0, "kg_caja": 10.0, "primera_real": None}
    sin_precio = dict(guia, importes=[None])
    sin_pesada = dict(guia, consumos=[(10.0, None)])
    r = resumen_de_rendimientos([guia, sin_precio, sin_pesada], [
        {"kg_por_cajon": 18.6, "bultos": 10.0, "kg_enviados": 180.0, "costo_por_kg": None,
         "bultos_merma": 0.0, "bultos_segunda": 0.0},
        {"kg_por_cajon": None, "bultos": 10.0, "kg_enviados": 85.0, "costo_por_kg": 1.0,
         "bultos_merma": 0.0, "bultos_segunda": 0.0}])
    g = r["guias_r"]
    assert (g["entraron"], g["perdida"], g["perdida_pct"], g["guias_sin_cuenta"]) == (200.0, 40.0, 20.0, 1)
    assert (g["costo"], g["perdida_pesos"], g["kg_sin_costo"]) == (300000.0, 60000.0, 100.0)
    s = r["sin_procesar"]
    assert (s["entraron"], s["perdida"], s["articulos_sin_dato"], s["kg_sin_costo"]) == (186.0, 6.0, 1, 186.0)
    assert s["costo"] == 0.0 and s["perdida_costo_pct"] is None
    assert r["total"]["entraron"] == 386.0 and r["total"]["kg_sin_costo"] == 286.0


def _renglones_del_resumen(texto):
    parte = texto.split("data-resumen>")[1].split('data-seccion="con-guia"')[0]
    bloques = re.split(r'<div class="renglon[^"]*" data-renglon-resumen="', parte)[1:]
    return {b.split('"')[0]: b for b in bloques}


def test_el_RESUMEN_arriba_ponderado_por_kilos(base):
    texto = base.get("/comercial/rendimientos", params=PERIODO).text
    assert texto.index("data-resumen>") < texto.index('data-seccion="con-guia"')
    r = _renglones_del_resumen(texto)
    assert list(r) == ["guias_r", "sin_procesar", "total"]
    esperado = {
        "guias_r": ("20,0%", "100", "80,0%", "10,0%", "0,0%", "10,0%", "$60.000", "20,0%"),
        "sin_procesar": ("16,6%", "318", "83,4%", "11,7%", "5,9%", "-0,9%", "$38.116", "8,6%"),
        "total": ("17,4%", "418", "82,6%", "11,3%", "4,5%", "1,7%", "$98.116", "13,2%"),
    }
    for clave, (perdida, kg, aprov, segunda, merma, falta, pesos, costo_pct) in esperado.items():
        b = r[clave]
        assert re.search(r'data-perdida-kg>' + re.escape(perdida) + " de pérdida<", b), clave
        assert f"Entraron <strong>{kg} kg</strong>" in b, clave
        assert f"Aprovechado <strong>{aprov}</strong>" in b, clave
        assert f"Segunda <strong>{segunda}</strong>" in b, clave
        assert f"Merma <strong>{merma}</strong>" in b, clave
        assert re.search(r"Falta sin registrar <strong[^>]*>" + re.escape(falta) + "<", b), clave
        assert re.search(r"data-perdida-pesos>" + re.escape(pesos) + "<", b), clave
        assert f"({costo_pct} del costo)" in b, clave
    sin = r["sin_procesar"].split("data-sin-procesar-kg>")[1]
    for dato in ("Comprado <strong>262 kg</strong>", "Enviado <strong>265 kg</strong>",
                 "Diferencia <strong>1,1%</strong>", "Merma <strong>19 kg</strong>",
                 "Segunda <strong>37 kg</strong>"):
        assert dato in sin, dato


def test_el_RESUMEN_respeta_el_filtro_por_articulo(base):
    r = _renglones_del_resumen(base.get("/comercial/rendimientos", params={**PERIODO, "articulo_id": "3"}).text)
    assert "Entraron <strong>0 kg</strong>" in r["guias_r"]
    for clave in ("sin_procesar", "total"):
        assert "Entraron <strong>242 kg</strong>" in r[clave], clave
        assert re.search(r"data-perdida-kg>25,6% de pérdida<", r[clave]), clave


def test_el_RESUMEN_sale_PRIMERO_en_el_PDF_y_el_EXCEL(base):
    from openpyxl import load_workbook
    import pypdfium2 as pdfium
    excel = base.get("/comercial/rendimientos/excel", params=PERIODO)
    valores = [[c.value for c in fila] for fila in load_workbook(io.BytesIO(excel.content)).active.iter_rows()]
    titulos = [f[0] for f in valores]
    assert titulos.index("Resumen del período") < titulos.index("Pasan por guía R")
    total = next(f for f in valores if f[0] == "TOTAL")
    assert total == ["TOTAL", "418", "82,6%", "11,3%", "4,5%", "1,7%", "17,4%", "$98.116", "13,2%"]
    guias = next(f for f in valores if f[0] == "Guías R (promedio ponderado)")
    assert guias[6:9] == ["20,0%", "$60.000", "20,0%"]
    pdf = base.get("/comercial/rendimientos/pdf", params=PERIODO)
    texto = " ".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(pdf.content))
    texto = " ".join(texto.split())
    assert texto.index("Resumen del período") < texto.index("Pasan por guía R")
    assert "TOTAL 418 82,6% 11,3% 4,5% 1,7% 17,4% $98.116 13,2%" in texto
