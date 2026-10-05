"""Resumen proveedores (dueño, 05/10): la MARCA del cajón en cada renglón; en el
celular sin el TOTAL de señas (la seña de cada compra sí); el PDF y el Excel
lo llevan todo.

Contra Postgres: la marca sale de la consulta (lo escrito por Recepción o, si
no hay texto, la marca de cajón pegada), y el renglón de la devolución la toma
de SU compra. Reusa el galpón de las devoluciones, con el RIVAL: la 12 no
tiene marca y tiene que decir "Sin marca".
"""
import io
import os
import re
import sys
from datetime import date

import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from tests.test_devoluciones_al_proveedor import base  # noqa: E402,F401

FILTROS = {"fecha_desde": "2026-09-01", "fecha_hasta": "2026-09-30", "proveedor_id": "1"}


def _cliente(monkeypatch):
    from fastapi.testclient import TestClient

    from app.main import PUERTA_ADMINISTRACION, app
    monkeypatch.setenv(PUERTA_ADMINISTRACION.env_var, "clave-adm")
    cliente = TestClient(app, base_url="https://testserver")
    cliente.cookies.set(PUERTA_ADMINISTRACION.cookie, PUERTA_ADMINISTRACION.firma("clave-adm"))
    return cliente


@pytest.fixture
def galpon(base, monkeypatch):
    """La 11 con la marca ESCRITA, la 13 (otro proveedor) solo con la marca de
    cajón pegada, y una devolución de la 11 desde depósito."""
    d, sql, _ = base
    westfalia = d.crear_marca_vacio(2, "Westfalia")
    sql("UPDATE compras SET marca = 'Camila' WHERE id = 11")
    sql("UPDATE compras SET marca_vacio_id = %s, sena = 300 WHERE id = 13", (westfalia,))
    d.crear_devolucion_deposito(11, 1, "EJ fea", date(2026, 9, 10), cargada_desde="administracion")
    return d, sql, _cliente(monkeypatch)


def _celdas_de_articulo(pagina: str, proveedor: str) -> list[str]:
    tarjeta = pagina.split("</style>")[-1].split(f'<p class="proveedor-encabezado">{proveedor}')[1]
    tarjeta = tarjeta.split('<div class="tarjeta">')[0]
    filas = re.findall(r"<tr[^>]*>(.*?)</tr>", tarjeta, re.S)[1:]
    return [re.findall(r"<td[^>]*>(.*?)</td>", f, re.S)[2] for f in filas]


def test_cada_RENGLON_dice_la_marca_del_cajon(galpon):
    _, _, cliente = galpon
    pagina = cliente.get("/administracion/ingresos/pagar", params={**FILTROS, "proveedor_id": ""}).text

    uno = [re.sub(r"<[^>]+>", " ", c).split() for c in _celdas_de_articulo(pagina, "EJ Uno")]
    # 05/09 la 11, 06/09 la 12 (el rival, sin marca), 10/09 la devolución de la 11.
    assert uno == [["EJEMPLO", "Fruta", "Marca", "Camila"], ["EJEMPLO", "Fruta", "Sin", "marca"],
                   ["EJEMPLO", "Fruta", "Marca", "Camila"]]
    dos, = _celdas_de_articulo(pagina, "EJ Dos")
    assert '<span class="marca-cajon">Marca Westfalia</span>' in dos
    assert pagina.count('<span class="marca-cajon">') == 4


def _visibles(html: str, ancho: int) -> dict:
    """Qué se ve, medido con getComputedStyle (corolario 32), al ancho dado."""
    pytest.importorskip("playwright")
    from playwright.sync_api import sync_playwright

    from scripts.medir_layout import CHROMIUM
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(executable_path=CHROMIUM)
        try:
            pagina = navegador.new_page(viewport={"width": ancho, "height": 800})
            pagina.set_content(html)
            return pagina.evaluate("""() => {
                const ve = el => getComputedStyle(el).display !== 'none' && el.offsetParent !== null;
                const senas = [...document.querySelectorAll('.desglose-senas')];
                const senaFila = [...document.querySelectorAll('tbody tr td:nth-child(6)')];
                const marcas = [...document.querySelectorAll('.marca-cajon')];
                return {senas: senas.length, senas_visibles: senas.filter(ve).length,
                        sena_fila: senaFila.length, sena_fila_visible: senaFila.filter(ve).length,
                        marcas: marcas.length, marcas_visibles: marcas.filter(ve).length,
                        desborde: document.documentElement.scrollWidth - innerWidth};
            }""")
        finally:
            navegador.close()


def test_en_el_CELULAR_no_va_el_total_de_senas_y_la_sena_de_cada_compra_SI(galpon):
    _, _, cliente = galpon
    respuesta = cliente.get("/administracion/ingresos/pagar", params=FILTROS)
    # La identidad: status y los renglones de la siembra (11, 12 y la devolución).
    assert respuesta.status_code == 200
    celular, compu = _visibles(respuesta.text, 313), _visibles(respuesta.text, 1024)

    # El desglose del subtotal y el del total: los dos con señas.
    assert celular["senas"] == compu["senas"] == 2
    assert (celular["senas_visibles"], compu["senas_visibles"]) == (0, 2)
    assert celular["sena_fila"] == celular["sena_fila_visible"] == 3
    assert celular["marcas"] == celular["marcas_visibles"] == 3
    assert celular["desborde"] == 0


def _texto_del_pdf(contenido: bytes) -> str:
    import pypdfium2 as pdfium
    return "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(contenido))


def test_el_PDF_lleva_la_marca_y_el_total_de_senas(galpon):
    _, _, cliente = galpon
    pytest.importorskip("pypdfium2")
    respuesta = cliente.get("/administracion/ingresos/exportar-pdf", params=FILTROS)
    assert respuesta.status_code == 200
    texto = " ".join(_texto_del_pdf(respuesta.content).split())

    assert texto.count("Marca Camila") == 2 and texto.count("Sin marca") == 1
    # 10 cajones × $500 que entraron, menos 1 × $500 que volvió.
    assert "señas de los cajones $4.500" in texto


def test_el_EXCEL_lleva_la_marca_y_el_total_de_senas(galpon):
    from openpyxl import load_workbook
    _, _, cliente = galpon
    respuesta = cliente.get("/administracion/ingresos/exportar-excel", params=FILTROS)
    assert respuesta.status_code == 200
    hoja = load_workbook(io.BytesIO(respuesta.content)).active
    filas = [[c.value for c in fila] for fila in hoja.iter_rows()]
    encabezado = next(f for f in filas if f[0] == "Fecha")
    columna = {nombre: i for i, nombre in enumerate(encabezado) if nombre}

    renglones = [f for f in filas if f[columna["Artículo"]] == "EJEMPLO Fruta"]
    assert [f[columna["Marca"]] for f in renglones] == ["Camila", "Sin marca", "Camila"]
    subtotal = next(f for f in filas if f[0] == "Subtotal")
    total = next(f for f in filas if f[0] == "Total a depositar")
    assert subtotal[columna["Total seña"]] == total[columna["Total seña"]] == 4500
    # Las columnas de plata siguen bajo SU encabezado después de correrlas.
    assert subtotal[columna["A depositar"]] == total[columna["A depositar"]] == 10 * 600 + 8 * 120 - 600
