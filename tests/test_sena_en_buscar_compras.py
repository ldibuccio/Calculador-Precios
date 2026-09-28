"""Buscar compras: la SEÑA en la pantalla, el Excel y el PDF (dueño, 28/09).

Las tres salidas tienen que decir lo mismo: la seña de cada compra al lado del
importe, vacía si no dejó, y el mismo total al pie. El total es seña por cajón
× cajones (la seña, como el importe, es por cajón).
"""
import re
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch

import pypdfium2 as pdfium
from openpyxl import load_workbook

import tests.test_app as ta
# La cookie de Compras, como la cruza una persona (fixture autouse de test_app).
from tests.test_app import _puerta_de_compras_abierta  # noqa: F401

# 40 cajones sin seña; 10 a $1.200; y una tercera de 4 a $300.
# Total: 10 × 1200 + 4 × 300 = $13.200.
COMPRAS = deepcopy(ta.COMPRAS_BUSQUEDA_DE_PRUEBA[:2])
COMPRAS[1]["sena"] = 1200.0
COMPRAS.append({**deepcopy(COMPRAS[1]), "id": 3, "articulo_nombre": "EJEMPLO Palta",
                "cantidad_cajones": 4, "sena": 300.0, "importe": 2000.0})
TOTAL = "$13.200"
# Importes: 40 × $45.000 + 4 × $2.000 = $1.808.000; el Mango no tiene precio.
TOTAL_IMPORTES = "Total de importes (importe por cajón × cajones): $1.808.000 · 1 sin precio, no suma"
RANGO = "fecha_desde=2026-08-01&fecha_hasta=2026-08-06"


def _pantalla():
    with (
        patch("app.main.listar_todos_los_proveedores", return_value=ta.PROVEEDORES_DE_PRUEBA),
        patch("app.main.listar_articulos", return_value=ta.ARTICULOS_CON_UNIDAD_COMPRA),
        patch("app.main.buscar_compras", return_value=deepcopy(COMPRAS)),
    ):
        respuesta = ta.cliente.get(f"/compras/buscar?{RANGO}")
    assert respuesta.status_code == 200
    return respuesta.text


def _senas_en_pantalla(html):
    marcado = html.split("</style>")[-1]
    celdas = re.findall(r'<td class="celda-sena">(.*?)</td>', marcado, re.S)
    return [re.sub(r"<[^>]+>|seña ", "", c).strip() for c in celdas]


def _excel():
    with patch("app.main.buscar_compras", return_value=deepcopy(COMPRAS)):
        contenido = ta.cliente.get(f"/compras/buscar/exportar-excel?{RANGO}").content
    hoja = load_workbook(BytesIO(contenido)).active
    filas = [[c.value for c in f] for f in hoja.iter_rows(max_col=4)]
    return filas


def _pdf():
    with patch("app.main.buscar_compras", return_value=deepcopy(COMPRAS)):
        contenido = ta.cliente.get(f"/compras/buscar/exportar-pdf?{RANGO}").content
    return "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(contenido))


def test_la_PANTALLA_muestra_la_sena_de_cada_compra_y_el_total():
    html = _pantalla()
    # Mismo orden que la tabla: sin seña va VACÍA, no "$0".
    assert _senas_en_pantalla(html) == ["", "$1.200", "$300"]
    assert f"Total de señas (seña por cajón × cajones): {TOTAL}" in html.split("</style>")[-1]


def test_el_EXCEL_trae_la_columna_Sena_al_lado_del_importe_y_el_total_al_pie():
    filas = _excel()
    encabezados = [f for f in filas if f[:4] == ["Artículo", "Cantidad", "Importe", "Seña"]]
    assert len(encabezados) == 2              # una tabla por fecha y proveedor
    por_articulo = {f[0]: f[3] for f in filas if f[0] in ("Tomate Cherry", "Mango", "EJEMPLO Palta")}
    assert por_articulo == {"Tomate Cherry": None, "Mango": 1200.0, "EJEMPLO Palta": 300.0}
    assert filas[-1][0] == "Total de señas (seña por cajón × cajones)"
    assert filas[-1][3] == 13200.0


def test_el_PDF_trae_la_columna_Sena_y_el_total_al_pie():
    texto = _pdf()
    assert texto.count("Seña") >= 2
    assert "$1.200" in texto and "$300" in texto
    assert f"Total de señas (seña por cajón × cajones): {TOTAL}" in texto


def test_las_TRES_salidas_dicen_la_MISMA_sena_por_compra():
    pantalla = dict(zip(("Tomate Cherry", "Mango", "EJEMPLO Palta"), _senas_en_pantalla(_pantalla())))
    excel = {f[0]: f[3] for f in _excel() if f[0] in pantalla}
    pdf = _pdf()
    for articulo, texto in pantalla.items():
        if texto:
            assert f"${excel[articulo]:,.0f}".replace(",", ".") == texto, articulo
            assert texto in pdf, articulo
        else:
            assert excel[articulo] is None, articulo


def test_la_columna_nueva_corrio_el_CSS_del_celular_y_no_quedo_desfasado():
    """El celular ubica cada celda por POSICIÓN (`td:nth-child`): una columna
    agregada corre todas las de la derecha. Tantas reglas como `<th>`."""
    partes = _pantalla().split("</style>")
    css = "".join(partes[:-1])
    ths = re.findall(r"<th[ >]", partes[-1][partes[-1].index("<thead>"):partes[-1].index("</thead>")])
    posiciones = {int(n) for n in re.findall(r"td:nth-child\((\d+)\)\s*\{\s*grid-area", css)}
    posiciones |= {int(n) for n in re.findall(r"td:nth-child\((\d+)\)\s*\{\s*\n\s*grid-area", css)}
    assert len(ths) == 8
    assert posiciones == set(range(1, 9))


def test_a_390px_la_sena_no_desborda():
    import pytest
    pytest.importorskip("playwright", reason="el desborde lo mide el navegador")
    from scripts.medir_layout import medir_sync
    medicion = medir_sync(_pantalla(), ancho=390)
    assert medicion["pares"] > 0
    # En tabla viene `desborde`; en tarjetas, `desborde_pagina` (corolario 53).
    assert medicion.get("desborde_pagina", medicion["desborde"]) == 0, medicion


def test_el_TOTAL_DE_IMPORTES_es_importe_por_cajones_y_dice_cuantas_no_suman():
    """El rival es sumar la columna: $47.000, que no es plata de nada."""
    assert TOTAL_IMPORTES in _pantalla().split("</style>")[-1]
    assert TOTAL_IMPORTES in _pdf()
    filas = _excel()
    fila, = [f for f in filas if str(f[0]).startswith("Total de importes")]
    assert fila[0] == TOTAL_IMPORTES.split(": $")[0] + " · 1 sin precio, no suma"
    assert fila[2] == 1808000.0


def test_sin_compras_sin_precio_el_total_no_lleva_la_cola():
    from core.exportar_compras import cola_sin_precio, total_de_importes
    con_precio = [c for c in COMPRAS if c["importe"] is not None]
    assert cola_sin_precio(con_precio) == ""
    assert total_de_importes(con_precio) == 1808000.0
    assert cola_sin_precio([{"importe": None}, {"importe": None}]) == " · 2 sin precio, no suman"
