"""Los VACÍOS DEL DEPÓSITO: los cajones del proveedor de COMPRAS, por PILA.

NO ES EL CIRCUITO DEL PUESTO, que vive en tests/test_app.py y en otras tablas
enteras. Acá el cajón llega CON la mercadería y se le devuelve al proveedor
que la vendió; allá un cliente del puesto lo trae y un proveedor del puesto lo
retira. No comparten una sola tabla.

DESDE EL 25/09 la cuenta es por PILA —proveedor y marca de cajón— y arranca
de una FOTO: el stock que el sistema mostraba ese día. Lo que la cuenta hace
con números se prueba CONTRA POSTGRES en
`tests/test_vacios_pilas_contra_la_base.py`; acá va lo que un mock SÍ puede
ver: el texto de la consulta, lo que se escribe, y las pantallas.
"""

import ast
import io
import os
import re
from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient
from fastapi.responses import HTMLResponse

from app.db import (
    COLUMNAS_PILAS_DE_VACIOS,
    _SQL_PILAS_DE_VACIOS,
    anular_ajuste_vacios_deposito,
    anular_devolucion_vacios,
    crear_devolucion_vacios,
)
from app.main import PUERTA_ADMINISTRACION, app

cliente = TestClient(app, base_url="https://testserver")


@pytest.fixture(autouse=True)
def _puerta_de_compras_abierta():
    """Vacíos vive bajo `/administracion`, que tiene puerta. La de Compras se
    abre igual: el test de la clave de ese sector entra por Cajas.

    ES UNA SEGUNDA COPIA de la fixture de `tests/test_app.py` y no se puede
    compartir: cada módulo tiene su propio `cliente`, y una fixture le pone
    la cookie a UNO. Lo que sí se comparte es de dónde sale la firma.
    """
    from app.main import PUERTA_COMPRAS

    # Sin arranque cargado, como Palmala: el cartel de origen dice "la foto".
    # Los tests del arranque corren contra Postgres (test_vacios_conteo_2809).
    with patch.dict(os.environ, {"CLAVE_COMPRAS": "compras-secreta",
                                 "CLAVE_ADMINISTRACION": "admin-secreta"}), \
            patch("app.main.arranque_de_vacios", return_value=None):
        cliente.cookies.set(PUERTA_COMPRAS.cookie, PUERTA_COMPRAS.firma("compras-secreta"))
        cliente.cookies.set(PUERTA_ADMINISTRACION.cookie,
                            PUERTA_ADMINISTRACION.firma("admin-secreta"))
        try:
            yield
        finally:
            cliente.cookies.delete(PUERTA_COMPRAS.cookie)
            cliente.cookies.delete(PUERTA_ADMINISTRACION.cookie)


FUENTE_DB = io.open("app/db.py", encoding="utf-8").read()


def _pila(marca_id, marca, stock):
    return {"proveedor_id": 7, "proveedor": "Puesto EJEMPLO",
            "marca_id": marca_id, "marca": marca, "arranque": 0, "recibidos": 0,
            "devueltos": 0, "ajustes": 0, "asignados": 0, "stock": stock}


UN_PROVEEDOR = [{
    "id": 7, "nombre": "Puesto EJEMPLO", "stock": 35,
    "pilas": [_pila(None, None, 12), _pila(71, "EJ Roja", 23)],
}]
MARCAS = [{"id": 71, "nombre": "EJ Roja"}, {"id": 72, "nombre": "EJ Azul"}]


def _conexion_falsa(filas_fetchone=None, filas_fetchall=None):
    cursor = MagicMock()
    if filas_fetchone is not None:
        cursor.fetchone.side_effect = filas_fetchone
    if filas_fetchall is not None:
        cursor.fetchall.return_value = filas_fetchall
    cursor.__enter__ = MagicMock(return_value=cursor)
    cursor.__exit__ = MagicMock(return_value=False)
    conexion = MagicMock()
    conexion.cursor.return_value = cursor
    return conexion, cursor


def _mov(tipo, id, cantidad, **resto):
    """Una fila de `movimientos_de_vacios`, con todas sus columnas."""
    fila = {"tipo": tipo, "id": id, "cantidad": cantidad, "fecha": date(2026, 9, 25),
            "marca": None, "marca_hasta": None, "motivo": None, "compra_id": None,
            "articulo": None, "foto_ruta": None, "importe": None, "anulada": False,
            "antes_del_arranque": False, "cargada_desde": None,
            "proveedor_id": 7, "proveedor": "Puesto EJEMPLO"}
    fila.update(resto)
    return fila


def _parches_del_detalle(proveedores=UN_PROVEEDOR, marcas=MARCAS, movimientos=(), senas=None):
    return [
        patch("app.main.stock_de_vacios_deposito", return_value=list(proveedores)),
        patch("app.main.proveedor_para_vacios", return_value=None),
        patch("app.main.listar_marcas_vacio", return_value=list(marcas)),
        patch("app.main.movimientos_de_vacios", return_value=list(movimientos)),
        patch("app.main.sena_por_cajon_de_la_ultima_recepcion", return_value=senas or {}),
    ]


def _con(parches):
    from contextlib import ExitStack
    pila = ExitStack()
    for parche in parches:
        pila.enter_context(parche)
    return pila


# ---------------------------------------------------------------------------
# La cuenta derivada: el TEXTO (los números van contra Postgres)
# ---------------------------------------------------------------------------

def test_los_NOMBRES_de_las_columnas_son_los_que_la_consulta_DEVUELVE():
    """La tupla y el SELECT final tienen que coincidir, en orden.

    En la cuenta de cajas esto reventó TODO guardado: la consulta perdió una
    columna, un lector quedó en `fila[8]`, y un índice no nombra ninguna
    columna — el `grep` del campo sacado no lo encuentra nunca.
    """
    seleccion = _SQL_PILAS_DE_VACIOS.rsplit("SELECT", 1)[1].split("FROM")[0]
    alias = re.findall(r"AS (\w+)", seleccion)
    assert tuple(alias) == COLUMNAS_PILAS_DE_VACIOS


def test_las_entradas_son_las_recepciones_CON_SENA_y_nada_mas():
    """Sin seña no hay cajón que devolver (dueño, 25/09); pendiente no llegó."""
    assert "co.estado = 'recepcionado'" in _SQL_PILAS_DE_VACIOS
    assert "COALESCE(co.sena, 0) > 0" in _SQL_PILAS_DE_VACIOS
    assert "vacios_deposito_entradas" not in FUENTE_DB, (
        "apareció una tabla de entradas: las entradas se derivan de las recepciones")


def test_el_corte_compara_contra_el_INSTANTE_de_la_foto_y_no_contra_su_dia():
    """Con la fecha, lo recibido el 25/09 después de sacar la foto se perdía.

    Encontrado corriendo la cuenta contra Postgres, no leyéndola: la foto es
    de las 14, la recepción de las 16 del mismo día, y `procesada_el::date >
    f.fecha` la descartaba. Se compara contra `f.creado_en`, en las DOS patas
    que la foto ya incluye.
    """
    # Con arranque, contra su instante; sin arranque, contra el de la foto.
    assert "co.procesada_el > COALESCE((SELECT creado_en FROM arr), f.creado_en," in _SQL_PILAS_DE_VACIOS
    assert "d.creado_en > COALESCE((SELECT creado_en FROM arr), f.creado_en," in _SQL_PILAS_DE_VACIOS
    assert "f.fecha" not in _SQL_PILAS_DE_VACIOS


def test_las_ANULADAS_no_cuentan_en_ninguna_de_las_tres_tablas_que_se_anulan():
    """Calificado por ALIAS (corolario 4): son tres tablas con la misma columna."""
    # CON EL CONTEO: la asignación entra dos veces (sale de una pila y entra
    # en otra), y un `in` pasa con una de las dos patas sin el filtro.
    for alias, veces in (("d", 1), ("a", 1), ("s", 2)):
        assert _SQL_PILAS_DE_VACIOS.count(f"{alias}.anulado_el IS NULL") == veces, alias


def test_la_ASIGNACION_resta_de_una_pila_y_suma_en_la_otra():
    """Una pata sola —o las dos con el mismo signo— cambia el total de un movimiento
    que por definición no lo cambia."""
    assert "s.marca_desde_id, 0, 0, 0, 0, -s.cantidad" in _SQL_PILAS_DE_VACIOS
    assert "s.marca_hasta_id, 0, 0, 0, 0, s.cantidad" in _SQL_PILAS_DE_VACIOS


def test_la_pila_se_lee_con_el_PROVEEDOR_BLOQUEADO_y_con_LA_MISMA_consulta():
    """Dos devoluciones a la vez no pueden leer el mismo stock y sacar de más.

    Y la cuenta no se escribe dos veces: el freno filtra la consulta de la
    pantalla, no una segunda versión.
    """
    cuerpo = ast.parse(FUENTE_DB)
    funcion = next(n for n in ast.walk(cuerpo)
                   if isinstance(n, ast.FunctionDef) and n.name == "_stock_de_la_pila")
    texto = ast.unparse(funcion)
    assert "FOR UPDATE" in texto
    assert "_SQL_PILAS_DE_VACIOS" in texto
    for escritor in ("crear_devolucion_vacios", "crear_asignacion_vacios"):
        nodo = next(n for n in ast.walk(cuerpo)
                    if isinstance(n, ast.FunctionDef) and n.name == escritor)
        llamadas = {c.func.id for c in ast.walk(nodo)
                    if isinstance(c, ast.Call) and isinstance(c.func, ast.Name)}
        assert "_stock_de_la_pila" in llamadas, escritor


# ---------------------------------------------------------------------------
# Lo que se escribe
# ---------------------------------------------------------------------------

def test_el_INSERT_de_la_devolucion_guarda_la_ESTRUCTURA_ENTERA():
    """La tupla completa, y `stock_sistema` es lo que la pila tenía al devolver."""
    # La devolución (88), la seña de la pila (ninguna) y el vale que deja (9).
    conexion, cursor = _conexion_falsa(filas_fetchone=[(88,), None, (9,)])
    with patch("app.db.obtener_conexion", return_value=conexion), \
         patch("app.db._stock_de_la_pila", return_value=30):
        devolucion_id = crear_devolucion_vacios(
            7, 71, 12, foto_ruta="vacios/2026-09-25/v.jpg", importe=5000.0,
            cargada_desde="administracion")

    assert devolucion_id == 88
    insert = next(ll for ll in cursor.execute.call_args_list
                  if "INSERT INTO vacios_deposito_devoluciones" in ll.args[0])
    assert insert.args[1] == (7, 71, 12, 5000.0, "vacios/2026-09-25/v.jpg", 30,
                              "administracion")
    for columna in ("proveedor_id", "marca_vacio_id", "cantidad", "importe",
                    "foto_ruta", "stock_sistema", "cargada_desde"):
        assert columna in insert.args[0], columna
    assert "compra_id" not in insert.args[0], "la devolución ya no va contra una compra"


def test_devolver_MAS_de_lo_que_hay_NO_escribe_nada():
    conexion, cursor = _conexion_falsa(filas_fetchone=[("EJ Roja",)])
    with patch("app.db.obtener_conexion", return_value=conexion), \
         patch("app.db._stock_de_la_pila", return_value=3):
        with pytest.raises(ValueError, match="hay 3 cajones: no se pueden devolver 4"):
            crear_devolucion_vacios(7, 71, 4, foto_ruta="v.jpg", cargada_desde="administracion")
    assert not any("INSERT" in ll.args[0] for ll in cursor.execute.call_args_list)
    conexion.commit.assert_not_called()


def test_sin_FOTO_no_llega_ni_a_abrir_la_base():
    """La guarda va donde se ESCRIBE, antes que nada: un POST armado a mano no la saltea."""
    with patch("app.db.obtener_conexion") as abrir:
        with pytest.raises(ValueError, match="ajuste"):
            crear_devolucion_vacios(7, None, 1, foto_ruta="   ", cargada_desde="deposito")
    abrir.assert_not_called()


def test_el_vale_NO_TOCA_compras_importe_ni_el_costeo():
    """El descuento vive SOLO en la fila de la devolución. Es plata de envase."""
    cuerpo = ast.parse(FUENTE_DB)
    funcion = next(n for n in ast.walk(cuerpo)
                   if isinstance(n, ast.FunctionDef) and n.name == "crear_devolucion_vacios")
    assert "UPDATE compras" not in ast.unparse(funcion)


def test_anular_pregunta_la_EXISTENCIA_con_un_select_SIN_AGREGADO():
    """Con un `count(*)` la guarda no se dispara nunca: la fila vuelve con 0."""
    conexion, cursor = _conexion_falsa(filas_fetchone=[None])
    with patch("app.db.obtener_conexion", return_value=conexion):
        with pytest.raises(ValueError, match="no existe"):
            anular_devolucion_vacios(9999)
    consulta = cursor.execute.call_args_list[0].args[0]
    assert "SELECT t.anulado_el" in consulta
    for agregado in ("count(", "COUNT(", "sum(", "SUM("):
        assert agregado not in consulta


def test_anular_dos_veces_NO_PISA_la_fecha_original_y_el_genero_viaja():
    """"Ese ajuste ya estaba anulada" salía igual de prolijo que la frase buena."""
    for anular, frase in ((anular_devolucion_vacios, "Esa devolución ya estaba anulada"),
                          (anular_ajuste_vacios_deposito, "Ese ajuste ya estaba anulado")):
        conexion, cursor = _conexion_falsa(filas_fetchone=[("2026-09-18",)])
        with patch("app.db.obtener_conexion", return_value=conexion):
            with pytest.raises(ValueError) as invalido:
                anular(5)
        assert str(invalido.value) == frase + "."
        assert not any("UPDATE" in ll.args[0] for ll in cursor.execute.call_args_list)


# ---------------------------------------------------------------------------
# Las rutas
# ---------------------------------------------------------------------------

def _foto():
    import base64
    # un JPEG de 1x1 de verdad: `_comprimir_foto_jpeg` lo tiene que poder abrir
    return ("vale.jpg", base64.b64decode(
        "/9j/4AAQSkZJRgABAQEASABIAAD/2wBDAP//////////////////////////////////////////////"
        "////////////////////////////////////////////wAALCAABAAEBAREA/8QAFAABAAAAAAAAAAAA"
        "AAAAAAAAA//EABQQAQAAAAAAAAAAAAAAAAAAAAD/2gAIAQEAAD8AN//Z"), "image/jpeg")


def test_la_DEVOLUCION_sin_foto_rebota_400_y_no_escribe():
    with _con(_parches_del_detalle()), \
         patch("app.main.crear_devolucion_vacios") as crear, \
         patch("app.main.subir_foto_comanda") as subir:
        respuesta = cliente.post("/administracion/vacios/7/devolucion",
                                 data={"marca_vacio_id": "71", "cantidad": "3"})
    assert respuesta.status_code == 400
    assert "Sin la foto del vale no es una devolución" in respuesta.text
    crear.assert_not_called()
    subir.assert_not_called()


def test_la_DEVOLUCION_con_foto_escribe_la_PILA_elegida_y_la_ruta_subida():
    with _con(_parches_del_detalle()), \
         patch("app.main._comprimir_foto_jpeg", return_value=b"jpg"), \
         patch("app.main.subir_foto_comanda", return_value="vacios/v.jpg"), \
         patch("app.main.crear_devolucion_vacios", return_value=1) as crear:
        respuesta = cliente.post("/administracion/vacios/7/devolucion",
                                 data={"marca_vacio_id": "", "cantidad": "3", "importe": "2.400"},
                                 files={"foto": _foto()}, follow_redirects=False)
    assert respuesta.status_code == 303
    assert respuesta.headers["location"].startswith("/administracion/vacios/7?")
    # La puerta es la del prefijo, no una que tipee alguien.
    crear.assert_called_once_with(7, None, 3, foto_ruta="vacios/v.jpg", importe=2400.0,
                                  cargada_desde="administracion", numero_vale="")


def test_si_la_foto_NO_SE_SUBE_no_se_guarda_nada_y_lo_dice():
    with _con(_parches_del_detalle()), \
         patch("app.main._comprimir_foto_jpeg", return_value=b"jpg"), \
         patch("app.main.subir_foto_comanda", side_effect=RuntimeError("storage caído")), \
         patch("app.main.crear_devolucion_vacios") as crear:
        respuesta = cliente.post("/administracion/vacios/7/devolucion",
                                 data={"marca_vacio_id": "71", "cantidad": "3"},
                                 files={"foto": _foto()})
    assert respuesta.status_code == 502
    assert "No se guardó nada" in respuesta.text
    crear.assert_not_called()


def test_el_FRENO_de_la_pila_se_ve_en_la_pantalla_con_el_numero():
    with _con(_parches_del_detalle()), \
         patch("app.main._comprimir_foto_jpeg", return_value=b"jpg"), \
         patch("app.main.subir_foto_comanda", return_value="vacios/v.jpg"), \
         patch("app.main.crear_devolucion_vacios",
               side_effect=ValueError("En la pila EJ Roja hay 23 cajones: no se pueden devolver 40.")):
        respuesta = cliente.post("/administracion/vacios/7/devolucion",
                                 data={"marca_vacio_id": "71", "cantidad": "40"},
                                 files={"foto": _foto()})
    assert respuesta.status_code == 400
    assert "hay 23 cajones: no se pueden devolver 40" in respuesta.text


def test_AJUSTE_y_ASIGNACION_existen_SOLO_bajo_administracion():
    """Decisión del dueño: los cierra la puerta de Administración, por prefijo.

    ENCONTRADO contra DECIDIDO: se barren las rutas de la app y no una lista
    escrita a mano, así que una tercera puerta bajo /compras falla acá.
    """
    rutas = {r.path for r in app.routes if hasattr(r, "path")}
    encontradas = {r for r in rutas if re.search(r"^/(compras|administracion)/vacios/.*(ajuste|asignacion|movimiento)", r)}
    assert encontradas == {
        "/administracion/vacios/{proveedor_id}/ajuste",
        "/administracion/vacios/{proveedor_id}/asignacion",
        "/administracion/vacios/movimiento/{tipo}/{movimiento_id}/anular",
        # La pantalla de Movimientos (29/09) también es solo de Administración.
        "/administracion/vacios/movimientos",
        "/administracion/vacios/movimientos/excel",
    }


def test_el_AJUSTE_lleva_el_signo_del_SENTIDO_y_no_uno_tipeado():
    for sentido, esperado in (("sobran", 4), ("faltan", -4)):
        with _con(_parches_del_detalle()), \
             patch("app.main.crear_ajuste_vacios_deposito") as crear:
            respuesta = cliente.post("/administracion/vacios/7/ajuste", data={
                "marca_vacio_id": "71", "sentido": sentido, "cantidad": "4",
                "motivo": "EJ se contaron"}, follow_redirects=False)
        assert respuesta.status_code == 303
        crear.assert_called_once_with(7, 71, esperado, "EJ se contaron")


def test_el_AJUSTE_sin_sentido_rebota_y_no_escribe():
    with _con(_parches_del_detalle()), \
         patch("app.main.crear_ajuste_vacios_deposito") as crear:
        respuesta = cliente.post("/administracion/vacios/7/ajuste", data={
            "marca_vacio_id": "71", "sentido": "", "cantidad": "4", "motivo": "x"})
    assert respuesta.status_code == 400
    crear.assert_not_called()


def test_la_ASIGNACION_pasa_de_la_pila_elegida_a_la_marca_elegida():
    with _con(_parches_del_detalle()), \
         patch("app.main.crear_asignacion_vacios") as crear:
        respuesta = cliente.post("/administracion/vacios/7/asignacion", data={
            "marca_desde_id": "", "marca_hasta_id": "72", "cantidad": "5"},
            follow_redirects=False)
    assert respuesta.status_code == 303
    crear.assert_called_once_with(7, None, 72, 5, marca_nueva="")


def test_el_CONTEO_va_al_COTEJO_y_no_arranca_ninguna_cuenta():
    with patch("app.main.crear_conteo_vacios_deposito") as crear:
        respuesta = cliente.post("/administracion/vacios/conteo", data={
            "proveedor_id": "7", "marca_vacio_id": "71", "cantidad": "0",
            "fecha": "2026-09-25"}, follow_redirects=False)
    assert respuesta.status_code == 303
    assert respuesta.headers["location"] == "/administracion/vacios/cotejo"
    # CERO VALE: contar cero es contar.
    crear.assert_called_once_with(7, 71, 0, date(2026, 9, 25))


# ---------------------------------------------------------------------------
# Las pantallas
# ---------------------------------------------------------------------------

def _indice(proveedores=UN_PROVEEDOR):
    with patch("app.main.stock_de_vacios_deposito", return_value=proveedores), \
         patch("app.main.listar_proveedores",
               return_value=[{"id": 7, "nombre": "Puesto EJEMPLO"}]), \
         patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: MARCAS}):
        return cliente.get("/administracion/vacios")


def test_el_indice_muestra_el_total_y_CADA_PILA_con_su_marca():
    respuesta = _indice()
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    assert "35 cajones" in marcado
    assert "Sin marca" in marcado
    assert "EJ Roja" in marcado
    # y la cuenta vieja no quedó dibujada en ningún lado
    for jerga in ("todavía no arrancó", "Contá a la mañana", "recibidos +"):
        assert jerga not in marcado, jerga


def _conteo(url="/administracion/vacios/conteo", metodo="get", datos=None):
    """La pantalla de contar. El stock se parchea para que EXPLOTE: si la
    pantalla lo pidiera, el que cuenta podría ver el número del sistema."""
    with patch("app.main.stock_de_vacios_deposito",
               side_effect=AssertionError("la pantalla de contar leyó el stock")), \
         patch("app.main.listar_proveedores",
               return_value=[{"id": 7, "nombre": "Puesto EJEMPLO"}]), \
         patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: MARCAS}):
        if metodo == "get":
            return cliente.get(url)
        return cliente.post(url, data=datos, follow_redirects=False)


def test_el_conteo_OFRECE_solo_proveedores_y_marcas_cargados():
    """Dueño, 25/09: solo lo cargado. Por eso son selectores y no texto libre."""
    marcado = _conteo().text.split("</style>")[-1]
    assert '<select id="proveedor_id" name="proveedor_id" required>' in marcado
    assert '<select id="marca_vacio_id" name="marca_vacio_id">' in marcado
    assert 'name="proveedor_nuevo"' not in marcado and 'name="marca_nueva"' not in marcado


@pytest.mark.parametrize("sector", ["/administracion", "/deposito"])
def test_la_pantalla_de_CONTAR_no_tiene_ningun_numero_del_sistema(sector):
    """Dueño, 27/09: el que cuenta no puede ver cuántos dice el sistema, o
    transcribe en vez de contar. Hasta ese día el formulario vivía en el
    índice, arriba de las tarjetas con el stock de cada pila.

    No alcanza con que no se DIBUJE: la pantalla no LEE el stock (el parche
    explota si lo pide), así que tampoco puede quedar en el HTML escondido."""
    respuesta = _conteo(f"{sector}/vacios/conteo")
    assert respuesta.status_code == 200, respuesta.text[:300]
    marcado = respuesta.text.split("</style>")[-1]
    assert f'action="{sector}/vacios/conteo"' in marcado
    assert 'name="cantidad"' in marcado
    assert "cajones</div>" not in marcado and 'class="pila-stock"' not in marcado


def test_el_INDICE_ya_no_tiene_el_formulario_de_contar_y_el_COTEJO_manda_a_su_pantalla():
    """El control del de arriba: el índice SÍ muestra el stock, así que el
    formulario no puede estar ahí. Desde el 29/09 el índice tiene Exportar,
    Cotejo y Movimientos, y el link a contar vive en el Cotejo."""
    marcado = _indice().text.split("</style>")[-1]
    assert "35 cajones" in marcado
    assert 'name="cantidad"' not in marcado
    assert "/vacios/conteo" not in marcado
    with patch("app.main.cotejo_de_vacios_deposito", return_value=[]):
        cotejo = cliente.get("/administracion/vacios/cotejo").text.split("</style>")[-1]
    assert 'href="/administracion/vacios/conteo"' in cotejo


def test_un_CONTEO_que_rebota_vuelve_a_la_pantalla_de_contar_y_no_al_indice():
    """El rebote es el camino que no se ve al abrir: si volviera al índice,
    el que se equivocó de número vería el stock justo antes de reintentar."""
    respuesta = _conteo(metodo="post", datos={"proveedor_id": "7", "cantidad": "x",
                                              "fecha": "2026-09-25"})
    assert respuesta.status_code == 400
    marcado = respuesta.text.split("</style>")[-1]
    assert "tienen que ser un número entero" in marcado
    assert 'name="cantidad"' in marcado


def test_VACIOS_ya_no_existe_bajo_COMPRAS():
    """Dueño, 29/09: Vacíos del depósito vive en Administración, y Depósito
    cuenta y devuelve sin clave. Bajo /compras no queda ninguna ruta, y ninguna
    plantilla manda ahí: un link viejo sería un 404 en el botón."""
    import app.main as m
    rutas = [r.path for r in m.app.routes if getattr(r, "path", "").startswith("/compras/vacios")]
    assert rutas == []
    assert cliente.get("/compras/vacios").status_code == 404
    linkean = [str(a) for a in Path("templates").glob("*.html")
               if "/compras/vacios" in a.read_text(encoding="utf-8")]
    assert linkean == []



def test_el_detalle_en_ADMINISTRACION_ofrece_ajuste_asignacion_y_anular():
    with _con(_parches_del_detalle(movimientos=[
            _mov("asignacion", 3, 2, marca_hasta="EJ Roja")])):
        respuesta = cliente.get("/administracion/vacios/7")
    marcado = respuesta.text.split("</style>")[-1]
    assert 'action="/administracion/vacios/7/ajuste"' in marcado
    assert 'action="/administracion/vacios/7/asignacion"' in marcado
    assert 'action="/administracion/vacios/movimiento/asignacion/3/anular"' in marcado
    assert "/compras/" not in marcado


def test_el_vale_se_PRECARGA_con_la_sena_de_CADA_pila():
    with _con(_parches_del_detalle(senas={None: 300.0, 71: 800.0})):
        respuesta = cliente.get("/administracion/vacios/7")
    marcado = respuesta.text.split("</style>")[-1]
    opciones = re.findall(r'<option value="(\d*)"\s+data-sena="([^"]*)"', marcado)
    # LAS TRES pilas, la azul en cero y sin seña: devolver de ahí lo frena el
    # server con el número, que es más claro que una opción que falta.
    assert opciones == [("", "300.0"), ("71", "800.0"), ("72", "")]
    assert "No se le descuenta a la" in marcado


# ── El detalle en tres zonas (dueño, 29/09) ─────────────────────────────────

def _abiertas(marcado):
    """Los <details> que vienen desplegados, por id."""
    return re.findall(r'<details class="(?:accion|historial)" id="(\w+)"[^>]*\bopen\b', marcado)


def _todas(marcado):
    return re.findall(r'<details class="(?:accion|historial)" id="(\w+)"', marcado)


def test_las_CUATRO_acciones_y_el_historial_llegan_CERRADOS_y_en_su_orden():
    with _con(_parches_del_detalle()):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    assert _todas(marcado) == ["devolver", "pasar", "corregir", "ajustar", "movimientos"]
    assert _abiertas(marcado) == []
    # Las cuatro comparten el `name`: el navegador las abre de a una.
    assert len(re.findall(r'<details class="accion" id="\w+" name="accion"', marcado)) == 4


@pytest.mark.parametrize("abrir", ["devolver", "pasar", "corregir", "ajustar", "movimientos"])
def test_abrir_despliega_ESA_y_ninguna_otra(abrir):
    with _con(_parches_del_detalle()):
        marcado = cliente.get(f"/administracion/vacios/7?abrir={abrir}").text.split("</style>")[-1]
    assert _abiertas(marcado) == [abrir]


def test_un_abrir_que_no_es_una_accion_no_abre_nada():
    with _con(_parches_del_detalle()):
        marcado = cliente.get('/administracion/vacios/7?abrir=x" onload="y').text.split("</style>")[-1]
    assert _abiertas(marcado) == [] and 'onload="y' not in marcado


def test_un_AJUSTE_que_rebota_vuelve_con_el_AJUSTE_abierto_y_el_error_arriba():
    """El error queda al lado de lo que hay que corregir: si el formulario
    volviera plegado, el que se equivocó tendría que ir a buscarlo."""
    with _con(_parches_del_detalle()):
        respuesta = cliente.post("/administracion/vacios/7/ajuste",
                                 data={"sentido": "", "cantidad": "2", "motivo": "x"})
    assert respuesta.status_code == 400
    marcado = respuesta.text.split("</style>")[-1]
    assert _abiertas(marcado) == ["ajustar"]
    assert "Decí si sobran o faltan" in marcado




def test_el_HISTORIAL_dice_cada_tipo_y_solo_ofrece_anular_lo_que_se_anula():
    movs = [_mov("arranque", 1, 40, marca="EJ Roja", motivo="Conteo físico 28/09"),
            _mov("entrada", 2, 10, marca="EJ Roja", compra_id=55, articulo="EJ Pera"),
            _mov("devolucion", 3, -5, foto_ruta="v.jpg", importe=4000.0),
            _mov("ajuste", 4, -2, motivo="EJ se rompieron"),
            _mov("asignacion", 5, 6, marca_hasta="EJ Roja"),
            _mov("ajuste", 6, 1, motivo="EJ anulado", anulada=True)]
    with _con(_parches_del_detalle(movimientos=movs)):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    historial = marcado[marcado.index('id="movimientos"'):]
    assert re.findall(r'data-tipo="(\w+)"', historial) == [m["tipo"] for m in movs]
    renglones = [" ".join(re.sub(r"<[^>]+>", " ", r).split())
                 for r in re.findall(r'<div class="vale-cabeza">(.*?)</div>', historial, re.S)]
    assert renglones == ["40 Conteo físico · EJ Roja",
                         "+10 Compra con seña · EJ Roja",
                         "-5 Devolución · Sin marca",
                         "-2 Ajuste · Sin marca",
                         "6 Pase · Sin marca → EJ Roja",
                         "+1 Ajuste · Sin marca anulado"]
    assert "compra #55 · EJ Pera · se corrige desde la compra" in historial
    assert "$4.000 · vale cargado" in historial and "EJ se rompieron" in historial
    # El signo va en color: verde lo que suma, rojo lo que resta.
    assert '<span class="vale-cantidad suma">+10' in historial
    assert '<span class="vale-cantidad resta">-5' in historial
    anular = re.findall(r'action="([^"]*/anular)"', historial)
    assert anular == ["/administracion/vacios/devolucion/3/anular",
                      "/administracion/vacios/movimiento/ajuste/4/anular",
                      "/administracion/vacios/movimiento/asignacion/5/anular"]


def test_en_el_NAVEGADOR_abrir_una_accion_CIERRA_la_que_estaba_abierta():
    """El efecto, no el atributo (corolario 32): con dos abiertas a la vez la
    pantalla se hace larguísima y el que carga no sabe cuál está llenando."""
    pytest.importorskip("playwright")
    from playwright.sync_api import sync_playwright
    from scripts.medir_layout import CHROMIUM
    with _con(_parches_del_detalle()):
        html = cliente.get("/administracion/vacios/7?abrir=pasar").text
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(executable_path=CHROMIUM)
        pagina = navegador.new_page(viewport={"width": 390, "height": 800})
        errores = []
        pagina.on("pageerror", lambda e: errores.append(str(e)))
        pagina.set_content(html)
        abiertas = lambda: pagina.eval_on_selector_all(
            "details.accion", "ds => ds.filter(d => d.open).map(d => d.id)")
        assert abiertas() == ["pasar"]
        pagina.click("#ajustar > summary")
        pagina.wait_for_timeout(100)
        assert abiertas() == ["ajustar"]
        alto = pagina.eval_on_selector("#ajustar > summary", "e => e.getBoundingClientRect().height")
        navegador.close()
    assert alto >= 44
    assert errores == []


def test_la_foto_del_vale_es_REQUIRED_en_el_formulario():
    with _con(_parches_del_detalle()):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    assert re.search(r'<input id="foto" name="foto" type="file"[^>]*required', marcado)


def test_un_proveedor_EN_CERO_se_abre_igual_para_cargarle_marcas():
    parches = _parches_del_detalle(proveedores=[], marcas=[])
    parches[1] = patch("app.main.proveedor_para_vacios", return_value={
        "id": 9, "nombre": "Puesto EJEMPLO DOS", "stock": 0, "pilas": []})
    with _con(parches):
        respuesta = cliente.get("/administracion/vacios/9")
    assert respuesta.status_code == 200
    assert 'action="/administracion/vacios/9/marca"' in respuesta.text


def test_un_proveedor_que_NO_existe_da_404():
    with _con(_parches_del_detalle(proveedores=[])):
        respuesta = cliente.get("/administracion/vacios/999")
    assert respuesta.status_code == 404


def test_el_COTEJO_dice_la_diferencia_y_cierra_en_cero():
    filas = [{"proveedor_id": 7, "proveedor": "Puesto EJEMPLO", "marca": "EJ Roja",
              "contado": 20, "fecha": date(2026, 9, 25), "sistema": 23, "diferencia": 3},
             {"proveedor_id": 7, "proveedor": "Puesto EJEMPLO", "marca": None,
              "contado": 12, "fecha": date(2026, 9, 25), "sistema": 12, "diferencia": 0}]
    with patch("app.main.cotejo_de_vacios_deposito", return_value=filas):
        respuesta = cliente.get("/administracion/vacios/cotejo")
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    assert "EJ Roja" in marcado and "cierra" in marcado


# ---------------------------------------------------------------------------
# AJUSTAR DESDE EL COTEJO (dueño, 29/09): el botón abre el proveedor con el
# ajuste desplegado y la diferencia precargada. La URL no lleva números: el
# server rehace la cuenta con la MISMA función del Cotejo.
# ---------------------------------------------------------------------------

COTEJO_CON_AJUSTES = [
    {"proveedor_id": 7, "proveedor": "Puesto EJEMPLO", "marca_id": 71, "marca": "EJ Roja",
     "contado": 20, "fecha": date(2026, 9, 25), "sistema": 23, "diferencia": 3},
    {"proveedor_id": 7, "proveedor": "Puesto EJEMPLO", "marca_id": None, "marca": None,
     "contado": 15, "fecha": date(2026, 9, 26), "sistema": 12, "diferencia": -3},
    {"proveedor_id": 7, "proveedor": "Puesto EJEMPLO", "marca_id": 72, "marca": "EJ Azul",
     "contado": 5, "fecha": date(2026, 9, 26), "sistema": 5, "diferencia": 0},
]


def test_la_PROPUESTA_sale_del_signo_de_la_diferencia():
    from core.movimientos_vacios import propuesta_desde_el_cotejo
    faltan = propuesta_desde_el_cotejo(COTEJO_CON_AJUSTES, 7, 71)
    assert (faltan["sentido"], faltan["cantidad"], faltan["marca_id"]) == ("faltan", 3, 71)
    sobran = propuesta_desde_el_cotejo(COTEJO_CON_AJUSTES, 7, None)
    assert (sobran["sentido"], sobran["cantidad"], sobran["marca_id"]) == ("sobran", 3, None)
    assert propuesta_desde_el_cotejo(COTEJO_CON_AJUSTES, 7, 72) is None, "cierra: nada que ajustar"
    assert propuesta_desde_el_cotejo(COTEJO_CON_AJUSTES, 8, 71) is None, "otro proveedor"
    assert propuesta_desde_el_cotejo(COTEJO_CON_AJUSTES, 7, 99) is None, "sin conteo"


def test_el_COTEJO_ofrece_AJUSTAR_solo_donde_no_cierra_y_sin_numeros_en_la_url():
    from app.db import SIN_MARCA
    with patch("app.main.cotejo_de_vacios_deposito", return_value=COTEJO_CON_AJUSTES):
        marcado = cliente.get("/administracion/vacios/cotejo").text.split("</style>")[-1]
    links = re.findall(r'<a class="ajustar" href="([^"]+)">Ajustar</a>', marcado)
    assert links == ["/administracion/vacios/7?abrir=ajustar&amp;pila=71",
                     f"/administracion/vacios/7?abrir=ajustar&amp;pila={SIN_MARCA}"]
    assert "Sin marca" in marcado and "sin asignar" not in marcado


@pytest.mark.parametrize("pila, marca, sentido", [("71", "71", "faltan"), ("sin", "", "sobran")])
def test_AJUSTAR_desde_el_cotejo_abre_el_ajuste_con_la_diferencia_PRECARGADA(pila, marca, sentido):
    with _con(_parches_del_detalle()), \
         patch("app.main.cotejo_de_vacios_deposito", return_value=COTEJO_CON_AJUSTES):
        respuesta = cliente.get(f"/administracion/vacios/7?abrir=ajustar&pila={pila}")
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    assert _abiertas(marcado) == ["ajustar"]
    ajuste = marcado[marcado.index('id="ajustar"'):marcado.index('id="movimientos"')]
    elegidas = re.findall(r'<option value="([^"]*)" selected>', ajuste)
    assert elegidas == [marca, sentido]
    assert 'name="cantidad" type="number" min="1" step="1" required value="3"' in ajuste
    # El motivo NO se precarga: es lo único que dice por qué no cerró.
    motivo = re.search(r'<input id="ajuste_motivo"[^>]*>', ajuste).group(0)
    assert "value=" not in motivo
    assert "Del Cotejo:" in ajuste


@pytest.mark.parametrize("url", [
    "/administracion/vacios/7?abrir=ajustar&pila=72",      # esa pila cierra
    "/administracion/vacios/7?abrir=ajustar&pila=basura",  # no es una pila
    "/administracion/vacios/7?abrir=pasar&pila=71",        # no pidió ajustar
])
def test_sin_DIFERENCIA_no_se_precarga_nada(url):
    with _con(_parches_del_detalle()), \
         patch("app.main.cotejo_de_vacios_deposito", return_value=COTEJO_CON_AJUSTES):
        marcado = cliente.get(url).text.split("</style>")[-1]
    ajuste = marcado[marcado.index('id="ajustar"'):marcado.index('id="movimientos"')]
    assert " selected>" not in ajuste and 'value="3"' not in ajuste
    assert "Del Cotejo:" not in ajuste


def test_el_detalle_SIN_pila_no_le_pregunta_al_cotejo():
    """El Cotejo recorre todas las pilas: pedírselo en cada apertura del detalle
    sería pagar esa cuenta para nada."""
    with _con(_parches_del_detalle()), \
         patch("app.main.cotejo_de_vacios_deposito", side_effect=AssertionError("no")) as cotejo:
        assert cliente.get("/administracion/vacios/7?abrir=ajustar").status_code == 200
    assert not cotejo.called


def test_la_pantalla_esta_LINKEADA_desde_el_hub_de_ADMINISTRACION():
    """Una ruta sin botón no la ve nadie, y la suite entera entra por la URL.
    Que Compras ya no linkee lo afirma `test_VACIOS_ya_no_existe_bajo_COMPRAS`."""
    admin = io.open("templates/administracion.html", encoding="utf-8").read()
    assert 'href="/administracion/vacios"' in admin



def test_el_boton_NO_se_llama_solo_Vacios_porque_el_del_puesto_ya_se_llama_asi():
    hub = io.open("templates/administracion.html", encoding="utf-8").read()
    etiqueta = re.search(r'href="/administracion/vacios">([^<]+)<', hub)
    assert etiqueta, "no encontré el botón"
    assert etiqueta.group(1).strip() != "Vacíos"
    assert "depósito" in etiqueta.group(1).lower()



# ---------------------------------------------------------------------------
# A 390px, con lo único cuyo largo NO controlamos
# ---------------------------------------------------------------------------

UN_NOMBRE_QUE_NO_SE_PUEDE_PARTIR = "PUESTODEEJEMPLOSINUNSOLOESPACIOPARAPARTIRLO"


def _medir(html, **opciones):
    pytest.importorskip("playwright", reason="la medición de layout necesita un navegador")
    from scripts.medir_layout import medir_sync
    return medir_sync(html, ancho=390, selector_filas=".tarjeta", **opciones)


def _pantallas_de_vacios(nombre):
    """Las dos pantallas con el nombre en TODOS los lugares que tipea una persona:
    el proveedor y la MARCA."""
    pilas = [dict(_pila(None, None, 12), proveedor=nombre),
             dict(_pila(71, nombre, 23), proveedor=nombre)]
    proveedores = [dict(UN_PROVEEDOR[0], nombre=nombre, pilas=pilas)]
    marcas = [{"id": 71, "nombre": nombre}]
    movimientos = [_mov("devolucion", 1, -25, importe=18500.0, marca=nombre,
                        foto_ruta="2026-09-12/vale.jpg", fecha=date(2026, 9, 12)),
                   _mov("ajuste", 3, 2, marca=nombre, motivo=nombre),
                   _mov("entrada", 9, 10, marca=nombre, compra_id=40, articulo=nombre),
                   _mov("asignacion", 4, 5, marca=nombre, marca_hasta=nombre)]

    with patch("app.main.stock_de_vacios_deposito", return_value=proveedores), \
         patch("app.main.listar_proveedores", return_value=[{"id": 7, "nombre": nombre}]), \
         patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: marcas}):
        indice = cliente.get("/administracion/vacios")
    with _con(_parches_del_detalle(proveedores=proveedores, marcas=marcas,
                                   movimientos=movimientos)):
        detalle = cliente.get("/administracion/vacios/7?abrir=movimientos")

    assert indice.status_code == 200 and detalle.status_code == 200
    return indice.text, detalle.text


def test_las_DOS_pantallas_aguantan_un_nombre_QUE_NO_SE_PUEDE_PARTIR():
    """Medido a 390px, y el desborde se lee de `desborde_pagina` (corolario 53)."""
    for pantalla, html in zip(("índice", "detalle"),
                              _pantallas_de_vacios(UN_NOMBRE_QUE_NO_SE_PUEDE_PARTIR)):
        medicion = _medir(html)
        desborde = medicion.get("desborde_pagina", medicion["desborde"])
        assert medicion["pares"] > 0, pantalla
        assert desborde == 0, f"{pantalla} desborda {desborde}px"
        assert medicion["solapes"] == [], f"{pantalla}: {medicion['solapes']}"


def test_las_DOS_pantallas_con_nombres_NORMALES_tampoco_se_pisan():
    for pantalla, html in zip(("índice", "detalle"),
                              _pantallas_de_vacios("Puesto EJEMPLO del Norte")):
        medicion = _medir(html)
        desborde = medicion.get("desborde_pagina", medicion["desborde"])
        assert medicion["pares"] > 0, pantalla
        assert desborde == 0, f"{pantalla} desborda {desborde}px"
        assert medicion["solapes"] == [], f"{pantalla}: {medicion['solapes']}"


def test_el_503_de_una_puerta_NOMBRA_SU_SECTOR_y_no_Gerencia():
    """La misma pantalla la comparten las tres zonas con clave.

    Hasta el 18/09 su texto era entero de Gerencia, así que un POST de
    Compras sin la variable cargada contestaba "Falta la clave de Gerencia"
    y explicaba que corregir una recepción mueve la cotización. El que lo
    leía se iba a pedir la clave equivocada: la url era una y los sectores,
    tres (corolario 56).

    Se descubrió acá porque Vacíos era de Compras y su primer POST dio 503.
    Desde el 29/09 Vacíos no está bajo /compras, así que prueba con Cajas.
    """
    from app.main import PUERTA_COMPRAS

    # Sin la variable cargada: es el único caso que dibuja esta pantalla.
    with patch.dict(os.environ, {"CLAVE_COMPRAS": ""}):
        respuesta = cliente.post("/compras/cajas/conteo-inicial",
                                 data={"envase_id": "7", "cantidad": "30",
                                       "fecha": "2026-09-18"},
                                 follow_redirects=False)

    assert respuesta.status_code == 503
    assert PUERTA_COMPRAS.env_var in respuesta.text, "tiene que decir QUÉ variable falta"
    marcado = respuesta.text.split("</style>")[-1]
    assert PUERTA_COMPRAS.titulo in marcado
    # Y la jerga del sector AJENO no puede aparecer: afirmar el texto bueno
    # pasa igual si la frase vieja quedó tres líneas más abajo.
    assert "Gerencia" not in marcado
    assert "Corregir una recepción" not in marcado


# ---------------------------------------------------------------------------
# El TIPO DE CAJÓN por proveedor se fue (dueño, 29/09): lo reemplazan las marcas
# ---------------------------------------------------------------------------

# Lo que decía cuando existía: los rótulos, los campos y el aviso. Se busca la
# JERGA y no solo que falte un campo: afirmar lo bueno pasa igual si la frase
# vieja quedó tres líneas más abajo.
JERGA_DEL_TIPO_DE_CAJON = ("En qué cajón entrega", "Tipo de cajón", "Guardar cajón",
                           "sin declarar", 'name="tipo_cajon_id"', 'name="cajon_nombre_nuevo"',
                           "/cajon\"")


def test_el_TIPO_DE_CAJON_no_aparece_en_ninguna_pantalla():
    """El índice, la lista de stock, el detalle y el alta de proveedores.

    Cada una lleva su IDENTIDAD al lado (corolario 53): sin eso, la pantalla
    de la clave tampoco dice "Tipo de cajón" y el test pasaría sobre ella.
    """
    pantallas = {"índice": (_indice(), "Puesto EJEMPLO")}
    with patch("app.main.stock_de_vacios_deposito", return_value=UN_PROVEEDOR):
        pantallas["stock"] = (cliente.get("/administracion/vacios/stock"), "Puesto EJEMPLO")
    with _con(_parches_del_detalle()):
        pantallas["detalle"] = (cliente.get("/administracion/vacios/7"), 'action="/administracion/vacios/7/marca"')
    with patch("app.main.listar_proveedores_para_abm", return_value=[]):
        pantallas["alta de proveedores"] = (cliente.get("/compras/proveedores"),
                                            'action="/compras/proveedores/nuevo"')
    for nombre, (respuesta, identidad) in pantallas.items():
        assert respuesta.status_code == 200, nombre
        assert identidad in respuesta.text, nombre
        for jerga in JERGA_DEL_TIPO_DE_CAJON:
            assert jerga not in respuesta.text, (nombre, jerga)


def test_los_EXPORTES_del_stock_no_llevan_la_columna_del_tipo_de_cajon():
    """El Excel lo mira test_cajas_y_vacios_en_administracion; acá, el PDF."""
    import pypdfium2 as pdfium
    from core.exportar_vacios_deposito import generar_pdf_stock_vacios_deposito
    documento = pdfium.PdfDocument(generar_pdf_stock_vacios_deposito(date(2026, 9, 29), UN_PROVEEDOR))
    texto = "\n".join(pagina.get_textpage().get_text_range() for pagina in documento)
    assert "Puesto EJEMPLO" in texto and "EJ Roja" in texto and "Cajones" in texto
    assert "Tipo de cajón" not in texto and "sin declarar" not in texto


def test_NINGUN_codigo_lee_ni_escribe_el_tipo_de_cajon():
    """Es lo que deja borrar la columna y la tabla (el paso 2) sin que nada explote.

    Por POSICIÓN y no por el nombre suelto (corolario 59): el comentario de
    app/db.py nombra `tipos_cajon` para contar que se fue, y eso no es leerla.
    `tipo_cajon_id` no aparece en prosa en ningún lado, así que ése sí se busca
    entero. Y ninguna ruta termina en /cajon.
    """
    import glob
    posicion = re.compile(r"(?:FROM|JOIN|INTO|UPDATE)\s+tipos_cajon\b", re.I)
    ofensores = []
    archivos = (glob.glob("app/**/*.py", recursive=True) + glob.glob("core/**/*.py", recursive=True)
                + glob.glob("scripts/**/*.py", recursive=True) + glob.glob("templates/**/*.html", recursive=True))
    for archivo in archivos:
        texto = io.open(archivo, encoding="utf-8").read()
        if posicion.search(texto) or "tipo_cajon_id" in texto:
            ofensores.append(archivo)
    assert len(archivos) > 100, "el barrido no encontró los archivos que tenía que mirar"
    assert ofensores == []
    from app.main import app as la_app
    assert [r.path for r in la_app.routes if getattr(r, "path", "").endswith("/cajon")] == []


def test_la_ASIGNACION_con_marca_ESCRITA_llega_a_la_escritura():
    with _con(_parches_del_detalle()), \
         patch("app.main.crear_asignacion_vacios") as crear:
        respuesta = cliente.post("/administracion/vacios/7/asignacion", data={
            "marca_desde_id": "", "marca_hasta_id": "", "marca_nueva": "EJ Verde",
            "cantidad": "5"}, follow_redirects=False)
    assert respuesta.status_code == 303
    crear.assert_called_once_with(7, None, None, 5, marca_nueva="EJ Verde")


def test_sin_marca_elegida_NI_escrita_rebota_sin_escribir():
    with _con(_parches_del_detalle()), \
         patch("app.main.crear_asignacion_vacios") as crear:
        respuesta = cliente.post("/administracion/vacios/7/asignacion", data={
            "marca_desde_id": "", "marca_hasta_id": "", "cantidad": "5"})
    assert respuesta.status_code == 400
    crear.assert_not_called()


def test_la_tarjeta_de_ASIGNAR_sale_AUNQUE_el_proveedor_no_tenga_marcas():
    """Iba adentro de un `if marcas` y ningún proveedor tenía: no salió nunca."""
    with _con(_parches_del_detalle(marcas=[])):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    tarjeta = marcado[marcado.index('id="pasar"'):marcado.index('id="corregir"')]
    assert 'action="/administracion/vacios/7/asignacion"' in tarjeta
    assert 'name="marca_hasta_id"' not in tarjeta           # no hay de dónde elegir
    assert re.search(r'name="marca_nueva"[^>]*required', tarjeta)
    assert '<p class="sin-marca-total hay">Sin marca: 12 cajones</p>' in tarjeta


def test_con_marcas_se_ELIGE_o_se_ESCRIBE():
    with _con(_parches_del_detalle()):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    tarjeta = marcado[marcado.index('id="pasar"'):marcado.index('id="corregir"')]
    opciones = re.findall(r'<option value="(\d*)">([^<]+)</option>',
                          tarjeta[tarjeta.index('name="marca_hasta_id"'):])
    assert opciones[:3] == [("", "Una marca nueva (escribila abajo)"), ("71", "EJ Roja"),
                            ("72", "EJ Azul")]
    assert 'name="marca_nueva"' in tarjeta
    assert not re.search(r'name="marca_nueva"[^>]*required', tarjeta)


def test_los_SIN_MARCA_con_cajones_van_RESALTADOS_y_los_en_cero_no():
    en_cero = [{**UN_PROVEEDOR[0], "id": 8, "stock": 23,
                "pilas": [_pila(None, None, 0), _pila(71, "EJ Roja", 23)]}]
    for url in ("/administracion/vacios",):
        with patch("app.main.stock_de_vacios_deposito", return_value=UN_PROVEEDOR + en_cero), \
             patch("app.main.listar_proveedores", return_value=[]), \
             patch("app.main.listar_marcas_vacio_por_proveedor", return_value={}):
            respuesta = cliente.get(url)
        partes = respuesta.text.split("</style>")
        assert ".pila.sin-marca { background:" in partes[0]
        marcado = partes[-1]
        assert marcado.count('class="pila sin-marca"') == 1, url
    # El link a asignar sale solo en Administración, y solo donde hay sin marca.
    assert respuesta.text.count('href="/administracion/vacios/7?abrir=pasar"') == 1
    assert 'href="/administracion/vacios/8?abrir=pasar"' not in respuesta.text


def test_RENOMBRAR_bien_vuelve_con_aviso():
    with _con(_parches_del_detalle()), \
         patch("app.main.renombrar_marca_vacio") as renombrar:
        respuesta = cliente.post("/administracion/vacios/7/marca/71/renombrar",
                                 data={"nombre": "EJ Colorada"}, follow_redirects=False)
    assert respuesta.status_code == 303
    renombrar.assert_called_once_with(7, 71, "EJ Colorada")


def test_RENOMBRAR_contra_una_que_YA_EXISTE_ofrece_juntarlas():
    from app.db import MarcaQueYaExiste
    with _con(_parches_del_detalle()), \
         patch("app.main.renombrar_marca_vacio", side_effect=MarcaQueYaExiste(72, "EJ Azul")):
        respuesta = cliente.post("/administracion/vacios/7/marca/71/renombrar",
                                 data={"nombre": "ej azul"})
    assert respuesta.status_code == 409
    marcado = respuesta.text.split("</style>")[-1]
    tarjeta = marcado[marcado.index('id="juntar"'):marcado.index('/renombrar"')]
    assert "¿Juntar «EJ Roja» con «EJ Azul»?" in tarjeta
    assert 'action="/administracion/vacios/7/marca/71/juntar"' in tarjeta
    assert '<input type="hidden" name="queda_id" value="72">' in tarjeta


def test_JUNTAR_llama_a_la_escritura_y_dice_lo_que_borro():
    with _con(_parches_del_detalle()), \
         patch("app.main.juntar_marcas_vacio",
               return_value={"movidos": 4, "asignaciones_entre_ellas": 1}) as juntar:
        respuesta = cliente.post("/administracion/vacios/7/marca/71/juntar",
                                 data={"queda_id": "72"}, follow_redirects=False)
    assert respuesta.status_code == 303
    juntar.assert_called_once_with(7, 71, 72)
    aviso = respuesta.headers["location"]
    assert "Se+borraron+1+asignaciones" in aviso or "Se%20borraron%201" in aviso


def test_la_tarjeta_de_MOVER_ofrece_salir_de_una_MARCA_y_no_solo_de_sin_marca():
    with _con(_parches_del_detalle()):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    tarjeta = marcado[marcado.index('id="pasar"'):marcado.index('id="corregir"')]
    desde = tarjeta[tarjeta.index('name="marca_desde_id"'):tarjeta.index("</select>")]
    assert re.findall(r'<option value="(\d*)">([^<]+)</option>', desde) == [
        ("", "Sin marca · hay 12"), ("71", "EJ Roja · hay 23"), ("72", "EJ Azul · hay 0")]


# ---------------------------------------------------------------------------
# EL ARRANQUE DESDE EL CONTEO FÍSICO (dueño, 28/09): la pantalla dice de dónde
# sale el número, y lo de antes del conteo no ofrece anularse.
# ---------------------------------------------------------------------------
ARRANQUE = {"motivo": "Conteo físico 28/09", "creado_en": datetime(2026, 9, 28, 21, 40),
            "total": 1086}


def test_el_INDICE_y_el_DETALLE_dicen_de_que_conteo_arranca_la_cuenta():
    with patch("app.main.arranque_de_vacios", return_value=ARRANQUE), \
            patch("app.main.stock_de_vacios_deposito", return_value=list(UN_PROVEEDOR)):
        indice = cliente.get("/administracion/vacios").text.split("</style>")[-1]
    with patch("app.main.arranque_de_vacios", return_value=ARRANQUE), _con(_parches_del_detalle()):
        detalle = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    for marcado in (indice, detalle):
        origen, = re.findall(r'<p class="origen">(.*?)</p>', marcado, re.S)
        texto = " ".join(origen.split())
        assert "<strong>Conteo físico 28/09</strong>" in texto
        assert "1086 cajones contados, cargado el 28/09 a las 21:40" in texto
        assert "foto del 25/09" not in texto


def test_SIN_arranque_el_cartel_dice_la_foto_del_25_09():
    with patch("app.main.stock_de_vacios_deposito", return_value=list(UN_PROVEEDOR)):
        marcado = cliente.get("/administracion/vacios").text.split("</style>")[-1]
    origen, = re.findall(r'<p class="origen">(.*?)</p>', marcado, re.S)
    assert "foto del 25/09" in origen


def test_lo_de_ANTES_del_conteo_no_ofrece_ANULAR_y_lo_de_despues_si():
    movs = [_mov("ajuste", 5, 2, motivo="viejo", antes_del_arranque=True),
            _mov("ajuste", 6, 3, motivo="nuevo"),
            _mov("devolucion", 4, -30, foto_ruta="x.jpg", antes_del_arranque=True),
            _mov("devolucion", 8, -30, foto_ruta="x.jpg")]
    with patch("app.main.arranque_de_vacios", return_value=ARRANQUE), \
            _con(_parches_del_detalle(movimientos=movs)):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    assert "/vacios/devolucion/4/anular" not in marcado
    assert "/vacios/devolucion/8/anular" in marcado
    assert "/vacios/movimiento/ajuste/5/anular" not in marcado
    assert "/vacios/movimiento/ajuste/6/anular" in marcado
    assert marcado.count("antes del conteo: no cuenta") == 2
    # El vale viejo se sigue pudiendo ver: la historia queda.
    assert "/vacios/devolucion/4/foto" in marcado


# ── Depósito cuenta vacíos sin la clave (dueño, 29/09) ──────────────────────

def test_el_menu_de_DEPOSITO_dice_Stock_Mercaderia_y_abajo_Stock_Vacios():
    """El botón se renombró y el nuevo va JUSTO abajo: ningún botón en el medio.

    Y la jerga vieja no puede quedar en ningún lado del menú."""
    marcado = cliente.get("/deposito").text.split("</style>")[-1]
    botones = re.findall(r'<a class="boton[^"]*" href="([^"]+)"', marcado)
    i = botones.index("/deposito/stock/fisico")
    assert botones[i + 1] == "/deposito/vacios/conteo", botones
    assert "<span>Stock Mercadería</span>" in marcado
    assert "<span>Stock Vacíos</span>" in marcado
    assert "Contar el stock" not in marcado and "Contar stock" not in marcado


def test_la_jerga_Contar_stock_no_aparece_en_ninguna_plantilla_ni_en_el_codigo():
    """El conjunto ENCONTRADO tiene que estar vacío: un link o una ayuda que
    siga diciendo el nombre viejo manda a buscar un botón que no existe."""
    raiz = Path(__file__).resolve().parent.parent
    encontrados = [
        str(a.relative_to(raiz))
        for carpeta in ("templates", "app", "core")
        for a in (raiz / carpeta).rglob("*")
        if a.suffix in (".html", ".py")
        and re.search(r"contar (el )?stock", a.read_text(encoding="utf-8"), re.I)
    ]
    assert encontrados == []


def test_DEPOSITO_abre_el_conteo_de_vacios_SIN_ninguna_cookie():
    """Sin la clave de Administración ni la de Compras: la barra es la de
    Depósito, el formulario manda a Depósito y el atrás vuelve a su menú."""
    limpio = TestClient(app)
    with patch.dict(os.environ, {"CLAVE_COMPRAS": "compras-secreta",
                                 "CLAVE_ADMINISTRACION": "admin-secreta"}), \
         patch("app.main.stock_de_vacios_deposito",
               side_effect=AssertionError("la pantalla de contar leyó el stock")), \
         patch("app.main.listar_proveedores",
               return_value=[{"id": 7, "nombre": "Puesto EJEMPLO"}]), \
         patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: MARCAS}):
        respuesta = limpio.get("/deposito/vacios/conteo")
        # EL CONTROL: con las claves puestas, el MISMO cliente rebota en las
        # otras dos puertas. Sin esto, el 200 de arriba se daría también con
        # las puertas apagadas y no probaría nada.
        cerradas = [limpio.get(f"{s}/vacios/conteo").status_code
                    for s in ("/compras", "/administracion")]
    assert cerradas == [401, 401], cerradas
    assert respuesta.status_code == 200, respuesta.text[:300]
    marcado = respuesta.text
    assert 'action="/deposito/vacios/conteo"' in marcado
    assert 'name="cantidad"' in marcado
    # Ningún camino de esta pantalla lleva a una puerta con clave.
    assert "/administracion/" not in marcado and "/compras/" not in marcado


def test_el_conteo_de_DEPOSITO_vuelve_a_contar_y_NO_al_cotejo():
    """El Cotejo muestra el número del sistema: el que sigue contando lo vería."""
    with patch("app.main.crear_conteo_vacios_deposito") as escritor:
        respuesta = TestClient(app).post(
            "/deposito/vacios/conteo",
            data={"proveedor_id": "7", "cantidad": "0", "fecha": "2026-09-20"},
            follow_redirects=False)
    assert respuesta.status_code == 303, respuesta.text[:300]
    assert escritor.called
    destino = respuesta.headers["location"]
    assert destino.startswith("/deposito/vacios/conteo?aviso="), destino
    assert "cotejo" not in destino


def test_DEPOSITO_no_tiene_ninguna_otra_pantalla_de_vacios():
    """Solo cuenta y devuelve: el índice, el stock, el cotejo y el detalle de un
    proveedor muestran el número del sistema, y no existen bajo /deposito."""
    rutas = {r.path for r in app.routes if getattr(r, "path", "").startswith("/deposito/vacios")}
    assert rutas == {"/deposito/vacios/conteo", "/deposito/vacios/devolucion"}


# ── Depósito DEVUELVE sin ver el stock (dueño, 29/09) ───────────────────────

CLAVES = {"CLAVE_COMPRAS": "compras-secreta", "CLAVE_ADMINISTRACION": "admin-secreta"}


def _devolucion_deposito(metodo="get", datos=None, archivo=True, escritor=None):
    """La pantalla y el POST de Depósito con un cliente SIN cookie y las claves
    PUESTAS: si alguna ruta pasara por una puerta, contestaría 401."""
    limpio = TestClient(app)
    parches = [
        patch.dict(os.environ, CLAVES),
        patch("app.main.stock_de_vacios_deposito",
              side_effect=AssertionError("la devolución de Depósito leyó el stock")),
        patch("app.main.listar_proveedores",
              return_value=[{"id": 7, "nombre": "Puesto EJEMPLO"}]),
        patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: MARCAS}),
        patch("app.main.senas_por_cajon_de_todas_las_pilas", return_value={7: {71: 800.0}}),
        patch("app.main._comprimir_foto_jpeg", return_value=b"jpg"),
        patch("app.main.subir_foto_comanda", return_value="vacios/vale.jpg"),
        patch("app.main.crear_devolucion_vacios", **(escritor or {"return_value": 1})),
    ]
    from contextlib import ExitStack
    with ExitStack() as pila:
        mocks = {p.attribute: pila.enter_context(p) for p in parches[1:]}
        pila.enter_context(parches[0])
        if metodo == "get":
            return limpio.get("/deposito/vacios/devolucion"), mocks
        files = {"foto": ("vale.jpg", b"xx", "image/jpeg")} if archivo else None
        return limpio.post("/deposito/vacios/devolucion", data=datos, files=files,
                           follow_redirects=False), mocks


def test_la_DEVOLUCION_de_Deposito_abre_sin_clave_y_no_muestra_el_stock():
    """El selector de marca de Administración dice "Marca · hay 420": éste no.
    Y el stock ni se lee — el parche explota si la pantalla lo pide."""
    respuesta, _ = _devolucion_deposito()
    assert respuesta.status_code == 200, respuesta.text[:300]
    marcado = respuesta.text.split("</style>")[-1]
    assert 'action="/deposito/vacios/devolucion"' in marcado
    for campo in ('name="proveedor_id"', 'name="marca_vacio_id"', 'name="cantidad"',
                  'name="importe"', 'name="foto"'):
        assert campo in marcado, campo
    assert " · hay " not in marcado
    assert "/administracion/" not in marcado and "/compras/" not in marcado


def test_la_puerta_de_Depósito_es_la_UNICA_abierta_EL_CONTROL():
    """Con las mismas claves, las otras dos puertas rebotan. Sin esto, el 200 de
    arriba se daría también con las puertas apagadas."""
    limpio = TestClient(app)
    with patch.dict(os.environ, CLAVES):
        cerradas = [limpio.post(f"{s}/vacios/7/devolucion", data={}).status_code
                    for s in ("/compras", "/administracion")]
    assert cerradas == [401, 401], cerradas


def test_devolver_DE_MAS_en_Deposito_manda_a_Administracion_SIN_el_numero():
    from app.db import DevolucionDeMas
    respuesta, _ = _devolucion_deposito(
        "post", {"proveedor_id": "7", "marca_vacio_id": "71", "cantidad": "500"},
        escritor={"side_effect": DevolucionDeMas("En la pila EJ Roja hay 420 cajones: "
                                                  "no se pueden devolver 500.")})
    assert respuesta.status_code == 400
    marcado = respuesta.text.split("</style>")[-1]
    assert "avisale a Administración" in marcado
    assert "420" not in marcado and "hay 420" not in respuesta.text


def test_devolver_DE_MAS_en_Administracion_SI_dice_el_numero():
    """El control del de arriba: la misma excepción por la otra puerta muestra el
    número, que es lo que distingue el mensaje de Depósito de un mensaje roto."""
    from app.db import DevolucionDeMas
    with patch.dict(os.environ, CLAVES), \
         patch("app.main._comprimir_foto_jpeg", return_value=b"jpg"), \
         patch("app.main.subir_foto_comanda", return_value="vacios/vale.jpg"), \
         patch("app.main.crear_devolucion_vacios",
               side_effect=DevolucionDeMas("En la pila EJ Roja hay 420 cajones: "
                                           "no se pueden devolver 500.")), \
         patch("app.main._renderizar_vacios_proveedor",
               side_effect=lambda req, pid, **kw: HTMLResponse(kw["error"], kw["status_code"])):
        cliente.cookies.set(PUERTA_ADMINISTRACION.cookie,
                            PUERTA_ADMINISTRACION.firma("admin-secreta"))
        try:
            respuesta = cliente.post("/administracion/vacios/7/devolucion",
                                     data={"marca_vacio_id": "71", "cantidad": "500"},
                                     files={"foto": ("vale.jpg", b"xx", "image/jpeg")})
        finally:
            cliente.cookies.clear()
    assert respuesta.status_code == 400
    assert "hay 420 cajones" in respuesta.text


def test_una_DEVOLUCION_de_Deposito_guarda_por_la_misma_escritura_y_vuelve_a_devolver():
    respuesta, mocks = _devolucion_deposito(
        "post", {"proveedor_id": "7", "marca_vacio_id": "71", "cantidad": "5",
                 "importe": "4.000"})
    assert respuesta.status_code == 303, respuesta.text[:300]
    assert respuesta.headers["location"].startswith("/deposito/vacios/devolucion?aviso=")
    mocks["crear_devolucion_vacios"].assert_called_once_with(
        7, 71, 5, foto_ruta="vacios/vale.jpg", importe=4000.0, cargada_desde="deposito",
        numero_vale="")


def test_una_DEVOLUCION_de_Deposito_SIN_FOTO_no_se_guarda():
    respuesta, mocks = _devolucion_deposito(
        "post", {"proveedor_id": "7", "marca_vacio_id": "", "cantidad": "5"}, archivo=False)
    assert respuesta.status_code == 400
    assert "Sin la foto del vale" in respuesta.text
    mocks["crear_devolucion_vacios"].assert_not_called()


def test_una_DEVOLUCION_de_Deposito_sin_PROVEEDOR_no_se_guarda():
    respuesta, mocks = _devolucion_deposito("post", {"proveedor_id": "", "cantidad": "5"})
    assert respuesta.status_code == 400
    assert "Elegí un proveedor" in respuesta.text
    mocks["crear_devolucion_vacios"].assert_not_called()


# ── Parte 3 (dueño, 29/09): por dónde entró, y la pantalla de Movimientos ────

from core.movimientos_vacios import (  # noqa: E402
    TEXTO_DE_LA_PUERTA, TEXTO_DEL_TIPO, generar_excel_movimientos_vacios_deposito,
    texto_de_cantidad, ventana,
)


def test_las_PUERTAS_de_la_devolucion_son_las_del_CHECK_de_la_migracion():
    """La lista se LEE del .sql, no se copia: copiada envejece en silencio."""
    from app.db import PUERTAS_DE_LA_DEVOLUCION
    for archivo in ("db/vacios_origen_devolucion_1.sql", "db/esquema_completo.sql"):
        texto = io.open(archivo, encoding="utf-8").read()
        lista = re.search(r"cargada_desde in \(([^)]*)\)", texto).group(1)
        assert tuple(re.findall(r"'(\w+)'", lista)) == PUERTAS_DE_LA_DEVOLUCION, archivo
    # Y cada puerta tiene su palabra en la pantalla y el Excel.
    assert set(PUERTAS_DE_LA_DEVOLUCION) <= set(TEXTO_DE_LA_PUERTA)


def test_una_PUERTA_desconocida_no_llega_ni_a_abrir_la_base():
    with patch("app.db.obtener_conexion") as abrir:
        with pytest.raises(ValueError, match="Puerta desconocida"):
            crear_devolucion_vacios(7, None, 1, foto_ruta="v.jpg", cargada_desde="gerencia")
    abrir.assert_not_called()


def test_la_puerta_NO_tiene_default_en_la_escritura():
    """Un camino nuevo que se olvide de decirla revienta, no guarda un NULL que
    se lee igual que una devolución vieja."""
    import inspect
    parametro = inspect.signature(crear_devolucion_vacios).parameters["cargada_desde"]
    assert parametro.default is inspect.Parameter.empty
    assert parametro.kind is inspect.Parameter.KEYWORD_ONLY


HOY_MOVIMIENTOS = date(2026, 3, 20)


@pytest.mark.parametrize("desde, hasta, esperado", [
    ("", "", (date(2026, 2, 19), HOY_MOVIMIENTOS, None)),              # 30 días, los dos incluidos
    ("2026-01-01", "2026-03-31", (date(2026, 1, 1), date(2026, 3, 31), None)),   # 90 justos
])
def test_la_VENTANA_por_defecto_son_30_dias_y_entran_90(desde, hasta, esperado):
    assert ventana(desde, hasta, HOY_MOVIMIENTOS) == esperado


@pytest.mark.parametrize("desde, hasta, error", [
    ("2026-01-01", "2026-04-01", "Hasta 90 días"),                     # 91: el de al lado del corte
    ("2026-03-10", "2026-03-01", "posterior"),
    ("ayer", "", "no se entiende"),
])
def test_la_VENTANA_rechaza_lo_que_no_se_puede_pedir(desde, hasta, error):
    assert error in ventana(desde, hasta, HOY_MOVIMIENTOS)[2]


def test_el_SIGNO_de_cada_tipo():
    assert texto_de_cantidad(_mov("entrada", 1, 10)) == "+10"
    assert texto_de_cantidad(_mov("devolucion", 1, -5)) == "-5"
    assert texto_de_cantidad(_mov("ajuste", 1, 2)) == "+2"
    assert texto_de_cantidad(_mov("ajuste", 1, -2)) == "-2"
    assert texto_de_cantidad(_mov("asignacion", 1, 6)) == "6"
    assert texto_de_cantidad(_mov("arranque", 1, 40)) == "40"
    assert set(TEXTO_DEL_TIPO) == set(__import__("app.db").db.TIPOS_DE_MOVIMIENTO_DE_VACIOS)


MOVS_PANTALLA = [
    _mov("devolucion", 3, -20, marca="EJ Roja", foto_ruta="v.jpg", importe=30000.0,
         cargada_desde="deposito"),
    _mov("asignacion", 5, 15, marca_hasta="EJ Azul", cargada_desde="administracion"),
    _mov("entrada", 2, 10, marca="EJ Roja", compra_id=1234, articulo="EJ Pera",
         cargada_desde="recepcion"),
    _mov("ajuste", 4, -3, motivo="EJ rotos", proveedor_id=8, proveedor="Puesto EJEMPLO Dos",
         cargada_desde="administracion"),
    _mov("devolucion", 6, -1, cargada_desde=None),
]


def _movimientos(url, movimientos=MOVS_PANTALLA):
    with patch("app.main.movimientos_de_vacios", return_value=[dict(m) for m in movimientos]) as leer, \
         patch("app.main.listar_proveedores", return_value=[{"id": 7, "nombre": "Puesto EJEMPLO"},
                                                             {"id": 8, "nombre": "Puesto EJEMPLO Dos"}]), \
         patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: MARCAS}), \
         patch("app.main._hoy_argentina", return_value=HOY_MOVIMIENTOS):
        respuesta = cliente.get(url)
    return respuesta, leer


def test_MOVIMIENTOS_arranca_en_los_ultimos_30_dias_y_pide_eso_a_la_base():
    respuesta, leer = _movimientos("/administracion/vacios/movimientos")
    assert respuesta.status_code == 200
    marcado = respuesta.text.split("</style>")[-1]
    assert 'value="2026-02-19"' in marcado and 'value="2026-03-20"' in marcado
    kwargs = leer.call_args.kwargs
    assert (kwargs["desde"], kwargs["hasta"], kwargs["marca"]) == (date(2026, 2, 19), HOY_MOVIMIENTOS, None)
    assert "5 movimientos" in marcado


def test_MOVIMIENTOS_cada_fila_dice_proveedor_marca_signo_y_POR_DONDE_ENTRO():
    marcado = _movimientos("/administracion/vacios/movimientos")[0].text.split("</style>")[-1]
    filas = re.findall(r'<div class="vale[^"]*" data-tipo="\w+">(.*?)</div>\s*(?:<a|</div>|<div class="vale)',
                       marcado, re.S)
    cabezas = [" ".join(re.sub(r"<[^>]+>", " ", c).split())
               for c in re.findall(r'<div class="vale-cabeza">(.*?)</div>', marcado, re.S)]
    assert cabezas == ["-20 Devolución · Puesto EJEMPLO · EJ Roja",
                       "15 Pase · Puesto EJEMPLO · Sin marca → EJ Azul",
                       "+10 Compra con seña · Puesto EJEMPLO · EJ Roja",
                       "-3 Ajuste · Puesto EJEMPLO Dos · Sin marca",
                       "-1 Devolución · Puesto EJEMPLO · Sin marca"], filas
    datos = [" ".join(re.sub(r"<[^>]+>", " ", d).split())
             for d in re.findall(r'<div class="vale-dato">(.*?)</div>', marcado, re.S)]
    assert [d.split(" · ")[1] for d in datos] == [
        "Depósito", "Administración", "Recepción", "Administración", "sin dato"]
    assert "compra #1234" in datos[2]


def test_MOVIMIENTOS_linkea_el_REGISTRO_y_nunca_una_pantalla_de_Compras():
    """La compra no se linkea: su detalle vive bajo la clave de Compras (corolario 56)."""
    marcado = _movimientos("/administracion/vacios/movimientos")[0].text.split("</style>")[-1]
    hrefs = re.findall(r'<a class="ir" href="([^"]+)"', marcado)
    assert hrefs == ["/administracion/vacios/devolucion/3/foto?proveedor_id=7",
                     "/administracion/vacios/7?abrir=movimientos",
                     "/administracion/vacios/8?abrir=movimientos",
                     "/administracion/vacios/7?abrir=movimientos"]
    assert "/compras/" not in marcado


def test_MOVIMIENTOS_con_mas_de_90_dias_no_consulta_y_lo_dice():
    respuesta, leer = _movimientos("/administracion/vacios/movimientos?desde=2026-01-01&hasta=2026-04-01")
    assert respuesta.status_code == 400
    assert "Hasta 90 días por búsqueda" in respuesta.text
    leer.assert_not_called()
    assert "/movimientos/excel" not in respuesta.text.split("</style>")[-1]


def test_MOVIMIENTOS_filtra_por_proveedor_y_marca_y_el_EXCEL_lleva_los_MISMOS_filtros():
    url = "/administracion/vacios/movimientos?desde=2026-03-01&hasta=2026-03-15&proveedor_id=7&marca=71"
    respuesta, leer = _movimientos(url)
    assert leer.call_args.args[0] == 7
    assert leer.call_args.kwargs["marca"] == 71
    marcado = respuesta.text.split("</style>")[-1]
    assert '<option value="71" selected>EJ Roja</option>' in marcado
    excel = re.search(r'href="(/administracion/vacios/movimientos/excel\?[^"]+)"', marcado).group(1)
    with patch("app.main.movimientos_de_vacios", return_value=[dict(m) for m in MOVS_PANTALLA]) as leer_excel, \
         patch("app.main.listar_proveedores", return_value=[{"id": 7, "nombre": "Puesto EJEMPLO"}]), \
         patch("app.main.listar_marcas_vacio_por_proveedor", return_value={7: MARCAS}):
        archivo = cliente.get(excel.replace("&amp;", "&"))
    assert archivo.status_code == 200
    assert leer_excel.call_args.args == (7, None)          # el Excel no tiene tope
    assert {k: leer_excel.call_args.kwargs[k] for k in ("desde", "hasta", "marca")} == \
        {k: leer.call_args.kwargs[k] for k in ("desde", "hasta", "marca")}


def test_MOVIMIENTOS_sin_proveedor_no_ofrece_marcas_y_sin_marca_se_pide_aparte():
    marcado = _movimientos("/administracion/vacios/movimientos")[0].text.split("</style>")[-1]
    assert re.search(r'<select id="marca" name="marca" disabled>', marcado)
    _, leer = _movimientos("/administracion/vacios/movimientos?proveedor_id=7&marca=sin")
    assert leer.call_args.kwargs["marca"] == "sin"


def test_MOVIMIENTOS_avisa_cuando_CORTA_la_lista():
    muchos = [_mov("ajuste", i, 1) for i in range(501)]
    marcado = _movimientos("/administracion/vacios/movimientos", muchos)[0].text.split("</style>")[-1]
    assert "500+ movimientos" in marcado
    assert "Se muestran los 500 más nuevos" in marcado
    assert marcado.count('data-tipo="ajuste"') == 500


def test_el_EXCEL_de_movimientos_dice_lo_MISMO_que_la_pantalla():
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(generar_excel_movimientos_vacios_deposito(
        date(2026, 3, 1), date(2026, 3, 15), "Puesto EJEMPLO", [dict(m) for m in MOVS_PANTALLA])))
    hoja = libro.active
    filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=5)]
    assert [f[1] for f in filas] == ["Devolución", "Pase", "Compra con seña", "Ajuste", "Devolución"]
    assert [f[4] for f in filas] == [-20, 15, 10, -3, -1]
    assert [f[5] for f in filas] == ["Depósito", "Administración", "Recepción", "Administración", "sin dato"]
    assert filas[2][6] == 1234 and filas[3][7] == "EJ rotos" and filas[0][8] == 30000.0
    assert "Del 01/03/2026 al 15/03/2026 · Puesto EJEMPLO" == hoja.cell(row=2, column=1).value


def test_el_INDICE_de_Administracion_tiene_los_TRES_botones():
    with patch("app.main.stock_de_vacios_deposito", return_value=list(UN_PROVEEDOR)):
        admin = cliente.get("/administracion/vacios").text.split("</style>")[-1]
    botones = re.findall(r'<a class="boton-stock" href="([^"]+)">([^<]+)</a>', admin)
    assert botones == [("/administracion/vacios/stock", "Exportar"),
                       ("/administracion/vacios/cotejo", "Cotejo"),
                       ("/administracion/vacios/movimientos", "Movimientos")]


# ── El detalle sin lo que está en cero (dueño, 30/09) ──────────────────────

def _resumen(marcado):
    """La tarjeta del resumen: de la primera tarjeta hasta las acciones."""
    return marcado.split('<details class="accion"')[0].rsplit('<div class="tarjeta">', 1)[1]


def test_el_detalle_NO_muestra_las_pilas_en_cero_y_SI_la_sin_marca_con_cajones():
    proveedor = [{"id": 7, "nombre": "Puesto EJEMPLO", "stock": 17, "pilas": [
        _pila(None, None, 12), _pila(71, "EJ Roja", 0), _pila(72, "EJ Azul", 5)]}]
    with _con(_parches_del_detalle(proveedores=proveedor)):
        marcado = cliente.get("/administracion/vacios/7").text.split("</style>")[-1]
    resumen = _resumen(marcado)
    assert re.findall(r'class="pila-marca[^"]*">([^<]+)<', resumen) == ["Sin marca", "EJ Azul"]
    assert "Sin vacíos" not in resumen
    # Las listas para ELEGIR siguen con la roja en cero: es a dónde se pasan
    # los cajones sin marca.
    assert re.search(r'<option value="71"[^>]*>[^<]*EJ Roja', marcado)


def test_la_SIN_MARCA_en_cero_no_aparece_y_un_NEGATIVO_si_en_rojo():
    proveedor = [{"id": 7, "nombre": "Puesto EJEMPLO", "stock": -2, "pilas": [
        _pila(None, None, 0), _pila(71, "EJ Roja", -2)]}]
    with _con(_parches_del_detalle(proveedores=proveedor)):
        resumen = _resumen(cliente.get("/administracion/vacios/7").text.split("</style>")[-1])
    assert re.findall(r'class="pila-marca[^"]*">([^<]+)<', resumen) == ["EJ Roja"]
    assert '<span class="pila-stock negativo">-2</span>' in resumen


def test_un_proveedor_SIN_NADA_dice_Sin_vacios():
    proveedor = [{"id": 7, "nombre": "Puesto EJEMPLO", "stock": 0, "pilas": [
        _pila(None, None, 0), _pila(71, "EJ Roja", 0)]}]
    with _con(_parches_del_detalle(proveedores=proveedor)):
        resumen = _resumen(cliente.get("/administracion/vacios/7").text.split("</style>")[-1])
    assert '<div class="stock">Sin vacíos</div>' in resumen
    assert "pila-marca" not in resumen
