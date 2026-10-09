# -*- coding: utf-8 -*-
"""VALES POR PROVEEDOR (dueño, 09/10).

Un renglón por proveedor con los vales EN CARTERA (ni cobrados ni aplicados a
una liquidación): cuántos y cuánto, de mayor a menor importe, y el total
general abajo. En Administración, en Gerencia y en el detalle del cuadro del
Panel de control. La suma por proveedor tiene que dar el número del panel.

Con el RIVAL plantado: un vale COBRADO del proveedor que más tiene, que no
puede sumar en ningún lado.
"""
import io
import os
import re
import sys
from datetime import date
from unittest.mock import patch

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

from tests.test_vales_a_cobrar import HOY, _anterior, _cliente, _devolucion, base  # noqa: E402,F401


def _cartera(base):
    """EJ Uno: $1.500 y $2.000 (2 vales). EJ Dos: $5.000 (1 vale) y uno de
    $99.999 COBRADO, que es el rival: con él EJ Dos tendría 2 vales."""
    d, sql = base
    _anterior(sql, 1500, date(2026, 8, 1), proveedor=1, numero="P-1")
    _devolucion(d, importe=2000.0)
    _anterior(sql, 5000, date(2026, 11, 20), proveedor=2, numero="P-2")
    cobrado = _anterior(sql, 99999, date(2026, 11, 1), proveedor=2, numero="P-3")
    d.registrar_salida_de_vale(cobrado, "cobrado", date(2026, 11, 30), sector="administracion", hoy=HOY,
                               importe_cobrado=99999)
    return d, sql


def _marcado(respuesta):
    """El cuerpo de la página: la tarjeta agrupada trae su propio <style>,
    así que no se corta en el último </style> como en las otras pantallas."""
    return respuesta.text.split("</head>", 1)[1]


def _renglones(marcado):
    """(proveedor, cantidad, importe) de cada renglón, en el orden en que se ven."""
    bloque = marcado.split("data-por-proveedor>", 1)[1].split("data-total-general>", 1)
    renglones = re.findall(r'<span class="fila-que">([^<]+)</span><br>\s*<span class="fila-dato">(\d+) vales?</span>'
                           r'</span>\s*<span class="importe">([^<]+)</span>', bloque[0])
    total = re.search(r"<span>Total · (\d+) vales?</span><span>([^<]+)</span>", bloque[1])
    return renglones, (total.group(1), total.group(2))


def test_la_cuenta_AGRUPA_la_misma_cartera_de_mayor_a_menor(base):
    from core.vales import vales_por_proveedor
    d, _ = _cartera(base)
    resumen = d.resumen_de_la_cartera(HOY)
    agrupados = vales_por_proveedor(resumen["vales"])
    assert [(g["proveedor"], g["cantidad"], g["total"]) for g in agrupados["filas"]] == \
        [("EJ Dos", 1, 5000.0), ("EJ Uno", 2, 3500.0)]
    assert agrupados["total"] == resumen["total"] == 8500.0
    assert agrupados["cantidad"] == resumen["cantidad"] == 3


def test_la_SUMA_por_proveedor_da_el_NUMERO_del_panel(base, monkeypatch):
    _cartera(base)
    cliente = _cliente(monkeypatch, "gerencia")
    with patch("app.main._hoy_argentina", return_value=HOY):
        panel = cliente.get("/gerencia/panel")
        detalle = cliente.get("/gerencia/panel/vales")
    assert panel.status_code == 200 and detalle.status_code == 200
    cuadro = _marcado(panel).split('data-cuadro="vales">', 1)[1].split("</a>", 1)[0]
    numero_del_panel = re.search(r'class="cuadro-numero">([^<]+)</div>', cuadro).group(1)
    renglones, (cantidad, total) = _renglones(_marcado(detalle))
    assert renglones == [("EJ Dos", "1", "$5.000"), ("EJ Uno", "2", "$3.500")]
    suma = sum(int(importe.strip("$").replace(".", "")) for _, _, importe in renglones)
    assert numero_del_panel == total == "$8.500" and suma == 8500 and cantidad == "3"
    # Tocar un proveedor abre la lista común con su filtro, y vuelve al detalle del panel.
    assert 'href="/gerencia/vales?proveedor_id=2&amp;volver=%2Fgerencia%2Fpanel%2Fvales"' in _marcado(detalle)


def test_ADMINISTRACION_y_GERENCIA_ven_la_vista_agrupada_y_vuelven_a_la_lista(base, monkeypatch):
    _cartera(base)
    for sector in ("administracion", "gerencia"):
        cliente = _cliente(monkeypatch, sector)
        with patch("app.main._hoy_argentina", return_value=HOY):
            lista = cliente.get(f"/{sector}/vales")
            agrupada = cliente.get(f"/{sector}/vales/por-proveedor")
            de_uno = cliente.get(f"/{sector}/vales?proveedor_id=1&volver=/{sector}/vales/por-proveedor")
        assert lista.status_code == agrupada.status_code == de_uno.status_code == 200, sector
        assert f'<a href="/{sector}/vales/por-proveedor">Vales por proveedor</a>' in _marcado(lista)
        marcado = _marcado(agrupada)
        assert _renglones(marcado) == ([("EJ Dos", "1", "$5.000"), ("EJ Uno", "2", "$3.500")], ("3", "$8.500"))
        assert f'<a href="/{sector}/vales">Lista común</a>' in marcado
        assert f'href="/{sector}/vales?proveedor_id=1&amp;volver=%2F{sector}%2Fvales%2Fpor-proveedor"' in marcado
        # La lista de UN proveedor: sus dos vales, y "Atrás" vuelve a la vista agrupada.
        marcado = _marcado(de_uno)
        assert "2 vales · $3.500" in marcado and "EJ Dos" not in marcado.split("<h2>", 1)[1]
        assert f'<a class="barra-boton" href="/{sector}/vales/por-proveedor" aria-label="Volver atrás">' in de_uno.text
        assert f'<input type="hidden" name="volver" value="/{sector}/vales/por-proveedor">' in marcado


def _texto_del_pdf(contenido: bytes) -> str:
    import pypdfium2 as pdfium
    return "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(contenido))


def test_el_PDF_y_el_EXCEL_bajan_la_vista_agrupada_y_dicen_que_se_muestra(base, monkeypatch):
    import openpyxl
    _cartera(base)
    cliente = _cliente(monkeypatch, "administracion")
    with patch("app.main._hoy_argentina", return_value=HOY):
        excel = cliente.get("/administracion/vales/por-proveedor/excel")
        pdf = cliente.get("/administracion/vales/por-proveedor/pdf")
    assert excel.status_code == pdf.status_code == 200
    hoja = openpyxl.load_workbook(io.BytesIO(excel.content)).active
    assert hoja.cell(row=2, column=1).value.startswith("Al 01/12/2026 · En cartera: ni cobrados ni aplicados")
    filas = [tuple(c.value for c in fila) for fila in hoja.iter_rows(min_row=5)]
    assert filas == [("EJ Dos", 1, 5000), ("EJ Uno", 2, 3500), ("Total", 3, 8500)]
    texto = _texto_del_pdf(pdf.content)
    assert "En cartera: ni cobrados ni aplicados a una liquidación" in texto
    orden = [texto.index(x) for x in ("EJ Dos", "$5.000", "EJ Uno", "$3.500", "Total", "$8.500")]
    assert orden == sorted(orden) and "99.999" not in texto


def test_ningun_IMPORTE_se_parte_a_313px_ni_con_un_nombre_largo(base, monkeypatch):
    """Un número nunca se parte (dueño, 09/10). En la lista común, con un
    proveedor de nombre largo, "$8.379.000.000" salía en tres renglones."""
    import pytest
    pytest.importorskip("playwright", reason="lo que se ve lo decide el navegador")
    from playwright.sync_api import sync_playwright
    from scripts.medir_layout import CHROMIUM
    d, sql = base
    sql("UPDATE proveedores SET nombre = 'EJEMPLO Proveedor con un nombre bastante largo SRL' WHERE id = 2")
    _anterior(sql, 8379000000, date(2026, 11, 1), proveedor=2)
    _anterior(sql, 1500, date(2026, 11, 1), proveedor=1)
    cliente = _cliente(monkeypatch, "administracion", "gerencia")
    with patch("app.main._hoy_argentina", return_value=HOY):
        paginas = {r: cliente.get(r).text for r in
                   ("/administracion/vales", "/administracion/vales/por-proveedor", "/gerencia/panel/vales")}
    medir = """() => {
      const importes = [...document.querySelectorAll('.fila-cabeza .importe, [data-por-proveedor] .importe')];
      // Los renglones que ocupa el TEXTO: un Range sobre él da un rectángulo por renglón.
      const renglones = n => { const r = document.createRange(); r.selectNodeContents(n);
        return new Set([...r.getClientRects()].filter(x => x.width > 0).map(x => Math.round(x.top))).size; };
      return {cuantos: importes.length,
              partidos: importes.filter(n => renglones(n) > 1).map(n => n.textContent),
              scroll: document.documentElement.scrollWidth - document.documentElement.clientWidth};
    }"""
    with sync_playwright() as pw:
        navegador = pw.chromium.launch(executable_path=CHROMIUM)
        try:
            for ruta, html in paginas.items():
                pagina = navegador.new_page(viewport={"width": 313, "height": 900})
                pagina.set_content(html)
                m = pagina.evaluate(medir)
                pagina.close()
                assert m["cuantos"] >= 2, (ruta, m)
                assert m["partidos"] == [] and m["scroll"] <= 0, (ruta, m)
        finally:
            navegador.close()
